"""The single agent-facing entry point. All names and data below are synthetic.

Rendering is replaced by a fake page image here; real Office rendering is
exercised by run_acceptance.py. Review texts are unit-test records, not real checks.
"""

from __future__ import annotations
import contextlib, hashlib, io, json, sys, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import xaufe as X
import delivery as D
import intake as I
from run_acceptance import original_text

JOB_B = {
    "task": "format",
    "template": "B",
    "info": {
        "course_name": "数据结构",
        "title_zh": "排序算法比较",
        "student_name": "合成同学",
        "student_id": "20230001",
        "major": "软件工程",
        "class_name": "软工2301",
    },
}
CONTENT = {
    "mode": "paper",
    "metadata": {"title_zh": "排序算法比较"},
    "abstract_zh": "这是合成摘要。",
    "keywords_zh": ["排序", "算法"],
    "references": ["合成作者. 合成文献[M]. 北京: 合成出版社, 2020."],
    "blocks": [
        {"type": "heading", "level": 1, "text": "一、引言"},
        {"type": "paragraph", "text": "这是合成正文。"},
        {"type": "heading", "level": 2, "text": "（一）背景"},
        {"type": "paragraph", "text": "这是合成正文第二段。"},
    ],
}


def fake_render(docx, outdir):
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    png = outdir / "page-001.png"
    png.write_bytes(b"unit-test-only-not-a-real-page")
    result = {
        "source_sha256": hashlib.sha256(Path(docx).read_bytes()).hexdigest(),
        "pages": [{"page": 1, "png": png.name, "png_sha256": hashlib.sha256(png.read_bytes()).hexdigest()}],
    }
    D.write_json(outdir / "render-report.json", result)
    return result


class CLI(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        # Synthetic renders below write no PDF; the rendered-page audit has its own tests
        # (tests/test_render_audit.py) and runs for real in tests/run_acceptance.py.
        audit = patch("delivery.audit_rendered", return_value={"passed": True, "errors": [], "warnings": [], "pages": 1, "fonts": {}})
        audit.start()
        self.addCleanup(audit.stop)

    def tearDown(self):
        self.tmp.cleanup()

    def call(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            code = X.main([str(a) for a in args])
        return code, out.getvalue(), err.getvalue()

    def new(self, job=None, content=None, sources=()):
        extra = []
        if content is not None and not sources and (job or {}).get("task", "format") == "format":
            sources = [self.d / f"原稿-{len(list(self.d.glob('原稿-*')))}.txt"]
            sources[0].write_text(original_text(content), encoding="utf-8")
        for s in sources:
            extra += ["--source", s]
        with patch("tempfile.gettempdir", return_value=str(self.d)):
            code, out, err = self.call("new", "--out", self.d / "out" / "成品.docx", *extra)
        self.assertEqual(code, 0, err)
        work = Path(out.split("WORK = ")[1].splitlines()[0].strip())
        if job is not None:
            job = {"task": "format", **job}
            (work / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        if content is not None:
            (work / "content.json").write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        return work

    def fill_review(self, work, vision=True):
        review = json.loads((work / "review.json").read_text(encoding="utf-8"))
        review["vision"] = vision
        review["pages"] = {k: "单元测试合成记录：这一页正常" for k in review["pages"]} if vision else {}
        if not vision:
            review["no_vision_reason"] = "单元测试：模拟不能看图"
        for key in ("content_check", "layout_check", "delivery_check"):
            review[key] = "单元测试合成记录：已核对"
        review["warnings"] = {k: "单元测试合成记录：已处理" for k in review["warnings"]}
        (work / "review.json").write_text(json.dumps(review, ensure_ascii=False), encoding="utf-8")

    # ------------------------------------------------------------ choosing
    def test_compare_prints_user_table_without_internal_ids(self):
        code, out, _ = self.call("compare")
        self.assertEqual(code, 0)
        self.assertTrue(out.startswith("| 对比项目 |"))
        self.assertIn("请选择 **A（学年论文）** 或 **B（课程作业）**", out)
        for internal in ("green-paper", "red-course", "模板标识", "template"):
            self.assertNotIn(internal, out)

    def test_check_asks_only_missing_fields(self):
        job = {"template": "A", "logo": "red", "info": {"title_zh": "合成题目", "student_name": "合成同学"}}
        work = self.new(job)
        code, out, _ = self.call("check", "--work", work)
        self.assertEqual(code, X.EXIT_ASK_USER)
        self.assertIn("采用A（本科学年论文），配红色校徽。", out)
        self.assertIn("还缺：学号、专业、班级、指导教师。这些项也可以留空", out)
        self.assertNotIn("题目", out.split("还缺：")[1].split("。")[0])

    def test_blank_and_default_date(self):
        job = dict(JOB_B, blank=["student_id"])
        job["info"] = {k: v for k, v in JOB_B["info"].items() if k != "student_id"}
        work = self.new(job)
        code, out, _ = self.call("check", "--work", work)
        self.assertEqual(code, 0)
        self.assertIn("按要求留空：学号", out)
        self.assertIn("日期采用执行当日", out)

    def test_job_errors_are_plain_and_actionable(self):
        cases = [
            ({"template": "C"}, '"A" 或 "B"'),
            ({"template": "B", "info": {"supervisor": "某老师"}}, "cover_supervisor"),
            ({"template": "A", "format": {"logo": "red"}}, "顶层的 logo"),
            ({"template": "A", "colour": "red"}, "不认识的项"),
            ({"template": "A", "info": {"student_name": "甲"}, "blank": ["student_name"]}, "既填写又留空"),
        ]
        for job, message in cases:
            with self.subTest(message=message):
                work = self.new(job)
                code, _, err = self.call("check", "--work", work)
                self.assertEqual(code, 2)
                self.assertIn(message, err)

    def test_job_must_say_write_or_format(self):
        work = self.new()
        job = {k: v for k, v in JOB_B.items() if k != "task"}
        (work / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        code, _, err = self.call("check", "--work", work)
        self.assertEqual(code, 2)
        self.assertIn('"task": "write"', err)

    def test_job_written_as_utf16_or_gbk_is_read(self):
        for encoding in ("utf-16", "gbk"):
            with self.subTest(encoding=encoding):
                work = self.new()
                (work / "job.json").write_bytes(json.dumps(JOB_B, ensure_ascii=False).encode(encoding))
                code, out, err = self.call("check", "--work", work)
                self.assertEqual(code, 0, err)

    def test_null_in_format_means_default(self):
        state = I.from_job(
            dict(JOB_B, format={"styles": {"body": {"cn": "仿宋", "size": None}}}), self.d / "s.json"
        )
        self.assertEqual(state["format_request"], {"styles": {"body": {"cn": "仿宋"}}})

    def test_same_workspace_gives_same_task_id(self):
        a = I.from_job(JOB_B, self.d / "x.json")["job_id"]
        b = I.from_job(JOB_B, self.d / "x.json")["job_id"]
        self.assertEqual(a, b)

    # ------------------------------------------------------------ run / finish
    def test_run_refuses_when_info_missing(self):
        work = self.new({"template": "B"}, CONTENT)
        code, out, _ = self.call("run", "--work", work)
        self.assertEqual(code, X.EXIT_ASK_USER)
        self.assertIn("还缺", out)
        self.assertFalse((work / "draft.docx").exists())

    def test_static_run_and_finish_deliver_one_word(self):
        work = self.new(JOB_B, CONTENT)
        with patch("office_discovery.discover") as discover:
            code, out, err = self.call("run", "--work", work, "--office", "none")
        discover.assert_not_called()
        self.assertEqual(code, 0, err)
        self.assertIn("需要在最终回复里告诉用户", out)
        self.assertIn("正文约 15 字", out)
        self.assertIn("参考文献 1 条", out)
        review = json.loads((work / "review.json").read_text(encoding="utf-8"))
        self.assertEqual(review["pages"], {})
        code, out, err = self.call("finish", "--work", work)
        self.assertEqual(code, 2)
        self.assertIn("content_check", err)
        self.fill_review(work, vision=False)
        code, out, err = self.call("finish", "--work", work)
        self.assertEqual(code, 0, err)
        self.assertIn("必须在给用户的回复里说明", out)
        self.assertFalse(work.exists())
        self.assertEqual([p.name for p in (self.d / "out").iterdir()], ["成品.docx"])

    def test_rendered_run_lists_pages_and_needs_every_page(self):
        # No table of contents here: real TOC paging is covered by run_acceptance.py.
        document = {"mode": "document", "metadata": {"title_zh": "排序算法比较"}, "blocks": CONTENT["blocks"]}
        work = self.new(JOB_B, document)
        with patch("delivery.render", side_effect=fake_render), patch("office_discovery.discover", return_value={}):
            code, out, err = self.call("run", "--work", work)
        self.assertEqual(code, 0, err)
        self.assertIn("第1页：", out)
        review = json.loads((work / "review.json").read_text(encoding="utf-8"))
        self.assertIn("请填写", review["pages"]["1"])
        self.fill_review(work)
        with patch("delivery.check_preview_toc", return_value={"required": False}):
            code, out, err = self.call("finish", "--work", work)
        self.assertEqual(code, 0, err)
        self.assertIn("已完成排版、分页和逐页检查", out)

    def test_cover_overflow_retries_with_next_fit_level(self):
        from cover_date_audit import CoverOverflow

        work = self.new(JOB_B, CONTENT)
        levels = []
        real_build = X._build

        def build(root, level):
            levels.append(level)
            return real_build(root, level)

        calls = {"n": 0}

        def prepare(*args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 1:
                raise CoverOverflow("合成：日期被挤到第二页")
            return {"validation_scope": "static-only", "pages": [], "warnings": [], "user_notice": "合成说明"}

        with patch("xaufe._build", side_effect=build), patch("delivery.prepare", side_effect=prepare):
            code, out, err = self.call("run", "--work", work)
        self.assertEqual(code, 0, err)
        self.assertEqual(levels, [0, 1])
        self.assertIn("自动收紧封面空白", out)

    def test_cover_overflow_with_substituted_fonts_blames_the_fonts(self):
        from cover_date_audit import CoverOverflow

        work = self.new(JOB_B, CONTENT)

        def prepare(*args, **kwargs):
            exc = CoverOverflow("合成：学号挤到第二页")
            exc.missing_fonts, exc.rendered_fonts = ["宋体", "黑体"], ["hiraginosansgb"]
            raise exc

        with patch("delivery.prepare", side_effect=prepare):
            code, out, err = self.call("run", "--work", work)
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("COVER_FONTS_MISSING", out)
        self.assertIn("渲染时没有宋体、黑体", out)
        self.assertIn("run --office none", out)
        self.assertNotIn("同意缩小封面字号", out)
        issues = json.loads((work / "issues.json").read_text(encoding="utf-8"))
        self.assertEqual(issues["issues"][0]["code"], "COVER_FONTS_MISSING")

    def test_cover_too_long_is_a_counted_located_issue(self):
        from cover_date_audit import CoverOverflow

        work = self.new(JOB_B, CONTENT)
        with patch("delivery.prepare", side_effect=CoverOverflow("合成：学号挤到第二页")):
            code, out, err = self.call("run", "--work", work)
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("COVER_TOO_LONG", out)
        self.assertIn("format.cover_fonts", out)

    def test_both_content_and_mapping_is_refused(self):
        work = self.new(JOB_B, CONTENT)
        (work / "mapping.json").write_text("{}", encoding="utf-8")
        code, _, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 2)
        self.assertIn("只能保留一个", err)

    # ------------------------------------------------------------ existing Word
    def sample_docx(self):
        doc = Document()
        for text in [
            "排序算法比较",
            "摘要",
            "本文比较了几种排序算法。",
            "关键词：排序   算法",
            "一、引言",
            "这是正文。",
            "（一）背景",
            "1.问题",
            "正文内容很长，这一段不是标题。",
            "英文摘要待删除的合成段落",
            "参考文献",
            "[1] 合成作者. 合成文献[M]. 2020.",
        ]:
            doc.add_paragraph(text)
        path = self.d / "原稿.docx"
        doc.save(path)
        return path

    def test_read_prints_document_text_once(self):
        source = self.sample_docx()
        code, out, err = self.call("read", "--docx", source)
        self.assertEqual(code, 0, err)
        self.assertEqual(out.count("排序算法比较"), 1)
        self.assertIn("一、引言", out)
        self.assertIn("[1] 合成作者", out)

    def test_inspect_guesses_roles(self):
        work = self.new(JOB_B)
        source = self.sample_docx()
        code, out, err = self.call("inspect", "--work", work, "--docx", source)
        self.assertEqual(code, 0, err)
        roles = [r["role"] for r in json.loads((work / "mapping.json").read_text(encoding="utf-8"))["paragraphs"]]
        self.assertEqual(
            roles,
            ["title_zh", "abstract_label_zh", "abstract_zh", "keywords_zh", "h1", "body", "h2", "h3", "body", "body", "reference_title", "reference"],
        )
        self.assertIn("加上所选模板的封面", out)
        mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
        self.assertEqual(mapping["replace_cover_before"], 0)
        # The workspace records resolved paths (macOS temp dirs live behind /var -> /private/var).
        self.assertIn(str(source.resolve()), [x["path"] for x in D.load(work)[1]["source_documents"]])

    def test_mapping_delete_role_removes_only_marked_paragraph(self):
        work = self.new(JOB_B)
        source = self.sample_docx()
        self.call("inspect", "--work", work, "--docx", source)
        mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
        mapping["paragraphs"][9]["role"] = "delete"
        (work / "mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        # Format only: deleting original text needs the user's word, recorded in job.json.
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("SOURCE_MISSING", out)
        self.assertIn("英文摘要待删除", out)
        # A snippet does not approve a whole paragraph: the approved text must contain the change.
        job = {**JOB_B, "approved_changes": [{"text": "英文摘要待删除", "why": "用户说不要英文摘要"}]}
        (work / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        job = {**JOB_B, "approved_changes": [{"text": "英文摘要待删除的合成段落", "why": "用户说不要英文摘要"}]}
        (work / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, err)
        self.assertIn("删除了 1 段", out)
        text = "\n".join(p.text for p in Document(work / "final.docx").paragraphs)
        self.assertNotIn("英文摘要待删除", text)
        self.assertIn("正文内容很长", text)

    def test_format_only_blocks_rewritten_or_added_text_and_needs_an_original(self):
        content = json.loads(json.dumps(CONTENT))
        work = self.new(JOB_B, content)
        content["blocks"][1]["text"] = "这是改写过的合成正文。"
        (work / "content.json").write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("SOURCE_CHANGED", out)
        self.assertIn("实际：这是改写过的合成正文。；应为：这是合成正文。", out)
        bare = self.new(JOB_B, CONTENT, sources=[])
        D.load(bare)  # a workspace without a registered original
        state = D.load(bare)[1]
        state["source_documents"] = []
        D.write_json(bare / D.MARKER, state)
        code, out, err = self.call("run", "--work", bare, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("SOURCE_REQUIRED", out)

    def test_wrong_role_is_refused_not_silently_formatted(self):
        work = self.new(JOB_B)
        source = self.sample_docx()
        self.call("inspect", "--work", work, "--docx", source)
        mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
        row = next(r for r in mapping["paragraphs"] if r["role"] == "body" and len(r["text"]) > 4)
        row["role"] = "spacer"
        (work / "mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("ROLE_MISMATCH", out)
        self.assertIn(row["id"], out)

    def test_plain_existing_word_numbers_the_body_from_one_after_the_cover(self):
        from docx.oxml.ns import qn

        doc = Document()
        for i in range(3):
            doc.add_paragraph(f"第{i + 1}段合成正文，没有摘要和标题。")
        source = self.d / "普通文档.docx"
        doc.save(source)
        work = self.new({**JOB_B, "format": {"footer_page_numbers": True}})
        code, out, err = self.call("inspect", "--work", work, "--docx", source)
        self.assertEqual(code, 0, err)
        self.assertEqual(json.loads((work / "mapping.json").read_text(encoding="utf-8"))["mode"], "document")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        page = Document(work / "final.docx").sections[1]._sectPr.find(qn("w:pgNumType"))
        self.assertEqual((page.get(qn("w:fmt")), page.get(qn("w:start"))), ("decimal", "1"))

    def test_gbk_job_and_original_still_guard_the_text(self):
        content = {"mode": "document", "metadata": {"title_zh": "排序算法比较"},
                   "blocks": [{"type": "paragraph", "text": "本段用于格式校准，不提供结论。"}]}
        source = self.d / "原稿-gbk.txt"
        source.write_bytes("本段用于格式校准，不提供结论。".encode("gbk"))
        work = self.new(JOB_B, content, sources=[source])
        (work / "job.json").write_bytes(json.dumps(JOB_B, ensure_ascii=False).encode("gbk"))
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        content["blocks"][0]["text"] = "本段用来格式校准，不提供结论。"
        (work / "content.json").write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, out + err)
        self.assertIn("SOURCE_CHANGED", out)

    def test_heading_marked_as_body_is_refused(self):
        work = self.new(JOB_B)
        self.call("inspect", "--work", work, "--docx", self.sample_docx())
        mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
        row = next(r for r in mapping["paragraphs"] if r["text"] == "（一）背景")
        row["role"] = "body"
        (work / "mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("ROLE_MISMATCH", out)
        self.assertIn("像是标题", out)

    def test_reference_entry_marked_as_body_is_refused(self):
        def body_reference(mapping):
            next(r for r in mapping["paragraphs"] if r["text"].startswith("[1]"))["role"] = "body"

        work, _, code, out, err = self.paper_run("B", edit=body_reference)
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("参考文献条目应为 reference", out)

    def test_body_paragraph_marked_as_reference_is_refused(self):
        def body_as_reference(mapping):
            next(r for r in mapping["paragraphs"] if r["text"] == "这是正文。")["role"] = "reference"

        work, _, code, out, err = self.paper_run("B", edit=body_as_reference)
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("在参考文献标题前面", out)

    def test_body_paragraph_marked_as_abstract_is_refused(self):
        def body_as_abstract(mapping):
            next(r for r in mapping["paragraphs"] if r["text"] == "正文内容很长，这一段不是标题。")["role"] = "abstract_zh"

        work, _, code, out, err = self.paper_run("B", edit=body_as_abstract)
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("“正文内容很长，这一段不是标题。”（role=abstract_zh）", out)
        self.assertIn("role 可能标错了", out)

    def test_document_title_must_match_the_cover_title(self):
        def title(mapping):
            next(r for r in mapping["paragraphs"] if r["text"] == "排序算法比较")["role"] = "title_zh"

        info = dict(JOB_B["info"], title_zh="另一个题目")
        work = self.new({"template": "B", "info": info})
        code, out, err = self.call("inspect", "--work", work, "--docx", self.sample_docx())
        mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
        mapping["paragraphs"][9]["role"] = "body"
        title(mapping)
        (work / "mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("info.title_zh“另一个题目”不一致", out)
        self.assertIn("问用户哪个对", out)

    def test_dropped_original_footer_is_reported(self):
        source = self.sample_docx()
        doc = Document(source)
        doc.sections[0].footer.paragraphs[0].text = "第 1 页"
        page = OxmlElement("w:fldSimple")
        page.set(qn("w:instr"), " PAGE ")
        doc.sections[0].footer.paragraphs[0]._p.append(page)
        doc.sections[0].header.paragraphs[0].text = "合成学院课程作业"
        doc.save(source)
        work = self.new(JOB_B)
        code, out, err = self.call("inspect", "--work", work, "--docx", source)
        mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
        mapping["paragraphs"][9]["role"] = "body"
        (work / "mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        self.assertIn("ORIGINAL_HEADER_FOOTER_REPLACED", out)
        self.assertIn("合成学院课程作业", out)
        self.assertIn("ORIGINAL_PAGE_NUMBERS_DROPPED", out)
        self.assertIn("footer_page_numbers", out)

    def test_tiny_margins_are_flagged_for_the_review(self):
        document = {"mode": "document", "metadata": {"title_zh": "排序算法比较"}, "blocks": CONTENT["blocks"]}
        work = self.new({**JOB_B, "format": {"page_twips": {"left": 200}}}, document)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        self.assertIn("MARGIN_TOO_SMALL", out)

    def test_self_defeating_styles_are_flagged_for_the_review(self):
        document = {"mode": "document", "metadata": {"title_zh": "排序算法比较"}, "blocks": CONTENT["blocks"]}
        job = {**JOB_B, "format": {"styles": {"body": {"color": "FFFFFF"}, "h1": {"size": 1}}}}
        work = self.new(job, document)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        self.assertIn("STYLE_HARD_TO_READ", out)
        self.assertIn("FFFFFF", out)

    def test_format_only_keeps_links_and_figures_of_the_original(self):
        source = self.d / "带链接.md"
        source.write_text("# 一、引言\n\n详见[学校官网](https://www.example.edu.cn/a)。", encoding="utf-8")
        content = {"mode": "document", "metadata": {"title_zh": "排序算法比较"}, "blocks": [
            {"type": "heading", "level": 1, "text": "一、引言"},
            {"type": "paragraph", "runs": [{"text": "详见"}, {"text": "学校官网", "link": "https://www.example.edu.cn/a"}, {"text": "。"}]}]}
        work = self.new(JOB_B, content, sources=[source])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        content["blocks"][1] = {"type": "paragraph", "text": "详见学校官网。"}
        (work / "content.json").write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED)
        self.assertIn("SOURCE_OBJECTS", out)
        self.assertIn("https://www.example.edu.cn/a", out)

    def test_unreachable_reference_service_needs_the_users_consent(self):
        entry = {"text": "MCKINNEY W. Python for Data Analysis[M]. Sebastopol: O'Reilly Media, 2017.",
                 "isbn": "9781491957660", "title": "Python for Data Analysis", "year": 2017, "authors": ["McKinney W"],
                 "quote": "pandas is a library for data cleaning and analysis"}
        content = json.loads(json.dumps(CONTENT))
        content["references"] = [entry["text"]]
        content["blocks"].append({"type": "paragraph", "text": "pandas 常用于数据清洗[1]。"})
        work = self.new({**JOB_B, "task": "write"}, content)
        (work / "refs.json").write_text(json.dumps([entry], ensure_ascii=False), encoding="utf-8")
        import refcheck as R

        with patch("refcheck.fetch", side_effect=R.Offline("no route", down=True)):
            self.call("refs", "--work", work)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("accept_unverified_refs", out)
        job = {**JOB_B, "task": "write", "accept_unverified_refs": True}
        (work / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, err)

    def paper_run(self, template, add=None, edit=None):
        info = dict(JOB_B["info"])
        if template == "A":
            info.pop("course_name")
        job = {"template": template, "info": info}
        if template == "A":
            job["blank"] = ["supervisor"]
        work = self.new(job)
        source = self.sample_docx()
        code, out, err = self.call("inspect", "--work", work, "--docx", source)
        self.assertEqual(code, 0, err)
        mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
        self.assertEqual(mapping["mode"], "paper")
        mapping["paragraphs"][9]["role"] = "body"
        if add:
            mapping["add"] = add
        if edit:
            edit(mapping)
        (work / "mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        code, run_out, err = self.call("run", "--work", work, "--office", "none")
        return work, out, code, run_out, err

    def test_existing_paper_gets_template_structure_b(self):
        work, _, code, out, err = self.paper_run("B")
        self.assertEqual(code, 0, err)
        doc = Document(work / "final.docx")
        styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
        # cover, abstract page, TOC, body: abstract before TOC in B
        self.assertLess(styles.index("XAUFE Abstract Chinese"), styles.index("XAUFE TOC Title"))
        self.assertLess(styles.index("XAUFE TOC Title"), styles.index("Heading 1"))
        self.assertEqual(len(doc.sections), 4)
        toc = [p.text for p in doc.paragraphs if p.style.name.startswith("TOC ")]
        self.assertTrue(any("一、引言" in t for t in toc))
        self.assertTrue(any("参考文献" in t for t in toc))
        text = "\n".join(p.text for p in doc.paragraphs)
        for original in ("这是正文。", "正文内容很长，这一段不是标题。", "[1] 合成作者"):
            self.assertIn(original, text)

    def test_existing_paper_a_needs_english_abstract_decision(self):
        work, inspect_out, code, out, err = self.paper_run("A")
        self.assertIn("先问用户", inspect_out)
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED)
        self.assertIn("english_abstract", out)

    def test_existing_paper_a_with_translated_english_abstract(self):
        add = {"title_en": "Comparison of Sorting Algorithms", "abstract_en": "A synthetic English abstract.", "keywords_en": ["sorting", "algorithms"]}
        work, _, code, out, err = self.paper_run("A", add=add)
        self.assertEqual(code, 0, err)
        doc = Document(work / "final.docx")
        styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
        order = [styles.index(s) for s in ("XAUFE TOC Title", "XAUFE Abstract Chinese", "XAUFE Abstract English", "Heading 1")]
        self.assertEqual(order, sorted(order))
        self.assertEqual(len(doc.sections), 5)

    def test_old_manual_toc_is_replaced_by_generated_one(self):
        doc = Document()
        for text in ["排序算法比较", "摘要", "本文比较了几种排序算法。", "关键词：排序   算法", "目录",
                     "一、引言……………1", "二、方法……………2", "一、引言", "这是正文。", "二、方法", "方法正文。", "参考文献",
                     "[1] 合成作者. 合成文献[M]. 2020."]:
            doc.add_paragraph(text)
        source = self.d / "带目录.docx"
        doc.save(source)
        work = self.new(JOB_B)
        code, out, err = self.call("inspect", "--work", work, "--docx", source)
        self.assertEqual(code, 0, err)
        roles = [r["role"] for r in json.loads((work / "mapping.json").read_text(encoding="utf-8"))["paragraphs"]]
        self.assertEqual(roles[4:7], ["old_toc", "old_toc", "old_toc"])
        self.assertIn("旧目录", out)
        code, run_out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, err)
        final = Document(work / "final.docx")
        self.assertEqual(sum(1 for p in final.paragraphs if p.style.name == "XAUFE TOC Title"), 1)
        self.assertNotIn("……………", "\n".join(p.text for p in final.paragraphs))

    def test_english_title_and_spaced_keyword_labels_are_kept(self):
        doc = Document()
        for text in ["排序算法比较", "摘要", "本文比较了几种排序算法。", "关键词：排序 算法 ", "Comparison of Sorting Algorithms",
                     "Abstract", "This paper compares sorting algorithms.", "Key words: sorting; algorithms",
                     "一、引言", "这是正文。", "参考文献", "[1] 合成作者. 合成文献[M]. 2020."]:
            doc.add_paragraph(text)
        source = self.d / "英文题目.docx"
        doc.save(source)
        work = self.new({**JOB_B, "template": "A", "info": {k: v for k, v in JOB_B["info"].items() if k != "course_name"}, "blank": ["supervisor"]})
        code, out, err = self.call("inspect", "--work", work, "--docx", source)
        self.assertEqual(code, 0, err)
        roles = [r["role"] for r in json.loads((work / "mapping.json").read_text(encoding="utf-8"))["paragraphs"]]
        self.assertEqual(roles[4], "title_en")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, err)
        text = [p.text for p in Document(work / "final.docx").paragraphs]
        self.assertIn("Key words: sorting; algorithms", text)
        self.assertIn("关键词：排序   算法", text)

    # ------------------------------------------------------------ writing mode
    def written(self, content, refs=None):
        work = self.new({**JOB_B, "task": "write"}, content)
        if refs is not None:
            (work / "refs.json").write_text(json.dumps(refs, ensure_ascii=False), encoding="utf-8")
            book = {"title": "Python for Data Analysis", "publish_date": "2017", "by_statement": "Wes McKinney"}
            with patch("refcheck.fetch", return_value=book):
                code, out, err = self.call("refs", "--work", work)
            self.assertEqual(code, 0, err)
        return work

    def test_writing_mode_blocks_unchecked_refs_unsourced_numbers_and_numbering_gaps(self):
        content = json.loads(json.dumps(CONTENT))
        content["blocks"] += [
            {"type": "paragraph", "text": "调查显示，女生月均消费为 710 元。"},
            {"type": "heading", "level": 1, "text": "三、结论"},
        ]
        work = self.written(content)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED)
        self.assertIn("参考文献还没有核对", out)
        self.assertIn("710 元", out)
        self.assertIn("标题编号不连续", out)
        self.assertFalse((work / "draft.docx").exists())

    def test_writing_mode_passes_and_finish_prints_checklist(self):
        entry = {"text": "MCKINNEY W. Python for Data Analysis[M]. Sebastopol: O'Reilly Media, 2017.",
                 "isbn": "9781491957660", "title": "Python for Data Analysis", "year": 2017, "authors": ["McKinney W"],
                 "supports": "pandas 适合做数据清洗", "relevance": "高",
                 "quote": "pandas is a library for data cleaning and analysis"}
        content = json.loads(json.dumps(CONTENT))
        content["references"] = [entry["text"]]
        content["blocks"].append({"type": "paragraph", "runs": [
            {"text": "pandas 的使用者超过 100 万人"}, {"text": "[1]", "superscript": True}, {"text": "。"}]})
        work = self.written(content, [entry])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, err)
        self.fill_review(work, vision=False)
        code, out, err = self.call("finish", "--work", work)
        self.assertEqual(code, 0, err)
        self.assertIn("参考文献核实清单", out)
        self.assertIn("| [1] |", out)
        self.assertIn("已验证", out)
        self.assertIn("pandas 适合做数据清洗", out)
        self.assertIn("观点支撑", out)
        self.assertIn("取不到原文，未能对照", out)
        self.assertIn("pandas 的使用者超过 100 万人", out)
        self.assertIn("少 5 篇", out)
        self.assertIn("少于默认的 5000 字", out)

    def test_writing_mode_requires_a_quote_and_flags_unconfirmed_support(self):
        entry = {"text": "MCKINNEY W. Python for Data Analysis[M]. Sebastopol: O'Reilly Media, 2017.",
                 "isbn": "9781491957660", "title": "Python for Data Analysis", "year": 2017, "authors": ["McKinney W"]}
        content = json.loads(json.dumps(CONTENT))
        content["references"] = [entry["text"]]
        content["blocks"].append({"type": "paragraph", "text": "pandas 常用于数据清洗[1]。"})
        work = self.written(content, [entry])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("没有写 quote", out)
        self.assertIn("refs.json", out)
        work = self.written(content, [dict(entry, quote="pandas is a library for data cleaning and analysis")])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        self.assertIn("REF_SUPPORT_UNCONFIRMED", out)
        self.assertIn("引用句“pandas 常用于数据清洗[1]。”", out)
        self.assertIn("摘录“pandas is a library for data", out)
        self.assertNotIn("REF_CITATION_MISMATCH", out)
        review = json.loads((work / "review.json").read_text(encoding="utf-8"))
        self.assertIn("REF_SUPPORT_UNCONFIRMED", review["warnings"])

    def test_writing_on_a_word_draft_gets_the_same_reference_gate(self):
        doc = Document()
        doc.add_paragraph("一、引言")
        doc.add_paragraph("这是写作模式下的合成正文[1]。")
        doc.add_paragraph("参考文献")
        doc.add_paragraph("[1] 张三. 合成文献题目[J]. 合成期刊, 2020, 1(1): 1-2.")
        source = self.d / "草稿.docx"
        doc.save(source)
        work = self.new({**JOB_B, "task": "write"})
        code, out, err = self.call("inspect", "--work", work, "--docx", source)
        self.assertEqual(code, 0, err)
        mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
        ref = next(r for r in mapping["paragraphs"] if r["text"].startswith("[1]"))
        self.assertEqual(ref["role"], "reference")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, out + err)
        self.assertIn("参考文献还没有核对", out)

    def test_citation_number_on_an_unrelated_sentence_is_flagged(self):
        entry = {"text": "MCKINNEY W. Python for Data Analysis[M]. Sebastopol: O'Reilly Media, 2017.",
                 "isbn": "9781491957660", "title": "Python for Data Analysis", "year": 2017, "authors": ["McKinney W"],
                 "quote": "pandas is a library for data cleaning and analysis", "supports": "pandas 适合做数据清洗"}
        content = json.loads(json.dumps(CONTENT))
        content["references"] = [entry["text"]]
        content["blocks"].append({"type": "paragraph", "text": "卷积神经网络在图像识别上效果显著[1]。"})
        work = self.written(content, [entry])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        self.assertIn("REF_CITATION_MISMATCH", out)
        self.assertIn("卷积神经网络在图像识别上效果显著", out)
        review = json.loads((work / "review.json").read_text(encoding="utf-8"))
        self.assertIn("REF_CITATION_MISMATCH", review["warnings"])

    def test_unverifiable_reference_needs_the_users_consent(self):
        import refcheck as R

        entry = {"text": "某作者. 某题目[J]. 某刊, 2020, 1(1): 1-2.", "title": "某题目", "authors": ["某作者"], "year": 2020,
                 "url": "https://example.org/a", "quote": "某题目的原文里的一整句话。"}
        content = json.loads(json.dumps(CONTENT))
        content["references"] = [entry["text"]]
        content["blocks"].append({"type": "paragraph", "text": "有研究讨论了这个问题[1]。"})
        work = self.new({**JOB_B, "task": "write"}, content)
        (work / "refs.json").write_text(json.dumps([entry], ensure_ascii=False), encoding="utf-8")
        blocked = R.Offline("HTTP 403")
        blocked.code = 403
        with patch("refcheck.fetch", return_value=None), patch("refcheck.fetch_page", side_effect=blocked):
            self.call("refs", "--work", work)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        self.assertIn("无法验证", out)
        self.assertIn("accept_unverified_refs", out)

    def test_figures_from_the_users_material_need_no_reference(self):
        material = self.d / "调查数据.txt"
        material.write_text("问卷共回收 312 份，女生月均消费 710 元。", encoding="utf-8")
        content = json.loads(json.dumps(CONTENT))
        content["blocks"].append({"type": "paragraph", "text": "本次调查共回收问卷 312 份，女生月均消费为 710 元。"})
        work = self.new({**JOB_B, "task": "write"}, content, sources=[material])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotIn("具体数字", out)
        content["blocks"][-1]["text"] = "本次调查共回收问卷 312 份，男生月均消费为 820 元。"
        (work / "content.json").write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertIn("820 元", out)
        self.assertIn("blocks[", out)

    def test_user_provided_reference_needs_a_source(self):
        entry = {"text": "MCKINNEY W. Python for Data Analysis[M]. Sebastopol: O'Reilly Media, 2017.",
                 "isbn": "9781491957660", "title": "Python for Data Analysis", "year": 2017, "authors": ["McKinney W"],
                 "from_user": True}
        content = json.loads(json.dumps(CONTENT))
        content["references"] = [entry["text"]]
        content["blocks"].append({"type": "paragraph", "text": "pandas 常用于数据清洗[1]。"})
        work = self.written(content, [entry])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED)
        self.assertIn("user_said", out)
        result = json.loads((work / "refs-result.json").read_text(encoding="utf-8"))["results"][0]
        self.assertEqual((result["status"], result["checked"]), ("用户提供", "已验证"))
        (work / "refs.json").write_text(json.dumps([{**entry, "user_said": "用户：参考文献用这本 pandas 的书"}],
                                                   ensure_ascii=False), encoding="utf-8")
        with patch("refcheck.fetch", return_value={"title": "Python for Data Analysis", "publish_date": "2017",
                                                   "by_statement": "Wes McKinney"}):
            self.call("refs", "--work", work)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotIn("user_said", out + err)

    def test_writing_mode_checks_citation_numbers_against_references(self):
        entry = {"text": "MCKINNEY W. Python for Data Analysis[M]. Sebastopol: O'Reilly Media, 2017.",
                 "isbn": "9781491957660", "title": "Python for Data Analysis", "year": 2017, "authors": ["McKinney W"]}
        content = json.loads(json.dumps(CONTENT))
        content["references"] = [entry["text"]]
        content["blocks"].append({"type": "paragraph", "text": "这里引用了不存在的第三条[3]。"})
        work = self.written(content, [entry])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED)
        self.assertIn("正文引用了 [3]，但参考文献只有 1 条", out)
        self.assertIn("第 1 条在正文里没有被引用", out)

    def test_citation_order_follows_first_citation(self):
        content = {"blocks": [{"type": "paragraph", "text": "先引第二条[2]，再引第一条[1]，范围[1-2]。"}], "references": ["甲", "乙"]}
        problems = X._citation_problems(content)
        self.assertTrue(any("第一次引用的顺序" in p for p in problems), problems)
        self.assertEqual(X._cited_numbers("见[1,3]和[4–6]"), [1, 3, 4, 5, 6])
        ok = {"blocks": [{"type": "paragraph", "text": "引[1]，又引[2]。"}], "references": ["甲", "乙"]}
        self.assertEqual(X._citation_problems(ok), [])

    def test_each_sentence_with_a_figure_needs_its_own_citation(self):
        content = {"blocks": [{"type": "paragraph", "text": "市场规模增长了30%[1]。另有调查称用户超过5万人。"}], "references": []}
        problems = X._writing_problems(self.d, content)
        flagged = [p for p in problems if "具体数字" in p]
        self.assertEqual(len(flagged), 1, problems)
        self.assertIn("用户超过5万人", flagged[0])
        cited = {"blocks": [{"type": "paragraph", "runs": [{"text": "用户超过5万人"}, {"text": "[1]", "superscript": True},
                                                            {"text": "。"}]}], "references": []}
        self.assertFalse([p for p in X._writing_problems(self.d, cited) if "具体数字" in p])

    def test_failed_acceptance_prints_located_fixable_issues_and_counts_rounds(self):
        import feedback as F

        document = {"mode": "document", "metadata": {"title_zh": "排序算法比较"}, "blocks": CONTENT["blocks"]}
        work = self.new(JOB_B, document)
        failing = {"passed": False, "pages": 3, "warnings": [], "fonts": {}, "errors": [
            {"code": "BLANK_PAGE", "severity": "error", "page": 3, "message": "这一页正文区什么都没有（多余的空白页）。",
             "where": "紧跟在“这是合成正文第二段。”之后"}]}
        with patch("delivery.render", side_effect=fake_render), patch("office_discovery.discover", return_value={}), \
                patch("delivery.audit_rendered", return_value=failing):
            for round_ in range(F.MAX_STALLED + 1):
                code, out, err = self.call("run", "--work", work)
                self.assertEqual(code, X.EXIT_NOT_ACCEPTED, err)
        for word in ("第3页", "content.json 的 blocks", "可能原因", "怎么改", "停止自动重试"):
            self.assertIn(word, out)
        # After the stop another run is refused until the user asks for one; their words are kept.
        code, out, err = self.call("run", "--work", work)
        self.assertEqual(code, X.EXIT_ASK_USER, out + err)
        self.assertIn("--user-says-retry", out)
        with patch("delivery.render", side_effect=fake_render), patch("office_discovery.discover", return_value={}), \
                patch("delivery.audit_rendered", return_value=failing):
            code, out, err = self.call("run", "--work", work, "--user-says-retry", "再试一次吧")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED, out + err)
        self.assertNotIn("停止自动重试", out)
        self.assertEqual(json.loads((work / "attempts.json").read_text(encoding="utf-8"))["user_retries"], ["再试一次吧"])
        issues = json.loads((work / "issues.json").read_text(encoding="utf-8"))
        self.assertEqual(issues["issues"][0]["code"], "BLANK_PAGE")
        self.assertFalse((work / "review.json").exists())

    def test_examples_run_as_documented_in_readme(self):
        import shutil

        for template in ("A", "B"):
            code, out, err = self.call("new", "--out", self.d / f"example-{template}.docx",
                                       "--source", ROOT / "examples" / template / "原稿.txt")
            self.assertEqual(code, 0, err)
            work = Path(out.split("WORK = ")[1].splitlines()[0].strip())
            for name in ("job.json", "content.json"):
                shutil.copy2(ROOT / "examples" / template / name, work / name)
            shutil.copy2(ROOT / "examples/layout-test.png", work / "layout-test.png")
            code, out, err = self.call("run", "--work", work, "--office", "none")
            self.assertEqual(code, 0, err)
            self.call("cancel", "--work", work)

    def test_writing_mode_rejects_reference_not_in_refs_json(self):
        entry = {"text": "合成作者. 合成书[M]. 2017.", "isbn": "9781491957660", "title": "Python for Data Analysis", "year": 2017, "authors": ["McKinney W"]}
        content = json.loads(json.dumps(CONTENT))
        content["references"] = ["编造作者. 编造的书[M]. 2020."]
        work = self.written(content, [entry])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, X.EXIT_NOT_ACCEPTED)
        self.assertIn("参考文献第 1 条不在 refs 核对可用的条目里", out)

    def test_existing_paper_keyword_spaces_normalized(self):
        work, _, code, out, err = self.paper_run("B")
        self.assertEqual(code, 0, err)
        keywords = next(p.text for p in Document(work / "final.docx").paragraphs if p.text.startswith("关键词"))
        self.assertEqual(keywords, "关键词：排序   算法")

    def test_inspect_recognises_this_skills_own_output(self):
        work = self.new(JOB_B, CONTENT)
        self.call("run", "--work", work, "--office", "none")
        previous = self.d / "上一版.docx"
        previous.write_bytes((work / "final.docx").read_bytes())
        work2 = self.new(dict(JOB_B, logo="green"))
        code, out, err = self.call("inspect", "--work", work2, "--docx", previous)
        self.assertEqual(code, 0, err)
        mapping = json.loads((work2 / "mapping.json").read_text(encoding="utf-8"))
        roles = {r["text"]: r["role"] for r in mapping["paragraphs"] if r["text"]}
        self.assertEqual(roles["一、引言"], "h1")
        self.assertEqual(roles["这是合成正文。"], "body")
        self.assertEqual(roles["（一）背景"], "h2")
        self.assertGreater(mapping["replace_cover_before"], 0)
        self.assertIn("旧封面", out)

    # ------------------------------------------------------------ instructions
    def test_skill_md_keeps_the_core_rules(self):
        text = (ROOT / "SKILL.md").read_text(encoding="utf-8")
        for phrase in (
            "版式和校徽是两回事",
            "只排版",
            "默认写作",
            "先问清主题、方向和内容",
            "留空自己填",
            "今天（北京时间）",
            "逐页打开查看",
            "只改用户这次说的那一项",
            "scripts/xaufe.py compare",
            "finish",
        ):
            self.assertIn(phrase, text)


if __name__ == "__main__":
    unittest.main()
