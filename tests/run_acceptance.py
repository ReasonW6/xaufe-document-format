#!/usr/bin/env python3
"""Real-render acceptance through the agent-facing commands (synthetic data only).

Each case runs `xaufe.py new / run` with the Office software available on this
machine, copies the final Word and page images to --output-dir for a person (or
model) to inspect page by page, then cancels its own workspace. Nothing here
signs a visual review.

    python -X utf8 tests/run_acceptance.py --output-dir 技能目录外的空文件夹
"""

from __future__ import annotations
import argparse, contextlib, io, json, shutil, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import xaufe as X
import delivery as D

LONG = {
    "title_zh": "基于多源异构数据融合与深度强化学习的城市交通信号自适应协同控制方法及其在大规模路网中的仿真验证研究",
    "student_name": "欧阳合成",
    "student_id": "2023123456789",
    "major": "计算机科学与技术（人工智能与大数据方向）",
    "class_name": "计算机科学与技术2023级卓越工程师实验班",
    "supervisor": "合成导师",
    "course_name": "面向对象程序设计与软件工程综合实践（双语）",
}


def example(template):
    job = json.loads((ROOT / f"examples/{template}/job.json").read_text(encoding="utf-8"))
    content = json.loads((ROOT / f"examples/{template}/content.json").read_text(encoding="utf-8"))
    for block in content["blocks"]:
        if block["type"] == "image":
            block["path"] = str(ROOT / "examples/layout-test.png")
    return job, content


def long_job(template, logo=None):
    job, content = example(template)
    fields = {"A": ("title_zh", "student_name", "student_id", "major", "class_name", "supervisor"),
              "B": ("course_name", "title_zh", "student_name", "student_id", "major", "class_name")}[template]
    job = {"task": "format", "template": template, "info": {k: LONG[k] for k in fields}}
    if logo:
        job["logo"] = logo
    content["metadata"]["title_zh"] = LONG["title_zh"]
    return job, content


def existing_paper(path):
    """A plain student paper as it might arrive: Word headings, abstract, table, references."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("高校图书馆座位预约系统的设计与实现")
    doc.add_paragraph("摘要")
    doc.add_paragraph("本文设计并实现了一个基于Web的座位预约系统，用于合成排版验收。")
    doc.add_paragraph("关键词：座位预约 Web系统 图书馆管理")
    for heading, text in (("一、引言", "合成正文第一段。"), ("二、系统设计", "合成正文第二段。"), ("三、结论", "合成结论。")):
        doc.add_heading(heading, level=1)
        doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for i, row in enumerate((("表名", "说明"), ("seat", "座位信息"))):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 合成作者. 合成文献[M]. 北京: 合成出版社, 2020.")
    doc.save(path)
    return path


def long_document(template):
    """About twenty pages: three heading levels, wide tables that cross pages, formulas,
    code with a long line, a long URL and an image - the regression case for long and
    extreme layouts."""
    job, content = example(template)
    sentence = "本段为合成的长文档回归文字，用于检验分页、行距、首行缩进和两端对齐在多页中保持一致，"
    blocks = [{"type": "heading", "level": 1, "text": "引言", "preface": True},
              {"type": "paragraph", "text": sentence * 6 + "不代表任何研究结论。"}]
    numerals = "一二三四五六七八"
    for i, n in enumerate(numerals[:6], 1):
        blocks.append({"type": "heading", "level": 1, "text": f"{n}、合成章节{i}"})
        for j, m in enumerate("一二三", 1):
            blocks.append({"type": "heading", "level": 2, "text": f"（{m}）合成小节{i}.{j}"})
            blocks.append({"type": "heading", "level": 3, "text": f"{j}.合成要点"})
            blocks += [{"type": "paragraph", "text": sentence * 5 + f"第{i}章第{j}节第{k}段。"} for k in range(1, 4)]
        if i == 2:
            blocks.append({"type": "table", "caption": "表1 合成宽表（七列，跨页）",
                           "header": ["编号", "名称", "类别", "说明文字较长的一列", "数值甲", "数值乙", "备注"],
                           "rows": [[str(r), f"项目{r}", "合成类", "单元格里放一段较长的说明文字，检验自动换行和列宽分配。", f"{r * 3.5:.1f}", f"{r * 7}", "无"]
                                    for r in range(1, 31)]})
        if i == 3:
            blocks += [{"type": "equation", "text": "f(x) = a₀ + Σᵢ aᵢ·cos(iωx) + bᵢ·sin(iωx)"},
                       {"type": "equation", "text": "E = mc²"},
                       {"type": "code", "text": "def merge(a, b):\n    return sorted(a + b, key=lambda item: (item.priority, item.created_at, item.identifier_that_is_quite_long))"},
                       {"type": "paragraph", "text": "长链接示例：https://example.org/a/very/long/path/that/keeps/going/without/any/space/to/check/line/breaking/behaviour/in/narrow/columns.html"}]
        if i == 4:
            blocks.append({"type": "image", "path": str(ROOT / "examples/layout-test.png"), "caption": "图1 合成图片"})
    content["blocks"] = blocks
    return job, content


def complex_paper(path):
    """An existing paper with merged cells, a nested table, a nine-column table and an oversized image."""
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    doc.add_paragraph("复杂表格合成论文")
    doc.add_paragraph("摘要")
    doc.add_paragraph("本文仅用于复杂表格和图片的排版验收，不含真实研究内容。")
    doc.add_paragraph("关键词：表格 合并单元格 验收")
    doc.add_heading("一、合并单元格", level=1)
    doc.add_paragraph("下表含横向与纵向合并。")
    t = doc.add_table(rows=4, cols=4)
    t.style = "Table Grid"
    for r in range(4):
        for c in range(4):
            t.cell(r, c).text = f"r{r}c{c}"
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(1, 2).merge(t.cell(3, 2))
    t.cell(2, 3).tables  # keep python-docx happy before nesting
    t.cell(2, 3).add_table(rows=2, cols=2)
    doc.add_heading("二、宽表与大图", level=1)
    wide = doc.add_table(rows=3, cols=9)
    wide.style = "Table Grid"
    for r in range(3):
        for c in range(9):
            wide.cell(r, c).text = f"列{c + 1}较长文字"
    doc.add_paragraph().add_run().add_picture(str(ROOT / "examples/layout-test.png"), width=Cm(22))
    doc.add_paragraph("参考文献")
    doc.add_paragraph("[1] 合成作者. 合成文献[M]. 北京: 合成出版社, 2020.")
    doc.save(path)
    return path


def original_text(content):
    """The user's original material behind a format-only content.json (synthetic)."""
    lines = [content.get("metadata", {}).get("title_zh", ""), content.get("abstract_zh", ""),
             "关键词：" + "；".join(content["keywords_zh"]) if content.get("keywords_zh") else "",
             content.get("metadata", {}).get("title_en", "") or content.get("title_en", ""),
             content.get("abstract_en", ""), "Key words: " + "; ".join(content["keywords_en"]) if content.get("keywords_en") else ""]
    for b in content.get("blocks", []):
        lines += [b.get("caption", ""), b.get("text", ""), "".join(r.get("text", "") for r in b.get("runs", []))]
        lines += ["| " + " | ".join(row) + " |" for row in [b.get("header") or []] + (b.get("rows") or []) if row]
    lines += content.get("references", [])
    return "\n\n".join(x for x in lines if x and x.strip())


def blank_page_paper(path):
    """A plain document whose author left two page breaks in a row (an empty page in the middle)."""
    from docx import Document
    from docx.enum.text import WD_BREAK

    doc = Document()
    doc.add_paragraph("空白页合成文档")
    doc.add_paragraph("第一页的合成正文。")
    for _ in range(2):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("空白页之后的合成正文。")
    doc.save(path)
    return path


def plain_word(path):
    """An ordinary Word file with no abstract or headings, long enough for several pages."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("普通文档合成标题")
    for i in range(1, 41):
        doc.add_paragraph(f"第 {i} 段合成正文，用来检验没有论文结构的普通 Word 加上封面和页码之后，"
                          "正文从第 1 页开始连续编号，并且每一页的版面都在版心以内。")
    doc.save(path)
    return path


def call(*args):
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = X.main([str(a) for a in args])
    return code, out.getvalue(), err.getvalue()


def run_case(out_dir, name, job, content=None, docx=None, mapping_edit=None, office=None, expect=None, retry=None):
    """Run one case; with ``expect`` the case is a defect that must be refused with that code.

    ``retry=(code, fix)``: the first run must be refused with ``code``; ``fix(job)`` then
    applies the change the issue's hint asks for, and the second run must pass.
    """
    target = out_dir / name
    target.mkdir()
    if docx is None and job.get("task") == "format":
        # Format only: the content came from the user's material, registered as the original.
        docx = target / "原稿.txt"
        docx.write_text(original_text(content), encoding="utf-8")
    code, text, err = call("new", "--out", target / "成品.docx", *(["--source", docx] if docx else []))
    if code:
        raise RuntimeError(err)
    work = Path(text.split("WORK = ")[1].splitlines()[0].strip())
    try:
        (work / "job.json").write_text(json.dumps(job, ensure_ascii=False, indent=2), encoding="utf-8")
        if content is not None:
            (work / "content.json").write_text(json.dumps(content, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            code, text, err = call("inspect", "--work", work, "--docx", docx)
            if code:
                raise RuntimeError(err)
            if mapping_edit:
                mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
                mapping_edit(mapping)
                (work / "mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        code, text, err = call("run", "--work", work, *(["--office", office] if office else []))
        (target / "run-output.txt").write_text(text + err, encoding="utf-8")
        if retry:
            if code != X.EXIT_NOT_ACCEPTED or retry[0] not in text:
                raise RuntimeError(f"{name}: 第一次应被 {retry[0]} 拦下，实际退出码 {code}：{(err or text)[-800:]}")
            retry[1](job)
            (work / "job.json").write_text(json.dumps(job, ensure_ascii=False, indent=2), encoding="utf-8")
            code, text2, err = call("run", "--work", work, *(["--office", office] if office else []))
            (target / "run-output.txt").write_text(text + "\n\n---- 按提示修改 job.json 后重新 run ----\n\n" + text2 + err,
                                                   encoding="utf-8")
            text = text2
        if expect:
            if code != X.EXIT_NOT_ACCEPTED or expect not in text:
                raise RuntimeError(f"{name}: 应被 {expect} 拦下，实际退出码 {code}：{(err or text)[-800:]}")
            row = {"case": name, "refused": expect}
            print(json.dumps(row, ensure_ascii=False), flush=True)
            return row
        if code:
            raise RuntimeError(err or text)
        shutil.copy2(work / "final.docx", target / "final.docx")
        for page in sorted((work / "preview").glob("page-*.png")) if (work / "preview").exists() else []:
            shutil.copy2(page, target / page.name)
        state = D.load(work)[1]
        audit = json.loads((work / "render-audit.json").read_text(encoding="utf-8")) if (work / "render-audit.json").exists() else {}
        row = {"case": name, "pages": len(state.get("pages", [])), "scope": state["validation_scope"],
               **({"fixed_after": retry[0]} if retry else {}),
               "render_audit": "passed" if audit.get("passed") else ("not-run" if not audit else "failed"),
               "render_warnings": sorted({w["code"] for w in audit.get("warnings", [])}),
               "cover_fit": __import__("engine").get_vars(__import__("docx").Document(work / "final.docx")).get("XAUFE_COVER_FIT")}
        print(json.dumps(row, ensure_ascii=False), flush=True)
        return row
    finally:
        call("cancel", "--work", work)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", required=True)
    a = parser.parse_args()
    out = Path(a.output_dir).resolve()
    if out.is_relative_to(ROOT) or (out.exists() and any(out.iterdir())):
        raise ValueError("输出须在技能目录之外且为空。")
    out.mkdir(parents=True, exist_ok=True)
    rows = []
    job, content = example("A")
    rows.append(run_case(out, "01-A-default", job, content))
    job, content = example("B")
    rows.append(run_case(out, "02-B-green-logo-pagenum-blank-id", job, content))
    job, content = long_job("A", "red")

    def short_header(job):  # what the HEADER_OVERFLOW hint asks for (the user picks the short title)
        job.setdefault("format", {})["header"] = {"right": "城市交通信号自适应协同控制研究"}

    rows.append(run_case(out, "03-A-red-logo-long-cover", job, content, retry=("HEADER_OVERFLOW", short_header)))
    job, content = long_job("B")
    rows.append(run_case(out, "04-B-long-cover", job, content))
    job, content = example("B")
    job["format"] = {
        "footer_page_numbers": True,
        "styles": {"body": {"cn": "仿宋", "size": 12}, "h1": {"color": "0070C0"}, "toc1": {"size": 12}, "toc2": {"size": 12}},
    }
    rows.append(run_case(out, "05-B-custom-formats", job, content))
    previous = out / "02-B-green-logo-pagenum-blank-id" / "final.docx"
    job, _ = example("B")
    job["logo"] = "red"
    rows.append(run_case(out, "06-revise-previous-logo-red", job, docx=previous))
    job, content = example("A")
    rows.append(run_case(out, "07-A-static-no-office", job, content, office="none"))
    existing = existing_paper(out / "existing-paper.docx")
    job, _ = example("A")
    job["info"]["title_zh"] = "高校图书馆座位预约系统的设计与实现"

    def translated(mapping):
        mapping["add"] = {
            "title_en": "Design and Implementation of a Library Seat Reservation System",
            "abstract_en": "Synthetic English abstract used only for layout acceptance.",
            "keywords_en": ["seat reservation", "web system"],
        }

    rows.append(run_case(out, "08-A-existing-paper-structured", job, docx=existing, mapping_edit=translated))
    job, content = long_document("B")
    rows.append(run_case(out, "09-B-long-document-tables-formulas", job, content))
    complex_docx = complex_paper(out / "complex-paper.docx")
    job, _ = example("B")
    job["info"]["title_zh"] = "复杂表格合成论文"
    rows.append(run_case(out, "10-B-existing-complex-tables", job, docx=complex_docx))
    job, _ = example("B")
    job["info"]["title_zh"] = "普通文档合成标题"
    job["format"] = {"footer_page_numbers": True}
    rows.append(run_case(out, "11-B-plain-word-cover-page-numbers", job, docx=plain_word(out / "plain-word.docx")))
    # Defects that must never be delivered.
    job, content = example("B")
    job["format"] = {"styles": {"body": {"size": 15, "line": 200, "line_rule": "exact"}}}
    rows.append(run_case(out, "12-defect-lines-overlap", job, content, expect="LINE_TOO_TIGHT"))
    rows.append(run_case(out, "13-defect-blank-page", *example("B")[:1], docx=blank_page_paper(out / "blank-page.docx"),
                         expect="BLANK_PAGE"))
    (out / "summary.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    bad = [r for r in rows if "refused" not in r and (
        r["render_audit"] != ("not-run" if r["case"].startswith("07-") else "passed")
        or r["scope"] != ("static-only" if r["case"].startswith("07-") else "full"))]
    if bad:
        raise SystemExit("这些用例没有完成完整验收：" + json.dumps(bad, ensure_ascii=False))
    print("页图和成品已复制到", out, "；请逐页查看。工作区已全部清理。")


if __name__ == "__main__":
    main()
