"""Delivery and customization regressions; all identities, instructions and visual records are synthetic."""

from __future__ import annotations
import hashlib, json, os, sys, tempfile, unittest
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import engine as E, intake as I, delivery as D
from customization import effective_profile, profile_for
from check_format import check
from toc_format import enforce
import office_discovery as O
import office_backend as B


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def cand(value):
    return {"value": value, "source": "user", "evidence": "仅合成单元测试，不代表真实用户授权"}


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def build(self, template="B", request=None, paper=True, cover=True):
        state = self.d / "intake.json"
        I.start(state, "合成交付验证")
        I.select(state, template, "合成夹具选择" + template)
        if request:
            I.customize(state, request, "合成夹具：按本测试字典覆盖对应格式")
        fields = profile_for(I.read(state))["cover"]["required_fields"]
        meta = {
            k: cand(
                "2026-09-24"
                if k == "date"
                else (
                    "00123"
                    if k == "student_id"
                    else ("测试课程" if k == "course_name" else "合成测试")
                )
            )
            for k in fields
        }
        I.resolve(state, {"metadata": meta})
        spec = profile_for(I.read(state))
        data = {
            "mode": "paper" if paper else "document",
            "cover": cover,
            "metadata": {},
            "blocks": [
                {"type": "heading", "level": 1, "text": "一、测试总览"},
                {"type": "heading", "level": 2, "text": "（一）检查方法"},
                {"type": "heading", "level": 3, "text": "1.细节验证"},
                {"type": "paragraph", "text": "这是虚构测试内容。"},
            ],
        }
        if paper:
            data.update(
                abstract_zh="这是合成摘要。",
                keywords_zh=["合成数据", "格式检查"],
                references=["[1] 虚构参考文献，仅测试排版。"],
            )
            if spec["english_abstract"]:
                data.update(
                    abstract_en="Synthetic fixture only.",
                    keywords_en=["Synthetic test", "Word format"],
                )
                data["metadata"]["title_en"] = "Synthetic Formatting Test"
        src = self.d / "content.json"
        src.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        out = self.d / "draft.docx"
        E.build(src, out, intake_path=state)
        return out, state

    def assertPass(self, result):
        self.assertTrue(result["passed"], result["errors"])

    def codes(self, path, state):
        return {e["code"] for e in check(path, intake_path=state)["errors"]}


class Customizations(Base):
    def test_b_layout_green_logo_passes(self):
        out, state = self.build(request={"logo": "green"})
        d = Document(out)
        spec = profile_for(I.read(state))
        self.assertPass(check(out, intake_path=state))
        self.assertEqual(d.sections[0].left_margin.twips, 1701)
        rid = d.paragraphs[0]._p.xpath(".//a:blip/@r:embed")[0]
        self.assertEqual(
            hashlib.sha256(d.part.related_parts[rid].blob).hexdigest(),
            E.load_profile("A")["logo"]["sha256"],
        )
        self.assertEqual(spec["logo"]["width_emu"], 3212465)
        self.assertEqual(spec["logo"]["height_emu"], round(3212465 * 125 / 505))
        self.assertFalse(any(p.style.name == "XAUFE English Title" for p in d.paragraphs))
        self.assertNotIn("supervisor", I.read(state)["metadata"])

    def test_a_layout_red_logo_passes(self):
        out, state = self.build("A", {"logo": "red"})
        self.assertPass(check(out, intake_path=state))
        doc = Document(out)
        spec = profile_for(I.read(state))
        base = E.load_profile("A")
        source = E.load_profile("B")["logo"]
        rid = doc.paragraphs[0]._p.xpath(".//a:blip/@r:embed")[0]
        self.assertEqual(
            hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest(), source["sha256"]
        )
        self.assertEqual(spec["logo"]["width_emu"], base["logo"]["width_emu"])
        self.assertEqual(
            spec["logo"]["height_emu"],
            round(base["logo"]["width_emu"] * source["pixel_height"] / source["pixel_width"]),
        )
        self.assertEqual(doc.sections[0].left_margin.twips, base["page_twips"]["left"])
        self.assertEqual(spec["parts"], base["parts"])
        self.assertEqual(spec["header"], base["header"])
        self.assertTrue(any(p.style.name == "XAUFE English Title" for p in doc.paragraphs))
        self.assertIn("supervisor", I.read(state)["metadata"])
        self.assertNotIn("course_name", I.read(state)["metadata"])

    def test_all_four_layout_logo_combinations_only_change_logo(self):
        for template in ("A", "B"):
            baseline = effective_profile(template)
            for source_id, color in (("A", "green"), ("B", "red")):
                with self.subTest(template=template, logo=color):
                    actual = effective_profile(template, {"logo": color})
                    self.assertEqual(
                        {k: v for k, v in actual.items() if k not in ("logo", "_format_request")},
                        {k: v for k, v in baseline.items() if k not in ("logo", "_format_request")},
                    )
                    self.assertEqual(
                        actual["logo"]["sha256"], E.load_profile(source_id)["logo"]["sha256"]
                    )
                    self.assertEqual(actual["logo"]["width_emu"], baseline["logo"]["width_emu"])

    def test_cross_logo_survives_later_font_edit_for_either_layout(self):
        for template, color in (("A", "red"), ("B", "green")):
            with self.subTest(template=template):
                self.d = Path(self.tmp.name) / template
                self.d.mkdir()
                out, state = self.build(template, {"logo": color})
                before = profile_for(I.read(state))
                I.customize(
                    state,
                    {"styles": {"body": {"cn": "仿宋", "size": 12}}},
                    "合成第二轮：仅正文改仿宋小四",
                )
                after = profile_for(I.read(state))
                self.assertEqual(after["logo"], before["logo"])
                self.assertEqual(after["parts"], before["parts"])
                self.assertEqual(after["header"], before["header"])
                I.customize(state, {"logo": None}, "合成第三轮：仅校徽恢复所选版式默认")
                restored = profile_for(I.read(state))
                self.assertEqual(restored["logo"], effective_profile(template)["logo"])
                self.assertEqual(restored["style_overrides"]["XAUFE Body"]["cn"], "仿宋")
                self.assertEqual(restored["style_overrides"]["XAUFE Body"]["size"], 12)

    def test_centered_custom_logo(self):
        out, state = self.build(
            request={"logo": "green", "logo_alignment": "center", "logo_width_cm": 8}
        )
        self.assertPass(check(out, intake_path=state))
        self.assertEqual(Document(out).paragraphs[0]._p.xpath(".//wp:extent/@cx"), ["2880000"])

    def test_toc_three_levels_user_sizes(self):
        out, state = self.build(
            request={
                "styles": {
                    "toc_title": {"size": 18},
                    "toc1": {"size": 16},
                    "toc2": {"size": 12},
                    "toc3": {"size": 10.5},
                }
            }
        )
        self.assertPass(check(out, intake_path=state))
        d = Document(out)
        self.assertEqual(d.styles["TOC 2"].element.xpath("./w:rPr/w:sz/@w:val"), ["24"])
        p = next(p for p in d.paragraphs if p.style.name == "TOC 2")
        self.assertEqual(set(p._p.xpath(".//w:rPr/w:sz/@w:val")), {"24"})

    def test_user_latin_and_chinese_font(self):
        out, state = self.build(
            request={"styles": {"toc2": {"cn": "仿宋", "latin": "Arial", "bold": True}}}
        )
        self.assertPass(check(out, intake_path=state))

    def test_custom_body_font_and_size(self):
        out, state = self.build(
            request={"styles": {"body": {"cn": "仿宋", "size": 12, "line": 440}}}
        )
        self.assertPass(check(out, intake_path=state))

    def test_custom_indent_replaces_alternate_units(self):
        out, state = self.build(
            request={"styles": {"body": {"first": 480}, "h2": {"before": 240, "after": 240}}}
        )
        self.assertPass(check(out, intake_path=state))
        self.assertFalse(
            Document(out).styles["XAUFE Body"].element.xpath("./w:pPr/w:ind/@w:firstLineChars")
        )

    def test_red_explicit_english_abstract(self):
        out, state = self.build(request={"english_abstract": True})
        self.assertPass(check(out, intake_path=state))

    def test_green_explicit_no_english_abstract(self):
        out, state = self.build("A", {"english_abstract": False})
        self.assertPass(check(out, intake_path=state))

    def test_explicit_part_order(self):
        out, state = self.build(
            request={"parts": ["cover", "toc", "abstract_zh", "body", "references"]}
        )
        self.assertPass(check(out, intake_path=state))

    def test_red_explicit_header(self):
        out, state = self.build(
            request={
                "header": {
                    "enabled": True,
                    "left": "定制页眉",
                    "right": "仅合成测试",
                    "border": False,
                }
            }
        )
        self.assertPass(check(out, intake_path=state))

    def test_explicit_footer_page_numbers(self):
        out, state = self.build(request={"footer_page_numbers": True})
        self.assertPass(check(out, intake_path=state))
        self.assertFalse(Document(out).sections[0].footer._element.xpath(".//w:fldSimple"))
        self.assertEqual(
            Document(out).sections[1].footer._element.xpath(".//w:fldSimple/@w:instr"), [" PAGE "]
        )

    def test_red_explicit_supervisor(self):
        out, state = self.build(request={"cover_supervisor": True})
        self.assertPass(check(out, intake_path=state))
        self.assertTrue(Document(out).paragraphs[13].text.startswith("指导教师"))

    def test_green_explicit_remove_supervisor(self):
        out, state = self.build("A", {"cover_supervisor": False})
        self.assertPass(check(out, intake_path=state))
        self.assertFalse(Document(out).paragraphs[13].text.strip())

    def test_cover_label_and_fonts(self):
        out, state = self.build(
            request={
                "cover_label": "课程研究报告",
                "cover_title_label": "报告题目：",
                "cover_fonts": {"title": {"size": 14}, "date": {"size": 16}},
            }
        )
        self.assertPass(check(out, intake_path=state))

    def test_margin_change_updates_date_width(self):
        out, state = self.build(request={"page_twips": {"left": 1600, "right": 1600}})
        self.assertPass(check(out, intake_path=state))
        self.assertEqual(Document(out).paragraphs[14]._p.xpath("./w:pPr/w:framePr/@w:w"), ["8706"])

    def test_requires_true_request_evidence(self):
        state = self.d / "intake.json"
        I.start(state, "synthetic")
        I.select(state, "B", "合成选择")
        with self.assertRaises(ValueError):
            I.customize(state, {"logo": "green"}, "")

    def test_change_after_build_is_stale(self):
        out, state = self.build()
        I.customize(state, {"logo": "green"}, "合成新要求")
        self.assertIn("CUSTOMIZATION_STALE", self.codes(out, state))

    def test_reset_to_defaults(self):
        out, state = self.build(request={"logo": "green"})
        I.customize(state, {}, "合成恢复默认要求", reset=True)
        self.assertEqual(profile_for(I.read(state))["logo"]["color"], "red")

    def test_unauthorized_inline_config_rejected(self):
        out, state = self.build()
        doc = Document(out)
        E.set_vars(doc, {"XAUFE_FORMAT_REQUEST": json.dumps({"logo": "green"})})
        doc.save(out)
        self.assertIn("CUSTOMIZATION_STALE", self.codes(out, state))

    def test_bad_style_role_rejected(self):
        with self.assertRaises(ValueError):
            effective_profile("B", {"styles": {"made_up": {"size": 12}}})

    def test_hanging_and_first_not_combined(self):
        with self.assertRaises(ValueError):
            effective_profile("B", {"styles": {"toc2": {"hanging": 280, "first": 280}}})

    def test_unknown_request_not_silently_ignored(self):
        with self.assertRaisesRegex(ValueError, "不能忽略"):
            effective_profile("B", {"random": 1})

    def test_defaults_not_mutated(self):
        before = E.load_profile("B")
        effective_profile(
            "B", {"logo": "green", "styles": {"toc2": {"size": 12}}}
        )
        self.assertEqual(before, E.load_profile("B"))


class TOCGuards(Base):
    def mutate(self, fn):
        out, state = self.build()
        doc = Document(out)
        fn(doc)
        doc.save(out)
        return out, state

    def toc(self, d, level=2):
        return next(p for p in d.paragraphs if p.style.name == f"TOC {level}")

    def test_page_number_wrong_size(self):
        def change(d):
            self.toc(d)._p.xpath("./w:fldSimple//w:rPr/w:sz")[0].set(qn("w:val"), "18")

        out, s = self.mutate(change)
        self.assertIn("TOC_FONT_SIZE", self.codes(out, s))

    def test_link_text_wrong_size(self):
        def change(d):
            self.toc(d)._p.xpath("./w:hyperlink//w:rPr/w:sz")[0].set(qn("w:val"), "20")

        out, s = self.mutate(change)
        self.assertIn("TOC_FONT_SIZE", self.codes(out, s))

    def test_style_wrong_even_with_correct_direct_runs(self):
        out, s = self.mutate(
            lambda d: d.styles["TOC 2"].element.xpath("./w:rPr/w:sz")[0].set(qn("w:val"), "18")
        )
        self.assertIn("TOC_FONT_SIZE", self.codes(out, s))

    def test_leader_paragraph_mark_wrong_size(self):
        out, s = self.mutate(
            lambda d: self.toc(d)._p.xpath("./w:pPr/w:rPr/w:sz")[0].set(qn("w:val"), "18")
        )
        self.assertIn("TOC_FONT_SIZE", self.codes(out, s))

    def test_szcs_wrong_size(self):
        out, s = self.mutate(
            lambda d: self.toc(d)._p.xpath("./w:hyperlink//w:rPr/w:szCs")[0].set(qn("w:val"), "60")
        )
        self.assertIn("TOC_FONT_SIZE", self.codes(out, s))

    def test_unexpected_italic(self):
        out, s = self.mutate(
            lambda d: self.toc(d)._p.xpath("./w:hyperlink//w:rPr/w:i")[0].set(qn("w:val"), "1")
        )
        self.assertIn("TOC_EMPHASIS", self.codes(out, s))

    def test_unexpected_bold(self):
        out, s = self.mutate(
            lambda d: self.toc(d)._p.xpath("./w:hyperlink//w:rPr/w:b")[0].set(qn("w:val"), "1")
        )
        self.assertIn("TOC_EMPHASIS", self.codes(out, s))

    def test_wrong_tab_leader(self):
        out, s = self.mutate(
            lambda d: self.toc(d)._p.xpath("./w:pPr/w:tabs/w:tab")[0].set(qn("w:leader"), "hyphen")
        )
        self.assertIn("TOC_TABS", self.codes(out, s))

    def test_wrong_tab_position(self):
        out, s = self.mutate(
            lambda d: self.toc(d)._p.xpath("./w:pPr/w:tabs/w:tab")[0].set(qn("w:pos"), "5000")
        )
        self.assertIn("TOC_TABS", self.codes(out, s))

    def test_directory_repair_preserves_text_and_links(self):
        out, s = self.build()
        d = Document(out)
        before = E.content_signature(d)
        self.toc(d)._p.xpath("./w:hyperlink//w:rPr/w:sz")[0].set(qn("w:val"), "18")
        E.select_profile("B")
        enforce(d, E.SPEC)
        self.assertEqual(E.content_signature(d), before)
        d.save(out)
        self.assertPass(check(out, intake_path=s))

    def test_no_automatic_update_on_open(self):
        out, s = self.build()
        self.assertEqual(Document(out).settings.element.xpath("./w:updateFields/@w:val"), ["false"])


class Discovery(Base):
    def test_windows_32bit_and_clicktorun_paths(self):
        paths = [
            v[1].replace("\\", "/")
            for v in O.known_paths(
                "Windows",
                {
                    "ProgramFiles": "C:/Program Files",
                    "ProgramFiles(x86)": "C:/Program Files (x86)",
                    "LOCALAPPDATA": "C:/Users/test/AppData/Local",
                },
                self.d,
            )
        ]
        self.assertTrue(
            any("Program Files (x86)/LibreOffice/program/soffice.com" in p for p in paths)
        )
        self.assertTrue(any("Microsoft Office/root/Office16/WINWORD.EXE" in p for p in paths))

    def test_mac_word_and_libreoffice(self):
        engines = {v[0] for v in O.known_paths("Darwin", {}, self.d)}
        self.assertIn("word-manual", engines)
        self.assertIn("libreoffice", engines)

    def test_missing_path_does_not_hide_known_install(self):
        exe = self.d / "soffice"
        exe.write_text("synthetic")
        with (
            patch("office_discovery.shutil.which", return_value=None),
            patch(
                "office_discovery.known_paths",
                return_value=[("libreoffice", str(exe), "common-install")],
            ),
            patch("office_discovery.registry_paths", return_value=([], [])),
        ):
            r = O.discover()
            self.assertTrue(r["candidates"][0]["found"])

    def test_explicit_executable_with_spaces(self):
        folder = self.d / "LibreOffice Test"
        folder.mkdir()
        exe = folder / "soffice.exe"
        exe.write_text("synthetic")
        with patch.dict(os.environ, {"XAUFE_LIBREOFFICE": str(exe)}):
            r = O.discover()
            self.assertEqual(O.usable_candidates(r)[0]["path"], str(exe))

    def test_registry_install_outside_default_dirs(self):
        exe = self.d / "WINWORD.EXE"
        exe.write_text("synthetic")
        with (
            patch("office_discovery.platform.system", return_value="Windows"),
            patch("office_discovery.powershell", return_value="powershell.exe"),
            patch("office_discovery.shutil.which", return_value=None),
            patch("office_discovery.known_paths", return_value=[]),
            patch(
                "office_discovery.registry_paths",
                return_value=([("word", str(exe), "registry")], []),
            ),
        ):
            self.assertEqual(O.usable_candidates(O.discover())[0]["engine"], "word")

    def test_wsl_boundary_notice(self):
        with patch.dict(os.environ, {"WSL_DISTRO_NAME": "test"}):
            r = O.discover()
            self.assertTrue(r["runtime"]["wsl"])
            self.assertTrue(any(c.get("source") == "host-boundary" for c in r["checks"]))

    def test_found_but_probe_failed_is_not_missing(self):
        exe = self.d / "soffice"
        exe.write_text("synthetic")
        with (
            patch("office_discovery.shutil.which", return_value=None),
            patch("office_discovery.known_paths", return_value=[("libreoffice", str(exe), "test")]),
            patch("office_discovery.registry_paths", return_value=([], [])),
            patch(
                "office_discovery.subprocess.run", side_effect=PermissionError("synthetic denial")
            ),
        ):
            r = O.discover(probe=True)
            self.assertTrue(r["candidates"][0]["found"])
            self.assertFalse(r["candidates"][0]["probe"]["usable"])

    def test_ps_unicode_safe_encoded_command(self):
        with (
            patch("office_discovery.powershell", return_value="powershell.exe"),
            patch(
                "office_discovery.subprocess.run", return_value=SimpleNamespace(returncode=0)
            ) as mock,
        ):
            O.ps_run("Write-Output '合成' ")
            args = mock.call_args.args[0]
            self.assertIn("-EncodedCommand", args)
            self.assertNotIn("-ExecutionPolicy", args)

    def test_all_exports_fail_do_not_claim_uninstalled(self):
        candidate = {
            "engine": "libreoffice",
            "path": "synthetic",
            "found": True,
            "automation": True,
        }
        with (
            patch("office_backend.discover", return_value={}),
            patch("office_backend.usable_candidates", return_value=[candidate]),
            patch("office_backend._lo", side_effect=RuntimeError("synthetic export error")),
        ):
            with self.assertRaisesRegex(RuntimeError, "不是已证实未安装"):
                B.export_pdf(self.d / "in.docx", self.d / "out.pdf")


class DeliveryGuards(Base):
    def fake_render(self, docx, outdir):
        outdir = Path(outdir)
        outdir.mkdir(exist_ok=True, parents=True)
        p = outdir / "page-001.png"
        p.write_bytes(b"unit-test-only-not-real-render")
        result = {
            "source_sha256": sha(docx),
            "pages": [{"page": 1, "png": p.name, "png_sha256": sha(p)}],
            "visual_review": "pending",
        }
        D.write_json(outdir / "render-report.json", result)
        return result

    def job(self, sources=()):
        source, state = self.build(paper=False, cover=False)
        info = D.create(self.d / "output/final.docx", parent=self.d, sources=sources)
        with patch("delivery.render", side_effect=self.fake_render):
            r = D.prepare(info["workspace"], source, state)
        return Path(info["workspace"]), source, state, r

    def review(self, w, r, vision=True, **extra):
        # Unit-test record only; never used for real specimen acceptance.
        data = {
            "vision": vision,
            "pages": (
                {str(p["page"]): f"单元测试合成记录：第{p['page']}页无异常" for p in r["pages"]}
                if vision
                else {}
            ),
            "content_check": "单元测试合成记录：内容核对完成",
            "layout_check": "单元测试合成记录：版面核对完成",
            "delivery_check": "单元测试合成记录：交付核对完成",
            "warnings": {},
        }
        if not vision:
            data["no_vision_reason"] = "单元测试：模拟不能看图"
        data.update(extra)
        return D.attest(w, data)

    def test_cannot_publish_before_review(self):
        w, src, s, r = self.job()
        with self.assertRaises(ValueError):
            D.publish(w)
        self.assertFalse((self.d / "output/final.docx").exists())

    def test_vision_requires_every_page(self):
        w, src, s, r = self.job()
        with self.assertRaisesRegex(ValueError, "还缺第 1 页"):
            self.review(w, r, pages={})

    def test_placeholder_text_rejected(self):
        w, src, s, r = self.job()
        with self.assertRaisesRegex(ValueError, "content_check"):
            self.review(w, r, content_check="请填写：需求与内容保全")
        with self.assertRaisesRegex(ValueError, "第 1 页"):
            self.review(w, r, pages={"1": "请填写：第1页实际看到的情况"})

    def test_unknown_review_key_rejected(self):
        w, src, s, r = self.job()
        with self.assertRaises(ValueError):
            self.review(w, r, passed=True)

    def test_no_vision_skips_only_visual(self):
        w, src, s, r = self.job()
        self.review(w, r, vision=False)
        result = D.publish(w)
        self.assertEqual(result["visual_review"], "skipped_no_vision")
        self.assertIn("没有逐页看图", result["user_notice"])
        self.assertFalse(w.exists())
        self.assertTrue(src.exists())
        self.assertEqual(list((self.d / "output").iterdir()), [self.d / "output/final.docx"])

    def test_capable_publish_removes_owned_intermediates(self):
        w, src, s, r = self.job()
        self.review(w, r)
        result = D.publish(w)
        self.assertEqual(result["user_notice"], "")
        self.assertFalse(w.exists())
        self.assertTrue(src.exists())
        self.assertTrue(s.exists())
        self.assertEqual(len(list((self.d / "output").iterdir())), 1)

    def test_does_not_delete_unrelated_output_files(self):
        w, src, s, r = self.job()
        folder = self.d / "output"
        folder.mkdir()
        (folder / "user-notes.txt").write_text("keep")
        self.review(w, r, vision=False)
        D.publish(w)
        self.assertEqual((folder / "user-notes.txt").read_text(), "keep")

    def test_word_modified_after_review_is_blocked(self):
        w, src, s, r = self.job()
        self.review(w, r)
        d = Document(w / "final.docx")
        d.core_properties.subject = "modified"
        d.save(w / "final.docx")
        with self.assertRaises(ValueError):
            D.publish(w)

    def test_images_modified_before_review_are_rejected(self):
        w, src, s, r = self.job()
        (w / "preview/page-001.png").write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "页图"):
            self.review(w, r)

    def test_images_modified_after_review_is_blocked(self):
        w, src, s, r = self.job()
        self.review(w, r)
        (w / "preview/page-001.png").write_bytes(b"changed")
        with self.assertRaises(ValueError):
            D.publish(w)

    def test_existing_destination_not_overwritten(self):
        w, src, s, r = self.job()
        dest = self.d / "output/final.docx"
        dest.parent.mkdir()
        dest.write_bytes(b"existing user file")
        self.review(w, r, vision=False)
        with self.assertRaises(FileExistsError):
            D.publish(w)
        self.assertEqual(dest.read_bytes(), b"existing user file")

    def test_arbitrary_folder_not_cleaned(self):
        folder = self.d / "xaufe-job-unowned"
        folder.mkdir()
        (folder / "user.txt").write_text("keep")
        with self.assertRaises(ValueError):
            D.clean(folder)
        self.assertTrue((folder / "user.txt").exists())

    def test_cannot_create_over_existing_file(self):
        file = self.d / "source.docx"
        file.write_bytes(b"user")
        with self.assertRaises(ValueError):
            D.create(file)

    def test_no_vision_must_not_forge_seen_pages(self):
        w, src, s, r = self.job()
        with self.assertRaises(ValueError):
            self.review(w, r, vision=False, pages={"1": "声称看过的合成记录内容"})

    def test_no_vision_requires_reason(self):
        w, src, s, r = self.job()
        with self.assertRaisesRegex(ValueError, "no_vision_reason"):
            self.review(w, r, vision=False, no_vision_reason="")

    def test_modified_original_stops_delivery(self):
        original = self.d / "original.txt"
        original.write_text("原始内容", encoding="utf-8")
        w, src, s, r = self.job(sources=[original])
        self.review(w, r, warnings={"SOURCE_DIFFERENCES": "单元测试：合成原稿与成品本来就不同"})
        original.write_text("交付前被改动", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "原稿"):
            D.publish(w)

    def test_changed_original_stops_before_rendering(self):
        source, state = self.build(paper=False, cover=False)
        original = self.d / "original.txt"
        original.write_text("原始内容", encoding="utf-8")
        info = D.create(self.d / "output/final.docx", parent=self.d, sources=[original])
        original.write_text("登记后改变的内容", encoding="utf-8")
        with patch("delivery.render") as render, patch("office_discovery.discover") as discover:
            with self.assertRaisesRegex(ValueError, "原稿"):
                D.prepare(info["workspace"], source, state)
        render.assert_not_called()
        discover.assert_not_called()

    def test_reprepare_after_draft_change_needs_new_review(self):
        w, src, state, r = self.job()
        self.review(w, r)
        doc = Document(src)
        doc.core_properties.subject = "authorized draft revision"
        doc.save(src)
        with patch("delivery.render", side_effect=self.fake_render):
            r2 = D.prepare(w, src, state)
        with self.assertRaises(ValueError):
            D.publish(w)
        self.review(w, r2)
        self.assertTrue(Path(D.publish(w)["output"]).is_file())


class ExtraGuards(Base):
    def test_empty_non_dict_request_not_accepted(self):
        for value in ([], False, 0, ""):
            with self.assertRaises(ValueError):
                effective_profile("B", value)

    def test_keyword_font_override_applies_to_words_and_label(self):
        out, state = self.build(
            request={
                "styles": {
                    "keywords_zh": {"cn": "仿宋", "latin": "Arial", "size": 12, "bold": True}
                }
            }
        )
        self.assertPass(check(out, intake_path=state))
        p = next(p for p in Document(out).paragraphs if p.style.name == "XAUFE Keywords Chinese")
        self.assertEqual(p.runs[-1].font.size.pt, 12)
        self.assertEqual(p.runs[0].font.size.pt, 14)

    def test_keyword_english_explicit_unbold(self):
        out, state = self.build(
            "A", {"styles": {"keywords_en": {"bold": False, "size": 12}}}
        )
        self.assertPass(check(out, intake_path=state))

    def test_body_explicit_unbold_beats_preserved_emphasis(self):
        out, state = self.build(request={"styles": {"body": {"bold": False}}})
        doc = Document(out)
        p = next(p for p in doc.paragraphs if p.style.name == "XAUFE Body")
        p.runs[0].bold = True
        E.style_paragraph(p, "body")
        self.assertFalse(p.runs[0].bold)

    def test_customization_hash_corruption_detected(self):
        out, state = self.build()
        doc = Document(out)
        E.set_vars(doc, {"XAUFE_FORMAT_REQUEST_HASH": "wrong"})
        doc.save(out)
        self.assertIn("CUSTOMIZATION_HASH", self.codes(out, state))

    def test_word_path_containing_placeholder_is_not_corrupted(self):
        src = self.d / "SOURCE_PATH_PDF_PATH_合成.docx"
        pdf = self.d / "out.pdf"

        def run(script, timeout):
            import base64, re

            payload = json.loads(
                base64.b64decode(re.search(r"FromBase64String\('([^']+)'\)", script).group(1))
            )
            self.assertEqual(str(src), payload["source"])
            pdf.write_bytes(b"synthetic-pdf")
            return SimpleNamespace(returncode=0, stdout="synthetic Word", stderr="")

        with patch("office_backend.ps_run", side_effect=run):
            B._word(src, pdf, {}, self.d)


if __name__ == "__main__":
    unittest.main()
