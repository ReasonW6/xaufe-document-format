"""Synthetic cases for user-facing scope, choice prompts and effective settings.

These are program tests, not proof that another model followed the instructions.
"""

from pathlib import Path
import json, subprocess, sys, tempfile, unittest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import intake as I
import engine as E
from check_format import check


def value(text):
    return {
        "value": text,
        "source": "user",
        "evidence": "SYNTHETIC test statement, not actual personal data or consent",
    }


class InstructionClarity(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.state = self.folder / "intake.json"

    def tearDown(self):
        self.temp.cleanup()

    def choose(self, template="B"):
        I.start(self.state, "SYNTHETIC interaction clarity")
        return I.select(self.state, template, "SYNTHETIC explicit choice " + template)

    def complete(self, template="B", patch=None):
        self.choose(template)
        if patch:
            I.customize(self.state, patch, "SYNTHETIC explicit customization")
        fields = I.profile_for(I.read(self.state))["cover"]["required_fields"]
        return I.resolve(
            self.state,
            {
                "metadata": {
                    key: value("2030-01-02" if key == "date" else "测试文字") for key in fields
                }
            },
        )

    def build(self, data):
        source = self.folder / "content.json"
        source.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        out = self.folder / "document.docx"
        E.build(source, out, intake_path=self.state)
        return out

    def sample(self, **extra):
        data = {
            "mode": "document",
            "metadata": {},
            "blocks": [
                {"type": "heading", "level": 1, "text": "一、合成测试"},
                {"type": "paragraph", "text": "本段仅用于检验交互范围，不是真实论文。"},
            ],
        }
        data.update(extra)
        return data

    def test_a_red_summary_keeps_base_and_logo_separate(self):
        self.choose("A")
        state = I.customize(
            self.state, {"logo": "red"}, "SYNTHETIC A with red logo"
        )
        text = I.summary(state)
        self.assertIn("A（本科学年论文）", text)
        self.assertIn("配红色校徽", text)

    def test_b_green_summary_does_not_expose_customization_json(self):
        self.choose()
        state = I.customize(
            self.state,
            {"logo": "green", "styles": {"body": {"color": "006699"}}},
            "SYNTHETIC B green with color",
        )
        text = I.summary(state)
        self.assertIn("B（课程期末大作业）", text)
        self.assertIn("配绿色校徽", text)
        for token in ("logo", "006699", '"styles"', "状态："):
            self.assertNotIn(token, text)

    def test_only_unknown_fields_are_asked(self):
        self.choose()
        state = I.resolve(
            self.state, {"metadata": {"student_name": value("合成甲"), "student_id": value("0001")}}
        )
        text = I.summary(state)
        self.assertNotIn("学生姓名", text)
        self.assertNotIn("学号", text)
        self.assertIn("课程名称", text)

    def test_a_disabled_english_error_does_not_blame_red_template(self):
        self.complete("A", {"english_abstract": False})
        with self.assertRaises(ValueError) as cm:
            self.build(self.sample(abstract_en="Synthetic abstract."))
        self.assertIn("本次有效格式未启用", str(cm.exception))
        self.assertNotIn("红色模板", str(cm.exception))

    def test_b_enabled_english_is_accepted(self):
        self.complete("B", {"english_abstract": True})
        data = self.sample(
            mode="paper",
            metadata={"title_en": "Synthetic title"},
            abstract_zh="合成摘要。",
            abstract_en="Synthetic abstract.",
            keywords_zh=["合成"],
            keywords_en=["synthetic"],
            references=["合成参考条目，非真实来源。"],
        )
        out = self.build(data)
        self.assertIn("XAUFE English Title", [p.style.name for p in Document(out).paragraphs])
        self.assertTrue(check(out, intake_path=self.state)["passed"])

    def test_teacher_is_asked_when_b_explicitly_adds_it(self):
        self.choose()
        state = I.customize(self.state, {"cover_supervisor": True}, "SYNTHETIC add supervisor")
        self.assertIn("supervisor", state["missing"])
        self.assertIn("指导教师", I.summary(state))

    def test_teacher_is_not_asked_when_a_explicitly_removes_it(self):
        self.choose("A")
        state = I.customize(self.state, {"cover_supervisor": False}, "SYNTHETIC remove supervisor")
        self.assertNotIn("supervisor", state["missing"])
        self.assertNotIn("指导教师", I.summary(state))

    def test_static_english_error_uses_effective_scope(self):
        self.complete("A", {"english_abstract": False})
        out = self.build(self.sample())
        doc = Document(out)
        doc.add_paragraph("Synthetic abstract", E.ROLES["abstract_en"])
        doc.save(out)
        errors = check(out, intake_path=self.state)["errors"]
        text = " ".join(x["message"] for x in errors if x["code"] == "RED_ENGLISH_ABSTRACT")
        self.assertIn("english_abstract=false", text)
        self.assertNotIn("红色模板不应", text)

    def test_b_custom_header_error_does_not_call_it_green(self):
        self.complete(
            "B", {"header": {"enabled": True, "left": "合成左侧", "right": "合成右侧"}}
        )
        out = self.build(self.sample())
        doc = Document(out)
        doc.sections[-1].header.paragraphs[0].runs[0].text = "错误内容"
        doc.save(out)
        texts = " ".join(
            x["message"]
            for x in check(out, intake_path=self.state)["errors"]
            if x["code"] == "HEADER_TEXT"
        )
        self.assertIn("本次有效设置", texts)
        self.assertNotIn("绿色版页眉", texts)

    def test_scope_requires_actual_instruction(self):
        self.choose()
        before = self.state.read_bytes()
        with self.assertRaises(ValueError):
            I.document_scope(self.state, False, "")
        self.assertEqual(before, self.state.read_bytes())

    def test_scope_rejects_nonboolean_before_write(self):
        self.choose()
        before = self.state.read_bytes()
        with self.assertRaises(ValueError):
            I.document_scope(self.state, "no", "SYNTHETIC")
        self.assertEqual(before, self.state.read_bytes())

    def test_scope_requires_existing_real_choice(self):
        I.start(self.state, "SYNTHETIC")
        with self.assertRaises(ValueError):
            I.document_scope(self.state, False, "SYNTHETIC no cover")

    def test_no_cover_b_has_no_unused_personal_questions(self):
        self.choose()
        state = I.document_scope(self.state, False, "SYNTHETIC no cover")
        self.assertEqual(state["stage"], "ready")
        self.assertEqual(state["missing"], [])
        self.assertEqual(state["blank_fields"], [])
        self.assertNotIn("date", state["metadata"])
        self.assertIn("不含封面", I.summary(state))
        self.assertNotIn("校徽", I.summary(state))

    def test_no_cover_summary_does_not_present_unused_blank_fields(self):
        self.choose()
        I.resolve(self.state, {"blank_all": True, "blank_instruction": "SYNTHETIC blank cover"})
        state = I.document_scope(self.state, False, "SYNTHETIC no cover instead of blank cover")
        self.assertTrue(state["blank_fields"])
        self.assertNotIn("按要求留空", I.summary(state))

    def test_no_cover_a_only_needs_default_header_fields(self):
        self.choose("A")
        state = I.document_scope(self.state, False, "SYNTHETIC no cover")
        self.assertEqual(set(state["missing"]), {"title_zh", "major"})
        state = I.resolve(
            self.state, {"metadata": {"title_zh": value("测试题目"), "major": value("测试专业")}}
        )
        self.assertEqual(state["stage"], "ready")
        self.assertNotIn("date", state["metadata"])

    def test_no_cover_with_disabled_header_needs_no_fields(self):
        self.choose("A")
        I.customize(self.state, {"header": {"enabled": False}}, "SYNTHETIC no header")
        state = I.document_scope(self.state, False, "SYNTHETIC no cover")
        self.assertEqual(state["missing"], [])
        self.assertEqual(state["stage"], "ready")

    def test_no_cover_with_explicit_header_does_not_ask_major(self):
        self.choose("A")
        I.customize(
            self.state, {"header": {"left": "测试", "right": "固定标题"}}, "SYNTHETIC fixed header"
        )
        state = I.document_scope(self.state, False, "SYNTHETIC no cover")
        self.assertEqual(state["stage"], "ready")

    def test_restoring_cover_reuses_known_info_and_asks_new_fields_only(self):
        self.choose("A")
        I.document_scope(self.state, False, "SYNTHETIC no cover")
        I.resolve(
            self.state, {"metadata": {"title_zh": value("已知题目"), "major": value("已知专业")}}
        )
        state = I.document_scope(self.state, True, "SYNTHETIC restore cover")
        self.assertNotIn("title_zh", state["missing"])
        self.assertNotIn("major", state["missing"])
        self.assertIn("student_name", state["missing"])
        self.assertIn("date", state["metadata"])

    def test_no_cover_default_build_is_not_blank_cover(self):
        self.choose()
        I.document_scope(self.state, False, "SYNTHETIC no cover")
        out = self.build(self.sample())
        doc = Document(out)
        self.assertEqual(E.get_vars(doc)["XAUFE_HAS_COVER"], "0")
        self.assertEqual(doc.paragraphs[0].text, "一、合成测试")
        self.assertTrue(check(out, final=True, intake_path=self.state)["passed"])

    def test_no_cover_scope_cannot_silently_output_a_cover(self):
        self.choose()
        I.document_scope(self.state, False, "SYNTHETIC no cover")
        with self.assertRaises(ValueError) as cm:
            self.build(self.sample(cover=True))
        self.assertIn("明确范围不同", str(cm.exception))
        self.assertFalse((self.folder / "document.docx").exists())

if __name__ == "__main__":
    unittest.main()
