from pathlib import Path
import json, sys, tempfile, unittest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import source_compare as S, delivery as D


class SourceProseTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def docx(self, *paragraphs):
        out = self.root / f"out-{len(list(self.root.glob('out-*')))}.docx"
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.save(out)
        return out

    def diff(self, source_text, *paragraphs, **options):
        source = self.root / "raw.md"
        source.write_text(source_text, encoding="utf-8")
        return S.compare([{"path": str(source)}], self.docx(*paragraphs), **options)["differences"]

    def test_missing_added_and_approved_paragraphs(self):
        text = "# 标题\n\n保留的第一段。\n\n转换中不能漏掉的第二段。"
        found = self.diff(text, "保留的第一段。", "这一段是排版时编出来的。")
        self.assertEqual(sorted(d.get("original", "") for d in found if d["kind"] == "missing"), ["标题", "转换中不能漏掉的第二段。"])
        self.assertEqual([d["output"] for d in found if d["kind"] == "added"], ["这一段是排版时编出来的。"])
        approved = ["标题", "转换中不能漏掉的第二段。", "这一段是排版时编出来的。"]
        self.assertEqual(self.diff(text, "保留的第一段。", "这一段是排版时编出来的。", approved=approved), [])
        # A short snippet does not approve the whole paragraph.
        self.assertTrue(self.diff(text, "保留的第一段。", "这一段是排版时编出来的。", approved=["标题", "编出来的"]))

    def test_numbers_signs_punctuation_and_case_are_differences(self):
        cases = [("增长了2.5%。", "增长了25%。"), ("结果为-5。", "结果为5。"), ("提高0.3个百分点。", "提高03个百分点。"),
                 ("结论成立。", "结论成立！"), ("使用 Times New Roman 字体。", "使用 times new roman 字体。")]
        for before, after in cases:
            with self.subTest(before=before):
                found = self.diff("第一段。\n\n" + before, "第一段。", after)
                self.assertEqual([d["kind"] for d in found], ["changed"])
                self.assertNotEqual(found[0]["original"], found[0]["output"])

    def test_spacing_and_keyword_separators_are_not_differences(self):
        found = self.diff("English title: Room Reservation\n\nKey words: room reservation; data records\n\n保留**重点**，好。\n\n一段被硬换行\n的中文。",
                          "Room Reservation", "Key words：room reservation   data records", "保留 重点，好。", "一段被硬换行的中文。")
        self.assertEqual(found, [])

    def test_punctuation_width_and_spaces_between_words_are_differences(self):
        for changed in ("用于格式校准,不提供结论。", "用于格式校准，不提供结论（示例）.", "字体用TimesNewRoman。"):
            with self.subTest(changed=changed):
                original = {"用于格式校准,不提供结论。": "用于格式校准，不提供结论。",
                            "用于格式校准，不提供结论（示例）.": "用于格式校准，不提供结论（示例）。",
                            "字体用TimesNewRoman。": "字体用 Times New Roman。"}[changed]
                found = self.diff(original, changed)
                self.assertEqual([d["kind"] for d in found], ["changed"], found)
                self.assertNotEqual(found[0]["original"], found[0]["output"])
        shown = self.diff("用于格式校准，不提供结论。", "用于格式校准,不提供结论。")[0]
        self.assertIn("，", shown["original"])  # the user's own text is shown, not a normalised copy

    def test_moved_and_duplicated_paragraphs(self):
        text = "甲段落内容。\n\n乙段落内容。\n\n丙段落内容。"
        self.assertEqual([d["kind"] for d in self.diff(text, "乙段落内容。", "甲段落内容。", "丙段落内容。")], ["moved"])
        self.assertEqual([d["kind"] for d in self.diff(text, "甲段落内容。", "乙段落内容。", "乙段落内容。", "丙段落内容。")], ["added"])

    def test_objects_are_counted(self):
        found = self.diff("见[链接名](https://example.com)。\n\n![图](a.png)\n\n| 甲 | 乙 |\n|---|---|\n| 1 | 2 |", "见链接名。", "甲", "乙", "1", "2")
        kinds = {d["location"]: d for d in found if d["kind"] == "objects"}
        self.assertEqual(set(kinds), {"图片", "表格", "链接"})
        self.assertEqual(kinds["链接"]["original"], "https://example.com")

    def test_create_registers_actual_original_file(self):
        source = self.root / "original.txt"
        source.write_text("原始材料", encoding="utf-8")
        job = D.create(
            self.root / "out.docx",
            parent=self.root,
            sources=[source],
        )
        work, state = D.load(job["workspace"])
        self.assertEqual(state["source_documents"][0]["sha256"], D.sha(source))
        self.assertEqual(state["source_documents"][0]["role"], "original")
        D.clean(work)
        self.assertEqual(source.read_text(encoding="utf-8"), "原始材料")

    def test_missing_source_stops_before_workspace_creation(self):
        with self.assertRaises(ValueError):
            D.create(
                self.root / "out.docx",
                parent=self.root,
                sources=[self.root / "missing.md"],
            )
        self.assertEqual(list(self.root.iterdir()), [])

    def test_explicit_workspace_parent_is_respected(self):
        job = D.create(
            self.root / "out.docx",
            parent=self.root,
        )
        work, state = D.load(job["workspace"])
        self.assertEqual(work.parent, self.root.resolve())
        self.assertTrue(work.name.startswith("xaufe-job-"))
        D.clean(work)
        self.assertFalse(work.exists())

    def test_source_gap_becomes_review_warning(self):
        source = self.root / "raw.txt"
        source.write_text("需要保留的原始段落。", encoding="utf-8")
        doc = Document()
        doc.add_paragraph("只有其他文字。")
        doc.save(self.root / "final.docx")
        result = {"warnings": []}
        D.compare_originals(
            self.root, {"source_documents": [{"path": str(source), "role": "original"}]}, result
        )
        self.assertEqual(result["warnings"][0]["code"], "SOURCE_DIFFERENCES")
        self.assertTrue((self.root / "source-comparison.json").is_file())


if __name__ == "__main__":
    unittest.main()
