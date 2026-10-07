"""Opt-in native Windows Word checks, without touching existing documents.

Run with python -X utf8 tests/run_word_lifecycle.py. Requires a desktop user
session with callable Word COM. This is not part of offline unittest discovery.
"""

from pathlib import Path
import hashlib, json, os, subprocess, sys, tempfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from docx import Document
from office_discovery import ps_run
from office_backend import _word
import word_com as W


def processes():
    result = ps_run(
        "@(Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object {$_.Id}) | ConvertTo-Json -Compress",
        15,
    )
    if result.returncode:
        raise RuntimeError(result.stderr)
    value = json.loads(result.stdout) if result.stdout.strip() else []
    return sorted(value if isinstance(value, list) else [value])


def main():
    if os.name != "nt":
        raise RuntimeError("This integration test requires native Windows and callable Word COM.")
    import pymupdf as fitz

    before = processes()
    rows = []
    with tempfile.TemporaryDirectory(prefix="xaufe-word-test-") as folder:
        folder = Path(folder)
        source = folder / "中文“引号”和'单引号 $().docx"
        doc = Document()
        doc.add_heading("Portable Word export", 1)
        doc.add_paragraph("Synthetic lifecycle test.")
        doc.save(source)
        pdf = folder / "输出“引号”.pdf"
        original = hashlib.sha256(source.read_bytes()).hexdigest()
        version = _word(source, pdf, {}, folder)
        with fitz.open(pdf) as result:
            assert len(result) == 1 and result.get_toc(), "Missing PDF page or heading bookmark"
        assert hashlib.sha256(source.read_bytes()).hexdigest() == original
        assert processes() == before, "Unexpected process change after export"
        rows.append(
            {"case": "export with Unicode/quoted paths", "passed": True, "version": version}
        )
        try:
            _word(folder / "absent.docx", folder / "failed.pdf", {}, folder)
        except RuntimeError as exc:
            assert "Word.Quit" not in str(exc), "Cleanup obscured primary failure"
        else:
            raise AssertionError("Missing input must fail")
        assert processes() == before, "Unexpected process change after input failure"
        rows.append({"case": "failed export cleanup", "passed": True})
        original_script = W.script

        def stalled(*args):
            return original_script(*args).replace(
                "  $options=$word.Options", "  Start-Sleep -Seconds 30\n  $options=$word.Options"
            )

        try:
            with patch.object(W, "script", side_effect=stalled):
                W.run(ps_run, timeout=10)
        except subprocess.TimeoutExpired:
            pass
        else:
            raise AssertionError("Deliberate stall must time out")
        assert processes() == before, "Unexpected process change after timeout"
        rows.append({"case": "owned PID timeout cleanup", "passed": True})
    print(
        json.dumps(
            {"tests": rows, "initial_pids": before, "final_pids": processes()},
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
