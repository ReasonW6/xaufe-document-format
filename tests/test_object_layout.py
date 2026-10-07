"""Regressions found by the full-length synthetic document exercise."""

from test_layout_integrity import Fixture
from docx import Document
from style_model import paragraph_properties
import engine as E


class CaptionFlow(Fixture):
    def test_restyle_binds_table_caption_both_profiles(self):
        for template in ("A", "B"):
            with self.subTest(template=template):
                d = Document()
                d.add_paragraph("一、图表测试")
                d.add_paragraph("表1 合成数据")
                t = d.add_table(rows=2, cols=2)
                t.cell(0, 0).text = "编号"
                t.cell(1, 0).text = "001"
                out, result, _, _ = self.restyle(d, template, roles={"p0": "h1", "p1": "caption"})
                self.assertPass(result)
                self.assertEqual(
                    paragraph_properties(out.paragraphs[1], "keepNext").get("val"), "1"
                )

    def test_standalone_caption_does_not_bind_unrelated_prose(self):
        d = Document()
        d.add_paragraph("一、测试")
        d.add_paragraph("图1 合成图题")
        d.add_paragraph("普通正文。")
        out, result, _, _ = self.restyle(d, roles={"p0": "h1", "p1": "caption", "p2": "body"})
        self.assertPass(result)
        self.assertEqual(paragraph_properties(out.paragraphs[1], "keepNext").get("val"), "0")


from pathlib import Path
import json
import struct, zlib
from docx.shared import Cm
from docx.oxml.ns import qn
from check_format import check
import intake as I
from test_layout_integrity import candidate


class CellImageSpacing(Fixture):
    def image(self):
        path = self.path("png")

        def chunk(kind, data):
            return (
                struct.pack(">I", len(data))
                + kind
                + data
                + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
            )

        raw = (b"\x00" + b"\xff" * 1200) * 200
        path.write_bytes(
            b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", struct.pack(">IIBBBBB", 400, 200, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw))
            + chunk(b"IEND", b"")
        )
        return path

    def prepared(self, template, width=10):
        d = Document()
        d.add_paragraph("一、单元格图片")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).paragraphs[0].add_run().add_picture(str(self.image()), width=Cm(width))
        t.cell(0, 1).text = "测试"
        return self.restyle(d, template, roles={"p0": "h1"})

    def test_wide_cell_images_explicit_zero_gaps_both_profiles(self):
        for template in ("A", "B"):
            with self.subTest(template=template):
                d, result, _, _ = self.prepared(template)
                self.assertPass(result)
                for inline in d.tables[0]._tbl.xpath(".//wp:inline"):
                    self.assertEqual(
                        [inline.get(k) for k in ("distT", "distB", "distL", "distR")], ["0"] * 4
                    )

    def test_small_image_also_has_zero_gaps(self):
        d, result, _, _ = self.prepared("A", 2)
        self.assertPass(result)
        self.assertEqual(d.tables[0]._tbl.xpath(".//wp:inline")[0].get("distL"), "0")

    def test_missing_gap_rejected(self):
        d, _, _, path = self.prepared("A")
        del d.tables[0]._tbl.xpath(".//wp:inline")[0].attrib["distL"]
        d.save(path)
        self.assertCode(check(path), "CELL_IMAGE_SPACING")

    def test_nonzero_gap_rejected(self):
        d, _, _, path = self.prepared("B")
        d.tables[0]._tbl.xpath(".//wp:inline")[0].set("distR", "114300")
        d.save(path)
        self.assertCode(check(path), "CELL_IMAGE_SPACING")

    def test_image_payload_and_aspect_ratio_preserved(self):
        image = self.image()
        d = Document()
        p = d.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        p.add_run().add_picture(str(image), width=Cm(20))
        out, result, _, _ = self.restyle(d)
        self.assertPass(result)
        inline = out.tables[0]._tbl.xpath(".//wp:inline")[0]
        ext = inline.find(qn("wp:extent"))
        self.assertAlmostEqual(int(ext.get("cx")) / int(ext.get("cy")), 2, places=4)
        rid = inline.xpath(".//a:blip")[0].get(qn("r:embed"))
        self.assertEqual(out.part.related_parts[rid].blob, image.read_bytes())

    def test_new_body_picture_has_explicit_zero_spacing(self):
        path, state = self.build(
            blocks=[
                {"type": "image", "path": str(self.image()), "width_cm": 5, "caption": "图1 测试"}
            ]
        )
        self.assertPass(check(path, intake_path=state))
        self.assertEqual(
            [
                Document(path).inline_shapes[0]._inline.get(k)
                for k in ("distT", "distB", "distL", "distR")
            ],
            ["0"] * 4,
        )


class LongCoverLayout(Fixture):
    def make(self, template, title):
        state = self.state(template)
        I.resolve(state, {"metadata": {"title_zh": candidate(title, corrects_previous=True)}})
        data = {
            "mode": "document",
            "cover": True,
            "metadata": {},
            "blocks": [{"type": "paragraph", "text": "合成内容。"}],
        }
        src = self.path("json")
        src.write_text(json.dumps(data, ensure_ascii=False))
        out = self.path("docx")
        E.build(src, out, intake_path=state)
        return out, state

    def test_green_long_cover_continuation_aligns_to_title(self):
        out, state = self.make("A", "校园自习空间预约与服务质量分析系统的设计与排版测试")
        self.assertPass(check(out, intake_path=state))
        ind = Document(out).paragraphs[5]._p.pPr.find(qn("w:ind"))
        self.assertEqual(
            dict(ind.attrib), {qn("w:left"): "2508", qn("w:right"): "0", qn("w:hanging"): "1800"}
        )

    def test_red_long_cover_keeps_original_indents(self):
        out, state = self.make("B", "校园自习空间预约与服务质量分析系统的设计与排版测试")
        self.assertPass(check(out, intake_path=state))
        d = Document(out)
        baseline = Document(E.ROOT / E.SPEC["master_file"])
        self.assertEqual(
            dict(d.paragraphs[5]._p.pPr.find(qn("w:ind")).attrib),
            dict(baseline.paragraphs[5]._p.pPr.find(qn("w:ind")).attrib),
        )

    def test_short_green_cover_keeps_original_indents(self):
        out, state = self.make("A", "合成短题目")
        self.assertPass(check(out, intake_path=state))
        d = Document(out)
        base = Document(E.ROOT / E.SPEC["master_file"])
        self.assertEqual(
            dict(d.paragraphs[5]._p.pPr.find(qn("w:ind")).attrib),
            dict(base.paragraphs[5]._p.pPr.find(qn("w:ind")).attrib),
        )

    def test_tampered_long_cover_indent_rejected(self):
        out, state = self.make("A", "校园自习空间预约与服务质量分析系统的设计与排版测试")
        d = Document(out)
        d.paragraphs[5]._p.pPr.find(qn("w:ind")).set(qn("w:left"), "999")
        d.save(out)
        self.assertCode(check(out, intake_path=state), "COVER_LAYOUT")

    def test_short_topic_cannot_claim_long_topic_spareline_exception(self):
        from docx.shared import Pt

        out, state = self.make("A", "短题目")
        d = Document(out)
        d.paragraphs[6].paragraph_format.line_spacing = Pt(1)
        d.save(out)
        self.assertCode(check(out, intake_path=state), "COVER_LAYOUT")
