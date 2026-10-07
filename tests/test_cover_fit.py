"""Long cover text: wrapping, field zones, fit levels and the matching audit."""

from __future__ import annotations
import json, sys, tempfile, unittest
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import cover_fit as F
import engine as E
import intake as I
from check_format import check
from customization import apply_cover

LONG = {
    "title_zh": "基于多源异构数据融合与深度强化学习的城市交通信号自适应协同控制方法及其在大规模路网中的仿真验证研究",
    "student_name": "欧阳某某",
    "student_id": "2023123456789",
    "major": "计算机科学与技术（人工智能与大数据方向）",
    "class_name": "计算机科学与技术2023级卓越工程师实验班",
    "supervisor": "诸葛某某教授",
    "course_name": "面向对象程序设计与软件工程综合实践（双语）",
}


class Wrapping(unittest.TestCase):
    def test_short_text_is_one_line(self):
        self.assertEqual(F.balanced_lines("软件工程", 16, 3200), ["软件工程"])

    def test_wrap_keeps_every_character(self):
        lines = F.balanced_lines(LONG["title_zh"], 16, 5798)
        self.assertGreater(len(lines), 1)
        self.assertEqual("".join(lines), LONG["title_zh"])
        self.assertTrue(all(F.width(x, 16) <= 5798 for x in lines))

    def test_prefers_bracket_boundary(self):
        self.assertEqual(
            F.balanced_lines("计算机科学与技术（人工智能与大数据方向）", 16, 4791),
            ["计算机科学与技术", "（人工智能与大数据方向）"],
        )

    def test_closing_punctuation_never_starts_a_line(self):
        for text in ("数据结构与算法分析课程设计报告，第二部分。", "学年论文（设计）题目示例文字较长的情况"):
            for line in F.balanced_lines(text, 16, 3000)[1:]:
                self.assertNotIn(line[0], F.NO_LINE_START)

    def test_english_words_are_not_split(self):
        lines = F.balanced_lines("Design and Implementation of a Web-based Library Management System", 16, 5798)
        words = "Design and Implementation of a Web-based Library Management System".split()
        self.assertEqual(" ".join(lines).split(), words)


class Fixture(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def build(self, template, info, fmt=None, level=0):
        fields = E.load_profile(template)["cover"]["required_fields"]
        if fmt and fmt.get("cover_supervisor"):
            fields = fields + ["supervisor"]
        job = {"template": template, "info": {k: v for k, v in info.items() if k in fields}, "format": fmt or {}}
        state = self.d / "intake.json"
        I.from_job(job, state)
        content = self.d / "content.json"
        content.write_text(
            json.dumps({"mode": "document", "metadata": {}, "blocks": [{"type": "paragraph", "text": "合成正文。"}]}, ensure_ascii=False),
            encoding="utf-8",
        )
        out = self.d / f"{template}-{level}.docx"
        E.build(content, out, force=True, intake_path=state, cover_min_level=level)
        return out, state


class Geometry(Fixture):
    def test_default_values_leave_master_geometry_untouched(self):
        out, state = self.build("A", {k: "合成" for k in LONG})
        self.assertTrue(check(out, intake_path=state)["passed"])
        self.assertEqual(E.get_vars(Document(out))["XAUFE_COVER_FIT"], "0")
        master = Document(ROOT / "assets/A/format-master.docx")
        doc = Document(out)
        for i in range(15):
            for tag in ("ind", "tabs", "spacing"):
                a = doc.paragraphs[i]._p.pPr.find(qn("w:" + tag))
                b = master.paragraphs[i]._p.pPr.find(qn("w:" + tag))
                self.assertEqual(a is None, b is None)
                if a is not None and i != 14:
                    self.assertEqual(dict(a.attrib), dict(b.attrib), (i, tag))

    def test_long_values_wrap_inside_field_zone_and_pass_audit(self):
        for template in ("A", "B"):
            with self.subTest(template=template):
                out, state = self.build(template, LONG)
                result = check(out, intake_path=state)
                self.assertTrue(result["passed"], result["errors"])
                doc = Document(out)
                major = doc.paragraphs[11]
                self.assertIn("\n", major.runs[1].text)
                ind = major._p.pPr.find(qn("w:ind"))
                self.assertIsNotNone(ind.get(qn("w:hanging")))
                right = [t.get(qn("w:pos")) for t in major._p.xpath("./w:pPr/w:tabs/w:tab") if t.get(qn("w:val")) == "right"]
                spec = E.load_profile(template)["page_twips"]
                self.assertEqual(int(right[0]), spec["width"] - spec["left"] - spec["right"])
                self.assertGreaterEqual(int(E.get_vars(doc)["XAUFE_COVER_FIT"]), 1)

    def test_font_sizes_never_change(self):
        out, state = self.build("B", LONG, level=4)
        doc = Document(out)
        from style_model import run_properties

        master = Document(ROOT / "assets/B/format-master.docx")
        for i in (2, 5, 9, 10, 11, 12):
            for a, b in zip(doc.paragraphs[i].runs, master.paragraphs[i].runs):
                self.assertEqual(run_properties(doc.paragraphs[i], a._r).get("sz"), run_properties(master.paragraphs[i], b._r).get("sz"))
        self.assertTrue(check(out, intake_path=state)["passed"])

    def test_every_level_passes_the_audit(self):
        for level in range(F.MAX_LEVEL + 1):
            with self.subTest(level=level):
                out, state = self.build("A", LONG, level=level)
                # The requested level is a floor; long text may need a higher one.
                self.assertGreaterEqual(int(E.get_vars(Document(out))["XAUFE_COVER_FIT"]), level)
                self.assertTrue(check(out, intake_path=state)["passed"])

    def test_tampered_fitted_geometry_is_rejected(self):
        out, state = self.build("A", LONG)
        doc = Document(out)
        doc.paragraphs[11]._p.pPr.find(qn("w:ind")).set(qn("w:left"), "100")
        doc.save(out)
        codes = {e["code"] for e in check(out, intake_path=state)["errors"]}
        self.assertIn("COVER_LAYOUT", codes)

    def test_supervisor_added_to_b_follows_field_zone(self):
        info = dict(LONG, supervisor="诸葛某某教授（校外导师：某某研究院高级工程师）")
        out, state = self.build("B", info, {"cover_supervisor": True})
        result = check(out, intake_path=state)
        self.assertTrue(result["passed"], result["errors"])

    def test_impossible_cover_asks_the_user(self):
        info = {k: v * 4 for k, v in LONG.items()}
        with self.assertRaisesRegex(ValueError, "请问用户"):
            self.build("A", info)


if __name__ == "__main__":
    unittest.main()
