"""Originals the user gives (.txt, .md, Word): read as the reader sees them, numbers and all."""

from __future__ import annotations
import json, unittest
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import test_xaufe_cli as base  # a module import: its CLI tests are not collected here a second time
from test_xaufe_cli import CONTENT, JOB_B
import xaufe as X
from fixtures import ABSTRACT


class Originals(base.CLI):
    def run_docx(self, doc, name, edit=None):
        source = self.d / name
        doc.save(source)
        work = self.new(JOB_B)
        code, inspected, err = self.call("inspect", "--work", work, "--docx", source)
        self.assertEqual(code, 0, inspected + err)
        if edit:
            mapping = json.loads((work / "mapping.json").read_text(encoding="utf-8"))
            edit(mapping)
            (work / "mapping.json").write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        return work, code, inspected, out + err

    # ------------------------------------------------------------ .txt / .md
    def test_plain_text_keeps_every_character_and_its_line_breaks(self):
        source = self.d / "原稿.txt"
        source.write_text("\n".join([
            "排序算法比较", "", "摘要", ABSTRACT, "", "关键词：合成；测试", "", "一、引言",
            "第一行是一个段落", "第二行是另一个段落", "（一）背景", "需要的步骤：", "1. 需求分析", "2. 系统设计",
            "表1 合成表", "编号\t名称\t密码", "1\t甲\t***", "二、结论", "合成结论。", "", "参考文献",
            "[1] 合成作者. 合成文献[M]. 北京: 合成出版社, 2020.",
        ]), encoding="utf-8")
        work = self.new(JOB_B, sources=[source])
        code, out, err = self.call("inspect", "--work", work, "--text", source)
        self.assertEqual(code, 0, out + err)
        blocks = json.loads((work / "content.json").read_text(encoding="utf-8"))["blocks"]
        texts = [b.get("text") or "".join(r["text"] for r in b.get("runs", [])) for b in blocks]
        self.assertIn("第一行是一个段落", texts)  # a short line is its own paragraph, not glued to the next
        self.assertIn("第二行是另一个段落", texts)
        self.assertIn("1. 需求分析", texts)  # numbered lines in a row are a list in the text, not headings
        self.assertEqual([b["level"] for b in blocks if b["type"] == "heading"], [1, 2, 1])
        table = next(b for b in blocks if b["type"] == "table")
        self.assertIn(["1", "甲", "***"], [table["header"]] + table["rows"])  # *** is the user's text
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)

    def test_markdown_reads_as_shown(self):
        from textio import md_visible, continues
        from text_import import _pieces

        self.assertEqual(md_visible("用 `a*b*c` 和 **加粗** 以及 snake_case_name"), "用 a*b*c 和 加粗 以及 snake_case_name")
        self.assertEqual(md_visible("单独的 * 号和 2*3"), "单独的 * 号和 2*3")
        self.assertEqual(len(_pieces("- 第一项\n- 第二项", markdown=True)), 2)
        long = "这是一行被硬折行的很长的合成正文，它在固定宽度处断开了没有标点"
        self.assertTrue(continues(long, "下一行接着写。", len(long)))
        self.assertFalse(continues("短行", "下一行。", len(long)))
        self.assertFalse(continues(long + "。", "下一行。", len(long) + 1))

    # ------------------------------------------------------------ Word numbering
    def numbered_doc(self, fmt="chineseCounting", text="%1、", headings=("引言", "方法")):
        """Headings numbered by a Word list: the label is not in the paragraph text."""
        doc = Document()
        doc.add_paragraph("排序算法比较")
        doc.add_paragraph("摘要")
        doc.add_paragraph(ABSTRACT)
        doc.add_paragraph("关键词：编号；标题")
        numbering = doc.part.numbering_part.element
        numbering.insert(0, parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
            f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/><w:suff w:val="nothing"/></w:lvl></w:abstractNum>'))
        numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
        for title in headings:
            h = doc.add_heading(title, 1)
            h._p.get_or_add_pPr().append(parse_xml(
                f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="90"/></w:numPr>'))
            doc.add_paragraph(f"{title}的合成正文。")
        doc.add_paragraph("参考文献")
        doc.add_paragraph("[1] 合成作者. 合成文献[M]. 北京: 合成出版社, 2020.")
        return doc

    def test_list_labels_follow_word_counters(self):
        from numbering import _chinese, _roman, list_labels

        self.assertEqual([_chinese(n) for n in (1, 10, 12, 20, 105, 110)], ["一", "十", "十二", "二十", "一百零五", "一百一十"])
        self.assertEqual(_roman(14), "XIV")
        self.assertEqual(list(list_labels(self.numbered_doc()).values()), ["一、", "二、"])
        doc = self.numbered_doc(fmt="decimal", text="%1.", headings=("引言", "方法", "结论"))
        self.assertEqual(list(list_labels(doc).values()), ["1.", "2.", "3."])

    def test_auto_numbered_headings_keep_their_visible_number_as_text(self):
        work, code, _, out = self.run_docx(self.numbered_doc(), "自动编号.docx")
        self.assertEqual(code, 0, out)
        self.assertIn("自动编号转成了同样的文字", out)
        heads = [p for p in Document(work / "final.docx").paragraphs if p.style.name == "Heading 1"]
        self.assertEqual([p.text for p in heads], ["一、引言", "二、方法"])
        self.assertFalse(any(p._p.xpath("./w:pPr/w:numPr") for p in heads))

    def test_style_numbered_list_items_keep_their_numbers(self):
        doc = Document()
        doc.add_paragraph("合成正文，下面是列表：")
        doc.add_paragraph("第一项", style="List Number")
        doc.add_paragraph("第二项", style="List Number")
        work, code, _, out = self.run_docx(doc, "列表.docx")
        self.assertEqual(code, 0, out)
        items = [p for p in Document(work / "final.docx").paragraphs if p.text in ("第一项", "第二项")]
        self.assertEqual(len(items), 2)
        self.assertTrue(all(p._p.xpath("./w:pPr/w:numPr/w:numId") for p in items))

    # ------------------------------------------------------------ Word structure
    def test_inline_abstract_label_is_kept_and_the_generated_label_follows_the_title(self):
        doc = Document()
        doc.add_paragraph("排序算法比较")
        doc.add_paragraph("摘要：" + ABSTRACT)
        doc.add_paragraph("关键词：合成；摘要")
        doc.add_heading("一、绪论", 1)
        doc.add_paragraph("合成正文。")
        doc.add_paragraph("参考文献")
        doc.add_paragraph("[1] 合成作者. 合成文献[M]. 北京: 合成出版社, 2020.")
        work, code, inspected, out = self.run_docx(doc, "行内摘要.docx")
        self.assertIn("p1 以“摘要：”开头", inspected)
        self.assertEqual(code, 0, out)
        shown = [p.text for p in Document(work / "final.docx").paragraphs if p.text.strip()]
        self.assertLess(shown.index("排序算法比较"), shown.index("内 容 摘 要"))
        self.assertLess(shown.index("内 容 摘 要"), shown.index("摘要：" + ABSTRACT))

    def test_errors_after_a_delete_name_the_ids_in_mapping_json(self):
        doc = Document()
        for text in ("要删掉的合成段落", "合成正文。", "另一段合成正文。"):
            doc.add_paragraph(text)

        def edit(mapping):
            mapping["paragraphs"][0]["role"] = "delete"
            mapping["paragraphs"][2]["role"] = "preserve"
        _, code, _, out = self.run_docx(doc, "删段.docx", edit)
        self.assertNotEqual(code, 0)
        self.assertIn("p2“另一段合成正文。”的 role 是 preserve，但它在 p1", out)
        self.assertIn("旧目录用 old_toc", out)

    def test_table_caption_below_its_table_is_told(self):
        doc = Document()
        doc.add_paragraph("合成正文，下表是合成数据。")
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, (("甲", "乙"), ("1", "2"))):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.add_paragraph("表1 合成表")
        doc.add_paragraph("合成正文结尾。")
        _, code, _, out = self.run_docx(doc, "表题在下.docx")
        self.assertEqual(code, 0, out)
        self.assertIn("表题“表1 合成表”在表格下面", out)

    def test_tracked_changes_are_located(self):
        from unsupported import unsupported_reasons

        doc = Document()
        p = doc.add_paragraph("合成正文，")
        ins = OxmlElement("w:ins")
        ins.set(qn("w:id"), "1")
        ins.set(qn("w:author"), "合成")
        run, t = OxmlElement("w:r"), OxmlElement("w:t")
        t.text = "修订加的字"
        run.append(t)
        ins.append(run)
        p._p.append(ins)
        self.assertIn("p0“合成正文，修订加的字”", unsupported_reasons(doc)[0])

    # ------------------------------------------------------------ job.json
    def test_json_error_shows_the_spot_and_says_escape_not_change_quotes(self):
        path = self.d / "job.json"
        path.write_text('{"info": {"title_zh": "关于"AI"的研究"}}', encoding="utf-8")
        with self.assertRaises(ValueError) as caught:
            X._read(path, "job.json")
        message = str(caught.exception)
        self.assertIn("第 1 行第", message)
        self.assertIn('\\"', message)
        self.assertIn("不要换成中文引号", message)

    def test_line_needs_line_rule_and_font_points_to_cn_latin(self):
        work = self.new({**JOB_B, "format": {"styles": {"body": {"line": 360}}}}, CONTENT)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotEqual(code, 0)
        self.assertIn("也要写 line_rule", out + err)
        from customization import locate_format
        with self.assertRaises(ValueError) as caught:
            locate_format({"styles": {"body": {"font": "宋体"}}})
        self.assertIn("cn（中文字体）", str(caught.exception))

    # ------------------------------------------------------------ format-only corrections
    def paper_doc(self, abstract, title="排序算法比较", references=True):
        doc = Document()
        doc.add_paragraph(title)
        doc.add_paragraph("摘要")
        doc.add_paragraph(abstract)
        doc.add_paragraph("关键词：排序；算法")
        doc.add_heading("一、引言", 1)
        doc.add_paragraph("合成正文。")
        if references:
            doc.add_paragraph("参考文献")
            doc.add_paragraph("[1] 合成作者. 合成文献[M]. 北京: 合成出版社, 2020.")
        return doc

    def with_job(self, work, **extra):
        job = json.loads((work / "job.json").read_text(encoding="utf-8"))
        job.update(extra)
        (work / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")

    def test_short_abstract_blocks_until_corrected_and_the_user_is_told(self):
        source = self.d / "短摘要.docx"
        self.paper_doc("本文比较了几种排序算法。").save(source)
        work = self.new(JOB_B)
        self.call("inspect", "--work", work, "--docx", source)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotEqual(code, 0)
        self.assertIn("[ABSTRACT_LENGTH]", out)
        self.assertIn("学校要求 150～300 字", out)
        self.with_job(work, corrections=[{"part": "abstract", "before": "本文比较了几种排序算法。", "after": ABSTRACT,
                                          "why": "原摘要 12 字，学校要求 150～300 字"}])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        self.assertIn("原摘要 12 字", out)
        shown = [p.text for p in Document(work / "final.docx").paragraphs]
        self.assertIn(ABSTRACT, shown)
        self.assertNotIn("本文比较了几种排序算法。", shown)
        import delivery as D
        self.assertTrue(any("原摘要 12 字" in line for line in D.load(work)[1].get("tell_user", [])))

    def test_user_can_keep_an_off_limit_abstract(self):
        source = self.d / "保留.docx"
        self.paper_doc("本文比较了几种排序算法。").save(source)
        work = self.new(JOB_B)
        self.call("inspect", "--work", work, "--docx", source)
        self.with_job(work, keep_as_is="摘要就这样，别动")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        import delivery as D
        self.assertTrue(any("按你的意思保留了原样" in line for line in D.load(work)[1].get("tell_user", [])))

    def test_added_reference_must_be_verified_first(self):
        source = self.d / "无文献.docx"
        self.paper_doc(ABSTRACT, references=False).save(source)
        work = self.new(JOB_B)
        self.call("inspect", "--work", work, "--docx", source)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotEqual(code, 0)
        self.assertIn("只排版也要补", out)
        self.with_job(work, corrections=[{"part": "references", "after": "合成作者. 补的文献[M]. 北京: 合成出版社, 2020.",
                                          "why": "原稿没有参考文献"}])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotEqual(code, 0)
        self.assertIn("还没有核对", out)

    def test_correction_before_must_quote_the_original(self):
        source = self.d / "找不到.docx"
        self.paper_doc(ABSTRACT).save(source)
        work = self.new(JOB_B)
        self.call("inspect", "--work", work, "--docx", source)
        self.with_job(work, corrections=[{"part": "title", "before": "不存在的题目", "after": "排序算法比较", "why": "题目超长"}])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotEqual(code, 0)
        self.assertIn("before 在原稿里找不到这一段", out)
        self.with_job(work, corrections=[{"part": "title", "after": "x"}])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertIn("缺 why", out + err)

    def test_long_title_is_flagged_with_its_count(self):
        from audit_helpers import word_count

        self.assertEqual(word_count("基于BERT的文本分类研究"), 10)
        self.assertEqual(word_count("城市公共交通票价结构对居民通勤方式选择的影响研究"), 24)
        source = self.d / "长题目.docx"
        title = "城市公共交通票价结构对居民通勤方式选择影响的实证分析研究"
        self.paper_doc(ABSTRACT, title=title).save(source)
        work = self.new({**JOB_B, "info": {**JOB_B["info"], "title_zh": title}})
        self.call("inspect", "--work", work, "--docx", source)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotEqual(code, 0)
        self.assertIn("[TITLE_TOO_LONG]", out)
        self.assertIn("有 28 字", out)

    def test_off_template_heading_numbers_are_corrected_and_told(self):
        doc = self.paper_doc(ABSTRACT)
        body = doc.paragraphs[4]._p
        for text, level in (("1 绪论", 1), ("1.1 研究背景", 2)):
            h = doc.add_heading(text, level)
            body.addprevious(h._p)
        source = self.d / "十进制标题.docx"
        doc.save(source)
        work = self.new(JOB_B)
        self.call("inspect", "--work", work, "--docx", source)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        heads = [p.text for p in Document(work / "final.docx").paragraphs if p.style.name.startswith("Heading")]
        self.assertEqual(heads[:3], ["一、绪论", "（一）研究背景", "二、引言"])
        import delivery as D
        told = " ".join(D.load(work)[1].get("tell_user", []))
        self.assertIn("“1 绪论”改为“一、绪论”", told)

    def test_missing_abstract_and_keywords_are_drafted_into_the_paper(self):
        doc = Document()
        doc.add_paragraph("排序算法比较")
        doc.add_heading("一、引言", 1)
        doc.add_paragraph("合成正文。")
        doc.add_paragraph("参考文献")
        doc.add_paragraph("[1] 合成作者. 合成文献[M]. 北京: 合成出版社, 2020.")
        source = self.d / "缺摘要.docx"
        doc.save(source)
        work = self.new(JOB_B)
        code, inspected, _ = self.call("inspect", "--work", work, "--docx", source)
        self.assertIn("原稿缺中文摘要、关键词：只排版也要补全", inspected)
        self.assertEqual(json.loads((work / "mapping.json").read_text(encoding="utf-8"))["mode"], "paper")
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertNotEqual(code, 0)
        self.assertIn("只排版也要补", out)
        self.assertNotIn("目\u3000录", out)
        self.with_job(work, corrections=[
            {"part": "abstract", "after": ABSTRACT, "why": "原稿没有摘要，根据正文起草"},
            {"part": "keywords", "after": "关键词：排序；算法", "why": "原稿没有关键词，根据正文提取"}])
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        shown = [p.text for p in Document(work / "final.docx").paragraphs if p.text.strip()]
        self.assertLess(shown.index("内 容 摘 要"), shown.index(ABSTRACT))
        self.assertLess(shown.index(ABSTRACT), shown.index("关键词：排序；算法"))
        self.assertIn("补了摘要", out)

    def test_old_cover_ending_in_a_page_break_is_replaced_and_told(self):
        from docx.enum.text import WD_BREAK

        doc = Document()
        for text in ("合成大学", "学年论文", "题目：排序算法比较", "学生姓名：合成同学", "学号：20230001",
                     "指导教师：合成老师", "诚信承诺：本文为本人独立完成。"):
            doc.add_paragraph(text)
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        body = self.paper_doc(ABSTRACT)
        for p in body.paragraphs:
            doc.element.body.append(__import__("copy").deepcopy(p._p))
        doc.element.body.append(doc.element.body.find(qn("w:sectPr")))
        source = self.d / "旧封面.docx"
        doc.save(source)
        work = self.new(JOB_B)
        code, inspected, err = self.call("inspect", "--work", work, "--docx", source)
        self.assertIn("检测到原稿开头有旧封面", inspected)
        self.assertIn("学号：20230001", inspected)
        code, out, err = self.call("run", "--work", work, "--office", "none")
        self.assertEqual(code, 0, out + err)
        text = "\n".join(p.text for p in Document(work / "final.docx").paragraphs)
        self.assertNotIn("诚信承诺", text)
        import delivery as D
        self.assertIn("旧封面", " ".join(D.load(work)[1].get("tell_user", [])))


# The inherited CLI tests already run in test_xaufe_cli.
for _name in [n for n in vars(base.CLI) if n.startswith("test_")]:
    setattr(Originals, _name, None)

if __name__ == "__main__":
    unittest.main()
