"""Customization and review tests. All identities, COM responses and review receipts are synthetic.

Real rendering and real page inspection are in acceptance_cases.py, not here.
No test receipt is evidence of an actual subagent invocation or native WPS use.
"""

from __future__ import annotations
from copy import deepcopy
import hashlib, json, os, subprocess, sys, tempfile, unittest, zipfile
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import customization as C, engine as E, intake as I, delivery as D
import office_discovery as O, office_backend as B, wps_support as W, wps_exchange as X
from check_format import check
from style_model import run_properties


def save(p, data):
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


class Fixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def intake(self, template="B", request=None):
        state = self.root / "intake.json"
        I.start(state, "SYNTHETIC test")
        I.select(state, template, "合成测试选择")
        if request:
            I.customize(state, request, "合成格式要求")
        spec = C.profile_for(I.read(state))
        values = {
            "title_zh": "合成排版测试",
            "student_name": "示例同学",
            "student_id": "000123",
            "major": "软件工程",
            "class_name": "示例班",
            "supervisor": "示例教师",
            "course_name": "格式测试",
            "date": "2026-09-25",
        }
        I.resolve(
            state,
            {
                "metadata": {
                    k: {"value": values[k], "source": "user", "evidence": "仅合成数据"}
                    for k in spec["cover"]["required_fields"]
                }
            },
        )
        return state

    def build(self, request=None, template="B", extra=None):
        state = self.intake(template, request)
        data = json.loads(
            (
                ROOT / f"examples/{template}/content.json"
            ).read_text()
        )
        data["metadata"] = (
            {"title_en": "Synthetic Formatting Test"} if template == "A" else {}
        )
        for block in data["blocks"]:
            if block["type"] == "image":
                block["path"] = str(ROOT / "examples/layout-test.png")
        if extra:
            data["blocks"] += extra
        src = self.root / "content.json"
        save(src, data)
        out = self.root / "draft.docx"
        E.build(src, out, intake_path=state)
        return out, state

    def passed(self, out, state):
        result = check(out, intake_path=state)
        self.assertTrue(result["passed"], result["errors"])


class MergeTests(Fixture):
    def test_siblings_preserved(self):
        a = {"styles": {"body": {"cn": "仿宋", "size": 12}, "toc1": {"size": 14}}}
        b = C.merge_request(a, {"styles": {"toc1": {"size": 16}}})
        self.assertEqual(b["styles"]["body"], a["styles"]["body"])
        self.assertEqual(b["styles"]["toc1"]["size"], 16)

    def test_leaf_reset(self):
        self.assertEqual(
            C.merge_request(
                {"styles": {"body": {"cn": "仿宋", "size": 12}}},
                {"styles": {"body": {"size": None}}},
            ),
            {"styles": {"body": {"cn": "仿宋"}}},
        )

    def test_role_reset(self):
        self.assertEqual(
            C.merge_request({"styles": {"body": {"size": 12}}}, {"styles": {"body": None}}), {}
        )

    def test_section_reset(self):
        self.assertEqual(
            C.merge_request(
                {"styles": {"body": {"size": 12}}, "logo": "green"}, {"styles": None}
            ),
            {"logo": "green"},
        )

    def test_empty_noop(self):
        a = {"styles": {"body": {"cn": "仿宋"}}}
        self.assertEqual(C.merge_request(a, {}), a)

    def test_arrays_replace(self):
        self.assertEqual(C.merge_request({"parts": ["a", "b"]}, {"parts": ["c"]}), {"parts": ["c"]})

    def test_input_not_mutated(self):
        a = {"styles": {"body": {"size": 12}}}
        old = deepcopy(a)
        b = C.merge_request(a, {"styles": {"toc1": {"size": 18}}})
        b["styles"]["body"]["size"] = 10
        self.assertEqual(a, old)

    def test_unknown_null_rejected(self):
        with self.assertRaises(ValueError):
            C.merge_request({}, {"unknown": None})

    def test_unknown_role_null_rejected(self):
        with self.assertRaises(ValueError):
            C.merge_request({}, {"styles": {"madeup": None}})

    def test_unknown_attribute_null_rejected(self):
        with self.assertRaises(ValueError):
            C.merge_request({}, {"styles": {"body": {"madeup": None}}})

    def test_styles_must_be_object(self):
        with self.assertRaises(ValueError):
            C.merge_request({}, {"styles": ["body"]})

    def test_invalid_keyword_null_rejected(self):
        with self.assertRaises(ValueError):
            C.merge_request({}, {"keyword_labels": {"zh": {"madeup": None}}})

    def test_invalid_header_null_rejected(self):
        with self.assertRaises(ValueError):
            C.merge_request({}, {"header": {"madeup": None}})

    def test_invalid_cover_region_null_rejected(self):
        with self.assertRaises(ValueError):
            C.merge_request({}, {"cover_fonts": {"madeup": None}})

    def test_label_multiturn_merge(self):
        a = {"keyword_labels": {"zh": {"size": 12, "italic": True}}}
        self.assertEqual(
            C.merge_request(a, {"keyword_labels": {"zh": {"italic": None}}}),
            {"keyword_labels": {"zh": {"size": 12}}},
        )

    def test_multi_turn_state(self):
        s = self.intake(request={"styles": {"body": {"cn": "仿宋"}}})
        I.customize(s, {"styles": {"toc1": {"size": 18}}}, "再改目录")
        self.assertEqual(I.read(s)["format_request"]["styles"]["body"]["cn"], "仿宋")

    def test_explicit_reset(self):
        s = self.intake(request={"styles": {"body": {"cn": "仿宋"}}})
        I.customize(s, {}, "恢复全部默认", reset=True)
        self.assertEqual(I.read(s)["format_request"], {})

    def test_explicit_replace(self):
        s = self.intake(request={"logo": "green"})
        I.customize(s, {"styles": {"body": {"size": 12}}}, "替换全部定制", replace=True)
        self.assertNotIn("logo", I.read(s)["format_request"])

    def test_reset_nonempty_rejected(self):
        s = self.intake()
        with self.assertRaises(ValueError):
            I.customize(s, {"logo": "green"}, "矛盾", reset=True)

    def test_invalid_change_atomic(self):
        s = self.intake()
        before = s.read_bytes()
        with self.assertRaises(ValueError):
            I.customize(s, {"styles": {"body": {"size": 0}}}, "无效测试")
        self.assertEqual(before, s.read_bytes())

    def test_evidence_required(self):
        s = self.intake()
        with self.assertRaises(ValueError):
            I.customize(s, {}, "")

    def test_absolute_indent_replaces_old_chars(self):
        self.assertEqual(
            C.merge_request(
                {"styles": {"body": {"chars": 200}}}, {"styles": {"body": {"first": 480}}}
            )["styles"]["body"],
            {"first": 480},
        )

    def test_chars_replace_hanging(self):
        self.assertEqual(
            C.merge_request(
                {"styles": {"body": {"hanging": 400}}}, {"styles": {"body": {"chars": 200}}}
            )["styles"]["body"],
            {"chars": 200},
        )

    def test_spacing_units_replace(self):
        self.assertEqual(
            C.merge_request(
                {"styles": {"body": {"before_lines": 100}}}, {"styles": {"body": {"before": 200}}}
            )["styles"]["body"],
            {"before": 200},
        )

    def test_explicit_conflicting_units_rejected(self):
        with self.assertRaises(ValueError):
            C.effective_profile("B", {"styles": {"body": {"first": 200, "chars": 200}}})

    def test_new_required_cover_field_only(self):
        s = self.intake()
        v = I.customize(s, {"cover_supervisor": True}, "加指导教师")
        self.assertEqual(v["missing"], ["supervisor"])
        self.assertEqual(v["metadata"]["student_id"], "000123")

    def test_history_retains_diff(self):
        s = self.intake()
        I.customize(s, {"styles": {"body": {"size": 12}}}, "合成改字号")
        events = [e for e in I.read(s)["events"] if e["action"] == "user_format_override"]
        self.assertEqual(events[-1]["mode"], "merge")
        self.assertEqual(events[-1]["previous"], {})


class TypographyTests(Fixture):
    def test_body_color(self):
        o, s = self.build({"styles": {"body": {"color": "336699"}}})
        self.passed(o, s)

    def test_heading_italic_color(self):
        o, s = self.build({"styles": {"h1": {"color": "0066AA", "italic": True}}})
        self.passed(o, s)

    def test_body_explicit_italic_false(self):
        o, s = self.build(
            {"styles": {"body": {"italic": False}}},
            extra=[{"type": "paragraph", "runs": [{"text": "合成斜体", "italic": True}]}],
        )
        self.passed(o, s)

    def test_body_underline(self):
        o, s = self.build({"styles": {"body": {"underline": "double"}}})
        self.passed(o, s)

    def test_toc_custom_color_and_italic(self):
        o, s = self.build(
            {"styles": {"toc2": {"color": "225588", "italic": True, "underline": "single"}}}
        )
        self.passed(o, s)

    def test_wrong_custom_color_rejected(self):
        o, s = self.build({"styles": {"body": {"color": "0066AA"}}})
        d = Document(o)
        p = next(p for p in d.paragraphs if p.style.name == "XAUFE Body")
        E.element(p.runs[0]._r.get_or_add_rPr(), "w:color", val="FF0000")
        d.save(o)
        self.assertIn("FONT_COLOR", [v["code"] for v in check(o, intake_path=s)["errors"]])

    def test_wrong_custom_italic_rejected(self):
        o, s = self.build({"styles": {"body": {"italic": True}}})
        d = Document(o)
        p = next(p for p in d.paragraphs if p.style.name == "XAUFE Body")
        E.element(p.runs[0]._r.get_or_add_rPr(), "w:i", val=0)
        d.save(o)
        self.assertIn("FONT_ITALIC", [v["code"] for v in check(o, intake_path=s)["errors"]])

    def test_invalid_color_rejected(self):
        with self.assertRaises(ValueError):
            C.effective_profile("B", {"styles": {"body": {"color": "blue"}}})

    def test_invalid_underline_rejected(self):
        with self.assertRaises(ValueError):
            C.effective_profile("B", {"styles": {"body": {"underline": "invented"}}})

    def test_keyword_label_independent(self):
        o, s = self.build(
            {
                "keyword_labels": {
                    "zh": {
                        "cn": "仿宋",
                        "latin": "Arial",
                        "size": 12,
                        "bold": True,
                        "color": "884422",
                        "italic": True,
                    }
                },
                "styles": {"keywords_zh": {"size": 10.5}},
            }
        )
        self.passed(o, s)

    def test_english_keyword_label(self):
        o, s = self.build(
            {"keyword_labels": {"en": {"size": 12, "italic": True, "underline": "single"}}},
            "A",
        )
        self.passed(o, s)

    def test_invalid_label_language(self):
        with self.assertRaises(ValueError):
            C.effective_profile("B", {"keyword_labels": {"jp": {"size": 12}}})

    def test_label_must_be_font_only(self):
        with self.assertRaises(ValueError):
            C.effective_profile("B", {"keyword_labels": {"zh": {"align": "right"}}})

    def test_four_levels_opt_in(self):
        o, s = self.build(
            {"heading_levels": 4, "styles": {"h4": {"color": "444444"}, "toc4": {"size": 10.5}}},
            extra=[
                {"type": "heading", "level": 4, "text": "（1）第四级合成标题"},
                {"type": "paragraph", "text": "明确启用的额外层级。"},
            ],
        )
        self.passed(o, s)

    def test_four_levels_without_optin_rejected(self):
        with self.assertRaises(ValueError):
            self.build(extra=[{"type": "heading", "level": 4, "text": "四级"}])

    def test_h4_config_requires_optin(self):
        with self.assertRaises(ValueError):
            C.effective_profile("B", {"styles": {"h4": {"size": 12}}})

    def test_fifth_level_rejected(self):
        with self.assertRaises(ValueError):
            C.effective_profile("B", {"heading_levels": 5})

    def test_bool_depth_rejected(self):
        with self.assertRaises(ValueError):
            C.effective_profile("B", {"heading_levels": True})

    def test_character_indent_physical_fallback(self):
        o, s = self.build({"styles": {"body": {"chars": 200, "size": 12}}})
        self.passed(o, s)
        self.assertEqual(
            Document(o).styles["XAUFE Body"].element.xpath("./w:pPr/w:ind/@w:firstLine"), ["480"]
        )

    def test_right_indent(self):
        o, s = self.build({"styles": {"body": {"right": 280}}})
        self.passed(o, s)

    def test_cover_color_and_italic(self):
        o, s = self.build({"cover_fonts": {"title": {"color": "006699", "italic": True}}})
        self.passed(o, s)

    def test_added_header_gets_safe_default_distance(self):
        self.assertEqual(
            C.effective_profile("B", {"header": {"enabled": True}})["page_twips"][
                "header"
            ],
            C.load_profile("A")["page_twips"]["header"],
        )

    def test_added_header_explicit_distance_wins(self):
        self.assertEqual(
            C.effective_profile(
                "B", {"header": {"enabled": True}, "page_twips": {"header": 120}}
            )["page_twips"]["header"],
            120,
        )

    def test_default_red_header_distance_unchanged(self):
        self.assertEqual(
            C.effective_profile("B")["page_twips"]["header"],
            C.load_profile("B")["page_twips"]["header"],
        )

    def test_header_color(self):
        o, s = self.build(
            {
                "header": {"enabled": True, "left": "自定义", "right": "页眉"},
                "styles": {"header": {"color": "777777", "italic": True}},
            }
        )
        self.passed(o, s)

    def test_restyle_custom_format(self):
        o, s = self.build(
            {"styles": {"body": {"color": "227799", "italic": True, "underline": "single"}}}
        )
        inv = E.inventory(o)
        roles = {
            b["id"]: next(k for k, v in E.ROLES.items() if v == b["style"])
            for b in inv["blocks"]
            if b["kind"] == "paragraph" and b["block_index"] >= 0 and b["style"] in E.ROLES.values()
        }
        # Cover indices are preserved explicitly; no content inferred from style alone.
        for i in range(E.SPEC["cover"]["paragraph_count"]):
            roles["p" + str(i)] = "cover"
        mapping = self.root / "mapping.json"
        save(
            mapping,
            {
                "input_sha256": inv["input_sha256"],
                "roles": roles,
                "default_role": "spacer",
                "has_cover": True,
                "mode": "paper",
                "format_tables": True,
            },
        )
        out = self.root / "restyled.docx"
        E.restyle(o, mapping, out, intake_path=s)
        self.passed(out, s)


class WPSTests(Fixture):
    def test_windows_paths(self):
        self.assertTrue(any("office6" in x[1] for x in W.known_paths("Windows", {}, self.root)))

    def test_mac_paths(self):
        self.assertTrue(any(".app" in x[1] for x in W.known_paths("Darwin", {}, self.root)))

    def test_linux_paths(self):
        self.assertTrue(any(x[1] == "/usr/bin/wps" for x in W.known_paths("Linux", {}, self.root)))

    def test_explicit_path_discovered(self):
        p = self.root / "wps.exe"
        p.write_text("fixture")
        with patch.dict(os.environ, {"XAUFE_WPS_PATH": str(p)}):
            self.assertTrue(
                any(c["engine"] == "wps" and c["found"] for c in O.discover()["candidates"])
            )

    def test_preference_wps(self):
        report = {
            "runtime": {"system": "Windows"},
            "candidates": [
                {"engine": "wps", "path": "x", "source": "PATH", "found": True, "automation": True},
                {
                    "engine": "word",
                    "path": "y",
                    "source": "PATH",
                    "found": True,
                    "automation": True,
                },
            ],
        }
        with patch.dict(os.environ, {"XAUFE_RENDER_ENGINE": "wps"}):
            self.assertEqual([c["engine"] for c in O.usable_candidates(report)], ["wps"])

    def test_linux_install_not_automatic(self):
        p = self.root / "wps"
        p.write_text("fixture")
        with (
            patch("office_discovery.platform.system", return_value="Linux"),
            patch.dict(os.environ, {"XAUFE_WPS_PATH": str(p)}),
        ):
            self.assertFalse(
                next(c for c in O.discover()["candidates"] if c["path"] == str(p))["automation"]
            )

    def test_probe_fallback_progid(self):
        calls = []

        def call(script, timeout):
            calls.append(script)
            return subprocess.CompletedProcess([], 1 if len(calls) == 1 else 0, "WPS mock", "")

        result = W.probe(call)
        self.assertTrue(result["usable"])
        self.assertEqual(result["progid"], "Wps.Application")

    def test_probe_failure_not_uninstalled(self):
        self.assertFalse(
            W.probe(lambda *a: subprocess.CompletedProcess([], 1, "", "synthetic failure"))[
                "usable"
            ]
        )

    def test_script_escapes_apostrophe(self):
        import base64, re

        values = [
            base64.b64decode(x).decode("utf-8")
            for x in re.findall(
                r"FromBase64String\('([^']+)'\)",
                W.com_script("Kwps.Application", "C:/a'b.docx", "C:/a.pdf"),
            )
        ]
        self.assertIn("C:/a'b.docx", values)

    def test_script_opens_read_only(self):
        self.assertIn("$false,$true,$false", W.com_script("Kwps.Application", "x", "y"))

    def test_script_does_not_resave_source(self):
        self.assertNotIn(".Save(", W.com_script("Kwps.Application", "x", "y"))

    def test_script_exports_heading_bookmarks(self):
        self.assertIn(
            "17,$false,0,0,1,1,0,$false,$false,1", W.com_script("Kwps.Application", "x", "y")
        )

    def test_script_no_broad_kill(self):
        self.assertNotIn("Stop-Process", W.com_script("Kwps.Application", "x", "y"))
        self.assertIn("$hadProcess", W.com_script("Kwps.Application", "x", "y"))

    def test_script_unknown_progid_refused(self):
        with self.assertRaises(ValueError):
            W.com_script("anything'; evil")

    def test_wps_backend_dispatch(self):
        import pymupdf as fitz

        source = self.root / "a.docx"
        Document().save(source)
        target = self.root / "a.pdf"

        def fake(docx, pdf, candidate, tmp, ps):
            d = fitz.open()
            d.new_page().insert_text((40, 40), "synthetic")
            d.save(pdf)
            d.close()
            return {"version": "mock WPS"}

        candidate = {
            "engine": "wps",
            "path": "fake",
            "source": "fixture",
            "found": True,
            "automation": True,
        }
        with (
            patch("office_backend.discover", return_value={}),
            patch("office_backend.usable_candidates", return_value=[candidate]),
            patch("wps_support.export", side_effect=fake),
        ):
            result = B.export_pdf(source, target)
        self.assertEqual(result["engine"], "wps")
        self.assertTrue(target.is_file())

    def test_invalid_pdf_backend_rejected(self):
        source = self.root / "a.docx"
        Document().save(source)

        def fake(docx, pdf, *a):
            Path(pdf).write_text("invalid")
            return "mock"

        candidate = {
            "engine": "wps",
            "path": "fake",
            "source": "fixture",
            "found": True,
            "automation": True,
        }
        original_temp = tempfile.TemporaryDirectory
        before = set(self.root.iterdir())
        with (
            patch("office_backend.discover", return_value={}),
            patch("office_backend.usable_candidates", return_value=[candidate]),
            patch("wps_support.export", side_effect=fake),
            patch(
                "office_backend.tempfile.TemporaryDirectory",
                side_effect=lambda **kw: original_temp(dir=self.root, **kw),
            ),
        ):
            with self.assertRaises(B.OfficeExportFailed):
                B.export_pdf(source, self.root / "x.pdf")
        self.assertEqual(
            set(self.root.iterdir()),
            before,
            "Invalid PDF must not leak a locked temporary directory",
        )


class ExchangeTests(Fixture):
    def setup_exchange(self):
        session = D.create(
            self.root / "output.docx",
            parent=self.root,
        )
        work = Path(session["workspace"])
        source = work / "draft.docx"
        d = Document()
        d.add_paragraph("Synthetic bridge document")
        d.save(source)
        return work, source, work / "wps-exchange"

    def pending(self):
        work, source, exchange = self.setup_exchange()
        with self.assertRaises(X.ManualExportRequired) as cm:
            X.export_cached(source, work / "render.pdf", exchange)
        return work, source, exchange, cm.exception.request

    def pdf(self, path):
        import pymupdf as fitz

        d = fitz.open()
        d.new_page().insert_text((72, 72), "Synthetic bridge document")
        d.save(path)
        d.close()
        return path

    def test_pending_request_created(self):
        w, s, e, r = self.pending()
        self.assertTrue(r.is_file())
        self.assertFalse((w / "render.pdf").exists())

    def test_arbitrary_root_refused(self):
        with self.assertRaises(ValueError):
            X.exchange_root(self.root / "arbitrary")

    def test_valid_export_resumed(self):
        w, s, e, r = self.pending()
        p = self.pdf(w / "input.pdf")
        X.accept(r, p, "Synthetic test PDF, NOT native WPS", "unit fixture")
        out = w / "render.pdf"
        v = X.export_cached(s, out, e)
        self.assertEqual(v["engine"], "wps-manual")
        self.assertEqual(p.read_bytes(), out.read_bytes())

    def test_evidence_missing_refused(self):
        w, s, e, r = self.pending()
        with self.assertRaises(ValueError):
            X.accept(r, self.pdf(w / "input.pdf"), "", "")

    def test_changed_request_source_refused(self):
        w, s, e, r = self.pending()
        (r.parent / "source.docx").write_bytes(b"corrupt")
        with self.assertRaises(ValueError):
            X.accept(r, self.pdf(w / "input.pdf"), "fixture", "fixture")

    def test_changed_pdf_refused(self):
        w, s, e, r = self.pending()
        X.accept(r, self.pdf(w / "in.pdf"), "fixture", "fixture")
        (r.parent / "accepted.pdf").write_bytes(b"changed")
        with self.assertRaises(ValueError):
            X.export_cached(s, w / "out.pdf", e)

    def test_new_doc_content_new_request(self):
        w, s, e, r = self.pending()
        X.accept(r, self.pdf(w / "in.pdf"), "fixture", "fixture")
        d = Document(s)
        d.add_paragraph("changed")
        d.save(s)
        with self.assertRaises(X.ManualExportRequired) as cm:
            X.export_cached(s, w / "out.pdf", e)
        self.assertNotEqual(cm.exception.request, r)

    def test_zip_timestamp_ignored(self):
        w, s, e, r = self.pending()
        new = w / "repacked.docx"
        with zipfile.ZipFile(s) as z, zipfile.ZipFile(new, "w") as dst:
            for n in z.namelist():
                info = zipfile.ZipInfo(n, (2000, 1, 1, 1, 1, 1))
                dst.writestr(info, z.read(n))
        self.assertNotEqual(X.sha(s), X.sha(new))
        self.assertEqual(X.package_digest(s), X.package_digest(new))

    def test_repacked_cache_reuse(self):
        w, s, e, r = self.pending()
        X.accept(r, self.pdf(w / "in.pdf"), "fixture", "fixture")
        d = Document(s)
        repacked = w / "resaved.docx"
        d.save(repacked)
        # python-docx roundtrip can change XML. Exact member-equivalent repacks only.
        with zipfile.ZipFile(s) as z, zipfile.ZipFile(repacked, "w") as dst:
            for n in z.namelist():
                dst.writestr(n, z.read(n))
        X.export_cached(repacked, w / "out.pdf", e)
        self.assertTrue((w / "out.pdf").is_file())

    def test_duplicate_zip_member_refused(self):
        import warnings

        p = self.root / "duplicate.docx"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with zipfile.ZipFile(p, "w") as z:
                z.writestr("x", "a")
                z.writestr("x", "b")
        with self.assertRaises(ValueError):
            X.package_digest(p)

    def test_manual_configuration(self):
        w, s, e = self.setup_exchange()
        D.configure_office(w, "wps", "manual")
        self.assertEqual(D.load(w)[1]["office_engine"], "wps")

    def test_manual_requires_wps_engine(self):
        w, s, e = self.setup_exchange()
        with self.assertRaises(ValueError):
            D.configure_office(w, "word", "manual")


if __name__ == "__main__":
    unittest.main()
