"""Behavior regressions for portable defaults and template customization."""

import base64
import hashlib
import json
import re
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from test_layout_integrity import Fixture, candidate
import engine as E
import intake as I


class PortableBehavior(Fixture):
    def test_sections_continue_without_inheriting_restart(self):
        doc = Document()
        E.select_profile("A")
        E.install_styles(doc)
        E.add_section(doc, start=1)
        E.add_section(doc)
        self.assertEqual(doc.sections[-2]._sectPr.find(qn("w:pgNumType")).get(qn("w:start")), "1")
        self.assertIsNone(doc.sections[-1]._sectPr.find(qn("w:pgNumType")).get(qn("w:start")))

    def test_abstracts_use_continuous_roman_numbers(self):
        out, _ = self.build(paper=True, cover=True)
        actual = [
            (
                s._sectPr.find(qn("w:pgNumType")).get(qn("w:fmt")),
                s._sectPr.find(qn("w:pgNumType")).get(qn("w:start")),
            )
            for s in Document(out).sections[1:]
        ]
        self.assertEqual(
            actual, [("decimal", "1"), ("upperRoman", "1"), ("upperRoman", None), ("decimal", "1")]
        )

    def test_replacing_logo_drops_orphan_but_preserves_body_use(self):
        for body_uses_old in (False, True):
            with self.subTest(body_uses_old=body_uses_old):
                state = self.state()
                I.customize(state, {"logo": "red"}, "synthetic red logo")
                old = E.ROOT / "assets/logos/green.png"
                blocks = [{"type": "paragraph", "text": "保留正文"}]
                if body_uses_old:
                    blocks.append({"type": "image", "path": str(old)})
                data = {"mode": "document", "cover": True, "blocks": blocks}
                src = self.path("json")
                src.write_text(json.dumps(data), encoding="utf-8")
                out = self.path("docx")
                E.build(src, out, intake_path=state)
                with ZipFile(out) as z:
                    media = [
                        hashlib.sha256(z.read(n)).hexdigest()
                        for n in z.namelist()
                        if n.startswith("word/media/")
                    ]
                self.assertEqual(
                    hashlib.sha256(old.read_bytes()).hexdigest() in media, body_uses_old
                )

    def test_long_course_and_fields_respect_explicit_small_font(self):
        state = self.state("B")
        I.customize(
            state,
            {"cover_fonts": {"course": {"size": 12}, "fields": {"size": 10.5}}},
            "synthetic smaller cover text",
        )
        course = "一二三四五六七八九十一二三"
        I.resolve(
            state,
            {
                "metadata": {
                    "course_name": candidate(course, corrects_previous=True),
                    "major": candidate("一二三四五六七八九十一", corrects_previous=True),
                }
            },
        )
        src = self.path("json")
        src.write_text(
            json.dumps(
                {
                    "mode": "document",
                    "cover": True,
                    "blocks": [{"type": "paragraph", "text": "测试"}],
                }
            ),
            encoding="utf-8",
        )
        out = self.path("docx")
        E.build(src, out, intake_path=state)
        self.assertIn(course, "\n".join(p.text for p in Document(out).paragraphs))

    def test_default_timezone_is_beijing_time(self):
        state = self.path("json")
        self.assertEqual(I.start(state, "synthetic")["timezone"], "Asia/Shanghai")

    def test_twenty_character_title_wraps_without_changing_words(self):
        import cover_fit
        from customization import apply_cover, effective_profile

        title = "城市公共交通票价结构对通勤方式选择的影响"

        def lines(spec):
            master = Document(E.ROOT / spec["master_file"])
            apply_cover(master, spec)
            return cover_fit.plan(master, spec, {"title_zh": title})["title"]["lines"]

        wrapped = lines(E.select_profile("A"))
        self.assertGreater(len(wrapped), 1)
        self.assertEqual("".join(wrapped), title)
        self.assertGreaterEqual(min(map(len, wrapped)), 8)
        smaller = effective_profile("A", {"cover_fonts": {"title": {"size": 10.5}}})
        self.assertEqual(lines(smaller), [title])

    def test_word_paths_are_data_even_with_smart_quotes(self):
        from word_com import script

        path = "C:/测试/含“引号”和‘撇号’ $(not-code).docx"
        code = script(path, "C:/测试/out.pdf")
        encoded = re.search(r"FromBase64String\('([^']+)'\)", code).group(1)
        self.assertEqual(json.loads(base64.b64decode(encoded))["source"], path)
        self.assertNotIn(path, code)

    def test_date_measurement_ignores_only_edge_whitespace(self):
        from cover_date_audit import visible_lines

        chars = [
            {"c": c, "bbox": (i * 10, 100, (i + 1) * 10, 112)} for i, c in enumerate("  A B  ")
        ]

        class Page:
            def get_text(self, mode):
                return {"blocks": [{"lines": [{"spans": [{"chars": chars}]}]}]}

        line = visible_lines(Page())[0]
        self.assertEqual(line["bbox"], (20, 100, 50, 112))
        self.assertEqual(line["spans"][0]["text"], "A B")

    def test_native_word_postsave_policy_changes_only_settings(self):
        from word_com import manual_open_updates

        out, _ = self.build(paper=True, cover=True)
        with ZipFile(out) as archive:
            before = {n: archive.read(n) for n in archive.namelist()}
        manual_open_updates(out)
        with ZipFile(out) as archive:
            after = {n: archive.read(n) for n in archive.namelist()}
        self.assertEqual(before.keys(), after.keys())
        for name in before:
            if name != "word/settings.xml":
                self.assertEqual(before[name], after[name], name)
        self.assertEqual(Document(out).settings.element.xpath("./w:updateFields/@w:val"), ["false"])

    def test_native_word_helper_rejects_input_as_output(self):
        from word_com import updated_copy

        out, _ = self.build()
        with patch("word_com.run") as run:
            with self.assertRaisesRegex(ValueError, "路径必须不同"):
                updated_copy(out, out, self.path("pdf"))
        run.assert_not_called()

    def test_toc_format_map_reads_effective_style_and_word_color(self):
        from word_com import toc_format_map

        doc = Document()
        toc1 = doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        toc1.font.name = "Times New Roman"
        toc1.font.size = Pt(14)
        fonts = toc1.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), "黑体")
        bold = OxmlElement("w:b")
        toc1.element.rPr.append(bold)  # Word treats an omitted w:val as on.
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "auto")
        toc1.element.rPr.append(color)
        first = doc.add_paragraph("一级目录项", style=toc1)
        visible = OxmlElement("w:webHidden")
        visible.set(qn("w:val"), "0")
        first.runs[0]._r.get_or_add_rPr().append(visible)

        toc2 = doc.styles.add_style("目录 2", WD_STYLE_TYPE.PARAGRAPH)
        toc2.font.size = Pt(12)
        fonts = toc2.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), "宋体")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "123456")
        toc2.element.rPr.append(color)
        second = doc.add_paragraph(style=toc2)
        hidden = second.add_run("隐藏项")
        web_hidden = OxmlElement("w:webHidden")
        web_hidden.set(qn("w:val"), "true")
        hidden._r.get_or_add_rPr().append(web_hidden)
        visible = second.add_run("二级目录项")
        web_hidden = OxmlElement("w:webHidden")
        web_hidden.set(qn("w:val"), "0")
        visible._r.get_or_add_rPr().append(web_hidden)
        path = self.path("docx")
        doc.save(path)

        formats = toc_format_map(path)
        self.assertEqual(formats["TOC1"]["format"]["b"], True)
        self.assertEqual(formats["TOC1"]["format"]["sz"], "28")
        self.assertEqual(formats["TOC1"]["format"]["font_eastAsia"], "黑体")
        self.assertEqual(formats["TOC1"]["format"]["word_color"], -16777216)
        self.assertEqual(formats["目录 2"]["format"]["font_eastAsia"], "宋体")
        self.assertEqual(formats["TOC2"], formats["目录 2"])
        self.assertEqual(formats["TOC2"]["format"]["word_color"], 0x563412)

    def test_toc_format_map_parses_true_and_on_run_bold_values(self):
        from word_com import toc_format_map

        for value in ("1", "true", "on"):
            with self.subTest(value=value):
                doc = Document()
                style = doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
                paragraph = doc.add_paragraph(style=style)
                run = paragraph.add_run("目录项")
                bold = OxmlElement("w:b")
                if value != "1":
                    bold.set(qn("w:val"), value)
                run._r.get_or_add_rPr().append(bold)
                path = self.path("docx")
                doc.save(path)
                self.assertIs(toc_format_map(path)["TOC1"]["format"]["b"], True)

    def test_toc_format_map_tolerates_missing_paragraph_style(self):
        from word_com import toc_format_map

        class UnstyledParagraph:
            style = None

            class XML:
                @staticmethod
                def xpath(_expression):
                    return []

            _p = XML()

        with patch(
            "docx.Document",
            return_value=type("DocumentStub", (), {"paragraphs": [UnstyledParagraph()]})(),
        ):
            self.assertEqual(toc_format_map(self.path("docx")), {})


if __name__ == "__main__":
    import unittest

    unittest.main()
