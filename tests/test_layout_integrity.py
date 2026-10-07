"""Independent layout regressions; fixtures are synthetic, not user consent."""

from __future__ import annotations
import hashlib, json, tempfile, unittest
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from docx import Document
from docx.shared import Cm, Pt, Twips
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import engine as E, intake as I
from check_format import check
from style_model import paragraph_properties, run_properties
from audit_helpers import toc_entries
from keywords import normalize, visible
from refresh_toc import logical_policy, update, verify_existing
from render_preview import font_report


def candidate(value, source="user", **kwargs):
    c = {"value": value, "source": source, "evidence": "自动化合成夹具；并非真实用户声明"}
    if source == "memory":
        c["target_confirmed"] = True
    c.update(kwargs)
    return c


class Fixture(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.serial = 0

    def tearDown(self):
        self.tmp.cleanup()

    def path(self, ext):
        self.serial += 1
        return self.d / f"case-{self.serial}.{ext}"

    def state(self, template="A"):
        p = self.path("intake.json")
        I.start(p, "自动化合成测试夹具")
        I.select(p, template, "测试模拟选择；不代表用户本次排版选择")
        fields = E.load_profile(template)["cover"]["required_fields"]
        data = {
            k: candidate("测试" + str(i))
            for i, k in enumerate(fields)
            if k not in ("date", "student_id")
        }
        data["student_id"] = candidate("00001234")
        I.resolve(p, {"metadata": data})
        return p

    def build(self, template="A", paper=False, cover=False, blocks=None, extras=None):
        state = self.state(template)
        data = {
            "mode": "paper" if paper else "document",
            "cover": cover,
            "metadata": {},
            "blocks": blocks
            or [
                {"type": "heading", "level": 1, "text": "一、独立审查"},
                {"type": "paragraph", "text": "本段仅供格式回归测试，不是研究结论。"},
            ],
        }
        if paper:
            data.update(
                abstract_zh="中文排版测试摘要。",
                keywords_zh=["格式", "验证"],
                references=["合成参考条目，仅作格式测试。"],
            )
            if template == "A":
                data.update(
                    abstract_en="Formatting test only.", keywords_en=["format", "quality assurance"]
                )
                data["metadata"]["title_en"] = "Formatting Test"
        if extras:
            data.update(extras)
        source = self.path("json")
        source.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        out = self.path("docx")
        E.build(source, out, intake_path=state)
        return out, state

    def mutate(self, fn, template="A", **kwargs):
        out, state = self.build(template, **kwargs)
        doc = Document(out)
        fn(doc)
        doc.save(out)
        return check(out, intake_path=state), out, state

    def restyle(self, doc, template="A", roles=None, **options):
        state = self.state(template)
        src = self.path("docx")
        doc.save(src)
        mapping = {
            "input_sha256": hashlib.sha256(src.read_bytes()).hexdigest(),
            "metadata": {},
            "mode": "document",
            "has_cover": False,
            "default_role": "body",
            "format_tables": True,
            "roles": roles or {},
        }
        mapping.update(options)
        mp = self.path("map.json")
        mp.write_text(json.dumps(mapping, ensure_ascii=False))
        out = self.path("docx")
        result = E.restyle(src, mp, out, intake_path=state)
        return Document(out), check(out, final=True, intake_path=state), result, out

    def assertPass(self, result):
        self.assertTrue(result["passed"], result["errors"])

    def assertCode(self, result, code):
        self.assertIn(code, [e["code"] for e in result["errors"]], result)


class IntakeRepairs(Fixture):
    def start(self):
        p = self.path("json")
        I.start(p, "合成测试")
        I.select(p, "B", "测试夹具选择红色版")
        return p

    def test_user_then_document_does_not_downgrade(self):
        p = self.start()
        I.resolve(p, {"metadata": {"student_name": candidate("甲")}})
        s = I.resolve(p, {"metadata": {"student_name": candidate("乙", "document")}})
        self.assertEqual(s["metadata"]["student_name"], "甲")

    def test_user_then_memory_does_not_downgrade(self):
        p = self.start()
        I.resolve(p, {"metadata": {"student_name": candidate("甲")}})
        s = I.resolve(p, {"metadata": {"student_name": candidate("乙", "memory")}})
        self.assertEqual(s["metadata"]["student_name"], "甲")

    def test_document_then_user_upgrades(self):
        p = self.start()
        I.resolve(p, {"metadata": {"student_name": candidate("甲", "document")}})
        s = I.resolve(p, {"metadata": {"student_name": candidate("乙")}})
        self.assertEqual(s["metadata"]["student_name"], "乙")

    def test_memory_then_document_upgrades(self):
        p = self.start()
        I.resolve(p, {"metadata": {"student_name": candidate("甲", "memory")}})
        s = I.resolve(p, {"metadata": {"student_name": candidate("乙", "document")}})
        self.assertEqual(s["metadata"]["student_name"], "乙")

    def test_same_rank_cross_round_conflict_is_not_silent(self):
        p = self.start()
        I.resolve(p, {"metadata": {"student_name": candidate("甲")}})
        s = I.resolve(p, {"metadata": {"student_name": candidate("乙")}})
        self.assertIn("student_name", s["conflicts"])

    def test_explicit_user_correction_archives_old_value(self):
        p = self.start()
        I.resolve(p, {"metadata": {"student_name": candidate("甲")}})
        s = I.resolve(p, {"metadata": {"student_name": candidate("乙", corrects_previous=True)}})
        self.assertEqual(s["metadata"]["student_name"], "乙")
        self.assertTrue(any(e["action"] == "user_value_corrected" for e in s["events"]))

    def test_document_cannot_claim_correction(self):
        p = self.start()
        before = p.read_bytes()
        with self.assertRaises(ValueError):
            I.resolve(
                p,
                {"metadata": {"student_name": candidate("甲", "document", corrects_previous=True)}},
            )
        self.assertEqual(before, p.read_bytes())

    def test_blanks_survive_later_document_value(self):
        p = self.start()
        I.resolve(p, {"blank_fields": ["student_name"], "blank_instruction": "测试明确留空"})
        s = I.resolve(p, {"metadata": {"student_name": candidate("甲", "document")}})
        self.assertIn("student_name", s["blank_fields"])

    def test_unblank_requires_user_instruction(self):
        p = self.start()
        I.resolve(p, {"blank_fields": ["student_name"], "blank_instruction": "测试留空"})
        with self.assertRaises(ValueError):
            I.resolve(
                p,
                {
                    "unblank_fields": ["student_name"],
                    "metadata": {"student_name": candidate("甲", "document")},
                },
            )

    def test_user_unblank_with_new_value(self):
        p = self.start()
        I.resolve(p, {"blank_fields": ["student_name"], "blank_instruction": "测试留空"})
        s = I.resolve(
            p, {"unblank_fields": ["student_name"], "metadata": {"student_name": candidate("甲")}}
        )
        self.assertNotIn("student_name", s["blank_fields"])

    def test_error_in_late_field_does_not_partially_save(self):
        p = self.start()
        before = p.read_bytes()
        with self.assertRaises(ValueError):
            I.resolve(
                p, {"metadata": {"student_name": candidate("甲"), "date": candidate("不是日期")}}
            )
        self.assertEqual(p.read_bytes(), before)

    def test_empty_candidate_list_cannot_clear_user_value(self):
        p = self.start()
        I.resolve(p, {"metadata": {"student_name": candidate("甲")}})
        with self.assertRaises(ValueError):
            I.resolve(p, {"metadata": {"student_name": []}})
        self.assertEqual(I.read(p)["metadata"]["student_name"], "甲")

    def test_unknown_state_schema_rejected_without_changes(self):
        p = self.start()
        I.resolve(p, {"metadata": {"student_name": candidate("甲")}})
        state = I.read(p)
        state["version"] = "99.0.0"
        state["integrity"] = I.digest({k: v for k, v in state.items() if k != "integrity"})
        p.write_text(json.dumps(state))
        before = p.read_bytes()
        with self.assertRaises(ValueError):
            I.resolve(p, {"metadata": {"student_name": candidate("乙", "document")}})
        self.assertEqual(before, p.read_bytes())

    def test_naive_clock_is_rejected(self):
        p = self.start()
        with self.assertRaises(ValueError):
            I.resolve(p, {}, now=datetime(2030, 1, 1))

    def test_strict_blank_field_type(self):
        p = self.start()
        with self.assertRaises(ValueError):
            I.resolve(p, {"blank_fields": "student_name", "blank_instruction": "test"})

    def test_duplicate_same_value_does_not_make_conflict(self):
        p = self.start()
        v = candidate("甲")
        I.resolve(p, {"metadata": {"student_name": v}})
        s = I.resolve(p, {"metadata": {"student_name": v}})
        self.assertNotIn("student_name", s["conflicts"])
        self.assertEqual(len(s["candidates"]["student_name"]), 1)


class ParagraphRepairs(Fixture):
    def test_h3_source_and_contract_green(self):
        E.select_profile("A")
        doc = Document()
        E.install_styles(doc)
        p = doc.add_paragraph("1.三级测试", "Heading 3")
        ind = paragraph_properties(p, "ind")
        space = paragraph_properties(p, "spacing")
        source = Document(ROOT / "sources/A/信息学院学年论文细则.docx").paragraphs[215]
        self.assertEqual(ind["left"], paragraph_properties(source, "ind")["left"])
        self.assertEqual(ind["leftChars"], "200")
        self.assertEqual(space["after"], "120")
        self.assertEqual(paragraph_properties(p, "jc")["val"], "both")

    def test_h3_red_has_point_indent_no_character_override(self):
        E.select_profile("B")
        doc = Document()
        E.install_styles(doc)
        p = doc.add_paragraph("1.三级", "Heading 3")
        ind = paragraph_properties(p, "ind")
        self.assertEqual(ind["left"], "420")
        self.assertNotIn("leftChars", ind)

    def test_reference_both_alignment(self):
        for template in ("A", "B"):
            out, _ = self.build(template, paper=True)
            p = next(p for p in Document(out).paragraphs if p.style.name == "XAUFE Reference")
            self.assertEqual(paragraph_properties(p, "jc")["val"], "both")

    def test_keyword_hanging_indents(self):
        out, _ = self.build(paper=True)
        doc = Document(out)
        for style, n in [("XAUFE Keywords Chinese", "1154"), ("XAUFE Keywords English", "1435")]:
            p = next(p for p in doc.paragraphs if p.style.name == style)
            ind = paragraph_properties(p, "ind")
            self.assertEqual(ind["left"], n)
            self.assertEqual(ind["hanging"], n)

    def test_extra_before_lines_rejected(self):
        r, _, _ = self.mutate(
            lambda d: E.element(d.paragraphs[1]._p.get_or_add_pPr(), "w:spacing", beforeLines=1000)
        )
        self.assertCode(r, "PARAGRAPH_EXTRA")

    def test_extra_zero_hanging_can_override_first_indent(self):
        r, _, _ = self.mutate(
            lambda d: E.element(d.paragraphs[1]._p.get_or_add_pPr(), "w:ind", hanging=0)
        )
        self.assertCode(r, "PARAGRAPH_EXTRA")

    def test_extra_zero_character_left_overrides_h3(self):
        r, _, _ = self.mutate(
            lambda d: E.element(d.paragraphs[0]._p.get_or_add_pPr(), "w:ind", leftChars=0),
            template="B",
            blocks=[{"type": "heading", "level": 3, "text": "1.测试"}],
        )
        self.assertCode(r, "PARAGRAPH_EXTRA")

    def test_modern_start_end_indent_resolves_equivalently(self):
        def change(d):
            E.element(d.paragraphs[0]._p.get_or_add_pPr(), "w:ind", start=420, end=0)

        r, _, _ = self.mutate(
            change,
            template="B",
            blocks=[{"type": "heading", "level": 3, "text": "1.测试"}],
        )
        self.assertPass(r)

    def test_modern_start_zero_cannot_hide_nonzero_indent(self):
        def change(d):
            E.element(d.paragraphs[0]._p.get_or_add_pPr(), "w:ind", start=0)

        r, _, _ = self.mutate(
            change,
            template="B",
            blocks=[{"type": "heading", "level": 3, "text": "1.测试"}],
        )
        self.assertCode(r, "PARAGRAPH")

    def test_contextual_spacing_rejected(self):
        r, _, _ = self.mutate(
            lambda d: E.element(d.paragraphs[1]._p.get_or_add_pPr(), "w:contextualSpacing")
        )
        self.assertCode(r, "PARAGRAPH_EXTRA")

    def test_title_orphan_rejected(self):
        r, _, _ = self.mutate(
            lambda d: setattr(d.paragraphs[0].paragraph_format, "keep_with_next", False)
        )
        self.assertCode(r, "PARAGRAPH")

    def test_reference_pagebreak_rejected(self):
        def change(d):
            next(
                p for p in d.paragraphs if p.style.name == "XAUFE References Title"
            ).paragraph_format.page_break_before = False

        r, _, _ = self.mutate(change, paper=True)
        self.assertCode(r, "PARAGRAPH")

    def test_green_border_missing_rejected(self):
        def change(d):
            for el in d.sections[0].header._element.xpath(".//w:pBdr"):
                el.getparent().remove(el)

        r, _, _ = self.mutate(change)
        self.assertCode(r, "HEADER_BORDER")

    def test_independent_checker_does_not_call_style_writer(self):
        out, state = self.build()
        with patch.object(E, "install_styles", side_effect=AssertionError("checker called writer")):
            self.assertPass(check(out, intake_path=state))

    def test_shared_wrong_style_is_caught_by_independent_contract(self):
        r, _, _ = self.mutate(lambda d: setattr(d.styles["XAUFE Body"].font, "size", Pt(30)))
        self.assertCode(r, "FONT_SIZE")

    def test_char_style_inheritance_bad_size_caught(self):
        def change(d):
            parent = d.styles.add_style("AuditBig", WD_STYLE_TYPE.CHARACTER)
            parent.font.size = Pt(30)
            child = d.styles.add_style("AuditChild", WD_STYLE_TYPE.CHARACTER)
            child.base_style = parent
            d.paragraphs[1].runs[0].style = child

        r, _, _ = self.mutate(change)
        self.assertCode(r, "FONT_SIZE")

    def test_direct_size_beats_bad_character_style(self):
        def change(d):
            s = d.styles.add_style("AuditBig", WD_STYLE_TYPE.CHARACTER)
            s.font.size = Pt(30)
            r = d.paragraphs[1].runs[0]
            r.style = s
            r.font.size = Pt(10.5)

        r, _, _ = self.mutate(change)
        self.assertPass(r)

    def test_hidden_char_style_rejected(self):
        def change(d):
            s = d.styles.add_style("HiddenAudit", WD_STYLE_TYPE.CHARACTER)
            s.font.hidden = True
            d.paragraphs[1].runs[0].style = s

        r, _, _ = self.mutate(change)
        self.assertCode(r, "HIDDEN_TEXT")

    def test_cyclic_style_is_explicit_failure(self):
        def change(d):
            s = d.styles.add_style("LoopAudit", WD_STYLE_TYPE.CHARACTER)
            s.base_style = s
            d.paragraphs[1].runs[0].style = s

        r, _, _ = self.mutate(change)
        self.assertCode(r, "STYLE_INHERITANCE")

    def test_style_toggle_and_direct_absolute_bold(self):
        doc = Document()
        a = doc.styles.add_style("A", WD_STYLE_TYPE.CHARACTER)
        a.font.bold = True
        b = doc.styles.add_style("B", WD_STYLE_TYPE.CHARACTER)
        b.base_style = a
        b.font.bold = True
        p = doc.add_paragraph("x")
        r = p.runs[0]
        r.style = b
        self.assertFalse(run_properties(p, r._r)["b"])
        r.bold = True
        self.assertTrue(run_properties(p, r._r)["b"])

    def test_input_emphasis_is_preserved(self):
        d = Document()
        p = d.add_paragraph()
        r = p.add_run("强调")
        r.bold = True
        r.italic = True
        out, result, _, _ = self.restyle(d)
        self.assertPass(result)
        v = run_properties(out.paragraphs[0], out.paragraphs[0].runs[0]._r)
        self.assertTrue(v["b"])
        self.assertTrue(v["i"])

    def test_hidden_input_stops_instead_of_revealing(self):
        d = Document()
        d.add_paragraph("隐字").runs[0].font.hidden = True
        with self.assertRaises(ValueError):
            self.restyle(d)

    def test_cover_char_style_hidden_rejected(self):
        def change(d):
            s = d.styles.add_style("HiddenCoverAudit", WD_STYLE_TYPE.CHARACTER)
            s.font.hidden = True
            d.paragraphs[9].runs[1].style = s

        r, _, _ = self.mutate(change, cover=True)
        self.assertCode(r, "COVER_EFFECTIVE_FONT")

    def test_unnumbered_heading_is_warning_not_content_rewrite(self):
        out, state = self.build(blocks=[{"type": "heading", "level": 1, "text": "系统设计"}])
        r = check(out, intake_path=state)
        self.assertPass(r)
        self.assertIn("HEADING_NUMBERING", [w["code"] for w in r["warnings"]])
        self.assertEqual(Document(out).paragraphs[0].text, "系统设计")

    def test_inline_content_control_does_not_silently_pass(self):
        d = Document()
        p = d.add_paragraph()
        sdt = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "内容控件"
        r.append(t)
        content.append(r)
        sdt.append(content)
        p._p.append(sdt)
        with self.assertRaises(ValueError):
            self.restyle(d)


class TableRepairs(Fixture):
    def table_doc(self):
        d = Document()
        t = d.add_table(rows=2, cols=2)
        t.autofit = False
        for col in t.columns:
            col.width = Cm(9)
        for i, row in enumerate(t.rows):
            for j, cell in enumerate(row.cells):
                cell.width = Cm(9)
                cell.text = f"单元格{i}{j}，完整文字。"
        return d, t

    def test_wide_existing_green_fixed(self):
        d, _ = self.table_doc()
        out, r, _, _ = self.restyle(d)
        self.assertPass(r)
        self.assertLessEqual(sum(c.width.twips for c in out.tables[0].columns), 8306)

    def test_table_grid_centered_not_cell_text_left_aligned(self):
        d, _ = self.table_doc()
        out, r, _, _ = self.restyle(d)
        self.assertPass(r)
        self.assertEqual(out.tables[0]._tbl.tblPr.find(qn("w:jc")).get(qn("w:val")), "center")

    def test_wide_existing_red_fixed(self):
        d, _ = self.table_doc()
        out, r, _, _ = self.restyle(d, "B")
        self.assertPass(r)
        self.assertLessEqual(sum(c.width.twips for c in out.tables[0].columns), 8958)

    def test_exact_5pt_row_changed_to_minimum(self):
        d, t = self.table_doc()
        for row in t.rows:
            row.height = Pt(5)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        out, r, _, _ = self.restyle(d)
        self.assertPass(r)
        for row in out.tables[0].rows:
            self.assertEqual(row.height_rule, WD_ROW_HEIGHT_RULE.AT_LEAST)

    def test_multilevel_nested_tables_formatted(self):
        d, t = self.table_doc()
        inner = t.cell(1, 0).add_table(rows=1, cols=1)
        p = inner.cell(0, 0).paragraphs[0]
        p.add_run("内层").font.size = Pt(30)
        deep = inner.cell(0, 0).add_table(rows=1, cols=1)
        deep.cell(0, 0).text = "更深层"
        out, r, _, _ = self.restyle(d)
        self.assertPass(r)
        self.assertEqual(
            out.tables[0].cell(1, 0).tables[0].cell(0, 0).paragraphs[0].runs[0].font.size.pt, 10.5
        )
        self.assertEqual(r["stats"]["tables_including_nested"], 3)

    def test_horizontal_and_vertical_merge_preserved(self):
        d = Document()
        t = d.add_table(rows=3, cols=3)
        for i, row in enumerate(t.rows):
            for j, cell in enumerate(row.cells):
                cell.text = f"格{i}{j}"
        t.cell(0, 0).merge(t.cell(0, 2))
        t.cell(1, 0).merge(t.cell(2, 0))
        before = E.content_signature(d)
        out, r, _, _ = self.restyle(d)
        self.assertPass(r)
        after = E.content_signature(out)
        self.assertEqual(before["cell_spans"], after["cell_spans"])
        self.assertEqual(before["text"], after["text"])

    def test_nowrap_and_fittext_cleared(self):
        d, t = self.table_doc()
        E.element(t.cell(0, 0)._tc.get_or_add_tcPr(), "w:noWrap")
        E.element(t.cell(0, 0)._tc.get_or_add_tcPr(), "w:tcFitText")
        out, r, _, _ = self.restyle(d)
        self.assertPass(r)
        self.assertEqual(
            out.tables[0].cell(0, 0)._tc.tcPr.find(qn("w:noWrap")).get(qn("w:val")), "0"
        )

    def test_inner_image_scaled_within_cell(self):
        d, t = self.table_doc()
        t.cell(0, 0).paragraphs[0].text = ""
        t.cell(0, 0).paragraphs[0].add_run().add_picture(
            str(ROOT / "examples/layout-test.png"), width=Cm(12)
        )
        out, r, _, _ = self.restyle(d)
        self.assertPass(r)
        ext = out.tables[0].cell(0, 0).paragraphs[0]._p.xpath(".//wp:extent")[0]
        self.assertLess(int(ext.get("cx")), int(Cm(8)))

    def test_float_table_rejected(self):
        d, t = self.table_doc()
        E.element(t._tbl.tblPr, "w:tblpPr", horzAnchor="page")
        with self.assertRaises(ValueError):
            self.restyle(d)

    def test_horizontal_merge_marker_rejected(self):
        d, t = self.table_doc()
        E.element(t.cell(0, 0)._tc.tcPr, "w:hMerge", val="restart")
        with self.assertRaises(ValueError):
            self.restyle(d)

    def test_impossible_narrow_column_rejected(self):
        d, t = self.table_doc()
        t.columns[0].width = Twips(100)
        with self.assertRaises(ValueError):
            self.restyle(d)

    def test_corrupt_grid_span_rejected(self):
        d, t = self.table_doc()
        E.element(t.cell(0, 0)._tc.tcPr, "w:gridSpan", val=5)
        with self.assertRaises(ValueError):
            self.restyle(d)

    def test_existing_header_repeat_marker_preserved(self):
        d, t = self.table_doc()
        E.element(t.rows[0]._tr.get_or_add_trPr(), "w:tblHeader")
        out, r, _, _ = self.restyle(d)
        self.assertPass(r)
        self.assertTrue(out.tables[0].rows[0]._tr.xpath("./w:trPr/w:tblHeader"))

    def test_checker_rejects_mutated_nested_font(self):
        d, t = self.table_doc()
        t.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套"
        out, _, _, path = self.restyle(d)
        out.tables[0].cell(0, 0).tables[0].cell(0, 0).paragraphs[0].runs[0].font.size = Pt(30)
        out.save(path)
        self.assertCode(check(path), "FONT_SIZE")

    def test_checker_rejects_reintroduced_exact_height(self):
        d, _ = self.table_doc()
        out, _, _, path = self.restyle(d)
        out.tables[0].rows[0].height = Pt(5)
        out.tables[0].rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        out.save(path)
        self.assertCode(check(path), "TABLE_FIXED_HEIGHT")

    def test_checker_rejects_mutated_cell_width(self):
        d, _ = self.table_doc()
        out, _, _, path = self.restyle(d)
        out.tables[0].cell(0, 0).width = Cm(20)
        out.save(path)
        self.assertCode(check(path), "CELL_WIDTH")

    def test_table_hidden_conditional_style_requires_confirmation(self):
        d, t = self.table_doc()
        s = d.styles.add_style("ConditionalHide", WD_STYLE_TYPE.TABLE)
        cond = OxmlElement("w:tblStylePr")
        cond.set(qn("w:type"), "firstRow")
        E.element(E.element(cond, "w:rPr"), "w:vanish")
        s.element.append(cond)
        t.style = s
        with self.assertRaises(ValueError):
            self.restyle(d)


class KeywordAndPreservation(Fixture):
    def test_chinese_2_4_spaces_become_3(self):
        d = Document()
        d.add_paragraph("关键词：格式  检查    模板")
        out, r, info, _ = self.restyle(d, roles={"p0": "keywords_zh"})
        self.assertPass(r)
        self.assertEqual(out.paragraphs[0].text, "关键词：格式   检查   模板")
        self.assertEqual(len(info["keyword_spacing_changes"]), 1)

    def test_english_phrase_single_spaces_preserved(self):
        d = Document()
        d.add_paragraph("Key words：machine learning  software engineering    quality assurance")
        out, r, _, _ = self.restyle(d, roles={"p0": "keywords_en"})
        self.assertPass(r)
        self.assertEqual(
            out.paragraphs[0].text,
            "Key words：machine learning   software engineering   quality assurance",
        )

    def test_explicit_keyword_items_disambiguates_single_separator(self):
        d = Document()
        d.add_paragraph("Key words：machine learning quality assurance")
        out, r, _, _ = self.restyle(
            d,
            roles={"p0": "keywords_en"},
            keyword_items={"p0": ["machine learning", "quality assurance"]},
        )
        self.assertPass(r)
        self.assertEqual(out.paragraphs[0].text, "Key words：machine learning   quality assurance")

    def test_explicit_items_cannot_rewrite_words(self):
        d = Document()
        d.add_paragraph("关键词：甲  乙")
        with self.assertRaises(ValueError):
            self.restyle(d, roles={"p0": "keywords_zh"}, keyword_items={"p0": ["甲", "丙"]})

    def test_spaces_split_across_runs(self):
        d = Document()
        p = d.add_paragraph()
        for text in ["关", "键词：甲 ", " ", " 乙  ", "  丙"]:
            p.add_run(text)
        out, r, _, _ = self.restyle(d, roles={"p0": "keywords_zh"})
        self.assertPass(r)
        self.assertEqual(out.paragraphs[0].text, "关键词：甲   乙   丙")

    def test_hyperlink_target_preserved_while_keywords_repaired(self):
        d = Document()
        p = d.add_paragraph("关键词：甲  ")
        link = OxmlElement("w:hyperlink")
        rid = d.part.relate_to("https://example.org/synthetic", RT.HYPERLINK, is_external=True)
        link.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "乙"
        r.append(t)
        link.append(r)
        p._p.append(link)
        p.add_run("    丙")
        out, result, _, _ = self.restyle(d, roles={"p0": "keywords_zh"})
        self.assertPass(result)
        self.assertEqual(visible(out.paragraphs[0]), "关键词：甲   乙   丙")
        self.assertEqual(out.part.rels[rid].target_ref, "https://example.org/synthetic")

    def test_keyword_tab_requires_explicit_preservation_workflow(self):
        d = Document()
        d.add_paragraph("关键词：甲\t乙")
        with self.assertRaises(ValueError):
            self.restyle(d, roles={"p0": "keywords_zh"})

    def test_empty_keywords_are_not_hard_content_gate(self):
        out, state = self.build(
            paper=True,
            extras={
                "abstract_zh": "",
                "keywords_zh": [],
                "references": [],
                "abstract_en": "",
                "keywords_en": [],
            },
        )
        r = check(out, intake_path=state)
        self.assertPass(r)
        self.assertIn("EMPTY_CONTENT", [x["code"] for x in r["warnings"]])

    def test_input_image_blob_hash_unchanged(self):
        d = Document()
        d.add_paragraph().add_run().add_picture(
            str(ROOT / "examples/layout-test.png"), width=Cm(20)
        )
        before = E.content_signature(d)
        out, r, _, _ = self.restyle(d, roles={"p0": "figure"})
        self.assertPass(r)
        self.assertEqual(before["related_payloads"], E.content_signature(out)["related_payloads"])

    def test_equation_payload_unchanged(self):
        d = Document()
        p = d.add_paragraph()
        math = OxmlElement("m:oMath")
        r = OxmlElement("m:r")
        t = OxmlElement("m:t")
        t.text = "x+y=z"
        r.append(t)
        math.append(r)
        p._p.append(math)
        before = E.content_signature(d)
        out, result, _, _ = self.restyle(d, roles={"p0": "equation"})
        self.assertPass(result)
        self.assertEqual(before["math"], E.content_signature(out)["math"])

    def test_real_hyperlink_and_numbering_preserved(self):
        d = Document()
        p = d.add_paragraph("列表文本", style="List Bullet")
        E.element(E.element(p._p.get_or_add_pPr(), "w:numPr"), "w:numId", val=1)
        before = E.content_signature(d)
        out, result, _, _ = self.restyle(d)
        self.assertPass(result)
        self.assertEqual(before["preserved_parts"], E.content_signature(out)["preserved_parts"])
        self.assertTrue(out.paragraphs[0]._p.xpath("./w:pPr/w:numPr"))


class TocRepairs(Fixture):
    def test_real_abstract_factory_title_is_legal(self):
        out, state = self.build(
            paper=True,
            blocks=[{"type": "heading", "level": 1, "text": "一、Abstract Factory模式设计"}],
        )
        self.assertPass(check(out, intake_path=state))
        self.assertEqual(len(toc_entries(Document(out))), 2)

    def test_toc_cannot_point_at_abstract_role(self):
        def change(d):
            target = toc_entries(d)[0]["target"]
            target.style = "XAUFE Chinese Abstract Label"

        r, _, _ = self.mutate(change, paper=True)
        self.assertCode(r, "TOC_SCOPE")

    def test_toc_wrong_anchor_rejected(self):
        def change(d):
            toc_entries(d)[0]["field"].set(qn("w:instr"), "PAGEREF MissingBookmark \\h")

        r, _, _ = self.mutate(change, paper=True)
        self.assertCode(r, "TOC_SCOPE")

    def test_toc_wrong_cached_title_rejected(self):
        def change(d):
            toc_entries(d)[0]["paragraph"]._p.xpath("./w:hyperlink//w:t")[0].text = "另一个标题"

        r, _, _ = self.mutate(change, paper=True)
        self.assertCode(r, "TOC_SCOPE")

    def test_toc_numeric_without_proof_is_not_final(self):
        out, state = self.build(paper=True)
        d = Document(out)
        for e in toc_entries(d):
            e["field"].findall(".//" + qn("w:t"))[0].text = "1"
        d.save(out)
        self.assertCode(check(out, final=True, intake_path=state), "TOC_PROOF_REQUIRED")

    def test_toc_must_start_at_explicit_body_section(self):
        out, _ = self.build(paper=True)
        d = Document(out)
        entries = toc_entries(d)
        for s in d.sections:
            pg = s._sectPr.find(qn("w:pgNumType"))
            if pg is not None:
                pg.attrib.pop(qn("w:start"), None)
        with self.assertRaises(ValueError):
            logical_policy(d, entries)

    def test_cache_change_with_same_filename_invalidates_proof(self):
        out, state = self.build(paper=True)
        d = Document(out)
        entries = toc_entries(d)
        for e in entries:
            e["field"].findall(".//" + qn("w:t"))[0].text = "1"
        d.save(out)
        vv = E.get_vars(d)
        report = {
            "document_sha256": hashlib.sha256(out.read_bytes()).hexdigest(),
            "converged": True,
            "template": vv["XAUFE_TEMPLATE"],
            "job_id": vv["XAUFE_INTAKE_ID"],
            "entries": [
                {"anchor": e["anchor"], "text": "".join(e["text"].split()), "logical_page": 1}
                for e in entries
            ],
            "rendered_pdf_sha256": "synthetic-proof-unit-test",
            "font_resolution": font_report(),
        }
        proof = out.with_suffix(".toc-report.json")
        proof.write_text(json.dumps(report))
        self.assertPass(check(out, final=True, intake_path=state))
        entries[0]["field"].findall(".//" + qn("w:t"))[0].text = "999"
        d.save(out)
        self.assertCode(check(out, final=True, intake_path=state), "TOC_PROOF_STALE")

    def test_empty_structure_is_allowed_with_verified_toc(self):
        out, state = self.build(
            paper=True,
            extras={
                "abstract_zh": "",
                "abstract_en": "",
                "keywords_zh": [],
                "keywords_en": [],
                "references": [],
            },
        )
        target = self.path("docx")
        with patch(
            "refresh_toc.rendered_positions",
            return_value=(
                [1, 2],
                {"physical_pages": [5, 6], "logical_pages": [1, 2], "total_pages": 6},
                "synthetic-pdf-hash",
            ),
        ):
            report = update(out, target, intake_path=state)
        self.assertEqual(len(report["passes"]), 2)
        r = check(target, final=True, intake_path=state)
        self.assertPass(r)
        self.assertIn("EMPTY_CONTENT", [w["code"] for w in r["warnings"]])

    def test_font_environment_change_requires_new_toc_render(self):
        out, state = self.build(paper=True)
        target = self.path("docx")
        with patch(
            "refresh_toc.rendered_positions",
            return_value=(
                [1, 2],
                {"physical_pages": [5, 6], "logical_pages": [1, 2], "total_pages": 6},
                "synthetic",
            ),
        ):
            update(out, target, intake_path=state)
        with patch("render_preview.font_report", return_value={"宋体": "Different substitution"}):
            self.assertCode(check(target, final=True, intake_path=state), "TOC_RENDER_ENV_CHANGED")

    def test_verify_only_rejects_wrong_cache_without_resaving(self):
        out, state = self.build(paper=True)
        before = out.read_bytes()
        with patch(
            "refresh_toc.rendered_positions",
            return_value=(
                [1, 2],
                {"physical_pages": [5, 6], "logical_pages": [1, 2], "total_pages": 6},
                "synthetic",
            ),
        ):
            with self.assertRaises(ValueError):
                verify_existing(out, state)
        self.assertEqual(out.read_bytes(), before)
        self.assertFalse(out.with_suffix(".toc-report.json").exists())


if __name__ == "__main__":
    unittest.main()
