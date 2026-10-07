"""Date position is a user rule, not a by-product of fields above it. Synthetic tests."""

from test_layout_integrity import Fixture, candidate
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from check_format import check
import engine as E, intake as I
from cover_date import apply
from zipfile import ZipFile
import json


class CoverDate(Fixture):
    def with_cover(self, template="A"):
        return self.build(template=template, cover=True)

    def mutate_date(self, change, template="A"):
        path, state = self.with_cover(template)
        d = Document(path)
        change(d.paragraphs[14])
        d.save(path)
        return check(path, intake_path=state)

    def test_green_build_bottom_date(self):
        path, state = self.with_cover()
        self.assertPass(check(path, intake_path=state))
        self.assertEqual(
            Document(path).paragraphs[14]._p.pPr.find(qn("w:framePr")).get(qn("w:yAlign")), "bottom"
        )

    def test_red_build_bottom_date(self):
        path, state = self.with_cover("B")
        self.assertPass(check(path, intake_path=state))
        self.assertEqual(
            Document(path).paragraphs[14]._p.pPr.find(qn("w:framePr")).get(qn("w:w")), "8958"
        )

    def test_both_masters_have_bottom_date(self):
        for t in ("A", "B"):
            spec = E.select_profile(t)
            E.asset_check()
            d = Document(E.ROOT / spec["master_file"])
            self.assertEqual(
                d.paragraphs[14]._p.pPr.find(qn("w:framePr")).get(qn("w:yAlign")), "bottom"
            )

    def test_removed_position_rejected(self):
        self.assertCode(
            self.mutate_date(lambda p: p._p.pPr.remove(p._p.pPr.find(qn("w:framePr")))),
            "COVER_DATE_POSITION",
        )

    def test_top_alignment_rejected(self):
        self.assertCode(
            self.mutate_date(lambda p: p._p.pPr.find(qn("w:framePr")).set(qn("w:yAlign"), "top")),
            "COVER_DATE_POSITION",
        )

    def test_paper_edge_anchor_rejected(self):
        self.assertCode(
            self.mutate_date(lambda p: p._p.pPr.find(qn("w:framePr")).set(qn("w:vAnchor"), "page")),
            "COVER_DATE_POSITION",
        )

    def test_absolute_y_override_rejected(self):
        self.assertCode(
            self.mutate_date(lambda p: p._p.pPr.find(qn("w:framePr")).set(qn("w:y"), "999")),
            "COVER_DATE_POSITION",
        )

    def test_wrong_frame_width_rejected(self):
        self.assertCode(
            self.mutate_date(
                lambda p: p._p.pPr.find(qn("w:framePr")).set(qn("w:w"), "10000"), "B"
            ),
            "COVER_DATE_POSITION",
        )

    def test_missing_anchor_lock_rejected(self):
        self.assertCode(
            self.mutate_date(
                lambda p: p._p.pPr.find(qn("w:framePr")).attrib.pop(qn("w:anchorLock"))
            ),
            "COVER_DATE_POSITION",
        )

    def test_arbitrary_spacing_rejected(self):
        self.assertCode(
            self.mutate_date(lambda p: p._p.pPr.find(qn("w:spacing")).set(qn("w:before"), "6000")),
            "COVER_DATE_LAYOUT",
        )

    def test_indent_rejected(self):
        self.assertCode(
            self.mutate_date(lambda p: p._p.pPr.find(qn("w:ind")).set(qn("w:left"), "100")),
            "COVER_DATE_LAYOUT",
        )

    def test_pagebreak_rejected(self):
        self.assertCode(
            self.mutate_date(lambda p: E.element(p._p.pPr, "w:pageBreakBefore", val="1")),
            "COVER_DATE_LAYOUT",
        )

    def test_keepnext_rejected(self):
        self.assertCode(
            self.mutate_date(lambda p: p._p.pPr.find(qn("w:keepNext")).set(qn("w:val"), "1")),
            "COVER_DATE_LAYOUT",
        )

    def test_no_cover_no_date_frame(self):
        path, state = self.build()
        self.assertFalse(Document(path).element.body.xpath(".//w:framePr"))
        self.assertPass(check(path, intake_path=state))

    def test_position_idempotent(self):
        path, _ = self.with_cover()
        d = Document(path)
        before = d.element.xml
        apply(d, E.SPEC)
        apply(d, E.SPEC)
        self.assertEqual(before, d.element.xml)

    def test_text_and_font_preserved(self):
        path, _ = self.with_cover()
        d = Document(path)
        p = d.paragraphs[14]
        before = [r._r.xml for r in p.runs]
        apply(d, E.SPEC)
        self.assertEqual(before, [r._r.xml for r in p.runs])

    def test_explicit_blank_still_bottom(self):
        path, state = self.with_cover()
        I.resolve(state, {"blank_fields": ["date"], "blank_instruction": "合成测试明确留空日期"})
        src = self.path("json")
        src.write_text(
            json.dumps(
                {
                    "cover": True,
                    "mode": "document",
                    "blocks": [{"type": "paragraph", "text": "合成测试。"}],
                }
            )
        )
        out = self.path("docx")
        E.build(src, out, intake_path=state)
        self.assertPass(check(out, intake_path=state))
        self.assertNotRegex(Document(out).paragraphs[14].text, r"\d")

    def test_floating_body_not_implicitly_supported(self):
        path, state = self.with_cover()
        d = Document(path)
        d.paragraphs[-1]._p.get_or_add_pPr().append(
            deepcopy(d.paragraphs[14]._p.pPr.find(qn("w:framePr")))
        )
        d.save(path)
        self.assertCode(check(path, intake_path=state), "UNSUPPORTED_STRUCTURE")

    def test_restyle_repairs_unanchored_cover(self):
        path, state = self.with_cover()
        d = Document(path)
        p = d.paragraphs[14]
        p._p.pPr.remove(p._p.pPr.find(qn("w:framePr")))
        d.save(path)
        inv = E.inventory(path)
        rev = {v: k for k, v in E.ROLES.items()}
        roles = {
            x["id"]: ("cover" if int(x["id"][1:]) < 15 else rev[x["style"]])
            for x in inv["blocks"]
            if x["kind"] == "paragraph"
        }
        m = self.path("json")
        m.write_text(
            json.dumps({"input_sha256": inv["input_sha256"], "has_cover": True, "roles": roles})
        )
        out = self.path("docx")
        E.restyle(path, m, out, intake_path=state)
        self.assertPass(check(out, intake_path=state))

    def test_explicit_false_pagebreak_rejected_for_renderer_compatibility(self):
        self.assertCode(
            self.mutate_date(lambda p: E.element(p._p.pPr, "w:pageBreakBefore", val="0")),
            "COVER_DATE_LAYOUT",
        )
