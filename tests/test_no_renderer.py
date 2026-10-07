"""Software-availability regression fixtures. Absence and agent receipts below are synthetic.

Actual native Office rendering is tested separately by run_acceptance.py.
No mocked PNG, claim of an agent, or test attestation is real visual evidence.
"""

from __future__ import annotations
import json
import os
import sys
from pathlib import Path
from unittest.mock import patch
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from test_delivery import Base
import delivery as D, render_availability as F
import office_backend as B, intake as I, engine as E
from check_format import check

ABSENT = {
    "runtime": {"system": "synthetic-test-runtime"},
    "candidates": [],
    "checks": [{"source": "SYNTHETIC full discovery", "status": "no-visible-apps"}],
}
INSTALLED = {
    "runtime": {"system": "synthetic-test-runtime"},
    "candidates": [
        {
            "engine": "libreoffice",
            "path": "/synthetic/soffice",
            "found": True,
            "automation": True,
            "source": "fixture",
            "probe": {"usable": False},
        }
    ],
    "checks": [{"source": "fixture", "status": "found-but-failed"}],
}
GUI_ONLY = {
    "runtime": {"system": "synthetic-test-runtime"},
    "candidates": [
        {
            "engine": "wps",
            "path": "/synthetic/wps",
            "found": True,
            "automation": False,
            "source": "fixture",
        }
    ],
    "checks": [{"source": "fixture", "status": "found-no-automation"}],
}


class NoRendererDelivery(Base):
    def prepare_absent(
        self,
        template="B",
        vision="available",
        request=None,
        paper=True,
        cover=True,
        note=None,
    ):
        source, state = self.build(template, request, paper, cover)
        job = D.create(
            self.d / "output" / "稿件.docx",
            parent=self.d,
        )
        work = Path(job["workspace"])
        with (
            patch("office_discovery.discover", return_value=deepcopy(ABSENT)),
            patch("office_backend.discover", return_value=deepcopy(ABSENT)),
        ):
            result = D.prepare(work, source, state, note)
        return work, source, state, result

    def acknowledgement(self, w, r, **changes):
        # Synthetic gate record; this is NOT a claim that anything was verified.
        record = {
            "vision": False,
            "pages": {},
            "content_check": "Synthetic fixture: static content check only",
            "layout_check": "Synthetic fixture: static layout check only",
            "delivery_check": "Synthetic fixture: single Word delivery",
            "warnings": {
                x["code"]: "Synthetic fixture acknowledges this is NOT verified."
                for x in r["warnings"]
            },
        }
        record.update(changes)
        return record

    def review_all(self, w, r):
        D.attest(w, self.acknowledgement(w, r))

    def test_red_without_office_prepares(self):
        w, _, s, r = self.prepare_absent()
        self.assertEqual(r["validation_scope"], "static-only")
        self.assertEqual(r["pages"], [])
        self.assertTrue(check(w / "final.docx", True, s, validation_scope="static-only")["passed"])

    def test_green_without_office_prepares(self):
        w, _, s, r = self.prepare_absent("A")
        self.assertTrue(check(w / "final.docx", True, s, validation_scope="static-only")["passed"])

    def test_static_notice_lists_skipped_checks(self):
        w, _, s, r = self.prepare_absent()
        self.assertIn("本次未完成", r["user_notice"])
        self.assertEqual(r["pages"], [])

    def test_no_fake_preview_or_toc_proof(self):
        w, _, s, r = self.prepare_absent()
        self.assertFalse((w / "preview").exists())
        self.assertFalse((w / "final.toc-report.json").exists())

    def test_pending_toc_caches(self):
        w, _, s, r = self.prepare_absent()
        self.assertTrue(
            all(x["cache"] == F.PENDING for x in D.toc_entries(Document(w / "final.docx")))
        )

    def test_live_update_request(self):
        w, _, s, r = self.prepare_absent()
        doc = Document(w / "final.docx")
        self.assertEqual(doc.settings.element.xpath("./w:updateFields/@w:val"), ["false"])
        self.assertTrue(all(f.get(qn("w:dirty")) not in ("true", "1") for f in F.page_fields(doc)))

    def test_optional_footer_page_numbers_pending(self):
        w, _, s, r = self.prepare_absent(request={"footer_page_numbers": True})
        fields = list(F.page_fields(Document(w / "final.docx")))
        self.assertTrue(any(f.get(qn("w:instr")).strip() == "PAGE" for f in fields))
        self.assertTrue(
            all(
                "".join(x.text or "" for x in f.findall(".//" + qn("w:t"))) == F.PENDING
                for f in fields
            )
        )

    def test_document_without_toc_not_reported_as_toc_skip(self):
        w, _, s, r = self.prepare_absent(paper=False, cover=False)
        self.assertNotIn("TOC_PAGES_NOT_VERIFIED", [a["code"] for a in r["skipped_steps"]])
        self.assertNotIn("待更新", r["user_notice"])

    def test_fonts_and_user_overrides_remain(self):
        w, _, s, r = self.prepare_absent(
            request={
                "styles": {"body": {"cn": "仿宋", "size": 12}, "toc2": {"size": 12}},
                "logo": "green",
            }
        )
        doc = Document(w / "final.docx")
        self.assertEqual(doc.styles["XAUFE Body"].font.size.pt, 12)
        self.assertEqual(doc.styles["TOC 2"].font.size.pt, 12)

    def test_preserves_content_and_objects(self):
        w, source, s, r = self.prepare_absent()
        self.assertEqual(
            F.page_neutral_signature(Document(source)),
            F.page_neutral_signature(Document(w / "final.docx")),
        )

    def test_same_file_cannot_claim_full_audit(self):
        w, _, s, r = self.prepare_absent()
        self.assertFalse(check(w / "final.docx", True, s)["passed"])

    def test_static_wrong_body_font_still_fails(self):
        w, _, s, r = self.prepare_absent()
        doc = Document(w / "final.docx")
        next(p for p in doc.paragraphs if p.style.name == "XAUFE Body").runs[0].font.size = Pt(30)
        doc.save(w / "bad.docx")
        self.assertFalse(check(w / "bad.docx", True, s, validation_scope="static-only")["passed"])

    def test_static_wrong_toc_font_still_fails(self):
        w, _, s, r = self.prepare_absent()
        doc = Document(w / "final.docx")
        p = next(p for p in doc.paragraphs if p.style.name == "TOC 2")
        p._p.xpath(".//w:rPr/w:sz")[0].set(qn("w:val"), "18")
        doc.save(w / "bad.docx")
        self.assertFalse(check(w / "bad.docx", True, s, validation_scope="static-only")["passed"])

    def test_replaced_pending_numeric_cache_rejected(self):
        w, _, s, r = self.prepare_absent()
        doc = Document(w / "final.docx")
        next(F.page_fields(doc)).find(".//" + qn("w:t")).text = "999"
        doc.save(w / "bad.docx")
        self.assertIn(
            "UNVERIFIED_PAGE_CACHE",
            {
                x["code"]
                for x in check(w / "bad.docx", True, s, validation_scope="static-only")["errors"]
            },
        )

    def test_missing_pending_marker_rejected(self):
        w, _, s, r = self.prepare_absent()
        doc = Document(w / "final.docx")
        F.clear_scope(doc)
        doc.save(w / "bad.docx")
        self.assertFalse(check(w / "bad.docx", True, s, validation_scope="static-only")["passed"])

    def test_unknown_scope_cannot_disable_checks(self):
        w, _, s, r = self.prepare_absent()
        self.assertFalse(check(w / "final.docx", True, s, validation_scope="skip-all")["passed"])

    def test_cannot_publish_before_acknowledgement(self):
        w, _, s, r = self.prepare_absent()
        with self.assertRaises(ValueError):
            D.publish(w)

    def test_no_renderer_rejects_fake_seen_pages(self):
        w, _, s, r = self.prepare_absent()
        with self.assertRaises(ValueError):
            D.attest(w, self.acknowledgement(w, r, pages={"1": "声称看过页面的合成记录"}))

    def test_warnings_must_be_addressed(self):
        w, _, s, r = self.prepare_absent()
        # Static deliveries fold page warnings into the notice; any other
        # outstanding warning (e.g. missing source text) still needs a note.
        root, state = D.load(w)
        state["warning_codes"] = ["SOURCE_DIFFERENCES"]
        D.write_json(root / D.MARKER, state)
        with self.assertRaisesRegex(ValueError, "SOURCE_DIFFERENCES"):
            D.attest(w, self.acknowledgement(w, r, warnings={}))

    def test_publication_returns_mandatory_notice(self):
        w, source, s, r = self.prepare_absent()
        self.review_all(w, r)
        out = D.publish(w)
        self.assertEqual(out["validation_scope"], "static-only")
        self.assertEqual(out["user_notice"], r["user_notice"])
        self.assertTrue(out["skipped_steps"])

    def test_cleanup_single_file_preserves_source(self):
        w, source, s, r = self.prepare_absent()
        old = D.sha(source)
        self.review_all(w, r)
        out = D.publish(w)
        self.assertFalse(w.exists())
        self.assertEqual(D.sha(source), old)
        self.assertEqual(len(list((self.d / "output").iterdir())), 1)

    def test_changed_word_rejected(self):
        w, _, s, r = self.prepare_absent()
        self.review_all(w, r)
        doc = Document(w / "final.docx")
        doc.core_properties.subject = "changed"
        doc.save(w / "final.docx")
        with self.assertRaises(ValueError):
            D.publish(w)

    def test_changed_intake_rejected(self):
        w, _, s, r = self.prepare_absent()
        self.review_all(w, r)
        I.customize(s, {"styles": {"body": {"size": 12}}}, "Synthetic later requirement")
        with self.assertRaises(ValueError):
            D.publish(w)

    def test_changed_probe_evidence_rejected(self):
        w, _, s, r = self.prepare_absent()
        self.review_all(w, r)
        D.write_json(w / "office-probe.json", {})
        with self.assertRaises(ValueError):
            D.publish(w)

    def test_changed_skip_receipt_rejected(self):
        w, _, s, r = self.prepare_absent()
        self.review_all(w, r)
        D.write_json(w / "limited-validation.json", {})
        with self.assertRaises(ValueError):
            D.publish(w)

    def test_review_on_previous_prepare_cannot_be_reused(self):
        w, source, s, r = self.prepare_absent()
        self.review_all(w, r)
        with (
            patch("office_discovery.discover", return_value=ABSENT),
            patch("office_backend.discover", return_value=ABSENT),
        ):
            D.prepare(w, source, s)
        # The earlier review belonged to the previous output; it must not carry over.
        with self.assertRaises(ValueError):
            D.publish(w)

    def test_skill_name(self):
        self.assertIn("name: xaufe-document-format", (ROOT / "SKILL.md").read_text())

    def test_workspace_uses_correct_name(self):
        w, _, s, r = self.prepare_absent()
        self.assertTrue(w.name.startswith("xaufe-job-"))
        self.assertEqual(D.MARKER, ".xaufe-session.json")


class EnvironmentDecision(Base):
    def context(self):
        source, state = self.build(paper=False, cover=False)
        job = D.create(
            self.d / "output/final.docx",
            parent=self.d,
        )
        return Path(job["workspace"]), source, state

    def test_installed_gui_only_requires_real_alternative_note(self):
        w, src, state = self.context()
        with (
            patch("office_discovery.discover", return_value=GUI_ONLY),
            patch("office_backend.discover", return_value=GUI_ONLY),
        ):
            with self.assertRaisesRegex(ValueError, "RENDER_ALTERNATIVE_REQUIRED"):
                D.prepare(w, src, state)

    def test_installed_gui_only_note_allows_limited_delivery(self):
        w, src, state = self.context()
        with (
            patch("office_discovery.discover", return_value=GUI_ONLY),
            patch("office_backend.discover", return_value=GUI_ONLY),
        ):
            r = D.prepare(w, src, state, "SYNTHETIC: no GUI tool and no PDF available")
        self.assertIn("已找到办公软件", r["user_notice"])
        self.assertEqual(r["validation_scope"], "static-only")

    def test_document_export_failure_with_working_smoke_not_skipped(self):
        w, src, state = self.context()
        failure = B.OfficeExportFailed(
            "synthetic source failure", INSTALLED, [{"error": "synthetic"}]
        )
        with (
            patch("office_discovery.discover", return_value=INSTALLED),
            patch("delivery.render", side_effect=failure),
            patch("office_backend.export_pdf", return_value={}),
        ):
            with self.assertRaisesRegex(ValueError, "当前稿件导出失败"):
                D.prepare(w, src, state, "SYNTHETIC no GUI")

    def test_all_exports_and_smoke_failed_note_allows_skip(self):
        w, src, state = self.context()
        failure = B.OfficeExportFailed(
            "synthetic environment failure", INSTALLED, [{"error": "synthetic"}]
        )
        with (
            patch("office_discovery.discover", return_value=INSTALLED),
            patch("delivery.render", side_effect=failure),
            patch("office_backend.export_pdf", side_effect=failure),
        ):
            r = D.prepare(w, src, state, "SYNTHETIC: permissions failed, no GUI available")
        self.assertEqual(r["validation_scope"], "static-only")
        self.assertEqual(
            D.read_json(w / "office-unavailable.json")["smoke_test"]["status"], "failed"
        )

    def test_layout_error_never_converted_to_software_absence(self):
        w, src, state = self.context()
        with patch("delivery.render", side_effect=ValueError("synthetic clipped cover")):
            with self.assertRaisesRegex(ValueError, "clipped cover"):
                D.prepare(w, src, state, "do not force skip")

    def test_missing_dependency_not_downgraded(self):
        w, src, state = self.context()
        with patch(
            "delivery.render", side_effect=ImportError("synthetic missing renderer dependency")
        ):
            with self.assertRaises(ImportError):
                D.prepare(w, src, state)

    def test_ordinary_runtime_error_not_downgraded(self):
        w, src, state = self.context()
        with patch("delivery.render", side_effect=RuntimeError("synthetic page mismatch")):
            with self.assertRaises(RuntimeError):
                D.prepare(w, src, state)

    def test_initial_static_defect_not_skipped(self):
        w, src, state = self.context()
        doc = Document(src)
        next(p for p in doc.paragraphs if p.style.name == "XAUFE Body").runs[0].font.size = Pt(30)
        doc.save(src)
        with (
            patch("office_discovery.discover", return_value=ABSENT),
            patch("office_backend.discover", return_value=ABSENT),
        ):
            with self.assertRaisesRegex(ValueError, "格式初查未通过"):
                D.prepare(w, src, state)


class NamespaceAndCandidates(Base):
    def test_new_environment_prefix_works(self):
        import office_discovery as O

        exe = self.d / "soffice"
        exe.write_text("synthetic")
        with patch.dict(os.environ, {"XAUFE_LIBREOFFICE": str(exe)}):
            report = O.discover()
            self.assertEqual(O.usable_candidates(report)[0]["path"], str(exe))

    def test_native_environment_does_not_mutate_process(self):
        import office_discovery as O

        exe = self.d / "renderer"
        exe.write_text("synthetic")
        with patch.dict(os.environ, {"XAUFE_LIBREOFFICE": str(exe)}):
            before = dict(os.environ)
            O.discover()
            self.assertEqual(dict(os.environ), before)

    def test_invalid_engine_preserves_environment(self):
        import office_discovery as O

        with patch.dict(os.environ, {"XAUFE_RENDER_ENGINE": "not-a-renderer"}):
            before = dict(os.environ)
            with self.assertRaises(ValueError):
                O.usable_candidates({"runtime": {}, "candidates": []})
            self.assertEqual(dict(os.environ), before)

    def test_nonexistent_selected_engine_uses_other_renderer(self):
        import pymupdf as fitz

        candidate = {
            "engine": "libreoffice",
            "path": "/synthetic/soffice",
            "found": True,
            "automation": True,
            "source": "fixture",
        }
        report = {"runtime": {"system": "Linux"}, "candidates": [candidate], "checks": []}

        def export(src, pdf, c, tmp):
            doc = fitz.open()
            doc.new_page()
            doc.save(pdf)
            doc.close()

        with (
            patch.dict(os.environ, {"XAUFE_RENDER_ENGINE": "wps"}),
            patch("office_backend.discover", return_value=report),
            patch("office_backend._lo", side_effect=export),
        ):
            result = B.export_pdf(self.d / "source.docx", self.d / "result.pdf")
        self.assertEqual(result["engine"], "libreoffice")

    def test_unavailable_note_does_not_force_skip_available_render(self):
        from test_delivery import DeliveryGuards

        source, state = self.build(paper=False, cover=False)
        job = D.create(
            self.d / "output/final.docx",
            parent=self.d,
        )
        with (
            patch(
                "delivery.render",
                side_effect=lambda doc, out: DeliveryGuards.fake_render(self, doc, out),
            ),
            patch("office_discovery.discover", return_value=INSTALLED),
        ):
            result = D.prepare(job["workspace"], source, state, "This text cannot force a skip.")
        self.assertEqual(result["validation_scope"], "full")


if __name__ == "__main__":
    import unittest

    unittest.main()


class UserRequestedStatic(Base):
    def test_office_none_skips_discovery_and_says_why(self):
        source, state = self.build()
        job = D.create(self.d / "output/final.docx", parent=self.d)
        with patch("office_discovery.discover") as discover, patch("delivery.render") as render:
            r = D.prepare(job["workspace"], source, state, office="none")
        discover.assert_not_called()
        render.assert_not_called()
        self.assertEqual(r["validation_scope"], "static-only")
        self.assertIn("按用户要求", r["user_notice"])
        D.attest(
            job["workspace"],
            {
                "vision": True,
                "pages": {},
                "content_check": "Synthetic fixture: static content check only",
                "layout_check": "Synthetic fixture: static layout check only",
                "delivery_check": "Synthetic fixture: single Word delivery",
                "warnings": {x["code"]: "Synthetic fixture note." for x in r["warnings"]},
            },
        )
        out = D.publish(job["workspace"])
        self.assertTrue(Path(out["output"]).is_file())
        self.assertIn("按用户要求", out["user_notice"])
