"""Rendered-page audit on synthetic PDFs (no office suite needed) and, when
LibreOffice is installed, on a really rendered defective document."""

from pathlib import Path
import json, shutil, sys, tempfile, unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import pymupdf as fitz
import engine as E
import intake as I
import feedback as F
from render_audit import audit

A4 = (595.3, 841.9)


def value(text):
    return {"value": text, "source": "user", "evidence": "SYNTHETIC test statement"}


class RenderAudit(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def docx(self, template="B", patch=None, cover=True, blocks=None):
        folder = Path(tempfile.mkdtemp(dir=self.folder))
        state = folder / "intake.json"
        I.start(state, "SYNTHETIC render audit")
        I.select(state, template, "SYNTHETIC choice")
        if patch:
            I.customize(state, patch, "SYNTHETIC customization")
        if not cover:
            I.document_scope(state, False, "SYNTHETIC no cover")
        fields = I.profile_for(I.read(state))["cover"]["required_fields"] if cover else []
        meta = {k: value("2030-01-02" if k == "date" else "测试") for k in fields}
        if not cover and template == "A":
            meta = {"title_zh": value("测试"), "major": value("测试")}
        I.resolve(state, {"metadata": meta})
        content = folder / "content.json"
        content.write_text(json.dumps({"mode": "document", "metadata": {}, "blocks": blocks or [
            {"type": "heading", "level": 1, "text": "一、合成"},
            {"type": "paragraph", "text": "合成正文。"}]}, ensure_ascii=False), encoding="utf-8")
        out = folder / "doc.docx"
        E.build(content, out, intake_path=state)
        return out

    def pdf(self, pages):
        """pages: list of lists of (kind, x, y, payload)."""
        path = self.folder / "doc.pdf"
        doc = fitz.open()
        for items in pages:
            page = doc.new_page(width=A4[0], height=A4[1])
            for kind, x, y, payload in items:
                if kind == "text":
                    page.insert_text((x, y), payload, fontname="china-s", fontsize=10.5)
                elif kind == "small":
                    page.insert_text((x, y), payload, fontname="china-s", fontsize=6)
                elif kind == "tall":
                    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4), 0)
                    page.insert_image(fitz.Rect(x, y, x + 200, y + payload), pixmap=pix)
                elif kind == "image":
                    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 4, 4), 0)
                    page.insert_image(fitz.Rect(x, y, x + payload, y + 40), pixmap=pix)
                elif kind == "line":
                    page.draw_line((x, y), (x + payload, y))
        doc.save(path)
        return path

    def codes(self, result, key="errors"):
        return [e["code"] for e in result[key]]

    # B: left 3.00 cm (85 pt), right 2.20 cm -> text column 85..533 pt, no header, no page numbers.
    def body(self, text="合成正文。"):
        return [("text", 106, 200, text)]  # 2-character first-line indent

    def cover(self):
        return [("text", 120, 300, "封面")]

    def test_clean_pages_pass(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body()]))
        self.assertTrue(result["passed"], result["errors"])

    def test_text_beyond_right_margin(self):
        result = audit(self.docx(), self.pdf([self.cover(), [("text", 480, 200, "伸出版心的文字文字文字")]]))
        self.assertIn("TEXT_OUTSIDE_MARGINS", self.codes(result))
        issue = result["errors"][0]
        self.assertEqual(issue["page"], 2)
        self.assertIn("actual", issue)

    def test_body_text_pushed_past_the_bottom_margin(self):
        # B: bottom margin 2.54 cm -> the text column ends at 769.9 pt.
        result = audit(self.docx(), self.pdf([self.cover(), [("text", 106, 780, "合成正文。")]]))
        self.assertTrue(any("正文伸进了上下页边距" in e["message"] for e in result["errors"]), result["errors"])

    def test_cover_underline_past_the_margin_is_named_as_the_cover(self):
        result = audit(self.docx(), self.pdf([self.cover() + [("line", 400, 330, 150)], self.body()]))
        self.assertEqual(self.codes(result), ["COVER_OUTSIDE_MARGINS"])

    def test_one_trailing_punctuation_may_hang(self):
        # The last glyph is a comma just past the column: 行末标点悬挂 is not a defect.
        x = 533 - 10.5 * 3
        result = audit(self.docx(), self.pdf([self.cover(), [("text", x, 200, "文字字，")]]))
        self.assertTrue(result["passed"], result["errors"])

    def test_blank_page_is_an_error_with_context(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body("前一页最后一句。"), []]))
        blank = [e for e in result["errors"] if e["code"] == "BLANK_PAGE"]
        self.assertEqual([e["page"] for e in blank], [3])
        self.assertIn("前一页最后一句", blank[0]["where"])

    def test_page_with_only_a_page_number_is_blank(self):
        doc = self.docx(patch={"footer_page_numbers": True})
        result = audit(doc, self.pdf([self.cover(), self.body() + [("text", 295, 805, "1")], [("text", 295, 805, "2")]]))
        self.assertIn("BLANK_PAGE", self.codes(result))

    def test_unexpected_header_on_b(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body() + [("text", 90, 50, "多余页眉")]]))
        self.assertIn("HEADER_UNEXPECTED", self.codes(result))

    def test_missing_header_on_a(self):
        result = audit(self.docx("A"), self.pdf([self.cover(), [("text", 95, 200, "正文")]]))
        self.assertIn("HEADER_MISSING", self.codes(result))

    def test_cover_must_not_show_footer(self):
        result = audit(self.docx(), self.pdf([self.cover() + [("text", 295, 805, "1")], self.body()]))
        self.assertIn("COVER_FOOTER_VISIBLE", self.codes(result))

    def test_requested_page_numbers_must_render(self):
        doc = self.docx(patch={"footer_page_numbers": True})
        result = audit(doc, self.pdf([self.cover(), self.body()]))
        self.assertIn("PAGE_NUMBER_MISSING", self.codes(result))
        result = audit(doc, self.pdf([self.cover(), self.body() + [("text", 295, 805, "1")]]))
        self.assertTrue(result["passed"], result["errors"])

    def test_page_number_at_the_paper_edge_is_refused(self):
        doc = self.docx(patch={"footer_page_numbers": True})
        result = audit(doc, self.pdf([self.cover(), self.body() + [("text", 295, 838, "1")]]))
        hit = [e for e in result["errors"] if e["code"] == "TEXT_AT_PAPER_EDGE"]
        self.assertEqual((hit[0]["page"], hit[0]["where"]), (2, "1"), result["errors"])
        self.assertIn("footer", F.HINTS["TEXT_AT_PAPER_EDGE"][1])

    def test_footer_text_when_page_numbers_are_off(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body() + [("text", 295, 805, "1")]]))
        self.assertIn("FOOTER_UNEXPECTED", self.codes(result))

    def test_image_and_table_rules_beyond_the_column(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body() + [("image", 300, 300, 260), ("line", 85, 500, 470)]]))
        self.assertIn("IMAGE_OUTSIDE_MARGINS", self.codes(result))
        self.assertIn("TABLE_OUTSIDE_MARGINS", self.codes(result))

    def test_off_page_text(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body() + [("text", 560, 300, "出界出界出界出界")]]))
        self.assertIn("OFF_PAGE", self.codes(result))

    def test_overlapping_lines_block_delivery(self):
        result = audit(self.docx(), self.pdf([self.cover(), [("text", 90, 200, "第一行文字"), ("text", 92, 202, "第二行文字")]]))
        self.assertFalse(result["passed"])
        overlap = [e for e in result["errors"] if e["code"] == "TEXT_OVERLAP"]
        self.assertEqual(overlap[0]["page"], 2)
        self.assertIn("第一行文字", overlap[0]["where"])

    def test_lines_one_line_apart_do_not_overlap(self):
        result = audit(self.docx(), self.pdf([self.cover(), [("text", 90, 200, "第一行文字"), ("text", 90, 216, "第二行文字")]]))
        self.assertTrue(result["passed"], result["errors"])

    def test_overlap_is_judged_against_the_font_size(self):
        # 10.5 pt text: baselines 8 pt apart print into each other, 10 pt apart only look tight.
        tight = audit(self.docx(), self.pdf([self.cover(), [("text", 90, 200, "第一行文字"), ("text", 90, 208, "第二行文字")]]))
        overlap = [e for e in tight["errors"] if e["code"] == "TEXT_OVERLAP"]
        self.assertTrue(overlap, tight["errors"])
        self.assertIn("0.9", overlap[0]["expected"])
        self.assertIn("基线相距 8.0pt", overlap[0]["actual"])
        snug = audit(self.docx(), self.pdf([self.cover(), [("text", 90, 200, "第一行文字"), ("text", 90, 210, "第二行文字")]]))
        self.assertTrue(snug["passed"], snug["errors"])

    def test_large_title_above_small_text_is_judged_by_the_room_left(self):
        doc = fitz.open()
        page = doc.new_page(width=A4[0], height=A4[1])
        page.insert_text((150, 120), "封面大字", fontname="china-s", fontsize=72)
        page.insert_text((160, 151), "下面的小字说明", fontname="china-s", fontsize=10.5)
        path = self.folder / "mixed.pdf"
        doc.save(path)
        from render_audit import _lines, _line_clash

        with fitz.open(path) as rendered:
            lines = _lines(rendered[0])
        self.assertIsNone(_line_clash(lines[0], lines[1]))

    def test_superscript_beside_a_word_is_not_an_overlap(self):
        result = audit(self.docx(), self.pdf([self.cover(), [("text", 90, 200, "引用文字"), ("small", 132.5, 195, "[1]"),
                                                            ("text", 145, 200, "后文")]]))
        self.assertTrue(result["passed"], result["errors"])

    def test_text_printed_on_an_image(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body() + [("image", 150, 300, 200), ("text", 160, 320, "压在图上的字")]]))
        hit = [e for e in result["errors"] if e["code"] == "TEXT_OVER_IMAGE"]
        self.assertEqual(hit[0]["page"], 2)
        self.assertIn("压在图上的字", hit[0]["where"])
        clear = audit(self.docx(), self.pdf([self.cover(), self.body() + [("image", 150, 300, 200), ("text", 160, 360, "图下的图题")]]))
        self.assertTrue(clear["passed"], clear["errors"])

    def test_image_taller_than_the_text_area(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body() + [("tall", 150, 100, 720)]]))
        tall = [e for e in result["errors"] if e["code"] == "IMAGE_OUTSIDE_MARGINS"]
        self.assertTrue(tall and "上下" in tall[0]["message"], result["errors"])

    def test_header_wrapping_into_the_text_area(self):
        # A: top margin 72 pt.  The header text ("测试") wraps to a third line that ends below it.
        doc = self.docx("A")
        header = [("text", 95, 50, "测试"), ("text", 95, 62.5, "测试")]
        ok = audit(doc, self.pdf([self.cover(), header + [("text", 95, 200, "正文")]]))
        self.assertNotIn("HEADER_OVERFLOW", self.codes(ok))
        bad = audit(doc, self.pdf([self.cover(), header + [("text", 95, 75, "测试"), ("text", 95, 200, "正文")]]))
        spill = [e for e in bad["errors"] if e["code"] == "HEADER_OVERFLOW"]
        self.assertEqual(spill[0]["page"], 2)
        self.assertIn("72.0pt", spill[0]["expected"])
        self.assertIn("format.header", F.HINTS["HEADER_OVERFLOW"][1].replace("\"header\"", "format.header"))

    def test_header_wrapped_to_two_lines_inside_the_margin(self):
        # Both rows stay above the 72 pt top margin, so nothing overflows, but the header is no longer one line.
        doc = self.docx("A")
        one = audit(doc, self.pdf([self.cover(), [("text", 95, 60, "测试"), ("text", 400, 60, "测试"), ("text", 95, 200, "正文")]]))
        self.assertNotIn("HEADER_WRAPPED", self.codes(one))
        two = audit(doc, self.pdf([self.cover(), [("text", 95, 50, "测试"), ("text", 95, 64, "测试"), ("text", 95, 200, "正文")]]))
        wrapped = [e for e in two["errors"] if e["code"] == "HEADER_WRAPPED"]
        self.assertEqual(wrapped[0]["page"], 2)
        self.assertEqual(wrapped[0]["actual"], "2 行")
        self.assertIn("header", F.HINTS["HEADER_WRAPPED"][1])

    def test_heading_left_alone_at_the_bottom_of_a_page(self):
        bottom = audit(self.docx(), self.pdf([self.cover(), self.body() + [("text", 90, 760, "一、合成")], self.body()]))
        hit = [e for e in bottom["errors"] if e["code"] == "HEADING_AT_PAGE_BOTTOM"]
        self.assertEqual(hit[0]["page"], 2)
        top = audit(self.docx(), self.pdf([self.cover(), [("text", 90, 120, "一、合成")] + self.body(), self.body()]))
        self.assertNotIn("HEADING_AT_PAGE_BOTTOM", self.codes(top))
        # The last page may end with anything.
        last = audit(self.docx(), self.pdf([self.cover(), self.body() + [("text", 90, 760, "一、合成")]]))
        self.assertNotIn("HEADING_AT_PAGE_BOTTOM", self.codes(last))

    def test_table_caption_left_at_the_bottom_of_a_page(self):
        doc = self.docx(blocks=[{"type": "heading", "level": 1, "text": "一、合成"},
                                {"type": "table", "caption": "表1 合成数据表", "header": ["甲", "乙"], "rows": [["1", "2"]]}])
        result = audit(doc, self.pdf([self.cover(), self.body() + [("text", 90, 760, "表1 合成数据表")], self.body()]))
        hit = [e for e in result["errors"] if e["code"] == "CAPTION_SEPARATED"]
        self.assertTrue(hit and "表题" in hit[0]["message"], result["errors"])

    def test_table_caption_behind_an_empty_paragraph_is_still_paired(self):
        from docx import Document

        doc = self.docx(blocks=[{"type": "heading", "level": 1, "text": "一、合成"},
                                {"type": "table", "caption": "表1 合成数据表", "header": ["甲", "乙"], "rows": [["1", "2"]]}])
        d = Document(doc)
        caption = next(p for p in d.paragraphs if p.text == "表1 合成数据表")
        from docx.oxml import OxmlElement

        caption._p.addnext(OxmlElement("w:p"))  # an empty paragraph between the caption and its table
        d.save(doc)
        result = audit(doc, self.pdf([self.cover(), self.body() + [("text", 90, 760, "表1 合成数据表")], self.body()]))
        self.assertIn("CAPTION_SEPARATED", self.codes(result))

    def test_body_rendered_at_the_wrong_size_is_refused(self):
        text = "这是一段用来检查正文实际显示字号的合成正文，长度足够统计，再多写几个字让它超过四十个字符。"
        doc = self.docx(blocks=[{"type": "heading", "level": 1, "text": "一、合成"},
                                {"type": "paragraph", "text": text}])
        right = audit(doc, self.pdf([self.cover(), [("text", 90, 200, text[:20]), ("text", 90, 220, text[20:])]]))
        self.assertNotIn("BODY_SIZE_RENDERED", self.codes(right))
        wrong = audit(doc, self.pdf([self.cover(), [("small", 90, 200, text[:20]), ("small", 90, 220, text[20:])]]))
        hit = [e for e in wrong["errors"] if e["code"] == "BODY_SIZE_RENDERED"]
        self.assertEqual((hit[0]["expected"], hit[0]["actual"]), ("10.5 磅", "6 磅"))

    def test_wrapped_cover_title_must_line_up(self):
        doc = self.docx("A")
        first = ("text", 125, 340, "论文题目：基于某某数据的")
        aligned = audit(doc, self.pdf([[first, ("text", 177.5, 356, "方法研究")], self.body()]))
        self.assertNotIn("COVER_TITLE_MISALIGNED", self.codes(aligned))
        drift = audit(doc, self.pdf([[first, ("text", 240, 356, "方法研究")], self.body()]))
        hit = [e for e in drift["errors"] if e["code"] == "COVER_TITLE_MISALIGNED"]
        self.assertEqual(hit[0]["page"], 1)
        self.assertIn("177.5", hit[0]["expected"])

    def test_every_overlap_on_a_page_is_reported(self):
        pairs = [("text", 90, 200, "第一行文字"), ("text", 90, 203, "第二行文字"),
                 ("text", 90, 400, "第三行文字"), ("text", 90, 403, "第四行文字")]
        result = audit(self.docx(), self.pdf([self.cover(), pairs]))
        self.assertEqual(len([e for e in result["errors"] if e["code"] == "TEXT_OVERLAP"]), 2)

    def test_page_with_only_a_drawing_is_not_blank(self):
        result = audit(self.docx(), self.pdf([self.cover(), [("line", 100, 300, 300)]]))
        self.assertNotIn("BLANK_PAGE", self.codes(result))

    def test_repeated_issues_are_shown_once_with_their_pages(self):
        result = audit(self.docx(), self.pdf([self.cover()] + [self.body() + [("text", 90, 50, "多余页眉")] for _ in range(3)]))
        issues = F.from_render(result)
        self.assertEqual(len(issues), 3)
        text = F.render_text(issues)
        self.assertEqual(text.count("[HEADER_UNEXPECTED]"), 1)
        self.assertIn("第2页、第3页、第4页（共 3 处）", text)

    def test_same_cause_and_fix_are_written_once(self):
        issues = [{"code": "TEXT_OUTSIDE_MARGINS", "message": "文字伸出了左右页边距。", "where": {"page": n, "text": t},
                   "cause": "长网址", "fix": "加断行"} for n, t in ((2, "甲"), (3, "乙"))]
        text = F.render_text(issues)
        self.assertEqual(text.count("怎么改：加断行"), 1)
        self.assertIn("同第 1 条", text)

    def test_blank_page_points_at_the_empty_break_paragraphs_of_the_mapping(self):
        from docx import Document
        from docx.enum.text import WD_BREAK

        doc = Document()
        doc.add_paragraph("前一页最后一句。")
        for _ in range(2):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph("后一页。")
        out = self.folder / "breaks.docx"
        doc.save(out)
        (self.folder / "mapping.json").write_text(json.dumps({"paragraphs": [
            {"id": "p0", "text": "前一页最后一句。", "role": "body"}, {"id": "p1", "text": "", "role": "spacer"},
            {"id": "p2", "text": "", "role": "spacer"}, {"id": "p3", "text": "后一页。", "role": "body"}]},
            ensure_ascii=False), encoding="utf-8")
        audit_result = {"errors": [{"code": "BLANK_PAGE", "page": 2, "message": "空白页",
                                    "where": "紧跟在“前一页最后一句。”之后"}]}
        issue = F.from_render(audit_result, self.folder, out)[0]
        self.assertEqual(issue["where"]["location"], "成品 p1、p2")
        self.assertIn("p1、p2", issue["where"]["edit"])
        self.assertIn("delete", issue["fix"])

    def test_blank_page_points_at_the_page_break_blocks_of_content_json(self):
        folder = Path(tempfile.mkdtemp(dir=self.folder))
        (folder / "content.json").write_text(json.dumps({"blocks": [
            {"type": "heading", "level": 1, "text": "一、引言"}, {"type": "paragraph", "text": "前一页最后一句。"},
            {"type": "page_break"}, {"type": "page_break"}, {"type": "paragraph", "text": "后一页。"}]},
            ensure_ascii=False), encoding="utf-8")
        audit_result = {"errors": [{"code": "BLANK_PAGE", "page": 2, "message": "空白页",
                                    "where": "紧跟在“前一页最后一句。”之后"}]}
        issue = F.from_render(audit_result, folder, self.folder / "missing.docx")[0]
        self.assertEqual(issue["where"]["edit"], "content.json 的 blocks[2]、blocks[3]（page_break）")

    def test_stop_rule_counts_runs_without_progress(self):
        def issue(page):
            return {"code": "BLANK_PAGE", "where": {"page": page}}

        rounds = [F.record_round(self.folder, [issue(3), issue(5)])]
        rounds.append(F.record_round(self.folder, [issue(5)]))  # one fixed: progress
        rounds += [F.record_round(self.folder, [issue(5)]) for _ in range(F.MAX_STALLED)]
        self.assertEqual([r[1] for r in rounds], [0, 0] + list(range(1, F.MAX_STALLED + 1)))
        self.assertTrue(rounds[-1][2])
        self.assertEqual(F.record_round(self.folder, []), (0, 0, False))
        # After a pass the first failure is new, not a stall; the same code at another block, or a
        # better measured value, is progress too.
        rejected = lambda block, actual="10 磅": {"code": "INPUT_REJECTED", "message": f"blocks[{block}] 不对",
                                                 "where": {"edit": f"content.json 的 blocks[{block}]"}, "actual": actual}
        self.assertEqual(F.record_round(self.folder, [rejected(5)])[1], 0)
        self.assertEqual(F.record_round(self.folder, [rejected(3)])[1], 0)
        self.assertEqual(F.record_round(self.folder, [rejected(3, "15 磅")])[1], 0)
        self.assertEqual(F.record_round(self.folder, [rejected(3, "15 磅")])[1:], (1, True))

    def numbered(self, *labels):
        doc = self.docx(patch={"footer_page_numbers": True})
        return audit(doc, self.pdf([self.cover()] + [self.body() + [("text", 295, 805, n)] for n in labels]))

    def test_page_numbers_roman_front_then_arabic_from_one(self):
        result = self.numbered("I", "II", "1", "2", "3")
        self.assertTrue(result["passed"], result["errors"])

    def test_page_number_restart_skip_and_bad_start(self):
        cases = {("I", "1", "1"): 4, ("1", "2", "4"): 4, ("I", "II", "2"): 4, ("2", "3"): 2, ("1", "II"): 3}
        for labels, page in cases.items():
            with self.subTest(labels=labels):
                result = self.numbered(*labels)
                wrong = [e for e in result["errors"] if e["code"] == "PAGE_NUMBER_SEQUENCE"]
                self.assertEqual([e["page"] for e in wrong], [page], result["errors"])
                self.assertIn("expected", wrong[0])

    def test_cover_may_use_its_own_measure_but_not_the_page_edge(self):
        ok = audit(self.docx(), self.pdf([[("text", 90, 300, "《很长的课程名称在封面上居中》")], self.body()]))
        self.assertTrue(ok["passed"], ok["errors"])
        bad = audit(self.docx(), self.pdf([[("text", 10, 300, "贴边")], self.body()]))
        self.assertIn("TEXT_OUTSIDE_MARGINS", self.codes(bad))

    def test_font_substitution_names_the_missing_families(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body()]))
        warning = [w for w in result["warnings"] if w["code"] == "FONT_SUBSTITUTED"]
        self.assertTrue(warning)
        self.assertIn("宋体", result["fonts"]["missing"])

    def test_issues_carry_location_cause_and_fix(self):
        result = audit(self.docx(), self.pdf([self.cover(), self.body(), []]))
        issues = F.from_render(result)
        self.assertEqual(issues[0]["where"]["page"], 3)
        self.assertTrue(issues[0]["cause"] and issues[0]["fix"])
        text = F.render_text(issues)
        for word in ("位置", "可能原因", "怎么改"):
            self.assertIn(word, text)


@unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "needs LibreOffice")
class RenderAuditRealOffice(unittest.TestCase):
    """A real export of a defective document must be caught from the PDF alone."""

    def setUp(self):
        # Workspaces made by `new` go here, so a test that stops half-way leaves nothing in /tmp.
        from unittest.mock import patch

        scratch = tempfile.TemporaryDirectory()
        self.addCleanup(scratch.cleanup)
        temp = patch("tempfile.gettempdir", return_value=scratch.name)
        temp.start()
        self.addCleanup(temp.stop)

    def test_blank_page_from_double_page_break(self):
        from office_backend import export_pdf
        from docx import Document
        from docx.enum.text import WD_BREAK

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            case = RenderAudit("test_clean_pages_pass")
            case.folder = folder
            out = case.docx(cover=False)
            doc = Document(out)
            for _ in range(2):
                empty = doc.add_paragraph("", style=E.ROLES["body"])
                empty.add_run().add_break(WD_BREAK.PAGE)
            doc.add_paragraph("尾页", style=E.ROLES["body"])
            doc.save(out)
            pdf = folder / "real.pdf"
            export_pdf(out, pdf)
            result = audit(out, pdf)
            self.assertIn("BLANK_PAGE", [e["code"] for e in result["errors"]])


    def test_typography_that_did_not_take_effect_is_caught_from_the_pdf(self):
        """Indent, centring and fixed line pitch, measured on a real export against the template."""
        from office_backend import export_pdf
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt

        text = "这是一段足够长的合成正文，用来检查渲染出来的首行缩进、行距和标题居中是否和模板要求一致，所以要写到两行以上才行，再多写几个字。"
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            case = RenderAudit("test_clean_pages_pass")
            case.folder = folder
            out = case.docx(cover=False, blocks=[{"type": "heading", "level": 1, "text": "一、合成标题 Results"},
                                                 {"type": "paragraph", "text": text}])
            pdf = folder / "good.pdf"
            export_pdf(out, pdf)
            self.assertTrue(audit(out, pdf)["passed"], audit(out, pdf)["errors"])
            doc = Document(out)
            heading = next(p for p in doc.paragraphs if p.text == "一、合成标题 Results")
            heading.paragraph_format.alignment = 0  # left
            for run in heading.runs:
                run.font.size, run.font.bold = Pt(18), True
            body = next(p for p in doc.paragraphs if p.text == text)
            body.paragraph_format.first_line_indent = 0
            body._p.pPr.find(qn("w:ind")).set(qn("w:firstLineChars"), "0")
            body.paragraph_format.line_spacing = Pt(15)
            doc.save(out)
            pdf = folder / "bad.pdf"
            export_pdf(out, pdf)
            codes = [e["code"] for e in audit(out, pdf)["errors"]]
            for code in ("NOT_CENTERED", "INDENT_RENDERED", "LINE_PITCH_RENDERED", "SIZE_RENDERED", "WEIGHT_RENDERED"):
                self.assertIn(code, codes)

    def test_page_numbers_on_b_stay_clear_of_the_paper_edge(self):
        from office_backend import export_pdf

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            case = RenderAudit("test_clean_pages_pass")
            case.folder = folder
            for patch, expect in (({"footer_page_numbers": True}, False),
                                  ({"footer_page_numbers": True, "page_twips": {"footer": 0}}, True)):
                out = case.docx(patch=patch)
                pdf = folder / "real.pdf"
                export_pdf(out, pdf)
                codes = [e["code"] for e in audit(out, pdf)["errors"]]
                self.assertEqual("TEXT_AT_PAPER_EDGE" in codes, expect, (patch, codes))

    def test_caption_split_from_its_table_by_a_page_break(self):
        """A stray page break between a caption and its table: caught in the Word file and in the PDF."""
        from office_backend import export_pdf
        from docx import Document
        from docx.enum.text import WD_BREAK
        import check_format as C

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            case = RenderAudit("test_clean_pages_pass")
            case.folder = folder
            out = case.docx(cover=False, blocks=[
                {"type": "heading", "level": 1, "text": "一、合成"},
                {"type": "table", "caption": "表1 站点统计", "header": ["甲", "乙"], "rows": [["1", "2"]]}])
            doc = Document(out)
            caption = next(p for p in doc.paragraphs if p.text == "表1 站点统计")
            blank = doc.add_paragraph("", style=E.ROLES["body"])
            blank.add_run().add_break(WD_BREAK.PAGE)
            caption._p.addnext(blank._p)
            doc.save(out)
            self.assertIn("CAPTION_DETACHED", [e["code"] for e in C.check(out)["errors"]])
            pdf = folder / "real.pdf"
            export_pdf(out, pdf)
            result = audit(out, pdf)
            self.assertIn("CAPTION_SEPARATED", [e["code"] for e in result["errors"]], result["errors"])

    def test_lowered_text_overlaps_although_the_static_rules_pass(self):
        """A run moved down by w:position keeps every style rule but prints into the next line."""
        from office_backend import export_pdf
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        import check_format as C

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            case = RenderAudit("test_clean_pages_pass")
            case.folder = folder
            out = case.docx(cover=False)
            doc = Document(out)
            first = doc.add_paragraph("", style=E.ROLES["body"])
            run = first.add_run("这一行的文字被整体降低了十二磅，落到下一行上。")
            position = OxmlElement("w:position")
            position.set(qn("w:val"), "-24")
            run._r.get_or_add_rPr().append(position)
            doc.add_paragraph("下一行正常的正文文字，用来检验两行是否叠在一起。", style=E.ROLES["body"])
            doc.save(out)
            static = [e["code"] for e in C.check(out)["errors"]]
            self.assertNotIn("AUDIT_ABORTED", static)
            self.assertNotIn("LINE_TOO_TIGHT", static)
            pdf = folder / "real.pdf"
            export_pdf(out, pdf)
            result = audit(out, pdf)
            self.assertIn("TEXT_OVERLAP", [e["code"] for e in result["errors"]], result["errors"])

    def test_hand_edited_copy_is_accepted_only_through_adopt(self):
        """A requirement the scripts cannot meet: edit a copy, adopt re-runs every check, the user is told."""
        import contextlib, io
        import xaufe as X
        from docx import Document
        from docx.shared import RGBColor

        def call(*args):
            out = io.StringIO()
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(out):
                code = X.main([str(a) for a in args])
            return code, out.getvalue()

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / "原稿.txt"
            source.write_text("一、引言\n\n这是合成正文。", encoding="utf-8")
            code, out = call("new", "--out", folder / "成品.docx", "--source", source)
            work = Path(out.split("WORK = ")[1].splitlines()[0].strip())
            job = {"task": "format", "template": "B", "cover": False, "info": {}}
            (work / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
            content = {"mode": "document", "metadata": {}, "blocks": [
                {"type": "heading", "level": 1, "text": "一、引言"}, {"type": "paragraph", "text": "这是合成正文。"}]}
            (work / "content.json").write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
            code, out = call("run", "--work", work)
            self.assertEqual(code, 0, out)

            edited = folder / "手工.docx"
            doc = Document(work / "final.docx")
            body = [p for p in doc.paragraphs if p.text == "这是合成正文。"][0]
            body.runs[0].font.color.rgb = RGBColor(0x00, 0x55, 0xAA)
            doc.save(edited)
            code, out = call("adopt", "--work", work, "--docx", edited, "--note", "正文按用户要求改成蓝色")
            self.assertEqual(code, 0, out)
            self.assertIn("MANUAL_DEVIATION_FONT_COLOR", out)
            self.assertIn("“这是合成正文。”的文字颜色和模板默认不同", out)
            self.assertNotIn("TOC_PROOF_STALE", out)
            review = json.loads((work / "review.json").read_text(encoding="utf-8"))
            self.assertIn("MANUAL_EDIT", review["warnings"])
            review.update(vision=True, pages={k: "一、引言下面这是合成正文，颜色为蓝色" for k in review["pages"]},
                          content_check="单元测试合成记录：已核对", layout_check="单元测试合成记录：已核对",
                          delivery_check="单元测试合成记录：已核对",
                          warnings={k: "单元测试合成记录：已处理" for k in review["warnings"]})
            (work / "review.json").write_text(json.dumps(review, ensure_ascii=False), encoding="utf-8")
            code, out = call("finish", "--work", work)
            self.assertEqual(code, 0, out)
            self.assertIn("正文按用户要求改成蓝色", out)
            self.assertIn("和模板默认格式不同", out)

    def test_adopted_body_size_is_not_held_to_the_template_size(self):
        """The user asked for a body size the scripts could not set: the rendered check follows the copy."""
        import contextlib, io
        import xaufe as X
        from docx import Document
        from docx.shared import Pt

        def call(*args):
            out = io.StringIO()
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(out):
                code = X.main([str(a) for a in args])
            return code, out.getvalue()

        text = "这是一段足够长的合成正文，用来检查手工改了正文字号以后，渲染后的字号检查以手工副本为准，而不是仍按模板默认字号拦下。"
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / "原稿.txt"
            source.write_text("一、引言\n\n" + text, encoding="utf-8")
            code, out = call("new", "--out", folder / "成品.docx", "--source", source)
            work = Path(out.split("WORK = ")[1].splitlines()[0].strip())
            job = {"task": "format", "template": "B", "cover": False, "info": {}}
            (work / "job.json").write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
            content = {"mode": "document", "metadata": {}, "blocks": [
                {"type": "heading", "level": 1, "text": "一、引言"}, {"type": "paragraph", "text": text}]}
            (work / "content.json").write_text(json.dumps(content, ensure_ascii=False), encoding="utf-8")
            code, out = call("run", "--work", work)
            self.assertEqual(code, 0, out)
            edited = folder / "手工.docx"
            doc = Document(work / "final.docx")
            for run in [p for p in doc.paragraphs if p.text == text][0].runs:
                run.font.size = Pt(14)
            doc.save(edited)
            code, out = call("adopt", "--work", work, "--docx", edited, "--note", "正文按用户要求改成四号")
            self.assertEqual(code, 0, out)
            self.assertIn("MANUAL_DEVIATION_FONT_SIZE", out)
            self.assertNotIn("BODY_SIZE_RENDERED", out)
            with contextlib.redirect_stdout(io.StringIO()):
                X.main(["cancel", "--work", str(work)])

    def test_named_renderer_that_cannot_run_is_reported(self):
        import contextlib, io
        import xaufe as X
        from office_discovery import discover

        if any(c["engine"] == "word" and c["found"] and c["automation"] for c in discover()["candidates"]):
            self.skipTest("Word can be driven here")
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / "原稿.txt"
            source.write_text("这是合成正文。", encoding="utf-8")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                X.main(["new", "--out", str(folder / "成品.docx"), "--source", str(source)])
            work = Path(out.getvalue().split("WORK = ")[1].splitlines()[0].strip())
            (work / "job.json").write_text(json.dumps({"task": "format", "template": "B", "cover": False}), encoding="utf-8")
            (work / "content.json").write_text(json.dumps({"mode": "document", "metadata": {}, "blocks": [
                {"type": "paragraph", "text": "这是合成正文。"}]}, ensure_ascii=False), encoding="utf-8")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                self.assertEqual(X.main(["run", "--work", str(work), "--office", "word"]), 0, out.getvalue())
            self.assertIn("RENDERER_FALLBACK", out.getvalue())
            self.assertIn("指定用 Word 渲染", out.getvalue())
            review = json.loads((work / "review.json").read_text(encoding="utf-8"))
            self.assertIn("RENDERER_FALLBACK", review["warnings"])
            with contextlib.redirect_stdout(io.StringIO()):
                X.main(["cancel", "--work", str(work)])

    def test_adopt_refuses_a_hand_edit_that_breaks_the_layout(self):
        import contextlib, io
        import xaufe as X
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_LINE_SPACING

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / "原稿.txt"
            source.write_text("这是合成正文。", encoding="utf-8")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                X.main(["new", "--out", str(folder / "成品.docx"), "--source", str(source)])
            work = Path(out.getvalue().split("WORK = ")[1].splitlines()[0].strip())
            (work / "job.json").write_text(json.dumps({"task": "format", "template": "B", "cover": False}), encoding="utf-8")
            (work / "content.json").write_text(json.dumps({"mode": "document", "metadata": {}, "blocks": [
                {"type": "paragraph", "text": "这是合成正文。"}]}, ensure_ascii=False), encoding="utf-8")
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(X.main(["run", "--work", str(work)]), 0)
            accepted = (work / "final.docx").read_bytes()
            doc = Document(work / "final.docx")
            body = [p for p in doc.paragraphs if p.text == "这是合成正文。"][0]
            body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            body.paragraph_format.line_spacing = Pt(5)
            doc.save(folder / "手工.docx")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                code = X.main(["adopt", "--work", str(work), "--docx", str(folder / "手工.docx"), "--note", "合成测试：行距改小"])
            self.assertEqual(code, X.EXIT_NOT_ACCEPTED, out.getvalue())
            self.assertIn("LINE_TOO_TIGHT", out.getvalue())
            self.assertIn("adopt 不放行", out.getvalue())
            self.assertEqual((work / "final.docx").read_bytes(), accepted)
            with contextlib.redirect_stdout(io.StringIO()):
                X.main(["cancel", "--work", str(work)])

    def test_adopt_refuses_changed_text(self):
        import contextlib, io
        import xaufe as X
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / "原稿.txt"
            source.write_text("这是合成正文。", encoding="utf-8")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                X.main(["new", "--out", str(folder / "成品.docx"), "--source", str(source)])
            work = Path(out.getvalue().split("WORK = ")[1].splitlines()[0].strip())
            (work / "job.json").write_text(json.dumps({"task": "format", "template": "B", "cover": False}), encoding="utf-8")
            (work / "content.json").write_text(json.dumps({"mode": "document", "metadata": {}, "blocks": [
                {"type": "paragraph", "text": "这是合成正文。"}]}, ensure_ascii=False), encoding="utf-8")
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(X.main(["run", "--work", str(work)]), 0)
            doc = Document(work / "final.docx")
            doc.add_paragraph("手工加的一句。")
            doc.save(folder / "手工.docx")
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                code = X.main(["adopt", "--work", str(work), "--docx", str(folder / "手工.docx"), "--note", "合成测试"])
            self.assertEqual(code, X.EXIT_NOT_ACCEPTED)
            self.assertIn("MANUAL_TEXT_CHANGED", out.getvalue())
            self.assertIn("手工文件的第", out.getvalue())
            self.assertIn("手工加的一句", out.getvalue())
            with contextlib.redirect_stdout(io.StringIO()):
                X.main(["cancel", "--work", str(work)])


if __name__ == "__main__":
    unittest.main()
