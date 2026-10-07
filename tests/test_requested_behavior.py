import json
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from test_layout_integrity import Fixture, candidate
import engine as E, intake as I, render_availability as RF
from check_format import check


class CoverAndManualFields(Fixture):
    def test_long_header_wraps_without_losing_metadata(self):
        state = self.state()
        I.customize(
            state,
            {
                "header": {
                    "left": "跨学科公共管理与数字化治理研究方向",
                    "right": "数字校园公共空间治理与信息服务协同机制研究" * 4,
                }
            },
            "synthetic long header",
        )
        source = self.path("json")
        source.write_text(
            json.dumps(
                {
                    "mode": "document",
                    "cover": False,
                    "blocks": [{"type": "paragraph", "text": "正文。"}],
                }
            ),
            encoding="utf-8",
        )
        output = self.path("docx")
        E.build(source, output, intake_path=state)
        doc = Document(output)
        header = doc.sections[0].header
        self.assertEqual(len(header.paragraphs), 2)
        self.assertEqual(
            header.paragraphs[1].text.replace(chr(10), ""), "数字校园公共空间治理与信息服务协同机制研究" * 4
        )
        self.assertEqual(len(header._element.xpath(".//w:pBdr/w:bottom")), 1)
        self.assertPass(check(output, intake_path=state))

    def test_toc_scan_skips_unstyled_section_carrier(self):
        from audit_helpers import toc_entries

        doc = Document()
        doc.styles["Normal"]._element.attrib.pop(qn("w:default"), None)
        doc.add_paragraph()
        self.assertIsNone(doc.paragraphs[0].style)
        self.assertEqual(toc_entries(doc), [])

    def test_added_teacher_uses_base_cover_columns(self):
        state = self.state("B")
        I.customize(
            state,
            {"cover_supervisor": True, "logo": "green"},
            "synthetic add teacher",
        )
        I.resolve(state, {"metadata": {"supervisor": candidate("陈测试")}})
        source = self.path("json")
        source.write_text(
            json.dumps(
                {
                    "mode": "document",
                    "cover": True,
                    "blocks": [{"type": "paragraph", "text": "合成正文。"}],
                }
            ),
            encoding="utf-8",
        )
        output = self.path("docx")
        E.build(source, output, intake_path=state)
        doc = Document(output)
        for tag in ("ind", "tabs", "jc", "spacing"):
            left = doc.paragraphs[9]._p.pPr.find(qn("w:" + tag))
            right = doc.paragraphs[13]._p.pPr.find(qn("w:" + tag))
            properties = lambda node: (
                dict(node.attrib),
                [(child.tag, dict(child.attrib)) for child in node],
            )
            self.assertEqual(properties(left), properties(right))
        self.assertEqual(doc.paragraphs[13].runs[1].text, "\t陈测试\t")
        self.assertPass(check(output, intake_path=state))

    def test_pending_fields_are_manual_and_editable(self):
        output, _ = self.build(paper=True)
        doc = Document(output)
        before = [f.get(qn("w:instr")) for f in RF.page_fields(doc)]
        RF.mark_pending(doc)
        self.assertEqual([f.get(qn("w:instr")) for f in RF.page_fields(doc)], before)
        self.assertEqual(doc.settings.element.xpath("./w:updateFields/@w:val"), ["false"])
        for field in RF.page_fields(doc):
            self.assertNotIn(field.get(qn("w:dirty")), ("true", "1"))
            self.assertNotIn(field.get(qn("w:fldLock")), ("true", "1"))
            self.assertEqual(
                "".join(t.text or "" for t in field.findall(".//" + qn("w:t"))), "待更新"
            )
        errors = []
        RF.pending_check(doc, lambda code, message: errors.append(code))
        self.assertEqual(errors, [])

    def test_open_time_update_request_is_rejected_by_format_audit(self):
        output, state = self.build()
        doc = Document(output)
        E.element(doc.settings.element, "w:updateFields", val="true")
        doc.save(output)
        self.assertCode(check(output, intake_path=state), "AUTO_FIELD_UPDATE")

    def test_document_title_mapped_as_chapter_is_identified(self):
        output, state = self.build(paper=True)
        doc = Document(output)
        old = next(p.text for p in doc.paragraphs if p.style.name == "Heading 1")
        for text in doc.element.body.xpath(".//w:t"):
            if text.text == old:
                text.text = doc.core_properties.title
        doc.save(output)
        warnings = check(output, intake_path=state)["warnings"]
        self.assertIn("DOCUMENT_TITLE_IN_TOC", {item["code"] for item in warnings})
        numbered = [item["message"] for item in warnings if item["code"] == "HEADING_NUMBERING"]
        self.assertTrue(any(doc.core_properties.title in item for item in numbered))
