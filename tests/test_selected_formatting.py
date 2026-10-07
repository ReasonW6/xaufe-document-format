from copy import deepcopy
from docx import Document
from docx.shared import RGBColor
from test_layout_integrity import Fixture
import engine as E
from style_model import run_properties, paragraph_properties


class SelectedFormatting(Fixture):
    def test_target_heading_clears_inherited_blue_and_bold_without_restyling_others(self):
        doc = Document()
        doc.add_heading("Keep original title", 0)
        doc.add_paragraph("Body text")
        doc.add_heading("一、Target heading", 1)
        doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string("365F91")
        doc.styles["Heading 1"].font.bold = True
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Unchanged table"
        original_title = doc.paragraphs[0]._p.xml
        original_table = table._tbl.xml
        styles = {name: doc.styles[name].element.xml for name in ("Normal", "Title", "Heading 1")}
        signature = E.content_signature(doc)
        E.select_profile("B")
        E.style_selected(doc, {"p1": "body", "p2": "h1"})
        actual = run_properties(doc.paragraphs[2], doc.paragraphs[2].runs[0]._r)
        self.assertEqual(actual["color"], "000000")
        self.assertFalse(actual.get("b", False))
        self.assertEqual(actual["fonts"]["ascii"], "Times New Roman")
        self.assertEqual(actual["fonts"]["eastAsia"], "黑体")
        self.assertEqual(actual["sz"], "28")
        self.assertEqual(doc.paragraphs[0]._p.xml, original_title)
        self.assertEqual(table._tbl.xml, original_table)
        self.assertEqual({name: doc.styles[name].element.xml for name in styles}, styles)
        self.assertEqual(E.content_signature(doc), signature)

    def test_selected_formatting_respects_effective_user_overrides(self):
        doc = Document()
        p = doc.add_paragraph("User override")
        p.add_run("[1]").font.superscript = True
        E.select_profile(
            "B", {"styles": {"body": {"cn": "仿宋", "size": 12, "color": "112233"}}}
        )
        E.style_selected(doc, {"p0": "body"})
        actual = run_properties(p, p.runs[0]._r)
        self.assertEqual(actual["fonts"]["eastAsia"], "仿宋")
        self.assertEqual(actual["sz"], "24")
        self.assertEqual(actual["color"], "112233")
        self.assertTrue(p.runs[1].font.superscript)
