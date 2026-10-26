"""Offline regression tests. No network; no external renderer needed for these unit tests."""

from __future__ import annotations
import hashlib, json, re, sys, tempfile, unittest
from pathlib import Path
from copy import deepcopy
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import engine
import intake
from check_format import check


class FormatTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        engine.select_profile("A")
        self.data = {
            "mode": "document",
            "cover": True,
            "metadata": {"title_zh": "格式回归测试"},
            "blocks": [
                {"type": "heading", "level": 1, "text": "一、测试标题"},
                {"type": "paragraph", "text": "保留原文和数字123，不写新内容。"},
                {"type": "heading", "level": 2, "text": "（一）二级标题"},
                {
                    "type": "paragraph",
                    "runs": [
                        {"text": "引用"},
                        {"text": "[1]", "superscript": True},
                        {"text": "及公式x²。"},
                    ],
                },
            ],
        }

    def tearDown(self):
        self.tmp.cleanup()

    def session_for(self, meta, template="A", name="intake.json"):
        path = self.d / name
        intake.start(path, "独立回归测试夹具")
        intake.select(path, template, "测试夹具：明确选择 " + template)
        spec = engine.load_profile(template)
        fields = spec["cover"]["required_fields"]
        known = {
            k: {"value": v, "source": "document", "evidence": "测试样本文档明确字段"}
            for k, v in meta.items()
            if k in fields and v
        }
        blanks = [k for k in fields if k not in known and k != "date"]
        intake.resolve(
            path,
            {
                "metadata": known,
                "blank_fields": blanks,
                "blank_instruction": "测试夹具：明确选择未提供项留空。",
            },
        )
        return path

    def make(self, data=None, name="out.docx"):
        src = self.d / (name + ".json")
        src.write_text(json.dumps(data or self.data, ensure_ascii=False), encoding="utf-8")
        out = self.d / name
        self.session = self.session_for(
            (data or self.data).get("metadata", {}), name=name + ".intake.json"
        )
        engine.build(src, out, intake_path=self.session)
        return out

    def changed(self, fn):
        p = self.make()
        d = Document(p)
        fn(d)
        d.save(p)
        return check(p)

    def assertCode(self, result, code):
        self.assertIn(code, {e["code"] for e in result["errors"]})

    def test_build_passes_static(self):
        self.assertTrue(check(self.make())["passed"])

    def test_without_cover(self):
        data = deepcopy(self.data)
        data["cover"] = False
        p = self.make(data)
        res = check(p)
        self.assertTrue(res["passed"])
        self.assertFalse(res["stats"]["cover"])
        self.assertEqual(len(Document(p).inline_shapes), 0)

    def test_original_logo_hash(self):
        engine.asset_check()
        self.assertEqual(
            hashlib.sha256((ROOT / "assets/logos/green.png").read_bytes()).hexdigest(),
            engine.SPEC["logo"]["sha256"],
        )

    def test_footer_default_rejected(self):
        self.assertCode(
            self.changed(lambda d: d.sections[-1].footer.paragraphs[0].add_run("1")), "FOOTER"
        )

    def test_footer_first_rejected(self):
        self.assertCode(
            self.changed(lambda d: d.sections[-1].first_page_footer.paragraphs[0].add_run("1")),
            "FOOTER",
        )

    def test_footer_even_rejected(self):
        self.assertCode(
            self.changed(lambda d: d.sections[-1].even_page_footer.paragraphs[0].add_run("1")),
            "FOOTER",
        )

    def test_footer_page_field_rejected(self):
        def mutate(d):
            fld = OxmlElement("w:fldSimple")
            fld.set(qn("w:instr"), " PAGE ")
            d.sections[-1].footer.paragraphs[0]._p.append(fld)

        r = self.changed(mutate)
        self.assertCode(r, "FOOTER")
        self.assertCode(r, "FOOTER_FIELD")

    def test_bold_heading_rejected(self):
        def heading(d):
            next(p for p in d.paragraphs if p.style.name == "Heading 1").runs[0].bold = True

        self.assertCode(self.changed(heading), "FONT_BOLD")

    def test_bold_words_in_body_allowed(self):
        def word(d):
            body = next(p for p in d.paragraphs if p.style.name == "XAUFE Body")
            body.add_run("强调").bold = True

        self.assertTrue(self.changed(word)["passed"])

    def test_wholly_bold_body_paragraph_flagged(self):
        def whole(d):
            body = next(p for p in d.paragraphs if p.style.name == "XAUFE Body")
            body.add_run("这是一整段都加粗的合成正文，用来检查整段加粗的提示。")
            for r in body.runs:
                r.bold = True

        r = self.changed(whole)
        self.assertIn("FONT_BOLD_PARAGRAPH", {w["code"] for w in r["warnings"]})

    def test_caption_cut_off_from_its_table_by_a_blank_rejected(self):
        data = deepcopy(self.data)
        data["blocks"].append({"type": "table", "caption": "表1 测试表", "header": ["甲", "乙"], "rows": [["1", "2"]]})

        def mutate(d):
            caption = next(p for p in d.paragraphs if p.text == "表1 测试表")
            caption._p.addnext(OxmlElement("w:p"))

        p = self.make(data)
        self.assertTrue(check(p)["passed"])
        d = Document(p)
        mutate(d)
        d.save(p)
        r = check(p)
        self.assertCode(r, "CAPTION_DETACHED")

    def make_own(self, data, name="own.docx"):
        """Built with the user's own heading numbering kept, so the checker sees the levels as set."""
        src = self.d / (name + ".json")
        src.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        session = self.session_for(data.get("metadata", {}), name=name + ".intake.json")
        intake.customize(session, {"heading_numbering": "own"}, "用户确认原稿就是这样编号")
        out = self.d / name
        engine.build(src, out, intake_path=session)
        return out

    def test_same_numbering_at_two_heading_levels_rejected(self):
        data = deepcopy(self.data)
        data["blocks"].append({"type": "heading", "level": 2, "text": "二、错设成二级的标题"})
        with self.assertRaises(ValueError) as caught:  # renumbering would hide the wrong level, so it refuses
            self.make(data)
        self.assertIn("“二、错设成二级的标题”设成了 h2，但它的编号是模板一级标题的写法", str(caught.exception))
        r = check(self.make_own(data))
        self.assertCode(r, "HEADING_LEVEL")
        self.assertIn("二、错设成二级的标题", " ".join(e["message"] for e in r["errors"]))

    def test_heading_at_a_level_owned_by_another_numbering_rejected(self):
        data = deepcopy(self.data)
        data["blocks"] = [{"type": "heading", "level": 1, "text": "一、测试标题"},
                          {"type": "paragraph", "text": "保留原文。"},
                          {"type": "heading", "level": 1, "text": "（二）错设成一级的标题"}]
        with self.assertRaises(ValueError) as caught:
            self.make(data)
        self.assertIn("应为 h2", str(caught.exception))
        r = check(self.make_own(data))
        self.assertIn("另一种编号", " ".join(e["message"] for e in r["errors"] if e["code"] == "HEADING_LEVEL"))

    def test_other_numbering_schemes_are_renumbered_the_template_way(self):
        data = deepcopy(self.data)
        data["blocks"] = [{"type": "heading", "level": 1, "text": "1 绪论"},
                          {"type": "heading", "level": 2, "text": "1.1 研究背景"},
                          {"type": "heading", "level": 3, "text": "1.1.1 问题来源"},
                          {"type": "paragraph", "text": "保留原文。"},
                          {"type": "heading", "level": 1, "text": "研究方法"},
                          {"type": "heading", "level": 2, "text": "2.1 数据"},
                          {"type": "paragraph", "text": "保留原文。"},
                          {"type": "heading", "level": 1, "text": "致谢"}]
        src = self.d / "renumber.json"
        src.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        result = engine.build(src, self.d / "renumber.docx", intake_path=self.session_for(data.get("metadata", {}), name="r.intake.json"))
        heads = [p.text for p in Document(self.d / "renumber.docx").paragraphs if p.style.name.startswith("Heading")]
        self.assertEqual(heads[:6], ["一、绪论", "（一）研究背景", "1.问题来源", "二、研究方法", "（一）数据", "致谢"])
        self.assertEqual(len(result["auto_corrections"]), 5)
        self.assertNotIn("HEADING_LEVEL", {e["code"] for e in check(self.d / "renumber.docx")["errors"]})

    def test_swapped_heading_levels_rejected_unless_user_owns_the_numbering(self):
        data = deepcopy(self.data)
        data["blocks"] = [{"type": "heading", "level": 2, "text": "一、本该是一级的标题"},
                          {"type": "paragraph", "text": "保留原文。"},
                          {"type": "heading", "level": 1, "text": "（一）本该是二级的标题"},
                          {"type": "paragraph", "text": "保留原文。"}]
        with self.assertRaises(ValueError) as caught:
            self.make(data)
        self.assertIn("本该是一级", str(caught.exception))
        self.assertIn("本该是二级", str(caught.exception))
        src = self.d / "own.json"
        src.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        session = self.session_for(data.get("metadata", {}), name="own.intake.json")
        intake.customize(session, {"heading_numbering": "own"}, "用户确认原稿就是这样编号")
        out = self.d / "own.docx"
        engine.build(src, out, intake_path=session)
        self.assertNotIn("HEADING_LEVEL", {e["code"] for e in check(out)["errors"]})

    def test_lost_heading_in_body_rejected_but_numbered_list_allowed(self):
        data = deepcopy(self.data)
        data["blocks"] += [{"type": "heading", "level": 3, "text": "1.三级标题"},
                           {"type": "paragraph", "text": "2.数据来源"},
                           {"type": "paragraph", "text": "这一段是正文，说明数据从哪里来。"}]
        r = check(self.make(data))
        self.assertIn("2.数据来源", " ".join(e["message"] for e in r["errors"] if e["code"] == "HEADING_LEVEL"))
        listed = deepcopy(self.data)
        listed["blocks"] += [{"type": "heading", "level": 3, "text": "1.三级标题"},
                             {"type": "paragraph", "text": "2.准确性"}, {"type": "paragraph", "text": "3.完整性"},
                             {"type": "paragraph", "text": "以上两点是要求。"}]
        r = check(self.make(listed, name="list.docx"))
        self.assertNotIn("HEADING_LEVEL", {e["code"] for e in r["errors"]})

    def test_wrong_body_size_rejected(self):
        def mutate(d):
            next(p for p in d.paragraphs if p.style.name == "XAUFE Body").runs[0].font.size = Pt(12)

        self.assertCode(self.changed(mutate), "FONT_SIZE")

    def test_wrong_body_chinese_font_rejected(self):
        def mutate(d):
            r = next(p for p in d.paragraphs if p.style.name == "XAUFE Body").runs[0]
            engine.element(r._r.get_or_add_rPr(), "w:rFonts", eastAsia="Arial")

        self.assertCode(self.changed(mutate), "FONT_CN")

    def test_center_logo_rejected(self):
        self.assertCode(
            self.changed(lambda d: setattr(d.paragraphs[0], "alignment", 1)), "LOGO_ALIGN"
        )

    def test_enlarged_logo_rejected(self):
        self.assertCode(
            self.changed(lambda d: setattr(d.inline_shapes[0], "width", Cm(14))), "LOGO_SIZE"
        )

    def test_wrong_margin_rejected(self):
        self.assertCode(
            self.changed(lambda d: setattr(d.sections[-1], "left_margin", Cm(2.54))), "PAGE"
        )

    def test_heading_four_stops(self):
        data = deepcopy(self.data)
        data["blocks"][0]["level"] = 4
        with self.assertRaises(ValueError):
            self.make(data)
        self.assertFalse((self.d / "out.docx").exists())

    def test_missing_image_stops(self):
        data = deepcopy(self.data)
        data["blocks"].append({"type": "image", "path": "absent.png"})
        with self.assertRaises(ValueError):
            self.make(data)
        self.assertFalse((self.d / "out.docx").exists())

    def test_refuses_in_place(self):
        p = self.make()
        with self.assertRaises(ValueError):
            engine.safe_out(p, p, True)

    def test_paper_toc_requires_refresh(self):
        data = json.loads((ROOT / "examples/A/content.json").read_text())
        data["blocks"] = [b for b in data["blocks"] if b["type"] != "image"]
        p = self.make(data)
        self.assertTrue(check(p)["passed"])
        self.assertCode(check(p, final=True), "TOC_UNRESOLVED")

    def test_abstract_tagged_as_body_named_as_abstract(self):
        data = json.loads((ROOT / "examples/A/content.json").read_text())
        data["blocks"] = [b for b in data["blocks"] if b["type"] != "image"]
        p = self.make(data)
        d = Document(p)
        abstract = next(q for q in d.paragraphs if q.style.name == "XAUFE Abstract Chinese")
        abstract.style = d.styles["XAUFE Body"]
        d.save(p)
        mismatches = [e["message"] for e in check(p)["errors"] if e["code"] in ("ROLE_MISMATCH", "PART_ORDER")]
        self.assertTrue(mismatches)
        self.assertTrue(all("abstract_zh" in m and "关键词之间" in m for m in mismatches), mismatches)

    def test_master_contains_no_filled_personal_fields(self):
        d = Document(ROOT / "assets/A/format-master.docx")
        for i in engine.SPEC["cover"]["fields"].values():
            self.assertEqual(d.paragraphs[i].runs[1].text.strip(), "")
        self.assertEqual(d.core_properties.author, "")

    def source_map(self, complex_math=False):
        d = Document()
        d.add_paragraph("一、已有标题")
        p = d.add_paragraph("原文保留，含链接")
        h = OxmlElement("w:hyperlink")
        h.set(qn("w:anchor"), "test")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "链接文字"
        r.append(t)
        h.append(r)
        p._p.append(h)
        d.add_paragraph("关键词：原词甲   原词乙")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "列甲"
        t.cell(0, 1).text = "列乙"
        t.cell(1, 0).text = "值甲"
        t.cell(1, 1).text = "值乙"
        p = d.add_paragraph()
        p.add_run().add_picture(str(ROOT / "examples/layout-test.png"), width=Cm(5))
        p = d.add_paragraph()
        m = OxmlElement("m:oMath")
        r = OxmlElement("m:r")
        t = OxmlElement("m:t")
        t.text = "x+1"
        r.append(t)
        m.append(r)
        p._p.append(m)
        src = self.d / "source.docx"
        d.save(src)
        inv = engine.inventory(src)
        mapping = {
            "input_sha256": inv["input_sha256"],
            "mode": "document",
            "has_cover": False,
            "metadata": {"title_zh": "保真测试"},
            "format_tables": True,
            "roles": {
                "p0": "h1",
                "p1": "body",
                "p2": "keywords_zh",
                "p3": "figure",
                "p4": "equation",
            },
        }
        mp = self.d / "map.json"
        mp.write_text(json.dumps(mapping, ensure_ascii=False), encoding="utf-8")
        self.session = self.session_for(mapping["metadata"])
        return src, mp

    def test_restyle_preserves_text_math_tables_images_fields(self):
        src, mp = self.source_map()
        before = engine.content_signature(Document(src))
        out = self.d / "restyled.docx"
        engine.restyle(src, mp, out, intake_path=self.session)
        after = engine.content_signature(Document(out))
        after["text"] = re.sub(r"（\d+）", "", after["text"])  # the template numbers each equation
        self.assertEqual(before, after)
        self.assertTrue(check(out)["passed"])

    def test_stale_map_stops(self):
        src, mp = self.source_map()
        d = Document(src)
        d.add_paragraph("改过")
        d.save(src)
        with self.assertRaises(ValueError):
            engine.restyle(src, mp, self.d / "result.docx", intake_path=self.session)

    def test_unmapped_paragraph_stops(self):
        src, mp = self.source_map()
        m = json.loads(mp.read_text())
        del m["roles"]["p1"]
        mp.write_text(json.dumps(m))
        with self.assertRaises(ValueError):
            engine.restyle(src, mp, self.d / "result.docx", intake_path=self.session)

    def test_tracked_change_stops(self):
        src, mp = self.source_map()
        d = Document(src)
        ins = OxmlElement("w:ins")
        d.paragraphs[1]._p.append(ins)
        d.save(src)
        m = json.loads(mp.read_text())
        m["input_sha256"] = engine.inventory(src)["input_sha256"]
        mp.write_text(json.dumps(m))
        with self.assertRaises(ValueError):
            engine.restyle(src, mp, self.d / "result.docx", intake_path=self.session)

    def test_top_level_content_control_stops(self):
        src, mp = self.source_map()
        d = Document(src)
        d.element.body.insert(0, OxmlElement("w:sdt"))
        d.save(src)
        m = json.loads(mp.read_text())
        m["input_sha256"] = engine.inventory(src)["input_sha256"]
        mp.write_text(json.dumps(m))
        with self.assertRaises(ValueError):
            engine.restyle(src, mp, self.d / "result.docx", intake_path=self.session)

    def test_figure_not_exact_line_height(self):
        d = Document()
        engine.install_styles(d)
        e = d.styles["XAUFE Figure"].element.pPr.find(qn("w:spacing"))
        self.assertEqual(e.get(qn("w:lineRule")), "auto")
        self.assertEqual(int(e.get(qn("w:before"))), 0)  # the blank line above the figure is the gap

    def test_figure_and_table_groups_get_a_blank_line_each_side(self):
        data = deepcopy(self.data)
        data["blocks"] = [{"type": "heading", "level": 1, "text": "一、图表"},
                          {"type": "paragraph", "text": "合成数据见表1。"},
                          {"type": "table", "caption": "表1 合成表", "unit": "单位：元", "note": "注：合成数据。",
                           "header": ["编号", "金额"], "rows": [["1", "30"]]},
                          {"type": "paragraph", "text": "由表1可见，金额不变。"}]
        out = self.make(data)
        texts = [p.text for p in Document(out).paragraphs]
        at = texts.index("表1 合成表")
        self.assertEqual(texts[at - 1], "")
        self.assertEqual(texts[at + 1], "元")
        self.assertEqual(texts[texts.index("合成数据。") + 1], "")
        r = check(out)
        self.assertNotIn("OBJECT_SPACING", {e["code"] for e in r["errors"]})
        self.assertFalse({"OBJECT_NOT_CITED", "OBJECT_NOT_DISCUSSED"} & {w["code"] for w in r["warnings"]})
        doc = Document(out)
        blank = next(p for p in doc.paragraphs if p.text == "表1 合成表")._p.getprevious()
        blank.getparent().remove(blank)
        doc.save(out)
        self.assertIn("OBJECT_SPACING", {e["code"] for e in check(out)["errors"]})

    def test_equations_are_numbered_in_order_at_the_right_margin(self):
        data = deepcopy(self.data)
        data["blocks"] = [{"type": "heading", "level": 1, "text": "一、公式"},
                          {"type": "paragraph", "text": "质能关系如下："},
                          {"type": "equation", "text": "E=mc2 (3)"},
                          {"type": "equation", "text": "F=ma"}]
        out = self.make(data)
        eqs = [p for p in Document(out).paragraphs if p.style.name == "XAUFE Equation"]
        self.assertEqual([p.text for p in eqs], ["\tE=mc2\t（1）", "\tF=ma\t（2）"])
        tabs = Document(out).styles["XAUFE Equation"].element.xpath("./w:pPr/w:tabs/w:tab")
        self.assertEqual({(t.get(qn("w:val")), t.get(qn("w:leader"))) for t in tabs}, {("center", "none"), ("right", "none")})
        self.assertFalse({"EQUATION_NUMBER", "EQUATION_LAYOUT"} & {e["code"] for e in check(out)["errors"]})
        doc = Document(out)
        doc.paragraphs[[p.text for p in doc.paragraphs].index("\tF=ma\t（2）")].runs[-1].text = "（5）"
        doc.save(out)
        self.assertIn("EQUATION_NUMBER", {e["code"] for e in check(out)["errors"]})

    def test_uncited_and_out_of_order_objects_are_told(self):
        data = deepcopy(self.data)
        data["blocks"] = [{"type": "heading", "level": 1, "text": "一、图表"},
                          {"type": "paragraph", "text": "这里没有提到任何表。"},
                          {"type": "table", "caption": "表2 合成表", "header": ["编号"], "rows": [["1"]]}]
        codes = {w["code"] for w in check(self.make(data))["warnings"]}
        self.assertTrue({"OBJECT_NOT_CITED", "OBJECT_NOT_DISCUSSED", "OBJECT_ORDER"} <= codes)

    def test_wrong_header_size_rejected(self):
        def mutate(d):
            d.sections[-1].header.paragraphs[0].runs[0].font.size = Pt(14)

        self.assertCode(self.changed(mutate), "FONT_SIZE")

    def test_wrong_cover_field_size_rejected(self):
        data = deepcopy(self.data)
        data["metadata"]["student_name"] = "测试姓名"
        p = self.make(data)
        d = Document(p)
        d.paragraphs[9].runs[1].font.size = Pt(12)
        d.save(p)
        self.assertCode(check(p), "COVER_FONT")

    def test_tampered_master_stops(self):
        (self.d / "assets/A").mkdir(parents=True)
        (self.d / "assets/logos").mkdir(parents=True)
        (self.d / "assets/logos/green.png").write_bytes(
            (ROOT / "assets/logos/green.png").read_bytes()
        )
        (self.d / "assets/A/format-master.docx").write_bytes(b"tampered")
        with patch.object(engine, "ROOT", self.d):
            with self.assertRaises(ValueError):
                engine.asset_check()

    def test_missing_original_logo_stops(self):
        with patch.object(engine, "ROOT", self.d):
            with self.assertRaises(ValueError):
                engine.asset_check()

    def test_add_cover_to_existing_keeps_original_text(self):
        src, mp = self.source_map()
        m = json.loads(mp.read_text())
        m["replace_cover_before"] = 0
        mp.write_text(json.dumps(m, ensure_ascii=False))
        before = engine.content_signature(Document(src))
        out = self.d / "covered.docx"
        engine.restyle(src, mp, out, intake_path=self.session)
        d = Document(out)
        self.assertTrue(re.sub(r"（\d+）", "", engine.content_signature(d)["text"]).endswith(before["text"]))
        self.assertTrue(check(out)["passed"])



class ColumnWidths(unittest.TestCase):
    def test_long_text_column_gets_the_room_and_total_is_exact(self):
        import engine

        rows = [["编号", "说明", "数值"]] + [[str(i), "一段较长的说明文字，需要更宽的列才不至于每行一个字。", "3.5"] for i in range(5)]
        widths = engine.column_widths(rows, 8958)
        self.assertEqual(sum(widths), 8958)
        self.assertGreater(widths[1], 2 * max(widths[0], widths[2]))
        self.assertTrue(all(w - 216 >= 210 for w in widths))

    def test_explicit_ratios_and_short_tables(self):
        import engine

        self.assertEqual(engine.column_widths([["a", "b", "c"]], 9000, [1, 1, 2]), [2250, 2250, 4500])
        even = engine.column_widths([["甲", "乙"], ["1", "2"]], 8958)
        self.assertEqual(sum(even), 8958)
        self.assertLessEqual(abs(even[0] - even[1]), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
