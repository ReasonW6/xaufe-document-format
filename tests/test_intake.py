"""Synthetic regression cases for the two-template conversation workflow.
No fixture represents a real user's consent or personal information.
"""

from __future__ import annotations
import json, sys, tempfile, unittest, subprocess
from pathlib import Path
from datetime import datetime, timezone
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import engine as E
import intake as I
from check_format import check


def candidate(v, source="user", **extra):
    return {
        "value": v,
        "source": source,
        "evidence": "独立测试夹具中的明确字段，不是实际用户资料。",
        **extra,
    }


class IntakeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.p = self.d / "intake.json"

    def tearDown(self):
        self.tmp.cleanup()

    def start(self, template=None):
        I.start(self.p, "测试任务：格式检查")
        if template:
            I.select(self.p, template, "测试夹具中的明确选择：" + template)

    def complete(self, template="A", **override):
        self.start(template)
        meta = {
            k: candidate("测试题目" if k == "title_zh" else "测试字段")
            for k in E.load_profile(template)["cover"]["required_fields"]
            if k != "date"
        }
        meta.update(override)
        return I.resolve(self.p, {"metadata": meta})

    def test_start_waits_for_template(self):
        self.start()
        self.assertEqual(I.read(self.p)["stage"], "awaiting_template")
        self.assertIsNone(I.read(self.p)["template"])

    def test_collect_before_selection_rejected(self):
        self.start()
        with self.assertRaises(ValueError):
            I.resolve(self.p, {})

    def test_missing_choice_quote_rejected(self):
        self.start()
        with self.assertRaises(ValueError):
            I.select(self.p, "B", "")

    def test_unknown_template_rejected(self):
        self.start()
        with self.assertRaises(ValueError):
            I.select(self.p, "default", "试图选择")

    def test_no_reselect_in_ready_state(self):
        self.complete()
        with self.assertRaises(ValueError):
            I.select(self.p, "B", "更换模板")

    def test_only_unknown_fields_are_missing(self):
        self.start("A")
        s = I.resolve(
            self.p,
            {
                "metadata": {
                    "title_zh": candidate("测试题目"),
                    "student_name": candidate("测试姓名"),
                }
            },
        )
        self.assertNotIn("title_zh", s["missing"])
        self.assertNotIn("student_name", s["missing"])
        self.assertNotIn("date", s["missing"])
        self.assertIn("student_id", s["missing"])

    def test_full_information_is_ready_without_another_confirmation(self):
        self.assertEqual(self.complete()["stage"], "ready")

    def test_red_never_requests_supervisor(self):
        self.start("B")
        s = I.resolve(self.p, {})
        self.assertIn("course_name", s["missing"])
        self.assertNotIn("supervisor", s["missing"])
        self.assertNotIn("date", s["missing"])

    def test_red_rejects_extra_supervisor(self):
        self.start("B")
        with self.assertRaises(ValueError):
            I.resolve(self.p, {"metadata": {"supervisor": candidate("测试教师")}})

    def test_blank_needs_actual_instruction(self):
        self.start("A")
        with self.assertRaises(ValueError):
            I.resolve(self.p, {"blank_all": True})

    def test_all_blank_includes_date(self):
        self.start("B")
        s = I.resolve(
            self.p, {"blank_all": True, "blank_instruction": "测试场景：全部留空，我自己填。"}
        )
        self.assertEqual(s["stage"], "ready")
        self.assertIn("date", s["blank_fields"])
        self.assertNotIn("date", s["metadata"])

    def test_explicit_blank_wins_known_cover_value(self):
        self.complete()
        s = I.resolve(
            self.p, {"blank_fields": ["student_name"], "blank_instruction": "测试场景：姓名留空。"}
        )
        self.assertNotEqual(s["metadata"]["student_name"], "")
        self.assertEqual(E.cover_metadata(s["metadata"], s)["student_name"], "")

    def test_empty_string_does_not_count_as_known(self):
        self.start("A")
        with self.assertRaises(ValueError):
            I.resolve(self.p, {"metadata": {"student_name": candidate("")}})

    def test_template_example_is_not_a_valid_source(self):
        self.start("A")
        with self.assertRaises(ValueError):
            I.resolve(self.p, {"metadata": {"student_name": candidate("示例", "template")}})

    def test_memory_requires_same_target(self):
        self.start("A")
        with self.assertRaises(ValueError):
            I.resolve(self.p, {"metadata": {"student_name": candidate("旧姓名", "memory")}})

    def test_confirmed_memory_avoids_reasking(self):
        self.start("A")
        s = I.resolve(
            self.p,
            {
                "metadata": {
                    "student_name": candidate("已核实名字", "memory", target_confirmed=True)
                }
            },
        )
        self.assertNotIn("student_name", s["missing"])

    def test_explicit_user_value_overrides_document_and_memory(self):
        self.start("A")
        s = I.resolve(
            self.p,
            {
                "metadata": {
                    "student_name": [
                        candidate("记忆名", "memory", target_confirmed=True),
                        candidate("稿件名", "document"),
                        candidate("本次更正名", "user"),
                    ]
                }
            },
        )
        self.assertEqual(s["metadata"]["student_name"], "本次更正名")

    def test_same_rank_conflict_is_not_guessed(self):
        self.start("A")
        s = I.resolve(
            self.p,
            {
                "metadata": {
                    "student_name": [candidate("甲", "document"), candidate("乙", "document")]
                }
            },
        )
        self.assertEqual(s["conflicts"]["student_name"], ["乙", "甲"])
        self.assertNotEqual(s["stage"], "ready")

    def test_default_date_uses_user_timezone_not_utc_day(self):
        self.start("A")
        now = datetime(2032, 1, 1, 18, 0, tzinfo=timezone.utc)
        s = I.resolve(self.p, {}, now=now)
        self.assertEqual(s["metadata"]["date"], "2032-01-02")

    def test_supplied_date_normalized_and_not_today(self):
        self.start("A")
        s = I.resolve(self.p, {"metadata": {"date": candidate("2024年2月29日")}})
        self.assertEqual(s["metadata"]["date"], "2024-02-29")
        self.assertEqual(s["provenance"]["date"]["source"], "user")

    def test_document_date_preserved(self):
        self.start("B")
        s = I.resolve(self.p, {"metadata": {"date": candidate("2024/02/29", "document")}})
        self.assertEqual(s["metadata"]["date"], "2024-02-29")

    def test_invalid_date_rejected(self):
        self.start("B")
        with self.assertRaises(ValueError):
            I.resolve(self.p, {"metadata": {"date": candidate("2025年2月29日")}})

    def test_missing_year_is_not_guessed(self):
        self.start("B")
        with self.assertRaises(ValueError):
            I.resolve(self.p, {"metadata": {"date": candidate("9月10日")}})

    def test_default_date_refreshes_after_midnight(self):
        self.complete()
        s = I.require(self.p, now=datetime(2032, 1, 2, 18, tzinfo=timezone.utc))
        self.assertEqual(s["metadata"]["date"], "2032-01-03")

    def test_user_date_does_not_refresh_after_midnight(self):
        self.complete(date=candidate("2024-02-29"))
        s = I.require(self.p, now=datetime(2032, 1, 2, 18, tzinfo=timezone.utc))
        self.assertEqual(s["metadata"]["date"], "2024-02-29")

    def test_state_tampering_is_rejected(self):
        self.start()
        s = I.load(self.p)
        s["stage"] = "ready"
        self.p.write_text(json.dumps(s))
        with self.assertRaises(ValueError):
            I.require(self.p)

    def test_new_task_does_not_reuse_old_selection(self):
        self.complete()
        p2 = self.d / "task2.json"
        I.start(p2, "另一个目标文档")
        self.assertEqual(I.read(p2)["stage"], "awaiting_template")

    def test_incomplete_cover_stops_write_gate(self):
        self.start("A")
        I.resolve(self.p, {})
        with self.assertRaises(ValueError):
            I.require(self.p)

    def test_readonly_inspection_can_follow_selection(self):
        self.start("A")
        self.assertEqual(I.require(self.p, ready=False)["stage"], "awaiting_metadata")


class ProfileTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def prepare(self, template="B", mode="document", blanks=None):
        state = self.d / (template + ".intake.json")
        I.start(state, "模板专属回归测试")
        I.select(state, template, "测试夹具：明确选择 " + template)
        spec = E.load_profile(template)
        vals = {
            "title_zh": "双模板测试题目",
            "student_name": "测试姓名",
            "student_id": "0012345678",
            "major": "测试专业",
            "class_name": "测试班级",
            "supervisor": "测试教师",
            "course_name": "测试课程",
        }
        known = {
            k: candidate(v, "document")
            for k, v in vals.items()
            if k in spec["cover"]["required_fields"]
        }
        ans = {"metadata": known}
        if blanks:
            ans.update({"blank_fields": blanks, "blank_instruction": "测试场景：指定字段留空。"})
        I.resolve(state, ans)
        data = {
            "mode": mode,
            "cover": True,
            "metadata": {},
            "blocks": [
                {"type": "heading", "level": 1, "text": "一、正文"},
                {"type": "paragraph", "text": "这是专供自动化测试使用的格式样本，不是研究结果。"},
            ],
        }
        if mode == "paper":
            data.update(
                {
                    "abstract_zh": "中文格式测试摘要。",
                    "keywords_zh": ["格式", "测试"],
                    "references": ["格式测试条目，不是真实参考文献。"],
                }
            )
            if template == "A":
                data.update(
                    {"abstract_en": "Formatting specimen only.", "keywords_en": ["format", "test"]}
                )
                data["metadata"]["title_en"] = "Formatting Specimen"
        return state, data

    def build(self, template="B", mode="document", blanks=None):
        state, data = self.prepare(template, mode, blanks)
        src = self.d / (template + ".json")
        src.write_text(json.dumps(data, ensure_ascii=False))
        out = self.d / (template + ".docx")
        E.build(src, out, intake_path=state)
        return out, state

    def mutate(self, fn, template="B", mode="document"):
        out, state = self.build(template, mode)
        d = Document(out)
        fn(d)
        d.save(out)
        return check(out, intake_path=state)

    def test_build_without_intake_stops_before_output(self):
        with self.assertRaises(ValueError):
            E.build(self.d / "missing.json", self.d / "out.docx")
        self.assertFalse((self.d / "out.docx").exists())

    def test_restyle_without_intake_stops(self):
        with self.assertRaises(ValueError):
            E.restyle(self.d / "missing.docx", self.d / "map.json", self.d / "out.docx")

    def test_red_document_passes(self):
        out, state = self.build()
        self.assertTrue(check(out, final=True, intake_path=state)["passed"])

    def test_green_document_passes(self):
        out, state = self.build("A")
        self.assertTrue(check(out, final=True, intake_path=state)["passed"])

    def test_red_complete_order_and_no_english_abstract(self):
        out, state = self.build("B", "paper")
        d = Document(out)
        names = [p.style.name for p in d.paragraphs]
        self.assertLess(names.index("XAUFE Chinese Title"), names.index("XAUFE TOC Title"))
        self.assertNotIn("XAUFE English Title", names)
        self.assertTrue(check(out, intake_path=state)["passed"])

    def test_green_complete_order_with_english(self):
        out, state = self.build("A", "paper")
        names = [p.style.name for p in Document(out).paragraphs]
        self.assertLess(names.index("XAUFE TOC Title"), names.index("XAUFE Chinese Title"))
        self.assertIn("XAUFE English Title", names)

    def test_red_rejects_supplied_english_instead_of_dropping_it(self):
        state, data = self.prepare("B", "paper")
        data["abstract_en"] = "Must not disappear."
        src = self.d / "in.json"
        src.write_text(json.dumps(data))
        out = self.d / "out.docx"
        with self.assertRaises(ValueError):
            E.build(src, out, intake_path=state)
        self.assertFalse(out.exists())

    def test_red_body_width_is_different_from_green(self):
        out, _ = self.build("B")
        s = Document(out).sections[-1]
        self.assertEqual(s.left_margin.twips, 1701)
        self.assertEqual(s.right_margin.twips, 1247)

    def test_red_all_headers_and_footers_empty(self):
        out, _ = self.build()
        for s in Document(out).sections:
            for part in [
                s.header,
                s.first_page_header,
                s.even_page_header,
                s.footer,
                s.first_page_footer,
                s.even_page_footer,
            ]:
                self.assertFalse(part._element.xpath(".//w:t|.//w:pBdr|.//w:instrText"))

    def test_red_unexpected_header_is_rejected(self):
        r = self.mutate(lambda d: d.sections[-1].header.paragraphs[0].add_run("错误页眉"))
        self.assertIn("RED_HEADER", [e["code"] for e in r["errors"]])

    def test_red_unexpected_header_border_is_rejected(self):
        def change(d):
            pp = d.sections[-1].header.paragraphs[0]._p.get_or_add_pPr()
            E.element(E.element(pp, "w:pBdr"), "w:bottom", val="single")

        r = self.mutate(change)
        self.assertIn("RED_HEADER", [e["code"] for e in r["errors"]])

    def test_red_has_no_supervisor_field(self):
        out, _ = self.build()
        self.assertNotIn("指导教师", "".join(p.text for p in Document(out).paragraphs[:15]))

    def test_red_has_runtime_date_not_source_example(self):
        out, state = self.build()
        date = I.read(state)["metadata"]["date"]
        y, m, d = map(int, date.split("-"))
        self.assertIn(f"{y}年{m}月{d}日", Document(out).paragraphs[14].text)

    def test_blank_date_is_respected(self):
        out, state = self.build(blanks=["date"])
        self.assertNotRegex(Document(out).paragraphs[14].text, r"\d")
        self.assertTrue(check(out, final=True, intake_path=state)["passed"])

    def test_cover_blank_does_not_delete_abstract_title(self):
        out, state = self.build("B", "paper", blanks=["title_zh"])
        d = Document(out)
        self.assertEqual(d.paragraphs[5].runs[1].text.strip(), "")
        self.assertIn(
            "双模板测试题目",
            [p.text for p in d.paragraphs if p.style.name == "XAUFE Chinese Title"],
        )

    def test_student_id_keeps_leading_zero(self):
        out, _ = self.build()
        self.assertIn("0012345678", Document(out).paragraphs[10].text)

    def test_supplied_metadata_cannot_override_confirmed_value(self):
        state, data = self.prepare()
        data["metadata"]["student_name"] = "未经确认的新名字"
        src = self.d / "in.json"
        src.write_text(json.dumps(data))
        with self.assertRaises(ValueError):
            E.build(src, self.d / "out.docx", intake_path=state)

    def test_abstract_uses_annotated_five_not_sample_four(self):
        out, _ = self.build("B", "paper")
        d = Document(out)
        st = d.styles["XAUFE Abstract Chinese"]
        self.assertEqual(st.font.size.pt, 10.5)

    def test_both_abstract_and_reference_have_one_point_five_spacing(self):
        for id in ["A", "B"]:
            E.select_profile(id)
            d = Document()
            E.install_styles(d)
            for name in ["XAUFE Abstract Chinese", "XAUFE Abstract English", "XAUFE Reference"]:
                e = d.styles[name].element.pPr.find(qn("w:spacing"))
                self.assertEqual(e.get(qn("w:line")), "360")
                self.assertEqual(e.get(qn("w:lineRule")), "auto")
            e = d.styles["XAUFE Body"].element.pPr.find(qn("w:spacing"))
            self.assertEqual(e.get(qn("w:line")), "400")
            self.assertEqual(e.get(qn("w:lineRule")), "exact")

    def test_red_wrong_abstract_size_rejected(self):
        def change(d):
            next(p for p in d.paragraphs if p.style.name == "XAUFE Abstract Chinese").runs[
                0
            ].font.size = Pt(14)

        r = self.mutate(change, mode="paper")
        self.assertIn("FONT_SIZE", [e["code"] for e in r["errors"]])

    def test_red_wrong_frontmatter_order_rejected(self):
        def change(d):
            p = next(p for p in d.paragraphs if p.style.name == "XAUFE TOC Title")._p
            p.getparent().remove(p)
            d.element.body.insert(15, p)

        r = self.mutate(change, mode="paper")
        self.assertIn("PART_ORDER", [e["code"] for e in r["errors"]])

    def test_wrong_template_logo_rejected(self):
        def change(d):
            rid, _ = d.part.get_or_add_image(str(ROOT / "assets/logos/green.png"))
            d.paragraphs[0]._p.xpath(".//a:blip")[0].set(qn("r:embed"), rid)

        r = self.mutate(change)
        self.assertIn("LOGO", [e["code"] for e in r["errors"]])

    def test_changed_cover_information_rejected(self):
        r = self.mutate(lambda d: setattr(d.paragraphs[9].runs[1], "text", "\t错误姓名\t"))
        self.assertIn("COVER_VALUE", [e["code"] for e in r["errors"]])

    def test_injected_english_abstract_role_rejected(self):
        r = self.mutate(lambda d: d.add_paragraph("Abstract", E.ROLES["abstract_label_en"]))
        self.assertIn("RED_ENGLISH_ABSTRACT", [e["code"] for e in r["errors"]])

    def test_final_audit_requires_intake(self):
        out, _ = self.build()
        self.assertIn("INTAKE_REQUIRED", [e["code"] for e in check(out, final=True)["errors"]])

    def test_blank_title_line_has_anchor_for_renderers(self):
        d = Document(ROOT / "assets/B/format-master.docx")
        self.assertEqual(d.paragraphs[5].runs[-1].text, "\u200b")

    def test_red_cover_fields_follow_song_font_annotation(self):
        d = Document(ROOT / "assets/B/format-master.docx")
        for idx in [9, 10, 11, 12]:
            r = d.paragraphs[idx].runs[1]
            self.assertEqual(r._r.rPr.find(qn("w:rFonts")).get(qn("w:eastAsia")), "宋体")
            self.assertEqual(r.font.size.pt, 16)

    def test_red_builds_table_without_builtin_style(self):
        state, data = self.prepare("B", "document")
        data["blocks"].append(
            {"type": "table", "header": ["列一", "列二"], "rows": [["值一", "值二"]]}
        )
        src = self.d / "table.json"
        src.write_text(json.dumps(data, ensure_ascii=False))
        out = self.d / "table.docx"
        E.build(src, out, intake_path=state)
        self.assertEqual(len(Document(out).tables), 1)
        self.assertTrue(check(out, final=True, intake_path=state)["passed"])

    def test_red_restyle_with_cover_preserves_existing_body_objects(self):
        state, _ = self.prepare("B", "document")
        d = Document()
        d.add_paragraph("一、原稿标题")
        d.add_paragraph("已有正文及数字0123保持不变。")
        table = d.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "原表甲"
        table.cell(0, 1).text = "原表乙"
        d.add_paragraph().add_run().add_picture(
            str(ROOT / "examples/layout-test.png"), width=E.Cm(5)
        )
        src = self.d / "source.docx"
        d.save(src)
        before = E.content_signature(d)
        mapping = {
            "input_sha256": E.inventory(src)["input_sha256"],
            "mode": "document",
            "metadata": {},
            "replace_cover_before": 0,
            "format_tables": True,
            "roles": {"p0": "h1", "p1": "body", "p2": "figure"},
        }
        mp = self.d / "mapping.json"
        mp.write_text(json.dumps(mapping, ensure_ascii=False))
        out = self.d / "restyled.docx"
        E.restyle(src, mp, out, intake_path=state)
        after = E.content_signature(Document(out))
        self.assertTrue(after["text"].endswith(before["text"]))
        self.assertEqual(after["tables"], before["tables"])
        self.assertEqual(after["images"], before["images"] + 1)
        self.assertTrue(check(out, final=True, intake_path=state)["passed"])

    def test_toc_update_also_requires_intake(self):
        from refresh_toc import update

        with self.assertRaises(ValueError):
            update(self.d / "missing.docx", self.d / "out.docx")

    def test_wrong_green_header_text_is_rejected(self):
        r = self.mutate(
            lambda d: setattr(
                d.sections[-1].header.paragraphs[0].runs[0], "text", "其他专业\t其他题目"
            ),
            template="A",
        )
        self.assertIn("HEADER_TEXT", [e["code"] for e in r["errors"]])


if __name__ == "__main__":
    unittest.main(verbosity=2)
