"""Independent static and real-PDF verification of the first-page bottom date."""

from __future__ import annotations
import json, re
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from style_model import paragraph_properties, on

ROOT = Path(__file__).resolve().parents[1]


class CoverOverflow(ValueError):
    """The rendered cover does not fit on page 1 (date moved, squeezed or split)."""


def baseline(template, request=None):
    data = json.loads((ROOT / "assets/checks/cover-date-baseline.json").read_text(encoding="utf-8"))
    result = {**data["common"], **data["templates"][template]}
    if request and request.get("page_twips"):
        from customization import effective_profile

        m = effective_profile(template, request)["page_twips"]
        result.update(
            width_twips=m["width"] - m["left"] - m["right"],
            height_twips=m["height"],
            left_twips=m["left"],
            bottom_twips=m["bottom"],
        )
    result["frame"] = {**data["common"]["frame"], "w": str(result["width_twips"])}
    return result


def static(doc, spec, err):
    """Check exact anchoring plus effective inherited paragraph properties."""
    rule = baseline(spec["id"], spec.get("_format_request"))
    index = rule["paragraph_index"]
    if len(doc.paragraphs) <= index:
        err("COVER_DATE_POSITION", "封面日期段落缺失。")
        return
    p = doc.paragraphs[index]
    actual = p._p.xpath("./w:pPr/w:framePr")
    expected = {qn("w:" + key): str(value) for key, value in rule["frame"].items()}
    if len(actual) != 1 or dict(actual[0].attrib) != expected:
        err(
            "COVER_DATE_POSITION",
            "完成日期必须单独锚定在封面版心底部；不得随信息栏流动、移到页边外或改成普通段落。",
        )
    for tag, want in rule["paragraph"].items():
        got = paragraph_properties(p, tag)
        if got != want:
            err("COVER_DATE_LAYOUT", f"日期的{tag}有效属性与独立基线不符：{got}。")
    inherited_break = paragraph_properties(p, "pageBreakBefore")
    if p._p.xpath("./w:pPr/w:pageBreakBefore") or (
        inherited_break and on(inherited_break.get("val"))
    ):
        err(
            "COVER_DATE_LAYOUT",
            "日期段落不得包含直接段前分页属性（包括显式0），也不得继承启用的分页。",
        )
    if p._p.xpath('./w:pPr/w:sectPr|.//w:br[@w:type="page"]'):
        err("COVER_DATE_BREAK", "日期段落不能携带分页或分节；须在第一页封面内。")
    for i, other in enumerate(doc.paragraphs):
        if i != index and paragraph_properties(other, "framePr"):
            err("COVER_FRAME", "本通用路径只允许封面日期段落使用定位框，其他浮动段落须专门处理。")


def norm(text):
    return re.sub(r"[\s\u200b\ufeff]+", "", text)


def visible_lines(page):
    """Exclude invisible edge spacing while retaining spaces within the text."""
    lines = []
    for block in page.get_text("rawdict")["blocks"]:
        for line in block.get("lines", []):
            chars = [c for span in line["spans"] for c in span["chars"]]
            visible = [i for i, c in enumerate(chars) if norm(c["c"])]
            if not visible:
                continue
            chars = chars[visible[0] : visible[-1] + 1]
            boxes = [c["bbox"] for c in chars]
            lines.append(
                {
                    "spans": [{"text": "".join(c["c"] for c in chars)}],
                    "bbox": (
                        min(b[0] for b in boxes),
                        min(b[1] for b in boxes),
                        max(b[2] for b in boxes),
                        max(b[3] for b in boxes),
                    ),
                }
            )
    return lines


def verify_pdf(docx, pdf_path):
    """Read rendered coordinates, not a cached assertion. Raises on cover overflow."""
    import pymupdf as fitz

    doc = Document(docx)
    vv = {
        e.get(qn("w:name")): e.get(qn("w:val"))
        for e in doc.settings.element.xpath("./w:docVars/w:docVar")
    }
    if vv.get("XAUFE_HAS_COVER") != "1":
        return {"required": False}
    template = vv.get("XAUFE_TEMPLATE")
    if template not in ("A", "B"):
        raise ValueError("无法确定封面日期的模板类型。")
    rule = baseline(template, json.loads(vv.get("XAUFE_FORMAT_REQUEST", "{}")))
    index = rule["paragraph_index"]
    if len(doc.paragraphs) <= index:
        raise ValueError("封面日期段落缺失。")
    expected = norm(doc.paragraphs[index].text)
    if not expected.startswith("完成日期："):
        raise ValueError("日期字段不是本模板字段。")
    with fitz.open(pdf_path) as pdf:
        if not len(pdf):
            raise ValueError("没有渲染出封面。")
        page = pdf[0]
        # Word may emit invisible edge spaces in a different font. Measure actual
        # date characters, preserving internal spacing, instead of the line box.
        lines = visible_lines(page)
        matches = [
            line for line in lines if norm("".join(s["text"] for s in line["spans"])) == expected
        ]
        if len(matches) != 1:
            raise CoverOverflow("完成日期未在第一页完整显示为独立单行，或同页出现重复日期。")
        box = matches[0]["bbox"]
        bottom = (rule["height_twips"] - rule["bottom_twips"]) / 20
        top = bottom - rule["frame_height_twips"] / 20
        tolerance = rule["render_tolerance_pt"]
        if box[1] < top - tolerance or box[3] > bottom + tolerance:
            raise CoverOverflow(
                f"完成日期没有落在封面底部保留区：实际y={box[1]:.2f}至{box[3]:.2f}磅，目标区{top:.2f}至{bottom:.2f}磅。"
            )
        left = rule["left_twips"] / 20
        right = left + rule["width_twips"] / 20
        if (
            box[0] < left - tolerance
            or box[2] > right + tolerance
            or abs((box[0] + box[2] - left - right) / 2) > tolerance
        ):
            raise ValueError("封面日期未在版心内居中，或超出左右边界。")
        first_text = norm(page.get_text())
        for p in doc.paragraphs[:index]:
            value = norm(p.text)
            if value and value not in first_text:
                raise CoverOverflow(
                    "封面信息未完整保留在第一页；不得挤到下一页或被日期遮住：" + value[:40]
                )
        other_bottom = max(
            (
                line["bbox"][3]
                for line in lines
                if line is not matches[0] and norm("".join(s["text"] for s in line["spans"]))
            ),
            default=0,
        )
        gap = box[1] - other_bottom
        if gap < rule["minimum_gap_pt"]:
            raise CoverOverflow(
                f"封面信息与底部日期相距仅{gap:.2f}磅，空间不足；请确认长字段布局，不能遮挡或缩字。"
            )
        # Cover is an independent first section. Reject a blank spillover cover page.
        if vv.get("XAUFE_MODE") != "template":
            following = next(
                (norm(p.text) for p in doc.paragraphs[rule["paragraph_count"] :] if norm(p.text)),
                None,
            )
            if following and (len(pdf) < 2 or following not in norm(pdf[1].get_text())):
                raise CoverOverflow(
                    "封面后首个正文/前置标题没有出现在第二页；须检查多余空白页或封面溢出。"
                )
        return {
            "required": True,
            "passed": True,
            "template": template,
            "physical_page": 1,
            "date_bbox_pt": list(box),
            "reserved_y_pt": [top, bottom],
            "distance_from_paper_bottom_pt": page.rect.height - box[3],
            "gap_to_information_pt": gap,
            "all_cover_fields_on_first_page": True,
            "following_content_starts_on_page_2": vv.get("XAUFE_MODE") != "template",
            "method": "independent rendered PDF text coordinates; not Word/WPS pixel parity",
        }
