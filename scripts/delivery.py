"""Task workspace, real pagination/preview, review gate and single-DOCX delivery.

The review is the agent's honest written record of what it checked; the script
binds it to the exact document and page images but cannot prove anyone looked.
Only the owned workspace is deleted; sources and unrelated files stay untouched.
"""

from __future__ import annotations
import hashlib, json, os, re, shutil, tempfile, uuid
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from textio import read_text_any
import engine as E
from intake import require
from check_format import check
from render_preview import render
from refresh_toc import update
from audit_helpers import toc_entries
from toc_format import enforce
from office_backend import OfficeUnavailable, OfficeExportFailed
import render_availability as RF
from feedback import AcceptanceFailed, explain, from_render, warning_key, _input_hint
from render_audit import audit as audit_rendered

MAGIC = "xaufe-owned-document-job-v2"
MARKER = ".xaufe-session.json"
REVIEW_TEXTS = {
    "content_check": "需求与内容保全",
    "layout_check": "格式与版面",
    "delivery_check": "最终交付",
}
PLACEHOLDER = "请填写"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(read_text_any(path))


def write_json(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def create(destination, *, parent=None, sources=()):
    dest = Path(destination).expanduser().resolve()
    if dest.suffix.lower() != ".docx":
        raise ValueError("最终交付文件必须是 .docx。")
    if dest.exists():
        raise ValueError("目标文件已存在；换一个文件名，不要覆盖用户已有文件。")
    if dest.is_relative_to(E.ROOT):
        raise ValueError("交付路径不能放在技能目录里。")
    originals = []
    for value in sources:
        path = Path(value).expanduser().resolve()
        if not path.is_file():
            raise ValueError("原始材料不是可读取的文件：" + str(path))
        if path.suffix.lower() in (".txt", ".md"):
            read_text_any(path)  # unreadable text must stop here, not after a full layout run
        if str(path) not in [x["path"] for x in originals]:
            originals.append({"path": str(path), "sha256": sha(path), "role": "original"})
    if os.name == "nt":
        # A plain directory inherits the parent's ACL, so an approved desktop-user
        # export can read it (mkdtemp would create a sandbox-private folder).
        folder = Path(parent) if parent is not None else Path(tempfile.gettempdir())
        root = (folder / ("xaufe-job-" + uuid.uuid4().hex[:16])).resolve()
        root.mkdir()
    else:
        root = Path(tempfile.mkdtemp(prefix="xaufe-job-", dir=parent)).resolve()
    state = {
        "magic": MAGIC,
        "workdir": str(root),
        "job_token": str(uuid.uuid4()),
        "destination": str(dest),
        "stage": "working",
        "source_documents": originals,
    }
    write_json(root / MARKER, state)
    return {"workspace": str(root), "destination": str(dest)}


def load(root):
    raw = Path(root).expanduser()
    if raw.is_symlink():
        raise ValueError("工作区不能是符号链接。")
    root = raw.resolve()
    marker = root / MARKER
    if (
        not root.name.startswith("xaufe-job-")
        or not root.is_dir()
        or not marker.is_file()
        or marker.is_symlink()
    ):
        raise ValueError("不是本工具创建的工作区；只能使用 xaufe.py new 返回的 WORK 路径。")
    state = read_json(marker)
    if (
        state.get("magic") != MAGIC
        or state.get("workdir") != str(root)
        or not state.get("job_token")
    ):
        raise ValueError("工作区所有权标记不匹配。")
    if root == E.ROOT or E.ROOT.is_relative_to(root) or root.is_relative_to(E.ROOT):
        raise ValueError("禁止把技能目录或其父目录当作工作区。")
    return root, state


def _job(root):
    """job.json as the agent wrote it; an unreadable one stops delivery instead of skipping checks."""
    path = Path(root) / "job.json"
    if not path.is_file():
        return {}
    try:
        job = read_json(path)
    except ValueError as exc:
        raise ValueError(f"job.json 读不出来（{exc}），请用 UTF-8 重新保存后重新 run。") from exc
    if not isinstance(job, dict):
        raise ValueError("job.json 必须是一个 JSON 对象。")
    return job


def previous_manual(path):
    """The hand edit an earlier output of this skill went through (adopt), or None."""
    if Path(path).suffix.lower() != ".docx":
        return None
    try:
        raw = E.get_vars(Document(path)).get("XAUFE_MANUAL")
        return json.loads(raw) if raw else None
    except Exception:
        return None


def _highlighted(path):
    """Texts of the runs an original .docx highlights or shades (the template has neither)."""
    if Path(path).suffix.lower() != ".docx":
        return []
    try:
        body = Document(path).element.body
    except Exception:
        return []
    found = []
    for p in body.iter(qn("w:p")):
        text = "".join(t for r in p.iter(qn("w:r")) if _painted(r) for t in r.xpath("./w:t/text()")).strip()
        if text:
            found.append(text)
    return found


def _painted(run):
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        return False
    mark = rpr.find(qn("w:highlight"))
    if mark is not None and mark.get(qn("w:val")) not in (None, "none"):
        return True
    shade = rpr.find(qn("w:shd"))
    return shade is not None and (shade.get(qn("w:fill")) or "auto").upper() not in ("AUTO", "FFFFFF")

def correction_line(c):
    """One correction as the user reads it: what changed and why."""
    part = E.CORRECTION_PARTS.get(c.get("part"), "文字")
    before, after = (c.get("before") or "").strip(), (c.get("after") or "").strip()
    short = lambda t: t[:24] + ("…" if len(t) > 24 else "")
    if before and before == after:  # moved, not reworded
        return f"{part}“{short(before)}”" + (c.get("why") or "换了位置")
    if before and after:
        what = f"{part}“{short(before)}”改为“{short(after)}”"
    elif after:
        what = f"补了{part}“{short(after)}”"
    elif c.get("part") == "cover":
        what = f"旧封面（“{short(before)}”）"
    else:
        what = f"删去{part}“{short(before)}”"
    return what + (f"（{c['why']}）" if c.get("why") else "")


def compare_originals(root, state, result):
    """Format-only (task "format"): any text added or dropped without the user's say blocks delivery.
    Writing: the originals are material, so text not reused is only a review note."""
    originals = [x for x in state["source_documents"] if x.get("role") == "original"]
    for item in originals:
        manual = previous_manual(item["path"])
        if manual:
            result["warnings"].append({"code": "PREVIOUS_MANUAL_EDIT", "message":
                f"原稿是本技能以前的成品，当时有一项是交付前手工处理的：“{manual.get('note', '')}”"
                + ("（" + "；".join(manual.get("summary") or []) + "）" if manual.get("summary") else "")
                + "。这次是按记录的设置重新排的，那项手工处理不会自动带上：写得进 job.json 的就写进去重新 run，"
                  "写不进的就在新成品上重新手工处理后 adopt；用户这次不要了，就在这里写明用户的原话。"})
    job = _job(root)
    strict = job.get("task") == "format"
    corrections = state.get("corrections") or []
    if corrections:
        result["warnings"].append({"code": "TEXT_CORRECTED", "message": "按学校版式要求改动了这些文字："
                                   + "；".join(f"{n}. {correction_line(c)}" for n, c in enumerate(corrections, 1))
                                   + "。要在回复里逐条告诉用户，用户不同意的就从 job.json 的 corrections 里去掉后重新 run。"})
    if not originals:
        if strict and job:
            raise AcceptanceFailed("原文保全", [{
                "code": "SOURCE_REQUIRED", "stage": "原文保全",
                "message": "只排版（task 为 format）却没有登记原稿，没法证明原文一个字没改。",
                "cause": "new 的时候没有用 --source 登记原稿，也没有用 inspect 读原稿。",
                "fix": "用 new --out 成品.docx --source 原稿文件 重新开始（已有 Word 用 inspect 也会自动登记）；"
                "用户是在对话里直接贴的文字，就先把原文原样存成 .txt 再用 --source 登记。",
            }])
        return
    from source_compare import compare, dropped_margins, dropped_page_numbers

    dropped = dropped_margins(originals, root / "final.docx")
    if dropped:
        result["warnings"].append({
            "code": "ORIGINAL_HEADER_FOOTER_REPLACED",
            "message": "原稿页眉/页脚里的“" + "”“".join(t[:20] for t in dropped[:3]) + "”没有保留：成品的页眉页脚按模板重新生成了。"
                       "要在回复里告诉用户；用户想保留原来的页眉文字，就写进 job.json 的 format.header（left / right）后重新 run。",
        })
    marked = [t for item in originals for t in _highlighted(item["path"])] if strict else []
    if marked:
        result["warnings"].append({
            "code": "ORIGINAL_HIGHLIGHT_REMOVED",
            "message": "原稿里“" + "”“".join(t[:15] for t in marked[:3]) + f"”等 {len(marked)} 处文字带高亮或底纹，"
                       "模板没有这种格式，成品里去掉了（文字没动）。要在回复里告诉用户：那可能是用户给自己做的待改标记。",
        })
    if dropped_page_numbers(originals, root / "final.docx"):
        result["warnings"].append({
            "code": "ORIGINAL_PAGE_NUMBERS_DROPPED",
            "message": "原稿页脚有页码，成品按模板默认不显示页码。要在回复里告诉用户；用户想要页码，就在 job.json 的 format 里写 "
                       "\"footer_page_numbers\": true 后重新 run。",
        })
    mapping = read_json(root / "mapping.json") if (root / "mapping.json").is_file() else {}
    toc_line = re.compile(r"^\s*目\s*录\s*$|\d+\s*$")
    old_toc = {r["id"] for r in mapping.get("paragraphs", []) if r.get("role") == "old_toc" and toc_line.search(r.get("text", ""))}
    # Cover values from job.json also appear on the abstract page and in the header.
    generated = list(require(root / "intake.json")["metadata"].values()) if (root / "intake.json").is_file() else []
    for value in (mapping.get("add") or {}).values():
        if isinstance(value, list):  # keywords: the template prints them after its own label
            generated += ["Key words: " + "; ".join(map(str, value)), "关键词：" + "；".join(map(str, value))]
        else:
            generated.append(value)
    approved = [c["text"] for c in job.get("approved_changes", []) if isinstance(c, dict) and c.get("text")]
    approved += [t for c in corrections for t in (c.get("before"), c.get("after")) if t]
    report = compare(originals, root / "final.docx", approved=approved, generated=[str(g) for g in generated],
                     replaced=mapping.get("replace_cover_before"), old_toc=old_toc)
    write_json(root / "source-comparison.json", report)
    found = report["differences"]
    if strict and found:
        fix = ("只排版不能改原文：已有 Word 时看 mapping.json 里这段的 role（delete 会删掉整段）；content.json 时照原稿逐字改回，"
               "数字、标点、大小写、段落顺序都要一样。用户明确同意的增删，把改动的完整文字写进 job.json 的 approved_changes"
               "（{\"text\": \"完整的那段文字\", \"why\": \"用户怎么说的\"}；图片、表格等对象写明种类和数量，如“删去图片 1 张”）。")
        what = {"changed": ("SOURCE_CHANGED", "原稿这段文字在成品里被改了。"),
                "missing": ("SOURCE_MISSING", "原稿这段文字在成品里找不到。"),
                "added": ("SOURCE_ADDED", "成品里这段文字在原稿里没有。"),
                "moved": ("SOURCE_MOVED", "原稿这段文字在成品里换了位置。"),
                "objects": ("SOURCE_OBJECTS", "成品里的图片/表格/公式/链接和原稿对不上。")}
        issues = []
        for d in found:
            code, message = what[d["kind"]]
            item = {"code": code, "stage": "原文保全", "message": message, "where": {"location": d["location"]},
                    "cause": "排版时漏掉、改写、挪动或新加了内容。", "fix": fix}
            if d["kind"] == "objects":
                item["fix"] = ("把原稿里的对象原样带上：content.json 里图片用 image 块、表格用 table 块（行列和原稿一样）、"
                               "链接写在 runs 的 link 里；已有 Word 时看 mapping.json 里图片、公式所在段落是不是被设成了 delete。"
                               "用户同意去掉的，在 job.json 的 approved_changes 里写明种类和数量（如 {\"text\": \"删去图片 1 张\", \"why\": \"用户的原话\"}）。")
            source = None
            for key in ("output_paragraph", "output", "original"):
                source = source or _input_hint(root, d.get(key) or "")
            if not source and d["kind"] == "missing" and d.get("previous"):
                before = _input_hint(root, d["previous"])
                source = before and f"补在 {before} 的后面"
            if source:
                item["where"]["edit"] = source
            if "original" in d:
                item["expected"] = d["original"]
                item["actual"] = "（成品里没有）"
            if "output" in d:
                item["actual"] = d["output"]
                item.setdefault("expected", "（原稿里没有）")
            issues.append(item)
        raise AcceptanceFailed("原文保全", issues)
    missing = [d for d in found if d["kind"] == "missing"]
    if missing:
        result["warnings"].append(
            {
                "code": "SOURCE_DIFFERENCES",
                "message": f"原稿里有 {len(missing)} 段文字没有用进成品（写作模式下原稿是素材，这不算错）。打开 WORK/source-comparison.json 核对："
                "用户要求保留的要补回；在 review.json 的 warnings 里写明怎么处理的。",
            }
        )


def check_preview_toc(root):
    """Bind page caches to the PDF the model actually sees, including backend fallback."""
    from refresh_toc import pdf_positions

    doc = Document(root / "final.docx")
    entries = toc_entries(doc)
    if not entries:
        return {"required": False}
    values, measurement, pdf_hash = pdf_positions(root / "preview/final.pdf", entries)
    if [entry["cache"] for entry in entries] != list(map(str, values)):
        raise ValueError("最终预览与目录页码不一致；可能切换了渲染软件。请重新运行 run。")
    proof = read_json(root / "final.toc-report.json")
    previous = proof["passes"][-1]
    if (
        measurement["physical_pages"] != previous["physical_pages"]
        or measurement["total_pages"] != previous["total_pages"]
    ):
        raise ValueError("预览版面与目录定位结果不一致，请重新运行 run。")
    return {
        "required": True,
        "entries": len(entries),
        "preview_pdf_sha256": pdf_hash,
        **measurement,
    }


def verify_originals(state):
    for source in state["source_documents"]:
        if source.get("role") != "original":
            continue
        path = Path(source["path"])
        if not path.is_file() or sha(path) != source["sha256"]:
            raise ValueError(
                "登记的原稿在处理期间被改动：" + str(path) + "。请确认原稿后用 new 重新开始。"
            )


def _reset_for_prepare(root, s, docx, intake_path):
    verify_originals(s)
    if docx == Path(s["destination"]):
        raise ValueError("原稿不能同时作为最终交付目标。")
    state = require(intake_path)
    doc = Document(docx)
    vv = E.get_vars(doc)
    if (
        vv.get("XAUFE_TEMPLATE") != state["template"]
        or vv.get("XAUFE_INTAKE_ID") != state["job_id"]
    ):
        raise ValueError("文档与本次任务记录不匹配；请重新运行 run。")
    if json.loads(vv.get("XAUFE_FORMAT_REQUEST", "{}")) != state.get("format_request", {}):
        raise ValueError("用户特别要求已变化，请重新运行 run。")
    local_intake = root / "active-intake.json"
    if intake_path != local_intake:
        shutil.copy2(intake_path, local_intake)
    s["active_request_source"] = str(intake_path)
    s["stage"] = "preparing"
    (root / "generated.docx").unlink(missing_ok=True)
    for key in (
        "manual",
        "review",
        "validation_scope",
        "limited_validation_sha256",
        "skipped_steps",
        "user_notice",
        "pages",
    ):
        s.pop(key, None)
    write_json(root / MARKER, s)
    E.select_profile(state["template"], state.get("format_request"))
    RF.clear_scope(doc)
    enforce(doc, E.SPEC)
    normalized = root / "toc-normalized.docx"
    doc.save(normalized)
    before = check(normalized, intake_path=local_intake)
    write_json(root / "initial-check.json", before)
    if not before["passed"]:
        # Usually a wrong role in mapping.json or a missing part in content.json.
        raise AcceptanceFailed("格式初查", explain(before["errors"], normalized, root))
    return doc, normalized, local_intake


def _prepare(root, s, docx, intake_path, unavailable_note=None, office="auto"):
    doc, normalized, local_intake = _reset_for_prepare(root, s, docx, intake_path)
    if office == "none":
        write_json(root / "office-probe.json", {"skipped": "user_request"})
        evidence = {
            "version": "1.0.0",
            "probe_sha256": sha(root / "office-probe.json"),
            "reason": "按用户要求，本次没有启动办公软件。",
            "status": "skipped_by_user_request",
        }
        write_json(root / "office-unavailable.json", evidence)
        return _prepare_static_only(root, s, normalized, local_intake, evidence)
    from office_discovery import discover

    # The real export below is the capability probe; no separate launch first.
    probe = discover()
    write_json(root / "office-probe.json", probe)
    final = root / "final.docx"
    try:
        has_toc = bool(toc_entries(doc))
        if has_toc:
            update(normalized, final, force=True, intake_path=local_intake)
        else:
            shutil.copy2(normalized, final)
        result = check(final, final=True, intake_path=local_intake)
        write_json(root / "final-check.json", result)
        if not result["passed"]:
            raise AcceptanceFailed("格式终查", explain(result["errors"], final, root))
        preview = render(final, root / "preview")
        write_json(root / "preview-toc-check.json", check_preview_toc(root))
        rendered = audit_rendered(final, root / "preview" / "final.pdf")
        write_json(root / "render-audit.json", rendered)
        if not rendered["passed"]:
            raise AcceptanceFailed("渲染后版面检查", from_render(rendered, root, final))
    except OfficeUnavailable as exc:
        evidence = _renderer_unavailable(root, exc, probe, unavailable_note)
        return _prepare_static_only(root, s, normalized, local_intake, evidence)
    if sha(final) != result["document_sha256"]:
        raise ValueError("预览时成品发生变化，请重新运行 run。")
    compare_originals(root, s, result)
    result["warnings"].extend(rendered["warnings"])
    wanted, used = s.get("office_engine", "auto"), _renderer_name(root)
    if wanted != "auto" and used and wanted != used.lower():
        names = dict(RENDERER_NAMES, word="Word")
        result["warnings"].append({"code": "RENDERER_FALLBACK", "message":
            f"指定用 {names.get(wanted, wanted)} 渲染，但它这次用不了，实际是用 {names.get(used.lower(), used)} 分页和检查的。"
            "用户点名要某个软件时要告诉用户；能修好那个软件就修好后重新 run。"})
    difference = renderer_difference(final, root / "preview" / "final.pdf", root / "second-render.pdf", used)
    if difference:
        result["warnings"].append(difference)
    write_json(root / "final-check.json", result)
    warnings = [w for w in result["warnings"] if w["code"] != "VISUAL_REQUIRED"]
    s.update(
        validation_scope="full",
        stage="awaiting_review",
        document_sha256=sha(final),
        render_report_sha256=sha(root / "preview/render-report.json"),
        render_audit_sha256=sha(root / "render-audit.json"),
        renderer=_renderer_name(root),
        has_toc=has_toc,
        pages=preview["pages"],
        warning_codes=sorted({warning_key(w) for w in warnings}),
        tell_user=tell_user(warnings),
    )
    write_json(root / MARKER, s)
    return {
        "workspace": str(root),
        "validation_scope": "full",
        "pages": [dict(p, path=str(root / "preview" / p["png"])) for p in preview["pages"]],
        "warnings": warnings,
    }


# Warnings about something the user will see differently from what they gave or asked for:
# finish puts them in the notice that must go into the reply, so they cannot be dropped.
TELL_USER = {"RENDERER_DIFFERENCE", "ORIGINAL_PAGE_NUMBERS_DROPPED", "ORIGINAL_HEADER_FOOTER_REPLACED", "ORIGINAL_HIGHLIGHT_REMOVED", "HEADING_NUMBERING", "MARGIN_TOO_SMALL",
             "STYLE_HARD_TO_READ", "FONT_MARKS", "RENDERER_FALLBACK", "REF_CITATION_MISMATCH", "REF_SUPPORT_UNCONFIRMED",
             "CAPTION_POSITION", "PAPER_STRUCTURE_SKIPPED", "TEXT_CORRECTED",
             "PART_LEFT_EMPTY", "TITLE_TOO_LONG", "ABSTRACT_LENGTH", "OBJECT_NOT_CITED", "OBJECT_NOT_DISCUSSED",
             "OBJECT_ORDER", "OBSOLETE_UNIT"}


def _user_line(w):
    """A warning as one plain sentence for the user: what happened to their document, no instructions
    meant for the agent, no internal codes, nothing cut off mid-word."""
    code = w["code"]
    message = re.sub(r"^[pt]\d+\S*:\s*", "", w.get("message", ""))
    quoted = re.findall(r"“([^”]+)”", message)
    first = quoted[0] if quoted else ""
    if code == "ORIGINAL_HEADER_FOOTER_REPLACED":
        return "原稿页眉页脚里的“" + "”“".join(quoted[:3]) + "”没有保留，页眉页脚按模板重新生成了；想保留原来的页眉文字可以告诉我。"
    if code == "RENDERER_DIFFERENCE":
        return message.split("。")[0] + "。用那个软件打开时以它的分页为准，目录页码在目录上右键“更新域”就能对上。"
    if code == "ORIGINAL_HIGHLIGHT_REMOVED":
        count = re.search(r"等 (\d+) 处", message)
        return (f"原稿里“{first}”等{count.group(1) + ' 处' if count else ''}文字带的高亮或底纹模板里没有，成品里去掉了（文字没动）；"
                "如果那是你给自己做的待改标记，记得在新文件里再找一下。")
    if code == "ORIGINAL_PAGE_NUMBERS_DROPPED":
        return "原稿页脚有页码，这次按模板默认没有显示页码；需要页码可以告诉我。"
    if code == "HEADING_NUMBERING":
        return f"标题“{first}”的编号和模板的写法（一级“一、”、二级“（一）”、三级“1.”）不一样，按原稿保留了。"
    if code in ("PART_LEFT_EMPTY", "TITLE_TOO_LONG", "ABSTRACT_LENGTH"):
        return message.split("用户说过保持原样")[0] + "按你的意思保留了原样。"
    if code in ("OBJECT_NOT_CITED", "OBJECT_NOT_DISCUSSED", "OBJECT_ORDER", "OBSOLETE_UNIT"):
        return message + "这是行文上的事，我没有替你改字，建议补上。"
    if code == "TEXT_CORRECTED":
        return message.split("。要在回复里")[0] + "。不同意哪一处告诉我，我改回去。"
    if code == "PAPER_STRUCTURE_SKIPPED":
        return "这次按普通文档排的：没有生成目录，摘要也没有单独成页（模板的完整论文结构里有这些）；需要的话可以告诉我。"
    if code == "CAPTION_POSITION":
        return message + "按原稿位置保留了，没有替你挪；需要的话可以告诉我。"
    if code == "MARGIN_TOO_SMALL":
        size = re.search(r"只有 ([\d.]+) 厘米", message)
        return f"按要求设的页边距只有 {size.group(1) if size else '不到 0.5'} 厘米，多数打印机印不到这么靠边，边上的文字可能被裁掉。"
    if code == "STYLE_HARD_TO_READ":
        return "按你的要求设置后，" + re.sub(r"（[a-z_0-9]+）", "", message.split("确认是用户要的")[0]).strip()
    if code == "FONT_MARKS":
        marks = re.search(r"带(.+?)，是原稿", message)
        return f"原稿里“{first}”等文字带有{marks.group(1) if marks else '特殊格式'}，模板没有这种格式，成品按原稿保留了。"
    if code == "RENDERER_FALLBACK":
        names = re.search(r"指定用 (.+?) 渲染.*实际是用 (.+?) 分页", message)
        if names:
            return f"指定的 {names.group(1)} 这次用不了，版面是用 {names.group(2)} 分页和检查的，和 {names.group(1)} 里的分页可能略有不同。"
    if code == "REF_CITATION_MISMATCH":
        numbers = sorted(set(re.findall(r"[\[［](\d+)[\]］]", message)), key=int)
        which = f"（第 {'、'.join(numbers)} 条）" if numbers else ""
        return f"有几处引用号所在的句子和所引文献{which}的主题看起来不太一致，我逐条核对过，建议你再看一眼。"
    if code == "REF_SUPPORT_UNCONFIRMED":
        import refcheck as R

        found = re.findall(r"[\[［](\d+)[\]］](?:（文献核对：([^）]+)）)?\s*引用句", message)
        real = [n for n, status in found if status == R.PASS]
        other = [n for n, status in found if status != R.PASS]
        parts = []
        if real:
            parts.append(f"第 {'、'.join(real)} 条文献确认真实存在，但它是否支持正文里引用它的那句话，没能在能读到的原文里确认")
        if other:
            parts.append(f"第 {'、'.join(other)} 条文献既没能核对是否真实存在，也没能确认它支持正文里引用它的那句话")
        return "；".join(parts or ["有几条文献是否支持正文里引用它的那句话没能确认"]) + "（核实清单里有标注），建议你核对。"
    return re.split(r"要在回复里告诉用户|。逐条核对|用户点名|确认是用户要的", message)[0].strip()


def tell_user(warnings):
    """One plain sentence per kind of warning the user must be told about (first of each kind)."""
    seen = {}
    for w in warnings:
        if w.get("code") in TELL_USER and w["code"] not in seen:
            seen[w["code"]] = _user_line(w)
    return list(seen.values())


def add_review_warnings(workspace, warnings):
    """Warnings found outside the layout checks (e.g. reference support) that review.json must answer too."""
    root, s = load(workspace)
    s["tell_user"] = list(dict.fromkeys(s.get("tell_user", []) + tell_user(warnings)))
    s["warning_codes"] = sorted(set(s.get("warning_codes", [])) | {warning_key(w) for w in warnings})
    write_json(root / MARKER, s)


def renderer_difference(final, pdf, second, used):
    """Pagination checked in LibreOffice or WPS while Word (or another engine) is installed too: render
    again there and compare page count and TOC page numbers.  A warning when they differ, else None."""
    import pymupdf as fitz
    from office_backend import second_opinion
    from refresh_toc import pdf_positions

    if not used or used.lower() == "word" or os.environ.get("XAUFE_SECOND_RENDER") == "0":
        return None
    other = second_opinion(final, second, used)
    if not other:
        return None
    names = dict(RENDERER_NAMES, word="Word")
    first_name, second_name = names.get(used.lower(), used), names.get(other["engine"], other["engine"])
    with fitz.open(pdf) as a, fitz.open(second) as b:
        counts = (len(a), len(b))
    found = [f"{first_name} 排出 {counts[0]} 页，{second_name} 排出 {counts[1]} 页"] if counts[0] != counts[1] else []
    entries = toc_entries(Document(final))
    if entries:
        try:
            values = list(map(str, pdf_positions(second, entries)[0]))
        except ValueError:
            values = []
        moved = [(e["text"], e["cache"], v) for e, v in zip(entries, values) if e["cache"] != v]
        found += [f"目录“{t[:15]}”在 {first_name} 里是第 {a} 页，在 {second_name} 里是第 {b} 页" for t, a, b in moved[:3]]
    if not found:
        return None
    return {"code": "RENDERER_DIFFERENCE", "message":
            f"同一份文件，{second_name} 和检查用的 {first_name} 分页不一样：" + "；".join(found) + "。"
            f"用户用 {second_name} 打开时以它为准；要在回复里告诉用户，目录页码在 {second_name} 里右键“更新域”即可。"}


def _renderer_name(root):
    try:
        backend = read_json(root / "preview" / "office-backend.json")
    except (OSError, ValueError):
        return ""
    return str(backend.get("engine") or backend.get("backend") or "")


def _renderer_unavailable(root, exc, probe, note):
    """Only real unavailability of the environment can skip rendering, never bad layout."""
    from office_backend import export_pdf

    installed = [c for c in probe.get("candidates", []) if c.get("found")]
    smoke_result = None
    if isinstance(exc, OfficeExportFailed):
        smoke = root / "renderer-smoke.docx"
        test = Document()
        test.add_paragraph("Renderer availability check.")
        test.save(smoke)
        try:
            export_pdf(smoke, root / "renderer-smoke.pdf")
        except OfficeUnavailable as smoke_exc:
            smoke_result = {
                "status": "failed",
                "error": str(smoke_exc),
                "attempts": smoke_exc.failures,
            }
        else:
            raise AcceptanceFailed("导出 PDF", [{
                "code": "EXPORT_FAILED", "stage": "导出 PDF",
                "message": "办公软件能导出一个测试文档，却导不出这份稿件。",
                "actual": str(exc)[:400],
                "cause": "问题在这份稿件本身（常见：超大或损坏的图片、嵌入对象、极长的表格或公式），不是缺软件。",
                "fix": "按上面的错误找出那一处：先去掉或换掉可疑的图片/对象（content.json 的对应块，或 mapping.json 里那一段）后重新 run；"
                       "找不到原因就停下来告诉用户哪一步失败了，不要改走静态交付。",
            }])
    elif any(c.get("automation") for c in installed):
        raise ValueError("还有可自动调用的办公软件没试过；请直接重新运行 run。")
    if installed and (not isinstance(note, str) or not note.strip()):
        raise ValueError(
            "RENDER_ALTERNATIVE_REQUIRED：电脑上找到了办公软件，但自动调用失败。"
            "如果你的运行环境支持经用户批准后在本机会话里执行命令，请用那种方式重跑同一条 run；"
            '否则用 run --office-note "失败的实际原因" 继续，进入静态交付（会告诉用户哪些检查没做）。'
            "不要绕过权限，也不要让用户另装软件。"
        )
    if installed:
        reason = (
            "已找到办公软件，但当前环境无法完成实际导出：" + note.strip().rstrip("。.;；") + "。"
        )
    else:
        reason = "当前运行环境里没有找到可用的 Word、WPS 或 LibreOffice（不代表用户电脑上没有）。"
    evidence = {
        "version": "1.0.0",
        "probe_sha256": sha(root / "office-probe.json"),
        "probe": probe,
        "export_discovery": exc.report,
        "export_attempts": exc.failures,
        "smoke_test": smoke_result,
        "alternative_note": note or "",
        "reason": reason,
        "status": "unavailable_in_current_runtime",
    }
    write_json(root / "office-unavailable.json", evidence)
    return evidence


def _prepare_static_only(root, s, normalized, intake, evidence):
    # Always start from the statically checked source, not a partly rendered output.
    doc = Document(normalized)
    preserved = RF.page_neutral_signature(doc)
    count = RF.mark_pending(doc)
    if RF.page_neutral_signature(doc) != preserved:
        raise ValueError("标记待更新页码时内容保全检查失败。")
    final = root / "final.docx"
    doc.save(final)
    for name in ("final.toc-report.json", "preview-toc-check.json", "limited-validation.json"):
        path = root / name
        if path.is_file():
            path.unlink()
    preview = root / "preview"
    if preview.exists():
        if preview.is_symlink():
            raise ValueError("预览目录不能是符号链接。")
        shutil.rmtree(preview)
    result = check(final, final=True, intake_path=intake, validation_scope="static-only")
    compare_originals(root, s, result)
    write_json(root / "final-check.json", result)
    if not result["passed"]:
        raise AcceptanceFailed("静态格式检查", explain(result["errors"], final, root))
    receipt = RF.create_receipt(root, final, evidence, count)
    warnings = [
        w
        for w in result["warnings"]
        # These are already part of the user notice for a static delivery.
        if w["code"] not in ("VISUAL_REQUIRED", "RENDER_NOT_VERIFIED", "TOC_UNRESOLVED")
    ]
    s.update(
        stage="awaiting_review",
        validation_scope="static-only",
        document_sha256=sha(final),
        render_report_sha256=None,
        pages=[],
        limited_validation_sha256=sha(root / "limited-validation.json"),
        warning_codes=sorted({warning_key(w) for w in warnings}),
        tell_user=tell_user(warnings),
        skipped_steps=receipt["skipped_steps"],
        user_notice=receipt["user_notice"],
    )
    write_json(root / MARKER, s)
    return {
        "workspace": str(root),
        "validation_scope": "static-only",
        "pages": [],
        "warnings": warnings,
        "skipped_steps": receipt["skipped_steps"],
        "user_notice": receipt["user_notice"],
    }


def configure_office(workspace, engine="auto", wps_mode="auto"):
    root, s = load(workspace)
    if engine not in ("auto", "word", "libreoffice", "wps"):
        raise ValueError("未知办公软件后端。")
    if wps_mode not in ("auto", "manual") or wps_mode == "manual" and engine != "wps":
        raise ValueError("手动界面导出只用于 WPS。")
    s.update(office_engine=engine, wps_mode=wps_mode, stage="working")
    s.pop("review", None)
    write_json(root / MARKER, s)
    return {"workspace": str(root), "engine": engine, "wps_mode": wps_mode}


def prepare(workspace, docx, intake_path, unavailable_note=None, office="auto"):
    root, s = load(workspace)
    docx = Path(docx).resolve()
    intake_path = Path(intake_path).resolve()
    keys = ("XAUFE_RENDER_ENGINE", "XAUFE_WPS_MODE", "XAUFE_WPS_EXCHANGE")
    saved = {key: os.environ.get(key) for key in keys}
    try:
        if "office_engine" in s:
            os.environ["XAUFE_RENDER_ENGINE"] = s["office_engine"]
        if "wps_mode" in s:
            os.environ["XAUFE_WPS_MODE"] = s["wps_mode"]
        if s.get("wps_mode") == "manual":
            os.environ["XAUFE_WPS_EXCHANGE"] = str(root / "wps-exchange")
        try:
            return _prepare(root, s, docx, intake_path, unavailable_note, office)
        except Exception as exc:
            from wps_exchange import ManualExportRequired

            if not isinstance(exc, ManualExportRequired) or not unavailable_note:
                raise
            # A note about an unusable GUI is not a skip switch: try every
            # callable automatic engine before entering the static-only branch.
            os.environ["XAUFE_WPS_MODE"] = "auto"
            os.environ["XAUFE_RENDER_ENGINE"] = "auto"
            root, s = load(workspace)
            return _prepare(root, s, docx, intake_path, unavailable_note, office)
    finally:
        for key, val in saved.items():
            if val is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = val


def _toc_matches_pdf(docx, pdf):
    """TOC page numbers printed in the Word file equal the pages the renderer put the headings on."""
    from refresh_toc import pdf_positions

    entries = toc_entries(Document(docx))
    if not entries:
        return True
    try:
        values, _, _ = pdf_positions(pdf, entries)
    except ValueError:
        return False
    return [entry["cache"] for entry in entries] == list(map(str, values))


# What a hand edit may change: how text looks.  Anything else it breaks
# (line spacing, objects past the margins, hidden text, page numbers, the
# cover, tables) blocks delivery exactly as it does for ``run``.
ADOPTABLE = {
    "FONT_SIZE": "字号", "FONT_CN": "中文字体", "FONT_EN": "西文字体", "FONT_BOLD": "加粗", "FONT_COLOR": "文字颜色",
    "FONT_ITALIC": "倾斜", "FONT_UNDERLINE": "下划线", "PARAGRAPH": "段落格式（对齐、缩进、段前段后或行距）",
    "PARAGRAPH_EXTRA": "段落格式（对齐、缩进、段前段后或行距）", "STYLE_INHERITANCE": "段落样式",
    "TOC_FONT": "目录字体", "TOC_FONT_SIZE": "目录字号", "TOC_EMPHASIS": "目录加粗", "TOC_COLOR": "目录颜色",
    "HEADER_TEXT": "页眉文字", "HEADER_BORDER": "页眉横线", "COVER_FONT": "封面字体", "COVER_EFFECTIVE_FONT": "封面字体",
    "LOGO_SIZE": "校徽大小", "LOGO_ALIGN": "校徽对齐",
}
# Proved again from the rendered PDF by adopt itself (the stored proof belongs to the generated copy).
REPROVED = {"TOC_PROOF_STALE", "TOC_UNRESOLVED"}


CN_SIZES = {84: "初号", 72: "小初", 52: "一号", 48: "小一", 44: "二号", 36: "小二", 32: "三号", 30: "小三",
            28: "四号", 24: "小四", 21: "五号", 18: "小五", 15: "六号"}


def _size_text(value):
    """'28' / '28半磅' (half-points, as the checks report sizes) -> '四号（14 磅）' for people."""
    def one(match):
        half = int(match.group(1))
        pt = f"{half / 2:g} 磅"
        return f"{CN_SIZES[half]}（{pt}）" if half in CN_SIZES else pt
    return re.sub(r"(\d+)\s*(?:半磅)?", one, str(value))


# A hand-made change in words: which property, and its values as Word's dialogs name them.
PROPERTY_NAMES = {
    "jc/val": "对齐", "ind/firstLine": "首行缩进", "ind/firstLineChars": "首行缩进", "ind/left": "左缩进",
    "ind/leftChars": "左缩进", "ind/hanging": "悬挂缩进", "ind/hangingChars": "悬挂缩进", "ind/right": "右缩进",
    "ind/rightChars": "右缩进", "spacing/before": "段前", "spacing/after": "段后", "spacing/beforeLines": "段前",
    "spacing/afterLines": "段后", "spacing/line": "行距", "spacing/lineRule": "行距方式", "keepNext/val": "与下段同页",
    "keepLines/val": "段中不分页", "pageBreakBefore/val": "段前分页", "widowControl/val": "孤行控制",
}
ALIGN_NAMES = {"both": "两端对齐", "center": "居中", "left": "左对齐", "start": "左对齐", "right": "右对齐",
               "end": "右对齐", "distribute": "分散对齐"}
RULE_NAMES = {"auto": "多倍行距", "exact": "固定值", "atLeast": "最小值"}
ROLE_NAMES = {
    "body": "正文", "h1": "一级标题", "h2": "二级标题", "h3": "三级标题", "h4": "四级标题", "preface": "序言标题",
    "title_zh": "中文题目", "title_en": "英文题目", "abstract_label_zh": "“摘要”标签", "abstract_label_en": "“Abstract”标签",
    "abstract_zh": "中文摘要", "abstract_en": "英文摘要", "keywords_zh": "中文关键词", "keywords_en": "英文关键词",
    "toc_title": "目录标题", "toc1": "目录一级", "toc2": "目录二级", "toc3": "目录三级", "toc4": "目录四级",
    "reference_title": "参考文献标题", "reference": "参考文献条目", "caption": "图表标题", "figure": "图片段",
    "table_text": "表格文字", "table_note": "表注", "table_unit": "表格单位", "code": "代码", "equation": "公式", "header": "页眉",
}
# How the same change is asked for in job.json (format.styles.<role>.<key>), when it is one.
STYLE_KEY = {"FONT_SIZE": "size", "FONT_CN": "cn", "FONT_EN": "latin", "FONT_BOLD": "bold", "FONT_COLOR": "color",
             "jc/val": "align", "ind/firstLine": "first", "ind/firstLineChars": "chars", "ind/left": "left",
             "ind/leftChars": "left_chars", "ind/hanging": "hanging", "spacing/before": "before",
             "spacing/after": "after", "spacing/beforeLines": "before_lines", "spacing/afterLines": "after_lines",
             "spacing/line": "line", "spacing/lineRule": "line_rule", "keepNext/val": "keep",
             "pageBreakBefore/val": "break_before"}


def color_name(value):
    """'0055AA' -> '蓝色（0055AA）': a colour as people name it, the hex value kept for exactness."""
    import colorsys

    raw = str(value).strip().upper()
    if raw in ("000000", "AUTO"):
        return "黑色"
    if not re.fullmatch(r"[0-9A-F]{6}", raw):
        return raw
    r, g, b = (int(raw[i:i + 2], 16) / 255 for i in (0, 2, 4))
    hue, light, sat = colorsys.rgb_to_hls(r, g, b)
    if sat < 0.15 or light < 0.12 or light > 0.95:
        name = "黑色" if light < 0.12 else "白色" if light > 0.95 else "灰色"
    else:
        name = next(n for top, n in ((15, "红色"), (45, "橙色"), (70, "黄色"), (160, "绿色"), (200, "青色"),
                                     (255, "蓝色"), (320, "紫色"), (361, "红色")) if hue * 360 < top)
    return f"{name}（{raw}）"


def _property(issue):
    """'jc/val' for a PARAGRAPH issue, else the issue code."""
    m = re.search(r"(?P<what>[\w/]+)=", issue.get("message", "")) if issue["code"] == "PARAGRAPH" else None
    return m.group("what") if m else issue["code"]


def _readable(what, value):
    raw = str(value).strip().strip("'\"")
    if raw in ("None", ""):
        return "未设置"
    if what in ("FONT_SIZE", "TOC_FONT_SIZE"):
        return _size_text(raw)
    if what == "jc/val":
        return ALIGN_NAMES.get(raw, raw)
    if what == "spacing/lineRule":
        return RULE_NAMES.get(raw, raw)
    if what == "FONT_COLOR":
        return color_name(raw)
    if what.startswith(("FONT_", "TOC_")):
        return raw
    if what.endswith("/val") and raw in ("0", "1"):
        return "是" if raw == "1" else "否"
    if not re.fullmatch(r"-?\d+", raw):
        return raw
    n = int(raw)
    if what.endswith("Chars"):
        return f"{n / 100:g} 字符"
    if what.endswith("Lines"):
        return f"{n / 100:g} 行"
    if what == "spacing/line":
        return f"{n / 240:g} 倍（固定值时 {n / 20:g} 磅）"
    return f"{n / 20:g} 磅"


def _role_of(issue):
    from customization import ROLE_STYLES

    style = (issue.get("where") or {}).get("style")
    return next((role for role, name in ROLE_STYLES.items() if name == style), None)


def _deviation_groups(explained):
    """Hand-made changes grouped by role and property: {(role, what, template, now): {paragraph, ...}}."""
    groups = {}
    for x in explained:
        what = _property(x)
        key = (_role_of(x), what, _readable(what, x.get("expected", "")) if "expected" in x else "",
               _readable(what, x.get("actual", "")) if "actual" in x else "")
        groups.setdefault(key, set()).add(re.sub(r"/r\d+$", "", (x.get("where") or {}).get("location", "")))
    return groups


def _group_line(key, places, totals):
    role, what, template, now = key
    name = PROPERTY_NAMES.get(what) or ADOPTABLE.get(what, what)
    total = totals.get(role, 0)
    who = ROLE_NAMES.get(role, "部分段落")
    who += f"（全部 {total} 段）" if total and len(places) >= total else f"（{len(places)} 段）"
    values = f"：模板是{template}，现在是{now}" if template and now else "和模板默认不同"
    return f"{who}的{name}{values}"


def _role_totals(docx):
    """Paragraphs with text per role in the document."""
    from collections import Counter
    from customization import ROLE_STYLES

    names = {name: role for role, name in ROLE_STYLES.items()}
    return Counter(names.get(p.style.name if p.style is not None else "") for p in Document(docx).paragraphs
                   if p.text.strip())


def _as_job_value(what, value):
    raw = str(value).strip().strip("'\"")
    if what == "FONT_SIZE":
        return int(raw) / 2 if raw.isdigit() else None
    if what == "FONT_BOLD":
        return raw == "是"
    if what.endswith("/val") and what != "jc/val":
        return raw == "1"
    if re.fullmatch(r"-?\d+", raw) and what not in ("FONT_CN", "FONT_EN", "FONT_COLOR"):
        return int(raw)
    return raw if raw not in ("None", "") else None


def _plain_deviation(issue):
    """One hand-made change in words (used when it cannot be grouped)."""
    where = issue.get("where") or {}
    what = _property(issue)
    place = f"“{where['text']}”" if where.get("text") else (where.get("location") or "")
    name = PROPERTY_NAMES.get(what) or ADOPTABLE.get(issue["code"], issue["code"])
    detail = (f"（模板：{_readable(what, issue['expected'])}；现在：{_readable(what, issue['actual'])}）"
              if "expected" in issue and "actual" in issue else "")
    return (f"{place}的{name}" if place else name) + "和模板默认不同" + detail


def adopt(workspace, docx, note):
    """Accept a hand-edited copy of the accepted output for a requirement the scripts cannot meet.

    The text must stay identical; the rendered-page audit, TOC check and source
    comparison run as for ``run``; every static check the edit now fails is
    recorded as a deviation that the review must address and the user is told,
    as long as it is a change of how text looks (``ADOPTABLE``); any other new
    static error is refused like a failed ``run``.
    """
    from source_compare import body_key, first_text_difference

    root, s = load(workspace)
    if s.get("stage") not in ("awaiting_review", "ready") or s.get("validation_scope") != "full":
        raise ValueError("先用 run 得到通过验收、有页图的成品，再在它的副本上手工处理；没有办公软件时不能验收手工修改。")
    if not isinstance(note, str) or len(note.strip()) < 4:
        raise ValueError("--note 要写清手工改了什么（这句话会告诉用户）。")
    source = Path(docx).expanduser().resolve()
    if source.suffix.lower() != ".docx" or not source.is_file():
        raise ValueError("--docx 要指向手工改好的 .docx。")
    base = root / "generated.docx"
    if not base.is_file():
        shutil.copy2(root / "final.docx", base)
    if body_key(source) != body_key(base):
        where, was, now = first_text_difference(base, source) or ("", "", "")
        raise AcceptanceFailed("手工修改", [{
            "code": "MANUAL_TEXT_CHANGED", "stage": "手工修改", "message": "手工处理的文件和脚本生成的成品文字不一样。",
            "where": {"location": f"手工文件的{where}"} if where else {},
            "expected": was or "（不应有）", "actual": now or "（没有）",
            "cause": "手工处理只能改格式，不能增删改文字（目录页码除外）。",
            "fix": "把这一处改回去，或从 WORK 里的 generated.docx 重新复制一份，只改用户要求的那一项格式；要改文字就改 content.json / mapping.json 后重新 run。",
        }])
    intake = root / "active-intake.json"
    known = {(e["code"], e["message"]) for e in check(base, final=True, intake_path=intake)["errors"]}
    final = root / "final.docx"
    # The copy accepted last (the run output, or an earlier adopted copy) comes back if this one fails.
    accepted = root / "accepted-before-adopt"
    shutil.rmtree(accepted, ignore_errors=True)
    accepted.mkdir()
    shutil.copy2(final, accepted / "final.docx")
    for name in ("preview", "render-audit.json"):  # its page images and audit too
        if (root / name).is_dir():
            shutil.copytree(root / name, accepted / name)
        elif (root / name).is_file():
            shutil.copy2(root / name, accepted / name)
    shutil.copy2(source, final)
    try:
        after = check(final, final=True, intake_path=intake)
        new = [e for e in after["errors"] if (e["code"], e["message"]) not in known]
        tolerated = [e for e in new if e["code"] in REPROVED]
        deviations = [e for e in new if e["code"] in ADOPTABLE]
        blocking = [e for e in new if e not in tolerated and e not in deviations]
        if blocking:
            issues = explain(blocking, final, root, stage="手工修改")
            for item in issues:
                item["cause"] = ("手工处理只能改文字的外观（字体、字号、颜色、加粗、对齐、缩进、间距等）；"
                                 "这一项是会让版面出错的问题，adopt 不放行。")
                item["fix"] = ("在手工处理的副本里把这一处改回去（位置见上），或换一种不引起这个问题的改法，再重新 adopt；"
                               "做不到就告诉用户这项要求会带来什么问题，问用户怎么取舍。")
            raise AcceptanceFailed("手工修改", issues)
        explained = explain(deviations, final, root, stage="手工修改")
        # One line per role and kind of change ("正文（全部 30 段）的字号：…"), not one per paragraph or run.
        groups, totals = _deviation_groups(explained), _role_totals(final)
        expressible = {}
        for (role, what, _, now), places in groups.items():
            raw = next((x.get("actual") for x in explained if _role_of(x) == role and _property(x) == what), None)
            value = _as_job_value(what, raw) if raw is not None else None
            if role in ROLE_NAMES and what in STYLE_KEY and value is not None and totals.get(role) \
                    and len(places) >= totals[role]:
                expressible.setdefault(role, {})[STYLE_KEY[what]] = value
        if expressible:
            styles = json.dumps(expressible, ensure_ascii=False)
            raise AcceptanceFailed("手工修改", [{
                "code": "ADOPT_USE_JOB", "stage": "手工修改",
                "message": "这次手工改的是整类段落的格式，job.json 就能做到，不需要手工处理："
                           + "；".join(_group_line(k, v, totals) for k, v in groups.items() if k[0] in expressible) + "。",
                "where": {"edit": "job.json 的 format.styles"},
                "cause": "adopt 只用于脚本做不到的要求；写进 job.json 的要求会记在成品里，以后改稿时自动保留，手工改的不会。",
                "fix": f"在 job.json 的 format.styles 里写 {styles}（和已有的项合并），重新 run，不要 adopt。",
            }])
        # The output remembers the hand edit, so revising it later (new --source) cannot silently drop it.
        marked = Document(final)
        E.set_vars(marked, {"XAUFE_MANUAL": json.dumps(
            {"note": note.strip(), "summary": [_group_line(k, v, totals) for k, v in groups.items()]},
            ensure_ascii=False)})
        marked.save(final)
        try:
            preview = render(final, root / "preview")
        except OfficeUnavailable as exc:
            raise ValueError("没有可用的办公软件，手工修改做不了渲染后验收，不能交付：" + str(exc)) from exc
        pdf = root / "preview" / "final.pdf"
        rendered = audit_rendered(final, pdf, adopted=True)
        write_json(root / "render-audit.json", rendered)
        if not rendered["passed"]:
            raise AcceptanceFailed("渲染后版面检查", from_render(rendered, root))
        if not _toc_matches_pdf(final, pdf):
            raise AcceptanceFailed("手工修改", [{
                "code": "TOC_PROOF_STALE", "stage": "手工修改", "message": "目录页码和实际分页对不上。",
                "cause": "手工修改改变了分页，目录没有更新。", "fix": "在 Word 里右键目录 →“更新域”→“更新整个目录”，保存后重新 adopt。",
            }])
        result = {"warnings": []}
        compare_originals(root, s, result)
    except Exception:
        # The copy accepted last stays in place, with its page images; the state still points at it.
        shutil.copy2(accepted / "final.docx", final)
        for name in ("preview", "render-audit.json"):
            if (accepted / name).is_dir():
                shutil.rmtree(root / name, ignore_errors=True)
                shutil.copytree(accepted / name, root / name)
            elif (accepted / name).is_file():
                shutil.copy2(accepted / name, root / name)
        raise
    finally:
        shutil.rmtree(accepted, ignore_errors=True)
    lines = [_group_line(k, v, totals) for k, v in groups.items()]
    warnings = rendered["warnings"] + result["warnings"] + [
        {"code": "MANUAL_EDIT", "message": "手工处理：" + note.strip() + "（交付时会告诉用户）"}
    ] + ([{"code": "MANUAL_DEVIATION", "message": "手工处理后和模板默认不同的地方：" + "；".join(lines) + "。"}]
         if lines else [])
    keys = sorted({warning_key(w) for w in warnings})
    s.pop("review", None)
    s.update(
        stage="awaiting_review",
        document_sha256=sha(final),
        render_report_sha256=sha(root / "preview/render-report.json"),
        render_audit_sha256=sha(root / "render-audit.json"),
        pages=preview["pages"],
        warning_codes=keys,
        manual={"note": note.strip(), "summary": lines,
                "deviations": [[e["code"], e["message"]] for e in deviations],
                "tolerated": [[e["code"], e["message"]] for e in tolerated]},
    )
    write_json(root / MARKER, s)
    return {
        "workspace": str(root),
        "validation_scope": "full",
        "pages": [dict(p, path=str(root / "preview" / p["png"])) for p in preview["pages"]],
        "warnings": warnings,
    }


def _real_text(value, minimum=4):
    return isinstance(value, str) and len(value.strip()) >= minimum and PLACEHOLDER not in value


def attest(workspace, review):
    """Accept the agent's written review of the current final document."""
    root, s = load(workspace)
    if s.get("stage") not in ("awaiting_review", "ready"):
        raise ValueError("还没有可检查的成品；先运行 run。")
    if sha(root / "final.docx") != s.get("document_sha256"):
        raise ValueError("成品在 run（或 adopt）之后被改动：不要直接改 WORK 里的 Word。手工处理过的就重新 adopt 那份副本，"
                         "否则重新运行 run，再检查。")
    if not isinstance(review, dict):
        raise ValueError("review.json 必须是 JSON 对象。")
    unknown = set(review) - {"vision", "no_vision_reason", "pages", *REVIEW_TEXTS, "warnings"}
    if unknown:
        raise ValueError("review.json 有不认识的项：" + "、".join(sorted(unknown)))
    for key, label in REVIEW_TEXTS.items():
        if not _real_text(review.get(key), 8):
            raise ValueError(f"review.json 的 {key}（{label}）没有写实际检查结果。")
    pages = review.get("pages", {})
    if not isinstance(pages, dict):
        raise ValueError('review.json 的 pages 必须写成 {"1": "第1页看到的情况", ...}。')
    vision = review.get("vision")
    if not isinstance(vision, bool):
        raise ValueError("review.json 的 vision 只能是 true（我能看图并看过了）或 false（我不能看图）。")
    if s.get("validation_scope") == "static-only":
        if any(_real_text(v) for v in pages.values()):
            raise ValueError("这次没有生成页图，pages 必须为空，不能写看图结果。")
        visual = "skipped_no_renderer"
    elif vision:
        expected = {p["page"] for p in s["pages"]}
        covered = {}
        for key in pages:
            m = re.fullmatch(r"\s*(\d+)\s*(?:[-–~]\s*(\d+)\s*)?", str(key))
            if not m or (m.group(2) and int(m.group(2)) <= int(m.group(1))):
                raise ValueError(f'review.json 的 pages 里“{key}”不是页码：写 "3" 或连续几页 "5-12"。')
            for n in range(int(m.group(1)), int(m.group(2) or m.group(1)) + 1):
                if n not in expected:
                    raise ValueError(f"pages 里有不存在的页码：{n}（“{key}”）。")
                if n in covered:
                    raise ValueError(f"第 {n} 页在 pages 里写了两次（“{covered[n]}”和“{key}”）。")
                covered[n] = key
        missing = sorted(expected - set(covered))
        if missing:
            raise ValueError("每一页都要打开看并写观察（连续几页情况相同可以合写，如 \"5-12\"），还缺第 "
                             + "、".join(map(str, missing)) + " 页。")
        weak = [k for k, v in pages.items() if not _real_text(v)]
        if weak:
            raise ValueError("这些页没有写具体看到的情况：" + "、".join(f"第 {k} 页" for k in weak) + "。")
        for page in s["pages"]:
            if sha(root / "preview" / page["png"]) != page["png_sha256"]:
                raise ValueError("页图已变化；请重新运行 run 后再看图。")
        visual = "passed_agent_review"
    else:
        if any(_real_text(v) for v in pages.values()):
            raise ValueError("vision=false 时 pages 必须为空，不能写看图结果。")
        if not _real_text(review.get("no_vision_reason")):
            raise ValueError("vision=false 时要在 no_vision_reason 写明为什么不能看图。")
        visual = "skipped_no_vision"
    notes = review.get("warnings", {})
    if not isinstance(notes, dict):
        raise ValueError("review.json 的 warnings 必须是 {\"提示（run 生成的 review.json 里已列好）\": \"怎么处理的\"}。")
    missing = [code for code in s.get("warning_codes", []) if not _real_text(notes.get(code))]
    if missing:
        raise ValueError("这些警告还没写处理说明：" + "、".join(missing))
    s["review"] = dict(review, document_sha256=s["document_sha256"], visual_review=visual)
    s["stage"] = "ready"
    write_json(root / MARKER, s)
    return {"visual_review": visual}


def clean(workspace):
    root, s = load(workspace)
    dest = Path(s["destination"]).resolve()
    if dest.is_relative_to(root):
        raise ValueError("最终 Word 不能位于将要清理的工作区内。")
    # rmtree does not follow interior symlinks; root ownership was checked above.
    shutil.rmtree(root)
    if root.exists():
        raise OSError("工作区没有清理干净，不能说只剩最终 Word。")


FONT_NOTICE = "检查用的页图是在没有宋体/黑体等原字体的环境里渲染的（用了替代字体），分页和字宽以装有这些字体的 Word 为准，建议在自己电脑上翻看一遍。"


RENDERER_NAMES = {"libreoffice": "LibreOffice", "wps": "WPS"}


def renderer_notice(s):
    """Pagination was checked in something other than Word: say which, and how to refresh the TOC."""
    name = RENDERER_NAMES.get(str(s.get("renderer", "")).lower())
    if not name:
        return ""
    text = f"分页和逐页检查是在 {name} 里做的，Word 的分页可能差一两行"
    if s.get("has_toc"):
        return text + "；用 Word 打开后如果目录页码和正文对不上，在目录上右键 →“更新域”→“只更新页码”。"
    return text + "，建议用 Word 打开翻看一遍。"


def user_notice(s):
    told = s.get("tell_user") or []
    extra = ("另外：" + "；".join(t.rstrip("。；") for t in told) + "。") if told else ""
    return _user_notice(s) + extra


def _user_notice(s):
    review = s["review"]
    if s.get("validation_scope") == "static-only":
        return s["user_notice"]
    fonts = renderer_notice(s)
    fonts += FONT_NOTICE if any(c.startswith("FONT_SUBSTITUTED") for c in s.get("warning_codes", [])) else ""
    manual = s.get("manual")
    if manual:
        fonts += "有一项是手工处理的：" + manual["note"].rstrip("。") + "。"
        summary = manual.get("summary") or []
        if summary:
            shown = summary[:6] + ([f"等共 {len(summary)} 处"] if len(summary) > 6 else [])
            fonts += "这样处理后，下面几项和模板默认格式不同：" + "；".join(shown) + "。"
    if review["visual_review"] == "skipped_no_vision":
        return (
            "已完成排版、格式检查和实际分页；没有逐页看图检查，原因："
            + review["no_vision_reason"].strip().rstrip("。")
            + "。建议打开文件快速翻看一遍。"
            + fonts
        )
    return fonts


def publish(workspace):
    root, s = load(workspace)
    if s.get("stage") != "ready" or not s.get("review"):
        raise ValueError("还没有通过 review.json 检查，不能交付。")
    final = root / "final.docx"
    if sha(final) != s.get("document_sha256") or s["review"]["document_sha256"] != sha(final):
        raise ValueError("检查之后 Word 又被改动；请重新运行 run 并重新检查。")
    limited = s.get("validation_scope") == "static-only"
    receipt = RF.verify_receipt(root, s) if limited else None
    if not limited:
        if sha(root / "preview/render-report.json") != s.get("render_report_sha256"):
            raise ValueError("预览已变化，请重新运行 run。")
        for page in s["pages"]:
            if sha(root / "preview" / page["png"]) != page["png_sha256"]:
                raise ValueError("页图已变化，请重新运行 run。")
        if s.get("render_audit_sha256") and sha(root / "render-audit.json") != s["render_audit_sha256"]:
            raise ValueError("渲染后版面检查结果被改动，请重新运行 run。")
        if s.get("manual"):
            if not _toc_matches_pdf(final, root / "preview" / "final.pdf"):
                raise ValueError("目录页码和实际分页对不上，请更新目录后重新 adopt。")
        else:
            check_preview_toc(root)
    verify_originals(s)
    source = s.get("active_request_source")
    if source and require(source)["integrity"] != require(root / "active-intake.json")["integrity"]:
        raise ValueError("用户要求或日期已变化，请重新运行 run。")
    result = check(
        final,
        final=True,
        intake_path=root / "active-intake.json",
        validation_scope="static-only" if limited else "full",
    )
    manual = s.get("manual") or {}
    allowed = {tuple(d[:2]) for d in manual.get("deviations", []) + manual.get("tolerated", [])}
    remaining = [e for e in result["errors"] if (e["code"], e["message"]) not in allowed]
    if remaining:
        raise AcceptanceFailed("交付前复查", explain(remaining, final, root))
    dest = Path(s["destination"])
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.is_relative_to(root) or dest.is_relative_to(E.ROOT):
        raise ValueError("交付位置不安全。")
    created = False
    try:
        # Exclusive creation: never overwrite an existing output or source file.
        with dest.open("xb") as out, final.open("rb") as src:
            created = True
            shutil.copyfileobj(src, out)
            out.flush()
            os.fsync(out.fileno())
        if sha(dest) != s["document_sha256"]:
            raise OSError("复制后的 Word 校验不一致。")
    except Exception:
        if created and dest.exists():
            dest.unlink()
        raise
    notice = user_notice(s)
    clean(root)
    return {
        "output": str(dest),
        "workspace_removed": not root.exists(),
        "validation_scope": "static-only" if limited else "full",
        "visual_review": s["review"]["visual_review"],
        "skipped_steps": receipt["skipped_steps"] if limited else [],
        "user_notice": notice,
    }
