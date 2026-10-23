#!/usr/bin/env python3
"""Rendered-page audit: reads only the exported PDF and the delivered DOCX.

It does not call the layout engine.  The page box and margins come from the
DOCX section properties (the static audit already ties those to the template),
and the header/footer expectations come from the recorded effective profile.
Everything else is measured on the PDF a real office suite produced, so a
wrong assumption in the generator cannot hide a visible defect here.

Errors are things a person would see as broken: content outside the page or
the text column, a page with nothing in the text area, lines printed on top of
each other, missing or unexpected header/footer text, page numbers that skip
or restart in the wrong place.  The one warning is a font the renderer had to
substitute, which changes the page images but not the Word file.
"""

from __future__ import annotations
import argparse, json, re, sys
from collections import Counter
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

PT = 20  # twips per point
EDGE_TOLERANCE = 2.0  # pt; drawing and image edges are exact
TEXT_TOLERANCE = 6.0  # pt; about half a 五号 glyph, absorbs justification and font metrics
COVER_SAFE = 28.0  # pt (~1 cm); the school cover centres 36 pt titles on its own measure
HEADER_ROOM = 24.0  # pt: a header or footer line and its leading, measured from the header/footer distance
EDGE_SAFE = 14.0  # pt (~0.5 cm); most printers cannot print closer to the paper edge
HANGING = set("，。、；：！？）》」』”’,.;:!?)")

# PDF font names that count as the requested family.
FONT_ALIASES = {
    "宋体": ("simsun", "nsimsun", "songti", "stsong", "宋体"),
    "黑体": ("simhei", "heiti", "stheiti", "黑体"),
    "仿宋": ("fangsong", "stfangsong", "仿宋"),
    "楷体": ("kaiti", "stkaiti", "楷体"),
    "微软雅黑": ("microsoftyahei", "msyh", "微软雅黑"),
    # Liberation/Tinos/Arimo are metric-compatible: same advance widths, same line breaks.
    "Times New Roman": ("timesnewroman", "timesnewromanps", "liberationserif", "tinos"),
    "Arial": ("arial", "arialmt", "liberationsans", "arimo"),
    "Courier New": ("couriernew", "couriernewps", "liberationmono", "cousine"),
    "Calibri": ("calibri",),
    "Consolas": ("consolas",),
}


def _key(name):
    # "ABCDEF+SimSun-Bold" -> "simsunbold"
    return re.sub(r"[^0-9a-z一-鿿]", "", name.split("+", 1)[-1].lower())


def requested_fonts(doc):
    """Font families the visible text of the DOCX resolves to (styles, runs, tables, headers)."""
    from style_model import run_properties
    from table_layout import walk_paragraphs

    paragraphs = list(doc.paragraphs)
    for t in doc.tables:
        paragraphs += [p for _, p in walk_paragraphs(t, "t")]
    for section in doc.sections:
        for part in (section.header, section.first_page_header, section.footer, section.first_page_footer):
            paragraphs += list(part.paragraphs)
    names = set()
    for p in paragraphs:
        for r in p._p.xpath(".//w:r"):
            text = "".join(r.xpath("./w:t/text()"))
            if not text.strip():
                continue
            try:
                fonts = run_properties(p, r).get("fonts", {})
            except ValueError:
                continue
            cjk = bool(re.search(r"[\u3400-\u9fff\uff00-\uffef]", text))
            latin = bool(re.search(r"[A-Za-z0-9]", text))
            wanted = ([fonts.get("eastAsia")] if cjk else []) + ([fonts.get("ascii")] if latin else [])
            for value in wanted:
                if value:
                    names.add(value.split(";")[-1].strip())
    return sorted(names)


def _font_present(family, used):
    aliases = FONT_ALIASES.get(family, (_key(family),))
    return any(alias and any(used_key.startswith(alias) for used_key in used) for alias in aliases)


def font_gap(doc, pdf_path):
    """(missing, rendered): requested families absent from the PDF, and what it used."""
    import pymupdf as fitz

    used = set()
    with fitz.open(pdf_path) as pdf:
        for page in pdf:
            used |= {_key(font[3]) for font in page.get_fonts()}
    return [f for f in requested_fonts(doc) if not _font_present(f, used)], sorted(used)


def _expectations(doc):
    vv = {e.get(qn("w:name")): e.get(qn("w:val")) for e in doc.settings.element.xpath("./w:docVars/w:docVar")}
    header = footer = None
    template = vv.get("XAUFE_TEMPLATE")
    if template in ("A", "B"):
        from customization import effective_profile

        from audit_helpers import contract

        request = json.loads(vv.get("XAUFE_FORMAT_REQUEST", "{}"))
        spec = effective_profile(template, request)
        header = bool(spec["header"]["enabled"])
        footer = bool(spec["footer"].get("page_numbers"))
        body_size = contract(template, request)["body"]["font"]["size_half_points"] / 2
    else:
        body_size, request = None, None
    return {"has_cover": vv.get("XAUFE_HAS_COVER") == "1", "header": header, "page_numbers": footer,
            "template": template, "body_size": body_size, "request": request}


def _margins(sections):
    """Page box and text column; when sections differ, the most generous edge of each."""
    get = lambda s, name: getattr(s, name).twips / PT
    return {
        "width": max(get(s, "page_width") for s in sections),
        "height": max(get(s, "page_height") for s in sections),
        **{k: min(get(s, k + "_margin") for s in sections) for k in ("left", "right", "top", "bottom")},
        **{k: max((getattr(s, k + "_distance").twips if getattr(s, k + "_distance") is not None else 0)
                  for s in sections) / PT for k in ("header", "footer")},
    }


def _issue(kind, code, page, message, *, where=None, expected=None, actual=None, bbox=None):
    item = {"code": code, "severity": kind, "page": page, "message": message}
    if where:
        item["where"] = where
    if expected is not None:
        item["expected"] = expected
    if actual is not None:
        item["actual"] = actual
    if bbox is not None:
        item["bbox_pt"] = [round(v, 1) for v in bbox]
    return item


def _lines(page):
    """Visible lines; ``ink_right`` ignores one trailing punctuation mark (行末标点悬挂)."""
    result = []
    for block in page.get_text("rawdict")["blocks"]:
        for line in block.get("lines", []):
            chars = [dict(c, size=span["size"], bold=bool(span["flags"] & 16) or "bold" in span["font"].lower())
                     for span in line["spans"] for c in span["chars"] if c["c"].strip()]
            if not chars:
                continue
            boxes = [c["bbox"] for c in chars]
            body = chars[:-1] if len(chars) > 1 and chars[-1]["c"] in HANGING else chars
            result.append(
                {
                    "spans": [{"text": "".join(c["c"] for c in chars)}],
                    "bbox": (min(b[0] for b in boxes), min(b[1] for b in boxes),
                             max(b[2] for b in boxes), max(b[3] for b in boxes)),
                    "ink_right": max(c["bbox"][2] for c in body),
                    "chars": [(c["bbox"][0], c["bbox"][2], c["origin"][1], c["size"]) for c in chars],
                    # Weight is only told apart for Latin letters: CJK fonts rarely ship a bold face.
                    "latin_bold": [c["bold"] for c in chars if c["c"].isascii() and c["c"].isalpha()],
                }
            )
    return result


def _ink(char):
    """Where a glyph actually puts ink: about 0.8 em above the baseline to 0.1 em below it."""
    x0, x1, base, size = char
    return x0, base - 0.8 * size, x1, base + 0.1 * size


def _line_clash(a, b):
    """(baseline gap, ink overlap) of the worst pair of glyphs printed into each other, or None.

    Each glyph inks roughly 0.8 em above to 0.1 em below its baseline.  Two glyphs
    clash when they share at least 30% of the narrower one's width and those ink
    bands overlap by more than 5% of the smaller size: for one size that is a
    baseline gap under about 0.85 of the size; a large title above small text is
    judged by the room actually left between them.  Superscripts beside a word
    never share its width.
    """
    if a["bbox"][1] > b["bbox"][3] + 40 or b["bbox"][1] > a["bbox"][3] + 40:
        return None
    worst = None
    for char in a["chars"]:
        x0, x1, base, size = char
        for other in b["chars"]:
            u0, u1, ubase, usize = other
            narrow = min(x1 - x0, u1 - u0)
            if narrow <= 0 or min(x1, u1) - max(x0, u0) < 0.3 * narrow:
                continue
            _, top, _, bottom = _ink(char)
            _, utop, _, ubottom = _ink(other)
            overlap = min(bottom, ubottom) - max(top, utop)
            if overlap > 0.05 * min(size, usize) and (worst is None or overlap > worst[1]):
                worst = (abs(base - ubase), overlap)
    return worst


def _covers(box, ink, share=0.2):
    width = min(box[2], ink[2]) - max(box[0], ink[0])
    height = min(box[3], ink[3]) - max(box[1], ink[1])
    area = (ink[2] - ink[0]) * (ink[3] - ink[1])
    return width > 0 and height > 0 and area > 0 and width * height > share * area


def _header_texts(doc):
    texts = set()
    for s in doc.sections:
        for part in (s.header, s.even_page_header):
            for p in part.paragraphs:
                t = re.sub(r"\s+", "", p.text)
                if t:
                    texts.add(t)
    return texts


def _header_paragraphs(doc):
    """Most non-empty paragraphs any one header has: how many rows a header may show."""
    return max((sum(1 for p in part.paragraphs if p.text.strip())
                for s in doc.sections for part in (s.header, s.even_page_header)), default=0)


def _header_rows(head, texts):
    """Distinct baselines the header text sits on (left and right parts share one when it fits)."""
    rows = []
    for line in head:
        text = re.sub(r"\s+", "", line["spans"][0]["text"])
        if not text or not any(text in t or t in text for t in texts):
            continue
        base, size = line["chars"][0][2], line["chars"][0][3]
        if not any(abs(base - b) < 0.5 * size for b, _ in rows):
            rows.append((base, line))
    return sorted(rows, key=lambda r: r[0])


def _header_overflow(lines, top, texts):
    """The last header line when the header wrapped past the top margin into the text area."""
    if not texts:
        return None
    chain = [l for l in lines if l["bbox"][3] <= top + 1]
    if not chain:
        return None
    last = max(chain, key=lambda l: l["bbox"][3])
    for line in sorted(lines, key=lambda l: l["bbox"][1]):
        if line in chain or line["bbox"][1] < last["bbox"][1]:
            continue
        height = last["bbox"][3] - last["bbox"][1]
        text = re.sub(r"\s+", "", line["spans"][0]["text"])
        if line["bbox"][1] - last["bbox"][3] < 0.6 * height and text and any(text in t for t in texts):
            last = line
        else:
            break
    return last if last["bbox"][3] > top + 1 else None


HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3", "Heading 4", "XAUFE Preface", "XAUFE References Title"}


def _flat(text):
    return re.sub(r"\s+", "", text)


def _keep_together(doc):
    """(headings, table captions above their table, figure captions below their figure) as plain text."""
    headings, above, below = [], [], []
    for p in doc.paragraphs:
        text = _flat(p.text)
        if not text:
            continue
        if p.style.name in HEADING_STYLES:
            headings.append(text)
        elif p.style.name == "XAUFE Caption":
            nxt, prev = p._p.getnext(), p._p.getprevious()
            # Look past empty paragraphs: a stray blank between them must not hide the pairing.
            while nxt is not None and nxt.tag == qn("w:p") and not _flat("".join(nxt.xpath(".//w:t/text()"))) \
                    and not nxt.xpath(".//w:drawing|.//w:pict"):
                nxt = nxt.getnext()
            while prev is not None and prev.tag == qn("w:p") and not _flat("".join(prev.xpath(".//w:t/text()"))) \
                    and not prev.xpath(".//w:drawing|.//w:pict"):
                prev = prev.getprevious()
            if nxt is not None and nxt.tag == qn("w:tbl"):
                above.append(text)
            elif prev is not None and prev.xpath(".//w:drawing|.//w:pict"):
                below.append(text)
    return headings, above, below


# Roles whose position on the page the rendered audit checks (first-line start, centring, pitch).
PLACED = {"body", "h1", "h2", "h3", "h4", "preface", "reference_title", "caption", "title_zh", "title_en",
          "abstract_label_zh", "abstract_label_en", "toc_title"}
INDENT_TOLERANCE = 3.0  # pt; well under the 10.5 pt a one-character indent error moves the text


def _twips(values, key, size):
    """An indent in points: the *Chars form (hundredths of a character) wins, as in Word."""
    chars = values.get(key + "Chars")
    if chars not in (None, "") and size:
        return int(chars) / 100 * size
    return int(values.get(key) or 0) / PT


def _starts(ind, sizes):
    """Where the first line may start (pt from the left text edge).  Office suites size a character
    unit differently (paragraph style size, run size, or the twip value), so each reading counts."""
    return sorted({round(_twips(ind, "left", z) + _twips(ind, "firstLine", z) - _twips(ind, "hanging", z), 1)
                   for z in sizes})


def _placements(doc, template, request, adopted):
    """[(flat text, role, {"jc", "start", "pitch"})] for paragraphs whose placement is checked.

    ``start``: where the first line should begin, from the left text edge; ``pitch``: the fixed
    line pitch (exact line spacing) or None.  From the template contract, or, for a hand-edited
    copy, from the copy's own paragraph settings (its deviations are reported statically)."""
    from audit_helpers import contract, FIXED_WEIGHT
    from customization import ROLE_STYLES
    from style_model import paragraph_properties, run_properties, style_chain

    rules = contract(template, request)
    roles = {style: role for role, style in ROLE_STYLES.items()}
    placed = []
    for p in doc.paragraphs:
        role = roles.get(p.style.name) if p.style is not None else None
        text = _flat(p.text)
        if role not in PLACED or len(text) < 2 or p._p.xpath(".//w:drawing|.//w:pict|.//m:oMath"):
            continue
        if adopted:
            ind, spacing = paragraph_properties(p, "ind"), paragraph_properties(p, "spacing")
            jc = paragraph_properties(p, "jc").get("val", "left")
            runs = [r for r in p._p.xpath(".//w:r") if "".join(r.xpath("./w:t/text()"))]
            styled = [v for st in style_chain(p.style) for v in st.element.xpath("./w:rPr/w:sz/@w:val")]
            first = run_properties(p, runs[0]) if runs else {}
            sizes = {int(first.get("sz") or 21) / 2 if runs else None, int(styled[-1]) / 2 if styled else None, None}
            font_size, bold = (int(first.get("sz") or 21) / 2 if runs else None), first.get("b")
        else:
            rule = rules[role]["pPr"]
            ind, spacing, jc = rule.get("ind", {}), rule.get("spacing", {}), rule.get("jc", {}).get("val", "left")
            sizes = {rules[role]["font"]["size_half_points"] / 2}
            font_size = rules[role]["font"]["size_half_points"] / 2
            bold = True if rules[role]["font"].get("bold_required") else (False if role in FIXED_WEIGHT else None)
        size = max((z for z in sizes if z), default=None)
        exact = spacing.get("lineRule") == "exact" and spacing.get("line")
        placed.append((text, role, {"jc": jc, "start": _starts(ind, sizes), "right": _twips(ind, "right", size),
                                    "size": font_size, "bold": bold,
                                    "pitch": int(spacing["line"]) / PT if exact else None}))
    return placed


def _visual_lines(lines):
    """PDF text lines on one baseline (a renderer splits a line where the script or font changes)
    joined left to right into the line a reader sees."""
    rows = []
    for line in sorted(lines, key=lambda l: (round(l["chars"][0][2]), l["bbox"][0])):
        row = rows[-1] if rows else None
        if row and abs(row["chars"][0][2] - line["chars"][0][2]) < 1 and line["bbox"][0] >= row["bbox"][2] - 1:
            row["spans"][0]["text"] += line["spans"][0]["text"]
            row["chars"] += line["chars"]
            row["latin_bold"] += line.get("latin_bold", [])
            row["bbox"] = (row["bbox"][0], min(row["bbox"][1], line["bbox"][1]), line["bbox"][2], max(row["bbox"][3], line["bbox"][3]))
        else:
            rows.append({"spans": [{"text": line["spans"][0]["text"]}], "chars": list(line["chars"]), "bbox": line["bbox"],
                         "latin_bold": list(line.get("latin_bold", []))})
    return rows


def _placement_issues(index, lines, placed, left, right):
    """First lines that start at the wrong place, centred lines off the axis, and body lines whose
    pitch differs from the fixed line spacing.  Lines are recognised by their text."""
    issues, pitches = [], []
    names = {"body": "正文", "caption": "图表标题", "reference_title": "参考文献标题", "preface": "序言标题",
             "title_zh": "摘要页题目", "title_en": "英文题目", "abstract_label_zh": "“摘要”标题",
             "abstract_label_en": "“Abstract”标题", "toc_title": "“目录”标题"}
    ordered = _visual_lines(lines)
    owner = []
    for line in ordered:
        key = _flat(line["spans"][0]["text"])
        hit = next(((t, r, w) for t, r, w in placed if key and (key == t or (len(key) >= 6 and key in t))), None)
        if hit and owner and owner[-1] and owner[-1] != hit and len(key) >= 6 and key in owner[-1][0]:
            hit = owner[-1]  # a continuation line whose words also occur in an earlier paragraph
        owner.append(hit)
        if not hit:
            continue
        text, role, want = hit
        name = names.get(role, "标题" if role.startswith("h") else role)
        x0, x1 = line["bbox"][0], line["bbox"][2]
        if role != "body" and want["size"]:  # body size is judged over the whole document (BODY_SIZE_RENDERED)
            size = Counter(round(c[3] * 2) / 2 for c in line["chars"]).most_common(1)[0][0]
            if abs(size - want["size"]) > 0.3:
                issues.append(_issue("error", "SIZE_RENDERED", index, f"{name}显示出来的字号不对。", where=line["spans"][0]["text"][:30],
                                     expected=f"{want['size']:g} 磅", actual=f"{size:g} 磅", bbox=line["bbox"]))
        weights = line.get("latin_bold", [])
        if want["bold"] is not None and len(weights) >= 2 and all(w != want["bold"] for w in weights):
            issues.append(_issue("error", "WEIGHT_RENDERED", index, f"{name}显示出来{'没有加粗' if want['bold'] else '被加粗了'}。",
                                 where=line["spans"][0]["text"][:30], expected="加粗" if want["bold"] else "不加粗",
                                 actual="加粗" if weights[0] else "不加粗", bbox=line["bbox"]))
        if want["jc"] == "center" and key == text:
            axis = (left + want["start"][0] + right - want["right"]) / 2
            if abs((x0 + x1) / 2 - axis) > INDENT_TOLERANCE:
                issues.append(_issue("error", "NOT_CENTERED", index, f"{name}应该居中，渲染出来偏了。", where=line["spans"][0]["text"][:30],
                                     expected=f"中线 {axis:.1f}pt", actual=f"中线 {(x0 + x1) / 2:.1f}pt", bbox=line["bbox"]))
        elif want["jc"] in ("left", "both", "start", "distribute") and text.startswith(key) and len(key) >= 4:
            if all(abs(x0 - left - start) > INDENT_TOLERANCE for start in want["start"]):
                start = want["start"][-1]
                chars = start / line["chars"][0][3] if line["chars"][0][3] else 0
                issues.append(_issue("error", "INDENT_RENDERED", index, f"{name}第一行的起点不对（首行缩进或左缩进没有按要求显示）。",
                                     where=line["spans"][0]["text"][:30], expected=f"从 {left + start:.1f}pt 起（缩进约 {chars:.1f} 个字）",
                                     actual=f"从 {x0:.1f}pt 起", bbox=line["bbox"]))
    for (a, ha), (b, hb) in zip(zip(ordered, owner), list(zip(ordered, owner))[1:]):
        if ha and hb and ha[0] == hb[0] and ha[1] == "body" and ha[2]["pitch"] and not hb[0].startswith(_flat(b["spans"][0]["text"])):
            pitches.append((b["chars"][0][2] - a["chars"][0][2], ha[2]["pitch"], b))
    return issues, pitches


def _is_part(line, texts):
    """The line is one of these texts, or the first/last line of one that wrapped."""
    key = _flat(line["spans"][0]["text"])
    return bool(key) and any(key == t or (len(key) >= 6 and (t.startswith(key) or t.endswith(key))) for t in texts)


def _cover_title_misaligned(lines, template):
    """A wrapped cover title whose continuation lines drift: A hangs them under the first
    character of the value, B keeps them on one centre axis.  None when they line up."""
    ordered = sorted(lines, key=lambda l: l["bbox"][1])
    head = next((i for i, l in enumerate(ordered) if re.match(r"^(论文)?题\s*目\s*[:：]", l["spans"][0]["text"])), None)
    if head is None:
        return None
    first = ordered[head]
    text = first["spans"][0]["text"]
    colon = max(text.find("："), text.find(":"))
    if colon < 0 or colon + 1 >= len(first["chars"]):
        return None
    value_x, size = first["chars"][colon + 1][0], first["chars"][-1][3]
    rest, bottom = [], first["bbox"][3]
    for line in ordered[head + 1:]:
        if re.search(r"[:：]", line["spans"][0]["text"]) or line["bbox"][1] - bottom > 1.5 * size:
            break
        rest.append(line)
        bottom = line["bbox"][3]
    if not rest:
        return None
    if template == "A":
        off = [l for l in rest if abs(l["bbox"][0] - value_x) > 3]
        return (off[0], f"续行左端 {off[0]['bbox'][0]:.1f}pt", f"和第一行题目文字左端对齐（{value_x:.1f}pt）") if off else None
    centres = [(l["bbox"][0] + l["bbox"][2]) / 2 for l in rest]
    if max(centres) - min(centres) > 0.5 * size:
        return rest[-1], f"续行中心 {min(centres):.1f}–{max(centres):.1f}pt", "各续行居中在同一条中线上"
    return None


ROMAN = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100}


def _page_number(text):
    """("decimal", 3) / ("roman", 3) for a footer like "3", "- 3 -", "第3页", "III"; None otherwise."""
    core = re.sub(r"[\s\-—第页]", "", text)
    if core.isdigit():
        return "decimal", int(core)
    if core and re.fullmatch(r"[IVXLC]+|[ivxlc]+", core):
        values = [ROMAN[c] for c in core.upper()]
        return "roman", sum(-v if i + 1 < len(values) and v < values[i + 1] else v for i, v in enumerate(values))
    return None


def _number_sequence(numbers):
    """Front matter I, II, ...; body 1, 2, ... to the end: no skip, repeat or stray restart."""
    errors, last = [], None
    for page, text in numbers:
        value = _page_number(text)
        if value is None:  # already reported as PAGE_NUMBER_MISSING; keep counting past it
            last = (last[0], last[1] + 1, text) if last else None
            continue
        kind, n = value
        first = "1" if kind == "decimal" else "I"
        if last is None and n != 1:
            errors.append(_issue("error", "PAGE_NUMBER_SEQUENCE", page, "第一个显示页码的页面没有从 1 开始。",
                                 expected=first, actual=text))
        elif last is not None and last[0] != kind and kind == "roman":
            errors.append(_issue("error", "PAGE_NUMBER_SEQUENCE", page, "正文之后又出现了罗马数字页码。",
                                 expected="阿拉伯数字，接着上一页", actual=text))
        elif last is not None and last[0] != kind and n != 1:
            errors.append(_issue("error", "PAGE_NUMBER_SEQUENCE", page, "换成阿拉伯数字页码时没有从 1 开始。",
                                 expected=first, actual=text))
        elif last is not None and last[0] == kind and n != last[1] + 1:
            errors.append(_issue("error", "PAGE_NUMBER_SEQUENCE", page, "页码不连续（跳号、重复或中途重新开始）。",
                                 expected=f"接着上一页（{last[2]}）往下数", actual=text))
        last = (kind, n, text)
    return errors


def _own_body_size(doc):
    """The body size the document itself sets (most body characters), in points."""
    from style_model import run_properties

    sizes = Counter()
    for p in doc.paragraphs:
        if p.style.name == "XAUFE Body":
            for r in p._p.xpath(".//w:r"):
                text = "".join(r.xpath("./w:t/text()"))
                size = run_properties(p, r).get("sz") if text else None
                if size and size.isdigit():
                    sizes[int(size) / 2] += len(text)
    return sizes.most_common(1)[0][0] if sizes else None


def audit(docx, pdf, adopted=False):
    """Return {"errors": [...], "warnings": [...], "pages": n, "fonts": {...}} for one rendered PDF.

    ``adopted``: a hand-edited copy (``xaufe.py adopt``) whose size changes are reported as
    deviations already, so the rendered body size is held to the size the document sets."""
    import pymupdf as fitz

    doc = Document(docx)
    want = _expectations(doc)
    if adopted and want["body_size"]:
        want["body_size"] = _own_body_size(doc) or want["body_size"]
    placed = _placements(doc, want["template"], want["request"], adopted) if want["template"] in ("A", "B") else []
    pitches = []
    sections = list(doc.sections)
    body_margins = _margins(sections[1:] if want["has_cover"] and len(sections) > 1 else sections)
    cover_margins = _margins(sections[:1])
    header_texts = _header_texts(doc)
    header_paragraphs = _header_paragraphs(doc)
    headings, table_captions, figure_captions = _keep_together(doc)
    # Body lines, recognised by their text, show the size the renderer really used for body text.
    body_texts = [t for t in (_flat(p.text) for p in doc.paragraphs if p.style is not None and p.style.name == "XAUFE Body")
                  if len(t) >= 6]
    body_sizes = Counter()
    errors, warnings = [], []
    used = set()
    previous = ""
    numbers = []
    with fitz.open(pdf) as rendered:
        pages = len(rendered)
        for index, page in enumerate(rendered, 1):
            for font in page.get_fonts():
                used.add(_key(font[3]))
            cover = want["has_cover"] and index == 1
            m = cover_margins if cover else body_margins
            w, h = page.rect.width, page.rect.height
            if abs(w - m["width"]) > 2 or abs(h - m["height"]) > 2:
                errors.append(
                    _issue("error", "PAGE_SIZE", index, "渲染出的纸张大小和文档设置不一致。",
                           expected=f"{m['width']:.0f}×{m['height']:.0f}pt", actual=f"{w:.0f}×{h:.0f}pt")
                )
            left, right = m["left"], w - m["right"]
            top, bottom = m["top"], h - m["bottom"]
            lines = _lines(page)
            images = [i["bbox"] for i in page.get_image_info() if i["bbox"][2] > i["bbox"][0]]
            drawings = [d["rect"] for d in page.get_drawings() if d["rect"].width > 0 or d["rect"].height > 0]

            if not cover and body_texts:
                for line in lines:
                    key = _flat(line["spans"][0]["text"])
                    if len(key) >= 6 and any(key in t for t in body_texts):
                        body_sizes.update(round(c[3] * 2) / 2 for c in line["chars"])
            if not cover and placed:
                column = [l for l in lines if l["bbox"][3] > top and l["bbox"][1] < bottom]
                found, pitch = _placement_issues(index, column, placed, left, right)
                errors += found[:10]
                pitches += [(index, *x) for x in pitch]
                for line in column:
                    key = _flat(line["spans"][0]["text"])
                    if (line["bbox"][1] < top - TEXT_TOLERANCE or line["bbox"][3] > bottom + TEXT_TOLERANCE) and len(key) >= 4 \
                            and any(key in t for t, role, _ in placed):
                        errors.append(_issue("error", "BODY_IN_MARGIN", index, "正文伸进了上下页边距（压到页眉页脚区）。",
                                             where=line["spans"][0]["text"][:30], expected=f"{top:.1f}–{bottom:.1f}pt",
                                             actual=f"{line['bbox'][1]:.1f}–{line['bbox'][3]:.1f}pt", bbox=line["bbox"]))
            if cover:
                drift = _cover_title_misaligned(lines, want["template"])
                if drift:
                    errors.append(_issue("error", "COVER_TITLE_MISALIGNED", index, "封面题目折行后，续行没有和第一行对齐。",
                                         where=drift[0]["spans"][0]["text"][:30], expected=drift[2], actual=drift[1],
                                         bbox=drift[0]["bbox"]))
            outside = None  # the issue for the run of lines now sticking out (one per paragraph, not per line)
            for line in lines:
                x0, y0, x1, y1 = line["bbox"]
                text = line["spans"][0]["text"]
                ink = line.get("ink_right", x1)
                if x0 < -0.5 or y0 < -0.5 or ink > w + 0.5 or y1 > h + 0.5:
                    errors.append(_issue("error", "OFF_PAGE", index, "文字超出纸张边缘，会被裁掉。",
                                         where=text[:30], bbox=line["bbox"]))
                elif cover and (x0 < COVER_SAFE or x1 > w - COVER_SAFE):
                    errors.append(_issue("error", "TEXT_OUTSIDE_MARGINS", index, "封面文字太靠近纸张边缘（不到 1 厘米）。",
                                         where=text[:30], expected=f"{COVER_SAFE:.0f}–{w - COVER_SAFE:.0f}pt",
                                         actual=f"{x0:.1f}–{x1:.1f}pt", bbox=line["bbox"]))
                elif not cover and (y0 < EDGE_SAFE or y1 > h - EDGE_SAFE):
                    errors.append(_issue("error", "TEXT_AT_PAPER_EDGE", index, "页眉、页脚或页码离纸张上下边缘太近，打印时可能被裁掉。",
                                         where=text[:30], expected=f"文字在 {EDGE_SAFE:.0f}–{h - EDGE_SAFE:.0f}pt 之间",
                                         actual=f"{y0:.1f}–{y1:.1f}pt", bbox=line["bbox"]))
                elif not cover and (x0 < left - TEXT_TOLERANCE or ink > right + TEXT_TOLERANCE):
                    if outside is not None:
                        outside["lines"] += 1
                        outside["message"] = f"文字伸出了左右页边距（从这里起连续 {outside['lines']} 行）。"
                        lo, hi = outside["actual"].rstrip("pt").split("–")
                        outside["actual"] = f"{min(float(lo), x0):.1f}–{max(float(hi), x1):.1f}pt"
                        continue
                    outside = _issue("error", "TEXT_OUTSIDE_MARGINS", index, "文字伸出了左右页边距。",
                                     where=text[:30], expected=f"{left:.1f}–{right:.1f}pt",
                                     actual=f"{x0:.1f}–{x1:.1f}pt", bbox=line["bbox"])
                    outside["lines"] = 1
                    errors.append(outside)
                    continue
                outside = None
            for box in images:
                x0, y0, x1, y1 = box
                if x0 < -0.5 or y0 < -0.5 or x1 > w + 0.5 or y1 > h + 0.5:
                    errors.append(_issue("error", "IMAGE_OFF_PAGE", index, "图片超出纸张边缘。", bbox=box))
                elif x0 < left - EDGE_TOLERANCE or x1 > right + EDGE_TOLERANCE:
                    errors.append(_issue("error", "IMAGE_OUTSIDE_MARGINS", index, "图片伸出了左右页边距。",
                                         expected=f"{left:.1f}–{right:.1f}pt", actual=f"{x0:.1f}–{x1:.1f}pt", bbox=box))
                elif not cover and (y0 < top - EDGE_TOLERANCE or y1 > bottom + EDGE_TOLERANCE):
                    errors.append(_issue("error", "IMAGE_OUTSIDE_MARGINS", index, "图片伸进了上下页边距（压到页眉页脚区）。",
                                         expected=f"{top:.1f}–{bottom:.1f}pt", actual=f"{y0:.1f}–{y1:.1f}pt", bbox=box))
                else:
                    text = next((l["spans"][0]["text"] for l in lines
                                 if any(_covers(box, _ink(c)) for c in l["chars"])), None)
                    if text:
                        errors.append(_issue("error", "TEXT_OVER_IMAGE", index, "文字压在图片上。",
                                             where=text[:30], bbox=box))
            for rect in drawings:
                if cover and (rect.x0 < left - EDGE_TOLERANCE or rect.x1 > right + EDGE_TOLERANCE):
                    errors.append(_issue("error", "COVER_OUTSIDE_MARGINS", index, "封面信息栏的下划线（或封面上的线条）伸出了左右页边距。",
                                         expected=f"{left:.1f}–{right:.1f}pt", actual=f"{rect.x0:.1f}–{rect.x1:.1f}pt",
                                         bbox=tuple(rect)))
                    break
                if rect.x0 < left - EDGE_TOLERANCE or rect.x1 > right + EDGE_TOLERANCE:
                    errors.append(_issue("error", "TABLE_OUTSIDE_MARGINS", index, "表格框线或图形伸出了左右页边距。",
                                         expected=f"{left:.1f}–{right:.1f}pt", actual=f"{rect.x0:.1f}–{rect.x1:.1f}pt",
                                         bbox=tuple(rect)))
                    break

            # Header and footer: outside the text column, or (margins narrowed below the header/footer
            # distance plus a line) a short line inside the column that is not the document's own text.
            band_top = max(top + 1, m["header"] + HEADER_ROOM) if m["header"] else top + 1
            band_bottom = min(bottom - 1, h - m["footer"] - HEADER_ROOM) if m["footer"] else bottom - 1
            own = lambda l: len(_flat(l["spans"][0]["text"])) >= 4 and any(
                _flat(l["spans"][0]["text"]) in t for t, _, _ in placed)
            head = [l for l in lines if l["bbox"][3] <= top + 1 or (l["bbox"][3] <= band_top and not own(l))]
            foot = [l for l in lines if l["bbox"][1] >= bottom - 1 or (l["bbox"][1] >= band_bottom and not own(l))]
            body = [l for l in lines if l["bbox"][3] > top and l["bbox"][1] < bottom and l not in head and l not in foot]
            # Rules, table borders and shapes count as content too; a lone header rule does not.
            body_objects = [b for b in images if b[3] > top and b[1] < bottom] + [
                r for r in drawings if r.y1 > top + 2 and r.y0 < bottom - 2 and (r.height > 2 or r.y0 > top + 30)]
            if not body and not body_objects:
                errors.append(_issue("error", "BLANK_PAGE", index, "这一页正文区什么都没有（多余的空白页）。",
                                     where=("紧跟在“" + previous + "”之后") if previous else None))
            elif body:
                previous = body[-1]["spans"][0]["text"][-20:]

            if want["header"] is not None:
                if cover and head:
                    errors.append(_issue("error", "COVER_HEADER_VISIBLE", index, "封面页出现了页眉文字。",
                                         where=head[0]["spans"][0]["text"][:30]))
                elif not cover and want["header"] and not head:
                    errors.append(_issue("error", "HEADER_MISSING", index, "这一页应该有页眉，渲染结果里没有。",
                                         expected="页眉文字", actual="（没有）"))
                elif not cover and want["header"]:
                    spill = _header_overflow(lines, top, header_texts)
                    if spill:
                        errors.append(_issue("error", "HEADER_OVERFLOW", index, "页眉文字折成多行，伸出上边距压进了正文区。",
                                             where=spill["spans"][0]["text"][:30], expected=f"页眉在 {top:.1f}pt 以上",
                                             actual=f"页眉底边 {spill['bbox'][3]:.1f}pt", bbox=spill["bbox"]))
                    else:
                        rows = _header_rows(head, header_texts)
                        if len(rows) > max(header_paragraphs, 1):
                            errors.append(_issue("error", "HEADER_WRAPPED", index, "页眉一行放不下，折成了两行（还没压进正文，但已经不是一行页眉）。",
                                                 where=rows[-1][1]["spans"][0]["text"][:30],
                                                 expected=f"页眉在一行内（{max(header_paragraphs, 1)} 行）",
                                                 actual=f"{len(rows)} 行", bbox=rows[-1][1]["bbox"]))
                elif not cover and not want["header"] and head:
                    errors.append(_issue("error", "HEADER_UNEXPECTED", index, "这一页不应该有页眉，渲染结果里有。",
                                         where=head[0]["spans"][0]["text"][:30]))
            if want["page_numbers"] is not None:
                texts = "".join(l["spans"][0]["text"] for l in foot).strip()
                if cover and texts:
                    errors.append(_issue("error", "COVER_FOOTER_VISIBLE", index, "封面页出现了页脚文字或页码。", where=texts[:30]))
                elif not cover and want["page_numbers"] and not re.fullmatch(r"[0-9IVXLCivxlc\-—\s第页]+", texts or "?"):
                    errors.append(_issue("error", "PAGE_NUMBER_MISSING", index, "这一页页脚应该显示页码，渲染结果里没有。",
                                         expected="页脚里的页码", actual=texts[:30] or "（页脚是空的）"))
                elif not cover and not want["page_numbers"] and texts:
                    errors.append(_issue("error", "FOOTER_UNEXPECTED", index, "页脚应该是空白，渲染结果里有文字。", where=texts[:30]))

            if body and not cover and index < pages:
                last = max(body, key=lambda l: l["bbox"][3])
                if _is_part(last, headings):
                    errors.append(_issue("error", "HEADING_AT_PAGE_BOTTOM", index, "标题落在页面最底下，它的正文到了下一页。",
                                         where=last["spans"][0]["text"][:30], bbox=last["bbox"]))
                elif _is_part(last, table_captions):
                    errors.append(_issue("error", "CAPTION_SEPARATED", index, "表题在这一页最底下，表格到了下一页。",
                                         where=last["spans"][0]["text"][:30], bbox=last["bbox"]))
            if body and not cover:
                first = min(body, key=lambda l: l["bbox"][1])
                if _is_part(first, figure_captions) and not any(b[3] <= first["bbox"][1] + 1 for b in images):
                    errors.append(_issue("error", "CAPTION_SEPARATED", index, "图题在这一页最上面，图片还在上一页。",
                                         where=first["spans"][0]["text"][:30], bbox=first["bbox"]))
            # Every clash on the page, so one repair round can fix them all.
            clashes = [(a, b, gap) for i, a in enumerate(lines) for b in lines[i + 1 :]
                       for gap in [_line_clash(a, b)] if gap]
            for a, b, (gap, overlap) in clashes[:10]:
                errors.append(_issue("error", "TEXT_OVERLAP", index, "两行文字在页面上叠在一起。",
                                     where=a["spans"][0]["text"][:20] + " / " + b["spans"][0]["text"][:20],
                                     expected="上下两行的字形互不重叠（同一字号时基线至少相距约 0.9 倍字号）",
                                     actual=f"基线相距 {gap:.1f}pt，字形重叠 {overlap:.1f}pt", bbox=a["bbox"]))
            if want["page_numbers"] and not cover:
                numbers.append((index, "".join(l["spans"][0]["text"] for l in foot).strip()))
    errors += _number_sequence(numbers)
    wrong = [(page, gap, pitch, line) for page, gap, pitch, line in pitches if abs(gap - pitch) > 0.6]
    if pitches and len(wrong) * 2 > len(pitches):
        page, gap, pitch, line = wrong[0]
        errors.append(_issue("error", "LINE_PITCH_RENDERED", page, "正文的行距显示出来和要求不一致（固定行距没有生效）。",
                             where=line["spans"][0]["text"][:30], expected=f"相邻两行基线相距 {pitch:g} 磅",
                             actual=f"{gap:.1f} 磅（{len(wrong)}/{len(pitches)} 处不对）", bbox=line["bbox"]))
    if want.get("body_size") and sum(body_sizes.values()) >= 40:
        size = body_sizes.most_common(1)[0][0]
        if abs(size - want["body_size"]) > 0.3:
            errors.append(_issue("error", "BODY_SIZE_RENDERED", None, "正文实际显示出来的字号和要求不一致（Word 文件里的样式设置没有生效）。",
                                 expected=f"{want['body_size']:g} 磅", actual=f"{size:g} 磅"))
    wanted = requested_fonts(doc)
    missing = [f for f in wanted if not _font_present(f, used)]
    if missing:
        warnings.append(
            _issue(
                "warning",
                "FONT_SUBSTITUTED",
                None,
                "渲染这台机器上没有文档要求的字体（"
                + "、".join(missing)
                + "），页图用了替代字体。字宽和分页可能和装有这些字体的 Word 不同，页图只能当近似参考；"
                "Word 文件本身的字体设置没有变。",
                expected="、".join(missing),
                actual="、".join(sorted(used))[:200],
            )
        )
    return {
        "passed": not errors,
        "pages": pages,
        "errors": errors,
        "warnings": warnings,
        "fonts": {"requested": wanted, "missing": missing, "rendered": sorted(used)},
    }


def main():
    p = argparse.ArgumentParser(description="检查实际渲染出的 PDF 版面（越界、空白页、页眉页脚、重叠、字体替代）")
    p.add_argument("docx")
    p.add_argument("pdf")
    p.add_argument("--report")
    a = p.parse_args()
    result = audit(a.docx, a.pdf)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if a.report:
        Path(a.report).write_text(text, encoding="utf-8")
    print(text)
    return 0 if result["passed"] else 2


if __name__ == "__main__":
    sys.exit(main())
