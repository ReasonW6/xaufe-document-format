"""Deterministic XAUFE Word layout. Only formatting; no research or prose generation."""

from __future__ import annotations
import hashlib, json, re, shutil
from copy import deepcopy
from pathlib import Path
from typing import Any
from zipfile import ZipFile
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.section import WD_SECTION_START
from docx.shared import Pt, Twips, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
from profiles import load_profile
from textio import read_text_any
from customization import (
    effective_profile,
    apply_cover,
    normalize_style_options,
    ROLE_STYLES,
    fingerprint,
)
from intake import require as require_intake, digest
from table_layout import fit_table, resize_inline
from keywords import normalize as normalize_keywords, split_label_runs
from style_model import run_properties, paragraph_properties, on
from cover_date import apply as position_cover_date

SPEC: dict = {}
WIDTH = 0


def select_profile(template: str, request=None) -> dict:
    """A process formats one template at a time; CLI always binds the explicit intake choice."""
    global WIDTH
    SPEC.clear()
    SPEC.update(effective_profile(template, request))
    WIDTH = SPEC["page_twips"]["width"] - SPEC["page_twips"]["left"] - SPEC["page_twips"]["right"]
    return SPEC


def require_profile():
    if not SPEC.get("id"):
        fail(
            "还没有记录用户选择的版式（A 或 B）。"
        )


FONT_CN, FONT_H, FONT_EN = "宋体", "黑体", "Times New Roman"
ROLES = ROLE_STYLES


def fail(message: str) -> None:
    raise ValueError(message)


def load_json(path: str | Path) -> dict:
    path = Path(path)
    try:
        return json.loads(read_text_any(path))
    except (OSError, json.JSONDecodeError) as exc:
        fail(f"无法读取 JSON {path}: {exc}")


def asset_check() -> None:
    require_profile()
    f = ROOT / SPEC["logo"]["file"]
    if not f.is_file() or hashlib.sha256(f.read_bytes()).hexdigest() != SPEC["logo"]["sha256"]:
        fail("原校徽文件缺失或哈希不匹配。禁止下载替代品、重绘或用文本代替。")
    master = ROOT / SPEC["master_file"]
    if not master.is_file():
        fail("格式母版缺失。请安装整个技能目录。")
    if (
        SPEC.get("master_sha256")
        and hashlib.sha256(master.read_bytes()).hexdigest() != SPEC["master_sha256"]
    ):
        fail("格式母版被改动；请恢复原母版，不凭记忆修复。")


def element(parent, tag, **attrs):
    e = parent.find(qn(tag))
    if e is None:
        e = OxmlElement(tag)
        parent.append(e)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def font(rpr, cn=FONT_CN, size=10.5, bold=False, latin=FONT_EN):
    rf = element(rpr, "w:rFonts", ascii=latin, hAnsi=latin, cs=latin, eastAsia=cn, hint="eastAsia")
    for key in list(rf.attrib):
        if key.lower().endswith("theme"):
            del rf.attrib[key]
    element(rpr, "w:sz", val=round(size * 2))
    element(rpr, "w:szCs", val=round(size * 2))
    element(rpr, "w:b", val=int(bold))
    element(rpr, "w:bCs", val=int(bold))
    color = element(rpr, "w:color")
    color.attrib.clear()
    color.set(qn("w:val"), "000000")
    element(rpr, "w:w", val=100)
    element(rpr, "w:spacing", val=0)
    element(rpr, "w:position", val=0)
    element(rpr, "w:vanish", val=0)
    element(rpr, "w:webHidden", val=0)


def apply_text_options(rp, options):
    if "italic" in options:
        element(rp, "w:i", val=int(options["italic"]))
        element(rp, "w:iCs", val=int(options["italic"]))
    if "color" in options:
        node = element(rp, "w:color")
        node.attrib.clear()
        node.set(qn("w:val"), options["color"].upper())
    if "underline" in options:
        element(rp, "w:u", val=options["underline"])


def run_font(run, cn=FONT_CN, size=10.5, bold=False):
    font(run._r.get_or_add_rPr(), cn, size, bold)


def configure_style(
    doc,
    name,
    *,
    cn=FONT_CN,
    size=10.5,
    bold=False,
    align="both",
    first=0,
    chars=None,
    left=0,
    before=0,
    after=0,
    line=400,
    before_lines=None,
    after_lines=None,
    outline=None,
    keep=False,
    break_before=False,
    line_rule="exact",
    hanging=None,
    left_chars=None,
    latin=FONT_EN,
    italic=None,
    color="000000",
    underline=None,
    right=0,
):
    st = (
        doc.styles[name]
        if name in doc.styles
        else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    )
    st.base_style = None
    ppr = st.element.get_or_add_pPr()
    for e in list(ppr):
        ppr.remove(e)
    element(ppr, "w:jc", val=align)
    ind = element(
        ppr,
        "w:ind",
        left=left if left_chars is None else round(size * 20 * left_chars / 100),
        right=right,
        firstLine=first if chars is None else round(size * 20 * chars / 100),
    )
    if chars is not None:
        ind.set(qn("w:firstLineChars"), str(chars))
    if hanging is not None:
        ind.attrib.pop(qn("w:firstLine"), None)
        ind.set(qn("w:hanging"), str(hanging))
    if left_chars is not None:
        ind.set(qn("w:leftChars"), str(left_chars))
    spacing = element(ppr, "w:spacing", before=before, after=after, line=line, lineRule=line_rule)
    if before_lines is not None:
        spacing.set(qn("w:beforeLines"), str(before_lines))
    if after_lines is not None:
        spacing.set(qn("w:afterLines"), str(after_lines))
    element(ppr, "w:keepNext", val=int(keep))
    element(ppr, "w:keepLines", val=int(keep))
    element(ppr, "w:widowControl", val=1)
    element(ppr, "w:pageBreakBefore", val=int(break_before))
    element(ppr, "w:snapToGrid", val=0)
    element(ppr, "w:kinsoku", val=1)
    element(ppr, "w:overflowPunct", val=0)
    # Avoid inherited outline levels bringing the cover/abstract into the TOC.
    element(ppr, "w:outlineLvl", val=outline if outline is not None else 9)
    font(st.element.get_or_add_rPr(), cn, size, bold, latin=latin)
    apply_text_options(
        st.element.get_or_add_rPr(),
        {
            "italic": False if italic is None else italic,
            "color": color,
            "underline": "none" if underline is None else underline,
        },
    )
    if name.startswith("TOC "):
        tabs = element(ppr, "w:tabs")
        t = OxmlElement("w:tab")
        t.set(qn("w:val"), "right")
        t.set(qn("w:leader"), "dot")
        t.set(qn("w:pos"), str(WIDTH))
        tabs.append(t)
    return st


def builtin_style_names(doc):
    """LibreOffice and WPS write built-in styles as "Heading 1"; Word and python-docx look them up as
    "heading 1".  Without this, the template's Heading 1 would be added a second time with the same
    styleId and Word would pick either one."""
    from docx.styles import BabelFish

    internal = {i for _, i in BabelFish.style_aliases}
    for name in doc.styles.element.xpath("./w:style/w:name"):
        value = name.get(qn("w:val")) or ""
        if value != value.lower() and value.lower() in internal:
            name.set(qn("w:val"), value.lower())


def install_styles(doc):
    require_profile()
    builtin_style_names(doc)
    # Some masters have no default paragraph style; then an unstyled paragraph
    # has style None. Make the existing Normal style the default.
    from docx.enum.style import WD_STYLE_TYPE as _T

    if doc.styles.default(_T.PARAGRAPH) is None and "Normal" in doc.styles:
        doc.styles["Normal"].element.set(qn("w:default"), "1")
    # Imported .doc templates may not contain Word's built-in Table Grid style.
    if "Table Grid" not in doc.styles:
        table_style = doc.styles.add_style("Table Grid", WD_STYLE_TYPE.TABLE)
        borders = element(element(table_style.element, "w:tblPr"), "w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element(borders, "w:" + edge, val="single", sz=4, color="000000")
    rules = load_json(ROOT / SPEC.get("paragraph_rules", "assets/paragraph-rules.json"))["styles"]
    for name, params in rules.items():
        options = normalize_style_options(params, SPEC.get("style_overrides", {}).get(name, {}))
        if options.get("left") == "$toc_indent":
            options["left"] = SPEC["toc"]["indents_twips"][int(name[-1]) - 1]
        configure_style(doc, name, **options)


def _default_role_format(role):
    if role in ("h1", "preface"):
        return FONT_H, 14, False
    if role == "h2":
        return FONT_H, 12, False
    if role in ("h3", "h4"):
        return FONT_H, 10.5, False
    if role in ("title_zh", "toc_title", "reference_title"):
        return FONT_H, 16, False
    if role == "title_en":
        return FONT_EN, 16, True
    if role == "abstract_label_zh":
        return FONT_H, 14, False
    if role == "abstract_label_en":
        return FONT_EN, 14, True
    if role == "abstract_en":
        return FONT_EN, 10.5, False
    if role == "keywords_zh":
        return FONT_H, 10.5, False
    if role == "keywords_en":
        return FONT_EN, 10.5, True
    if role.startswith("toc") and role[-1:].isdigit():
        return FONT_H if role == "toc1" else FONT_CN, 14, False
    if role == "header":
        return FONT_CN, 9, False
    return FONT_CN, 10.5, False


def role_format(role):
    cn, size, bold = _default_role_format(role)
    patch = SPEC.get("style_overrides", {}).get(ROLES.get(role, ""), {})
    return patch.get("cn", cn), patch.get("size", size), patch.get("bold", bold)


def role_latin(role):
    return SPEC.get("style_overrides", {}).get(ROLES.get(role, ""), {}).get("latin", FONT_EN)


def style_paragraph(p, role, *, preserve_num=False, style_name=None):
    if role not in ROLES:
        fail(f"不认识的段落角色：{role}")
    if role == "cover":
        return
    old_emphasis = {r: run_properties(p, r) for r in p._p.xpath(".//w:r")}
    ppr = p._p.get_or_add_pPr()
    # Preserve section breaks, existing numbering and tab stops (e.g. a real table of contents).
    keep_tags = {qn("w:sectPr"), qn("w:numPr"), qn("w:tabs")}
    for child in list(ppr):
        if child.tag not in keep_tags:
            ppr.remove(child)
    p.style = style_name or ROLES[role]
    if not preserve_num:
        for e in ppr.findall(qn("w:numPr")):
            ppr.remove(e)
    cn, size, bold = role_format(role)
    # Descend into hyperlinks without rebuilding text, fields, formulas or image relationships.
    for r in p._p.xpath(".//w:r"):
        rp = r.get_or_add_rPr()
        previous = old_emphasis.get(r, {})
        preserve_bold = (
            role in ("body", "reference", "table_text", "code")
            and previous.get("b", False)
            and "bold" not in SPEC.get("style_overrides", {}).get(ROLES[role], {})
        )
        if previous.get("vanish") or previous.get("webHidden") or previous.get("specVanish"):
            fail("原文含隐藏文字；必须先明确是否显示或保留，不自动揭示或删除隐藏内容。")
        font(rp, cn, size, bold or preserve_bold, latin=role_latin(role))
        # Look-only decorations are not part of the template; marks that can carry meaning
        # (strike, caps, emphasis dots) stay and are reported by the check instead.
        for tag in ("w:highlight", "w:shd", "w:emboss", "w:imprint", "w:outline", "w:shadow", "w:effect"):
            for node in rp.findall(qn(tag)):
                rp.remove(node)
        if (
            role in ("body", "reference", "table_text", "code")
            and previous.get("i")
            and "italic" not in SPEC.get("style_overrides", {}).get(ROLES[role], {})
        ):
            element(rp, "w:i", val=1)
        apply_text_options(rp, SPEC.get("style_overrides", {}).get(ROLES[role], {}))
    if role in ("keywords_zh", "keywords_en"):
        normalize_keywords(p)
        from customization import keyword_label_options

        for r, is_label in split_label_runs(p):
            opts = (
                keyword_label_options(SPEC, role)
                if is_label
                else dict(
                    cn=cn,
                    size=size,
                    bold=bold,
                    latin=role_latin(role),
                    **{
                        k: v
                        for k, v in SPEC.get("style_overrides", {}).get(ROLES[role], {}).items()
                        if k in ("italic", "color", "underline")
                    },
                )
            )
            rp = r.get_or_add_rPr()
            font(rp, opts["cn"], opts["size"], opts["bold"], latin=opts["latin"])
            apply_text_options(rp, opts)


def style_selected(doc, roles):
    """Apply exact role formats to selected top-level paragraphs only.

    Copy isolated role styles instead of overwriting the source's Normal or
    Heading styles, which could restyle unselected titles and table cells.
    """
    import uuid

    selected = []
    for key, role in roles.items():
        match = re.fullmatch(r"p(\d+)", str(key))
        if (
            not match
            or int(match.group(1)) >= len(doc.paragraphs)
            or role not in ROLES
            or role == "cover"
        ):
            fail("局部映射必须使用实际清单中的p索引与受支持的非封面角色。")
        selected.append((int(match.group(1)), role))
    before = content_signature(doc)
    reference = Document()
    install_styles(reference)
    aliases = {}
    for role in dict.fromkeys(role for _, role in selected):
        name = "XAUFE Local " + role + " " + uuid.uuid4().hex[:8]
        local = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        source = reference.styles[ROLES[role]].element
        for tag in ("pPr", "rPr"):
            node = source.find(qn("w:" + tag))
            if node is not None:
                local.element.append(deepcopy(node))
        aliases[role] = name
    for index, role in selected:
        style_paragraph(doc.paragraphs[index], role, preserve_num=True, style_name=aliases[role])
    if content_signature(doc) != before:
        fail("局部排版改变了原始内容或对象，不能保存该结果。")
    return {"formatted_paragraphs": ["p" + str(index) for index, _ in selected]}


def set_page(doc):
    m = SPEC["page_twips"]
    for s in doc.sections:
        for attr, key in [
            ("page_width", "width"),
            ("page_height", "height"),
            ("top_margin", "top"),
            ("bottom_margin", "bottom"),
            ("left_margin", "left"),
            ("right_margin", "right"),
            ("header_distance", "header"),
            ("footer_distance", "footer"),
            ("gutter", "gutter"),
        ]:
            setattr(s, attr, Twips(m[key]))
        s._sectPr.get_or_add_pgSz().set(
            qn("w:orient"), "landscape" if m["width"] > m["height"] else "portrait"
        )
        # No odd/even section breaks: no accidental blank sheets in single-sided printing.
        t = s._sectPr.find(qn("w:type"))
        if t is not None and t.get(qn("w:val")) in ("oddPage", "evenPage"):
            t.set(qn("w:val"), "nextPage")


def empty_part(part):
    for child in list(part._element):
        part._element.remove(child)
    part._element.append(OxmlElement("w:p"))


def drop_unused_images(doc, candidates):
    """Drop only replaced image relationships with no remaining document use."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    used = {
        value
        for el in doc.element.iter()
        for name, value in el.attrib.items()
        if name.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}")
    }
    for rid in candidates - used:
        rel = doc.part.rels.get(rid)
        if rel is not None and rel.reltype == RT.IMAGE:
            doc.part.drop_rel(rid)


def headers_footers(doc, meta, has_cover):
    left = SPEC["header"].get("left", (meta.get("major") or "＿＿") + "专业本科学年论文")
    right = SPEC["header"].get("right", meta.get("title_zh", ""))
    for i, s in enumerate(doc.sections):
        s.different_first_page_header_footer = False
        for f in [s.footer, s.first_page_footer, s.even_page_footer]:
            f.is_linked_to_previous = False
            empty_part(f)
            if SPEC["footer"].get("page_numbers") and not (i == 0 and has_cover):
                p = f.paragraphs[0]
                p.style = ROLES["header"]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                field = OxmlElement("w:fldSimple")
                field.set(qn("w:instr"), " PAGE ")
                r = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.text = "1"
                r.append(t)
                field.append(r)
                p._p.append(field)
        for h in [s.header, s.first_page_header, s.even_page_header]:
            h.is_linked_to_previous = False
            empty_part(h)
            if not SPEC["header"]["enabled"] or (i == 0 and has_cover):
                continue
            cn, size, bold = role_format("header")
            size = SPEC.get("style_overrides", {}).get(ROLES["header"], {}).get("size", size)
            import unicodedata

            # Estimate only the line choice; actual pagination remains authoritative.
            extent = lambda text: sum(
                size * 20 * (1 if unicodedata.east_asian_width(c) in ("W", "F") else 0.6)
                for c in text
            )
            stacked = extent(left) + extent(right) > WIDTH
            texts = [left, right] if stacked else [left + "\t" + right]
            if stacked:
                # A text longer than one line is split into even lines instead
                # of leaving a few characters alone on the last line.
                from cover_fit import balanced_lines

                texts = ["\n".join(balanced_lines(t, size, WIDTH)) for t in texts]
            for index, text in enumerate(texts):
                p = h.paragraphs[0] if index == 0 else h.add_paragraph()
                p.style = ROLES["header"]
                if stacked:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if index else WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.paragraph_format.tab_stops.add_tab_stop(Twips(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
                r = p.add_run(text)
                font(r._r.get_or_add_rPr(), cn, size, bold, latin=role_latin("header"))
                apply_text_options(
                    r._r.get_or_add_rPr(), SPEC.get("style_overrides", {}).get(ROLES["header"], {})
                )
            if SPEC["header"].get("border", True):
                bd = element(p._p.get_or_add_pPr(), "w:pBdr")
                element(bd, "w:bottom", val="single", sz=6, space=1, color="000000")
    # Prevent stale previous variant settings from selecting an unprocessed footer.
    for x in doc.settings.element.findall(qn("w:evenAndOddHeaders")):
        doc.settings.element.remove(x)


def set_vars(doc, values):
    vs = element(doc.settings.element, "w:docVars")
    for key, val in values.items():
        for old in list(vs):
            if old.get(qn("w:name")) == key:
                vs.remove(old)
        e = OxmlElement("w:docVar")
        e.set(qn("w:name"), key)
        e.set(qn("w:val"), str(val))
        vs.append(e)


def get_vars(doc):
    return {
        e.get(qn("w:name")): e.get(qn("w:val"))
        for e in doc.settings.element.xpath("./w:docVars/w:docVar")
    }


def cover_metadata(meta, state):
    result = dict(meta)
    for key in state.get("blank_fields", []):
        result[key] = ""
    return result


def merge_metadata(data_meta, state):
    result = dict(data_meta or {})
    allowed = set(SPEC["cover"]["required_fields"]) | (
        {"title_en"} if SPEC["english_abstract"] else set()
    )
    if set(result) - allowed:
        fail("元数据含不属于所选模板的字段：" + ",".join(sorted(set(result) - allowed)))
    for key, value in state["metadata"].items():
        if key in result and result[key] != value:
            fail(
                f"content.json 里的 {key} 和 job.json 的封面信息不一致；封面信息只写在 job.json 的 info 里。"
            )
        result[key] = value
    if "supervisor" not in SPEC["cover"]["fields"] and result.get("supervisor"):
        fail(
            "本次有效格式未启用指导教师栏。用户已要求增加时，先记录cover_supervisor=true；只换校徽不增减字段。"
        )
    return result


def fill_cover(doc, meta):
    """Fill the cover of a fresh master copy; long text is wrapped and fitted."""
    import cover_fit

    apply_cover(doc, SPEC)
    ps = doc.paragraphs
    if len(ps) < SPEC["cover"]["paragraph_count"]:
        fail("母版封面结构已损坏。")
    ps[SPEC["cover"]["label_paragraph"]].runs[0].text = SPEC["cover"]["document_label"]
    values = {key: meta.get(key, "") for key in SPEC["cover"]["fields"]}
    values["title_zh"] = meta.get("title_zh", "")
    values["course_name"] = meta.get("course_name", "")
    layout = cover_fit.plan(doc, SPEC, values, min_level=SPEC.get("_cover_min_level", 0))
    if not layout["fits"]:
        fail(
            "封面文字太长，第一页放不下（已收紧空白和行距，仍约差"
            + str(round(layout["extra_pt"] - cover_fit.capacity(SPEC, layout["level"])))
            + "磅）。请问用户：缩短题目、课程名或过长的封面项，或同意缩小封面字号。"
        )
    cover_fit.apply(doc, SPEC, layout)
    SPEC["_cover_fit"] = layout["level"]
    date = meta.get("date", "")
    if date and re.fullmatch(r"\d{4}-\d{2}-\d{2}", date):
        y, m, d = map(int, date.split("-"))
        date = f"{y}年{m}月{d}日"
    dp = ps[SPEC["cover"]["date_paragraph"]]
    if SPEC["cover"].get("date_value_run") is not None:
        dp.runs[0].text = "完成日期："
        dp.runs[SPEC["cover"]["date_value_run"]].text = date or "      年   月   日"
    else:
        dp.runs[0].text = "完成日期：" + (date or "      年   月   日")
    position_cover_date(doc, SPEC)


def safe_out(input_path, output, force=False):
    output = Path(output).resolve()
    if input_path and Path(input_path).resolve() == output:
        fail("禁止原地覆盖输入文件；请另存输出。")
    if output.exists() and not force:
        fail(f"输出已存在：{output}。另取文件名，或明确使用 --force。")
    output.parent.mkdir(parents=True, exist_ok=True)
    return output


FRONT_ROLES = {"title_zh", "title_en", "abstract_label_zh", "abstract_label_en", "abstract_zh", "abstract_en",
               "keywords_zh", "keywords_en", "toc_title", "toc1", "toc2", "toc3", "toc4"}
MAIN_ROLES = {"preface", "h1", "h2", "h3", "h4", "body", "reference_title", "reference"}


def number_sections(doc):
    """Front matter (abstracts, TOC) numbered I, II, ...; body from 1, continuing to the end."""
    names = {name: role for role, name in ROLES.items()}
    current, front_seen, main_seen = set(), False, False
    for el in doc.element.body:
        if el.tag == qn("w:p"):
            current.add(names.get(Paragraph(el, doc._body).style.name))
            sect = el.find(qn("w:pPr") + "/" + qn("w:sectPr"))
        else:
            sect = el if el.tag == qn("w:sectPr") else None
        if sect is None:
            continue
        if main_seen or current & MAIN_ROLES:
            page = element(sect, "w:pgNumType", fmt="decimal")
            page.attrib.pop(qn("w:start"), None)
            if not main_seen:
                page.set(qn("w:start"), "1")
            main_seen = True
        elif current & FRONT_ROLES:
            page = element(sect, "w:pgNumType", fmt="upperRoman")
            page.attrib.pop(qn("w:start"), None)
            if not front_seen:
                page.set(qn("w:start"), "1")
            front_seen = True
        current = set()


def number_plain_sections(doc, has_cover):
    """A document without paper structure: arabic numbers from 1 after the cover, continuing to the end."""
    for i, section in enumerate(doc.sections):
        if has_cover and i == 0:
            continue
        page = element(section._sectPr, "w:pgNumType", fmt="decimal")
        page.attrib.pop(qn("w:start"), None)
        if i == (1 if has_cover else 0):
            page.set(qn("w:start"), "1")


def finalize(doc, out, meta, has_cover, mode, state=None):
    if state and "include_cover" in state and state["include_cover"] != bool(has_cover):
        fail(
            "文档是否含封面与本次明确范围不同。无封面新稿设置cover=false，已有稿按明确范围处理并更新has_cover；不能只少填信息却保留空封面。"
        )
    set_page(doc)
    headers_footers(doc, meta, has_cover)
    front = {ROLES["toc_title"], ROLES["abstract_label_zh"], ROLES["abstract_label_en"]}
    if mode == "paper" or any(p.style is not None and p.style.name in front for p in doc.paragraphs):
        number_sections(doc)  # also an earlier output of this skill revised as a "document"
    else:
        number_plain_sections(doc, has_cover)
    if has_cover:
        apply_cover(doc, SPEC)
        position_cover_date(doc, SPEC)
    set_vars(
        doc,
        {
            "XAUFE_MODE": mode,
            "XAUFE_HAS_COVER": int(has_cover),
            "XAUFE_COVER_FIT": SPEC.get("_cover_fit", 0) if has_cover else 0,
            "XAUFE_FORMAT_VERSION": "1.0.0",
            "XAUFE_TEMPLATE": SPEC["id"],
        },
    )
    if state:
        set_vars(
            doc,
            {
                "XAUFE_WORKFLOW": "ready",
                "XAUFE_INTAKE_ID": state["job_id"],
                "XAUFE_INTAKE_HASH": state["integrity"],
                "XAUFE_COVER_METADATA_HASH": digest(state["metadata"]),
                "XAUFE_COVER_BLANK_FIELDS": json.dumps(
                    state.get("blank_fields", []), ensure_ascii=False
                ),
                "XAUFE_EFFECTIVE_DATE": state["metadata"].get("date", ""),
                "XAUFE_DATE_SOURCE": state.get("provenance", {})
                .get("date", {})
                .get("source", "blank"),
            },
        )
    set_vars(
        doc,
        {
            "XAUFE_FORMAT_REQUEST": json.dumps(
                SPEC.get("_format_request", {}), ensure_ascii=False, sort_keys=True
            ),
            "XAUFE_FORMAT_REQUEST_HASH": fingerprint(SPEC.get("_format_request", {})),
        },
    )
    element(doc.settings.element, "w:updateFields", val="false")
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.core_properties.title = meta.get("title_zh", "")
    for p in doc.paragraphs:
        role = next((k for k, v in ROLES.items() if v == p.style.name), None)
        if role is None or role == "cover":
            continue
        patch = SPEC.get("style_overrides", {}).get(ROLES[role], {})
        for r in p._p.xpath(".//w:r"):
            if "bold" in patch:
                element(r.get_or_add_rPr(), "w:b", val=int(patch["bold"]))
                element(r.get_or_add_rPr(), "w:bCs", val=int(patch["bold"]))
            apply_text_options(r.get_or_add_rPr(), patch)
        if role.startswith("keywords_"):
            from customization import keyword_label_options

            for r, is_label in split_label_runs(p):
                if is_label:
                    opts = keyword_label_options(SPEC, role)
                    rp = r.get_or_add_rPr()
                    font(rp, opts["cn"], opts["size"], opts["bold"], latin=opts["latin"])
                    apply_text_options(rp, opts)
    from toc_format import enforce

    enforce(doc, SPEC)
    drop_unused_images(doc, set(doc.part.rels))
    doc.save(out)


def add_section(doc, *, start=None):
    s = doc.add_section(WD_SECTION_START.NEW_PAGE)
    # The empty paragraph carrying the section break must not take a full body line.
    p = doc.paragraphs[-1]
    p.style = ROLES["spacer"]
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    # add_section clones the prior section, including a restart. None means
    # continue, so explicitly remove the copied start rather than inherit it.
    page = element(s._sectPr, "w:pgNumType", fmt="decimal")
    page.attrib.pop(qn("w:start"), None)
    if start is not None:
        page.set(qn("w:start"), str(start))
    return s


def bookmark(p, name, bid):
    a = OxmlElement("w:bookmarkStart")
    a.set(qn("w:id"), str(bid))
    a.set(qn("w:name"), name)
    b = OxmlElement("w:bookmarkEnd")
    b.set(qn("w:id"), str(bid))
    idx = 1 if p._p.pPr is not None else 0
    p._p.insert(idx, a)
    p._p.append(b)


def add_text(p, b):
    if "runs" in b:
        for item in b["runs"]:
            r = p.add_run(item["text"])
            if item.get("link"):
                # A real external hyperlink around the run; the run keeps the paragraph's formatting.
                from docx.opc.constants import RELATIONSHIP_TYPE as RT

                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), p.part.relate_to(item["link"], RT.HYPERLINK, is_external=True))
                r._r.addprevious(link)
                link.append(r._r)
            if item.get("superscript"):
                r.font.superscript = True
            if item.get("subscript"):
                r.font.subscript = True
            if item.get("italic"):
                r.italic = True
            if item.get("bold"):
                r.bold = True
    else:
        p.add_run(b.get("text", ""))


def add_toc(doc, entries):
    p = doc.add_paragraph("目　录", ROLES["toc_title"])
    for i, (name, text, level) in enumerate(entries):
        p = doc.add_paragraph(style=f"TOC {level}")
        if i == 0:
            r = p.add_run()._r
            a = OxmlElement("w:fldChar")
            a.set(qn("w:fldCharType"), "begin")
            r.append(a)
            r = p.add_run()._r
            a = OxmlElement("w:instrText")
            a.set(qn("xml:space"), "preserve")
            a.text = (
                ' TOC \\o "1-'
                + str(SPEC.get("heading_levels", 3))
                + '" \\h \\z \\t "XAUFE Preface,1,XAUFE References Title,1" '
            )
            r.append(a)
            r = p.add_run()._r
            a = OxmlElement("w:fldChar")
            a.set(qn("w:fldCharType"), "separate")
            r.append(a)
        # Real hyperlink + real PAGEREF field. Cached page values are filled by refresh_toc.py.
        h = OxmlElement("w:hyperlink")
        h.set(qn("w:anchor"), name)
        h.set(qn("w:history"), "1")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        h.append(r)
        p._p.append(h)
        p.add_run("\t")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), f" PAGEREF {name} \\h ")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "?"
        r.append(t)
        fld.append(r)
        p._p.append(fld)
        if i == len(entries) - 1:
            r = p.add_run()._r
            a = OxmlElement("w:fldChar")
            a.set(qn("w:fldCharType"), "end")
            r.append(a)


def format_table(table, *, repeat_header=False):
    return fit_table(
        table,
        WIDTH,
        style_paragraph,
        repeat_header=repeat_header,
        max_height=SPEC["page_twips"]["height"]
        - SPEC["page_twips"]["top"]
        - SPEC["page_twips"]["bottom"],
    )


def _display_units(text):
    """Width in half-em units at the body size: a CJK character is 2, a Latin letter 1."""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(text))


def column_widths(rows, total, ratios=None):
    """Split ``total`` twips between columns by content (or explicit ``widths`` ratios).

    Short columns get what their longest cell needs; long text columns share the
    rest, so a description column is not squeezed to one character per line while
    a numeric column wastes space.  Cell margins (2 × 108 twips) are included.
    """
    n = len(rows[0])
    if ratios:
        weights = [float(x) for x in ratios]
    else:
        unit, margin = 105, 216  # half of a 10.5 pt character, default left+right cell margins
        desired, minimum = [], []
        for c in range(n):
            cells = [str(row[c]) for row in rows]
            longest = max(_display_units(x) for x in cells)
            word = max((_display_units(w) for x in cells for w in re.findall(r"[A-Za-z0-9_.%-]+", x)), default=0)
            desired.append(min(longest, 60) * unit + margin)
            minimum.append(max(4, min(word, 16)) * unit + margin)
        if sum(desired) <= total:
            weights = desired
        elif sum(minimum) >= total:
            weights = minimum
        else:
            # Columns that fit in a fair share get their full width; long-text
            # columns share what is left, never below their minimum.
            weights, open_, remaining = [0.0] * n, set(range(n)), float(total)
            while open_:
                share = remaining / len(open_)
                fixed = [c for c in open_ if desired[c] <= share]
                if not fixed:
                    break
                for c in fixed:
                    weights[c] = desired[c]
                    remaining -= desired[c]
                    open_.discard(c)
            if open_:
                need = sum(desired[c] for c in open_)
                for c in open_:
                    weights[c] = max(minimum[c], remaining * desired[c] / need)
    scale = total / sum(weights)
    widths = [int(w * scale) for w in weights]
    widths[-1] += total - sum(widths)
    return widths


def add_blocks(doc, blocks, base, entries):
    heading_iter = iter([e for e in entries if e[0] != "_xaufe_references"])
    for index, b in enumerate(blocks):
        try:
            _add_block(doc, b, base, heading_iter)
        except ValueError as exc:
            what = str(b.get("caption") or b.get("text") or "")[:15]
            raise ValueError(f"content.json 的 blocks[{index}]（{b.get('type')}{'“' + what + '”' if what else ''}）：{exc}") from exc


def _add_block(doc, b, base, heading_iter):
    typ = b["type"]
    if typ == "heading":
        role = "preface" if b.get("preface") else "h" + str(b["level"])
        p = doc.add_paragraph(b["text"], ROLES[role])
        name, _, _ = next(heading_iter)
        bookmark(p, name, int(name.rsplit("_", 1)[-1]) + 10)
    elif typ in ("paragraph", "code", "equation"):
        p = doc.add_paragraph(style=ROLES["body" if typ == "paragraph" else typ])
        add_text(p, b)
    elif typ == "image":
        f = (base / b["path"]).resolve()
        if not f.is_file():
            fail(f"图片缺失：{f}。禁止用占位图替代。")
        if b.get("caption") and not isinstance(b["caption"], str):
            fail("图题必须是字符串。")
        p = doc.add_paragraph(style=ROLES["figure"])
        r = p.add_run()
        width = Cm(b.get("width_cm", WIDTH / 1440 * 2.54))
        if width > Twips(WIDTH):
            fail("图片宽度超过版心，必须等比缩小。")
        shape = r.add_picture(str(f), width=width)
        for key in ("distT", "distB", "distL", "distR"):
            shape._inline.set(key, "0")
        maxheight = Pt(580)
        if shape.height > maxheight:
            ratio = maxheight / shape.height
            shape.width = int(shape.width * ratio)
            shape.height = int(maxheight)
        shape._inline.docPr.set("descr", b.get("alt", b.get("caption", "文档插图")))
        if b.get("caption"):
            doc.add_paragraph(b["caption"], ROLES["caption"])
        else:
            p.paragraph_format.keep_with_next = False
    elif typ == "table":
        if b.get("caption"):
            p = doc.add_paragraph(b["caption"], ROLES["caption"])
            p.paragraph_format.keep_with_next = True
        rows = [b["header"]] + b["rows"]
        n = len(rows[0])
        t = doc.add_table(rows=len(rows), cols=n)
        t.style = "Table Grid"
        for col, width in zip(t.columns, column_widths(rows, WIDTH, b.get("widths"))):
            col.width = Twips(width)
        for ri, row in enumerate(rows):
            for ci, value in enumerate(row):
                t.cell(ri, ci).text = value
        format_table(t, repeat_header=True)
    elif typ == "page_break":
        p = doc.add_paragraph(style=ROLES["spacer"])
        p.paragraph_format.line_spacing = Pt(1)
        p.add_run().add_break(WD_BREAK.PAGE)
    else:
        fail(f"未实现的块类型：{typ}")


def _schema_message(schema, data, exc):
    """A content.json schema error as "where: what is wrong, what is allowed" (for a block, checked
    against the one shape its type asks for, not the list of all shapes)."""
    import jsonschema

    path = list(exc.absolute_path)
    if len(path) >= 2 and path[0] == "blocks" and isinstance(path[1], int):
        block = data["blocks"][path[1]]
        shapes = schema["properties"]["blocks"]["items"]["oneOf"]
        kinds = lambda shape: [shape["properties"]["type"]["const"]] if "const" in shape["properties"]["type"] \
            else shape["properties"]["type"]["enum"]
        where = f"content.json 的 blocks[{path[1]}]"
        if not isinstance(block, dict) or block.get("type") not in [k for sh in shapes for k in kinds(sh)]:
            got = block.get("type") if isinstance(block, dict) else block
            return f"{where}: type={got!r} 不认识，可用的 type：" + "、".join(k for sh in shapes for k in kinds(sh)) + "。"
        shape = next(sh for sh in shapes if block["type"] in kinds(sh))
        inner = jsonschema.Draft202012Validator(shape).iter_errors(block)
        first = min(inner, key=lambda e: len(list(e.absolute_path)), default=None)
        allowed = "、".join(shape["properties"])
        if first is None:
            return f"{where}: 写法不对（{block['type']} 可用的字段：{allowed}）。"
        field = "".join(f"[{p}]" if isinstance(p, int) else f".{p}" for p in first.absolute_path)
        if first.validator == "additionalProperties":
            extra = sorted(set(block) - set(shape["properties"])) if not first.absolute_path else []
            return f"{where}{field}: 多了不认识的字段 {'、'.join(extra) or first.message}；{block['type']} 可用的字段：{allowed}。"
        if first.validator == "oneOf" and not first.absolute_path:
            return f"{where}: text 和 runs 要写一个（只写一个）；{block['type']} 可用的字段：{allowed}。"
        if first.validator == "required":
            return f"{where}{field}: 缺少必填字段 {first.message.split()[0]}；{block['type']} 可用的字段：{allowed}。"
        if first.validator == "type":
            kind = {"boolean": "true 或 false", "string": "文字（加引号）", "integer": "整数", "number": "数字",
                    "array": "列表 [...]", "object": "{...} 对象"}.get(str(first.validator_value), first.validator_value)
            return f"{where}{field}: 现在是 {first.instance!r}，应为 {kind}。"
        return f"{where}{field}: {_schema_words(first)}；{block['type']} 可用的字段：{allowed}。"
    field = ".".join(map(str, path)) or "顶层"
    return f"content.json 的 {field}: {_schema_words(exc)}（content.json 的写法见 references/content.md）。"


def _schema_words(err):
    """A jsonschema error in Chinese (the common validators; others keep jsonschema's text)."""
    v, value, got = err.validator, err.validator_value, err.instance
    words = {
        "maximum": lambda: f"现在是 {got!r}，最大只能是 {value}",
        "minimum": lambda: f"现在是 {got!r}，最小是 {value}",
        "exclusiveMaximum": lambda: f"现在是 {got!r}，要小于 {value}",
        "exclusiveMinimum": lambda: f"现在是 {got!r}，要大于 {value}",
        "enum": lambda: f"现在是 {got!r}，只能是 " + "、".join(map(repr, value)),
        "const": lambda: f"现在是 {got!r}，应为 {value!r}",
        "minLength": lambda: "不能是空的" if value == 1 else f"至少要 {value} 个字",
        "maxLength": lambda: f"最多 {value} 个字",
        "minItems": lambda: "不能是空列表" if value == 1 else f"至少要有 {value} 项",
        "maxItems": lambda: f"最多 {value} 项",
        "pattern": lambda: f"写法不对：{got!r}",
    }
    return words[v]() if v in words else err.message


def validate_data(data, base=None):
    import jsonschema

    schema = load_json(ROOT / "assets/checks/content.schema.json")
    # Every broken block at once (one line each), so one repair round can fix them all.
    groups = {}
    for exc in jsonschema.Draft202012Validator(schema).iter_errors(data):
        groups.setdefault(tuple(exc.absolute_path)[:2], []).append(exc)
    problems = [_schema_message(schema, data, jsonschema.exceptions.best_match(found)) for _, found in sorted(
        groups.items(), key=lambda item: [str(x).zfill(6) for x in item[0]])]
    blocks = data.get("blocks") if isinstance(data, dict) and isinstance(data.get("blocks"), list) else []
    for i, b in enumerate(blocks):
        if isinstance(b, dict) and b.get("type") == "table" and isinstance(b.get("rows"), list) \
                and isinstance(b.get("header"), list) and ("blocks", i) not in groups:
            bad = [n for n, row in enumerate(b["rows"]) if isinstance(row, list) and len(row) != len(b["header"])]
            if bad:
                problems.append(f"content.json 的 blocks[{i}]（表格）: rows[{bad[0]}] 有 {len(b['rows'][bad[0]])} 格，"
                                f"表头有 {len(b['header'])} 格；每行格数要和表头相同（空格子写 \"\"）。")
        # Missing image files too, so they do not first show up in the run after the schema is fixed.
        if base is not None and isinstance(b, dict) and b.get("type") == "image" and isinstance(b.get("path"), str) \
                and not (Path(base) / b["path"]).resolve().is_file():
            problems.append(f"content.json 的 blocks[{i}]（图片）: 图片文件 {b['path']} 不存在（路径相对 content.json 所在的文件夹）；"
                            "请用户提供图片，不能用占位图代替。")
    if problems:
        fail("\n".join(dict.fromkeys(problems)))
    if not SPEC["english_abstract"]:
        if any(k in data for k in ("abstract_en", "keywords_en")) or data.get("metadata", {}).get(
            "title_en"
        ):
            fail(
                "content.json 里有英文摘要、英文关键词或英文题目（title_en），但本次版式不含英文摘要。问用户："
                "要保留就在 job.json 的 format 写 \"english_abstract\": true；不要就从 content.json 去掉 abstract_en、keywords_en 和 metadata.title_en。"
                "不要按校徽颜色猜。"
            )
    if data.get("mode", "paper") == "paper":
        required = ["abstract_zh", "keywords_zh", "references"]
        if SPEC["english_abstract"]:
            required += ["abstract_en", "keywords_en"]
        for key in required:
            if key not in data:
                fail(
                    f"mode=paper 需要 {key}。用户要求写作就写好这一部分；只排版且原文没有这部分，就把 content.json 的 mode 改成 document。"
                )
        if not data["blocks"] or data["blocks"][0]["type"] != "heading":
            fail("content.json 的 blocks 第一块要是标题（type heading）或序言：正文从标题开始。")
        if not data["metadata"].get("title_zh"):
            fail("content.json 的 metadata 要写 title_zh（论文题目，摘要页用）；封面可以不填。")
        if SPEC["english_abstract"] and not data["metadata"].get("title_en"):
            fail("本次包含英文摘要，需要在 content.json 的 metadata 里写 title_en（英文题目）。用户要求写作或翻译时由你写好；只排版且原文没有，就问用户。")
    for lang in ("zh", "en"):
        for word in data.get("keywords_" + lang, []):
            if not word.strip() or word.strip() != word or re.search(r" {2,}|[\t\r\n]", word):
                fail("关键词词项不能是空白或自带分隔符；英文短语内的单个空格可以保留。")
    for b in data["blocks"]:
        if b["type"] == "heading" and b["level"] > SPEC.get("heading_levels", 3):
            fail("content.json 有四级标题（level 4），模板只到三级：改成三级以内，或用户同意后在 job.json 的 format 写 \"heading_levels\": 4。")
        if b["type"] == "heading" and b.get("preface") and b["level"] != 1:
            fail("content.json 里序言（preface）的 level 要写 1。")
        if b["type"] == "table" and b.get("widths") is not None and len(b["widths"]) != len(b["header"]):
            fail("表格 widths 的个数要和列数相同（每列一个比例数字）。")
    if any(re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", str(x)) for x in data["metadata"].values()):
        fail("元数据包含无效控制字符。")


def add_abstract(doc, data, meta, english=False):
    if english and not SPEC["english_abstract"]:
        fail("本次版式不含英文摘要：用户要英文摘要就在 job.json 的 format 写 \"english_abstract\": true 再 run。")
    sp = ROLES["abstract_spacer"]
    doc.add_paragraph("", sp)
    doc.add_paragraph(
        meta["title_en" if english else "title_zh"], ROLES["title_en" if english else "title_zh"]
    )
    doc.add_paragraph("", sp)
    doc.add_paragraph(
        "Abstract" if english else "内 容 摘 要",
        ROLES["abstract_label_en" if english else "abstract_label_zh"],
    )
    doc.add_paragraph("", sp)
    for text in data["abstract_en" if english else "abstract_zh"].split("\n"):
        doc.add_paragraph(text, ROLES["abstract_en" if english else "abstract_zh"])
    doc.add_paragraph("", sp)
    p = doc.add_paragraph(style=ROLES["keywords_en" if english else "keywords_zh"])
    role = "keywords_en" if english else "keywords_zh"
    cn, size, bold = role_format(role)
    font(
        p.add_run("Key words：" if english else "关键词：")._r.get_or_add_rPr(),
        cn,
        14,
        bold,
        latin=role_latin(role),
    )
    font(
        p.add_run(
            "   ".join(data["keywords_en" if english else "keywords_zh"])
        )._r.get_or_add_rPr(),
        cn,
        size,
        bold,
        latin=role_latin(role),
    )


def build(input_path, output, force=False, intake_path=None, cover_min_level=0):
    state = require_intake(intake_path)
    select_profile(state["template"], state.get("format_request"))
    SPEC["_cover_min_level"] = cover_min_level
    asset_check()
    data = load_json(input_path)
    data["metadata"] = merge_metadata(data.get("metadata", {}), state)
    validate_data(data, Path(input_path).parent)
    out = safe_out(input_path, output, force)
    meta = data["metadata"]
    mode = data.get("mode", "paper")
    has_cover = data.get("cover", state.get("include_cover", True))
    doc = Document(ROOT / SPEC["master_file"])
    install_styles(doc)
    if has_cover:
        fill_cover(doc, cover_metadata(meta, state))
    else:
        for e in list(doc.element.body):
            if e.tag != qn("w:sectPr"):
                doc.element.body.remove(e)
    entries = []
    for b in data["blocks"]:
        if b["type"] == "heading":
            entries.append((f"_xaufe_heading_{len(entries)+1}", b["text"], b["level"]))
    if mode == "paper" or data.get("references"):
        entries.append(("_xaufe_references", "参 考 文 献", 1))
    if mode == "paper":
        parts = [x for x in SPEC["parts"] if x not in ("cover", "references")]
        for i, part in enumerate(parts):
            if has_cover or i > 0:
                add_section(doc, start=1 if part in ("toc", "body") else None)
            elif part == "body":
                element(doc.sections[0]._sectPr, "w:pgNumType", fmt="decimal", start=1)
            if part.startswith("abstract_"):
                page = element(doc.sections[-1]._sectPr, "w:pgNumType", fmt="upperRoman")
                if i == 0 or not parts[i - 1].startswith("abstract_"):
                    page.set(qn("w:start"), "1")
            if part == "toc":
                add_toc(doc, entries)
            elif part == "abstract_zh":
                add_abstract(doc, data, meta)
            elif part == "abstract_en":
                add_abstract(doc, data, meta, english=True)
            elif part == "body":
                add_blocks(doc, data["blocks"], Path(input_path).resolve().parent, entries)
    else:
        if has_cover:
            add_section(doc, start=1)
        add_blocks(doc, data["blocks"], Path(input_path).resolve().parent, entries)
    if mode == "paper" or data.get("references"):
        p = doc.add_paragraph("参 考 文 献", ROLES["reference_title"])
        bookmark(p, "_xaufe_references", 100000)
        for i, text in enumerate(data["references"], 1):
            # Already numbered ("[1]", "1.", "1、", "（1）"): keep the user's number, never add a second one.
            numbered = re.match(r"^\s*(?:[\[［]\d+[\]］]|\d+\s*[.．、](?!\d)|[(（]\d+[)）])", text)
            prefix = "" if numbered else f"[{i}] "
            doc.add_paragraph(prefix + text, ROLES["reference"])
    finalize(doc, out, meta, has_cover, mode, state)
    return {
        "output": str(out),
        "template": SPEC["id"],
        "mode": mode,
        "toc_needs_refresh": mode == "paper",
        "message": "排版初稿已生成，尚未交付。",
    }


def inventory(input_path):
    doc = Document(input_path)
    blocks = []
    pno = tno = 0
    from numbering import list_labels

    labels = list_labels(doc)
    for bi, el in enumerate(doc.element.body):
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc._body)
            text = "".join(el.xpath(".//w:t/text()"))
            blocks.append(
                {
                    "block_index": bi,
                    "id": f"p{pno}",
                    "kind": "paragraph",
                    "style": p.style.name,
                    "text": text,
                    "images": len(el.xpath(".//w:drawing")),
                    "equations": len(el.xpath(".//m:oMath")),
                    "section_break": bool(el.xpath("./w:pPr/w:sectPr")),
                    "has_fields": bool(el.xpath(".//w:instrText|.//w:fldSimple")),
                    "numbered": el in labels,  # an auto number or bullet Word shows before the text
                    "list_label": labels.get(el) or "",
                }
            )
            pno += 1
        elif el.tag == qn("w:tbl"):
            blocks.append(
                {
                    "block_index": bi,
                    "id": f"t{tno}",
                    "kind": "table",
                    "text": " | ".join(el.xpath(".//w:t/text()")),
                }
            )
            tno += 1
    return {
        "input_sha256": hashlib.sha256(Path(input_path).read_bytes()).hexdigest(),
        "path": str(Path(input_path).resolve()),
        "sections": len(doc.sections),
        "blocks": blocks,
        "instructions": "paragraph id 的角色映射必须根据完整清单确认；block_index 仅用于明确的封面替换范围。",
    }


# Roles whose first-line position the template sets: blanks typed in front of the text to fake an
# indent would push it further (code and table cells keep theirs).
INDENTED_ROLES = {"body", "abstract_zh", "abstract_en", "h1", "h2", "h3", "h4", "preface", "caption",
                  "reference", "title_zh", "title_en", "keywords_zh", "keywords_en", "equation"}
NEW_PAGE_SECTIONS = {None, "nextPage", "oddPage", "evenPage"}


def blank_page_breaks(doc):
    """Page starts with nothing visible since the page before started: Word shows an empty page at
    each, even where LibreOffice folds the breaks together.  Yields (paragraph index, what, the
    break that can go, text before it): two page breaks in a row, a page break and then a paragraph
    set to start a new page, a page break and then an empty paragraph ending the section, a page
    break at the very end.  ``None`` as the break means it is a style or section, not a w:br."""
    state, last, text, index = None, None, "", -1  # state "break"/"section": a page has just started
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:tbl"):
            state = None
            continue
        if el.tag != qn("w:p"):
            continue
        index += 1
        before = paragraph_properties(Paragraph(el, doc._body), "pageBreakBefore")
        if before and on(before.get("val")):
            if state == "break":
                yield index, "紧跟着设了“段前分页”的段落", last, text
            state, last = "break", None
        own = ""
        for node in el.iter(qn("w:t"), qn("w:br"), qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("m:oMath")):
            if node.tag == qn("w:br"):
                if node.get(qn("w:type")) == "page":
                    if state in ("break", "section"):
                        yield index, "又是一个分页符", node, own if own.strip() else text
                    state, last = "break", node
            elif node.tag != qn("w:t"):
                state = None
            else:
                own += node.text or ""
                state = None if own.strip() else state
        text = own if own.strip() else text
        sect = el.find(f"./{qn('w:pPr')}/{qn('w:sectPr')}")
        if sect is not None:
            kind = sect.find(qn("w:type"))
            if kind is None or kind.get(qn("w:val")) in NEW_PAGE_SECTIONS:
                if state == "break" and not el.xpath('.//w:br[@w:type="page"]'):
                    yield index, "紧跟着分节符（下一页）", last, text
                state = "section"
    if state == "break" and last is not None and not doc.paragraphs[-1]._p.xpath('.//w:br[@w:type="page"]'):
        yield index, "就是文档结尾", last, text


def drop_blank_page_breaks(doc):
    """Remove the page breaks that only make an empty page; [(paragraph index, text before)]."""
    removed = []
    while True:
        found = next(((i, node, text) for i, _, node, text in blank_page_breaks(doc) if node is not None), None)
        if not found:
            return removed
        i, node, text = found
        node.getparent().remove(node)
        removed.append((i, text))


def keep_list_numbering(p):
    """A bullet or number that comes from the paragraph's style ("List Bullet") would vanish when the
    template style replaces it: copy it onto the paragraph.  True when the paragraph is numbered."""
    from style_model import style_chain

    if p._p.xpath("./w:pPr/w:numPr"):
        return True
    for style in reversed(style_chain(p.style) if p.style is not None else []):
        num = style.element.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if num is not None:
            if num.xpath('./w:numId[@w:val="0"]'):
                return False
            ppr = p._p.get_or_add_pPr()
            ppr.insert(1 if ppr.find(qn("w:pStyle")) is not None else 0, deepcopy(num))
            return True
    return False


# Roles whose auto number becomes typed text: the template numbers headings and references in
# the text itself, and the TOC, the heading-number check and the citation check all read text.
LIVE_LIST_ROLES = {"body", "table_text", "code", "cover", "spacer", "abstract_spacer", "preserve"}


def freeze_list_labels(doc, roles, default, ids):
    """Turn the auto number of headings, references etc. into the same visible text.

    Returns {"p<i>": "p<i>“label text…”"} for the paragraphs changed; the reader sees the same
    characters as before (the source comparison counts list labels as text on both sides)."""
    from numbering import list_labels

    labels = list_labels(doc)
    frozen = {}
    for i, p in enumerate(doc.paragraphs):
        role = roles.get(ids[i], default)
        if role in LIVE_LIST_ROLES or p._p not in labels:
            continue
        label = labels[p._p]
        if label is None:
            fail(f"原稿 {ids[i]}“{p.text.strip()[:15]}”（role {role}）用的是自动编号，但编号格式认不出来，"
                 "转不成文字。请用户在 Word 里把自动编号转成文字（选中这些段落，复制后“选择性粘贴→只保留文本”），"
                 "或者在 mapping.json 里把它改成 body。")
        ppr = p._p.get_or_add_pPr()
        for num in ppr.findall(qn("w:numPr")):
            ppr.remove(num)
        if not label:
            continue
        run = OxmlElement("w:r")
        first = p._p.find(f".//{qn('w:r')}/{qn('w:rPr')}")
        if first is not None:
            run.append(deepcopy(first))
        t = OxmlElement("w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = label
        run.append(t)
        ppr.addnext(run)
        frozen[f"{ids[i]}"] = f"{ids[i]}“{label}{p.text[len(label):].strip()[:12]}”"
    return frozen


LEADING_BLANKS = " \t\u3000\u00a0"


def strip_typed_indent(p):
    """Remove spaces, full-width spaces and tabs typed before the first character; return how many."""
    removed = 0
    for node in list(p._p.iter()):
        if node.tag == qn("w:tab") and node.getparent() is not None and node.getparent().tag == qn("w:r"):
            node.getparent().remove(node)
            removed += 1
        elif node.tag == qn("w:t"):
            text = node.text or ""
            kept = text.lstrip(LEADING_BLANKS)
            removed += len(text) - len(kept)
            node.text = kept
            if kept:
                break
        elif node.tag in (qn("w:drawing"), qn("w:pict"), qn("w:fldChar"), qn("w:instrText"), qn("w:sym"),
                          qn("w:br"), qn("w:cr"), qn("m:oMath"), qn("m:oMathPara"), qn("w:object")):
            break
    return removed


def content_signature(doc, root=None):
    """Payload signature: runs may be reformatted, their content/targets cannot change."""
    root = doc.element.body if root is None else root

    def resource(rid):
        if not rid:
            return None
        rel = doc.part.rels.get(rid)
        if rel is None:
            fail("正文引用了缺失的关系：" + rid)
        if rel.is_external:
            return ("external", rel.target_ref)
        return ("internal", hashlib.sha256(rel.target_part.blob).hexdigest())

    related = []
    for node in root.iter():
        if node.tag in (qn("w:headerReference"), qn("w:footerReference")):
            continue
        for key, value in node.attrib.items():
            if key.startswith(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
            ):
                related.append((node.tag, key, resource(value)))
    preserved_parts = []
    for part in doc.part.package.parts:
        name = str(part.partname)
        if re.match(r"/word/(footnotes|endnotes|comments|numbering)\.xml$", name):
            preserved_parts.append((name, hashlib.sha256(part.blob).hexdigest()))
    return {
        "text": "".join(root.xpath(".//w:t/text()")),
        "math": tuple(
            __import__("lxml").etree.tostring(x, method="c14n") for x in root.xpath(".//m:oMath")
        ),
        "tables": len(root.xpath(".//w:tbl")),
        "images": len(root.xpath(".//a:blip")),
        "fields": tuple(root.xpath(".//w:instrText/text()|.//w:fldSimple/@w:instr")),
        "hyperlink_anchors": tuple(root.xpath(".//w:hyperlink/@w:anchor")),
        "related_payloads": tuple(related),
        "preserved_parts": tuple(sorted(preserved_parts)),
        "cell_spans": tuple(
            (
                tuple(tc.xpath("./w:tcPr/w:gridSpan/@w:val")),
                tuple(tc.xpath("./w:tcPr/w:vMerge/@w:val")),
                len(tc.xpath("./w:tcPr/w:vMerge")),
            )
            for tc in root.xpath(".//w:tc")
        ),
    }


def subtree_signature(doc, elements):
    root = OxmlElement("w:body")
    for el in elements:
        root.append(deepcopy(el))
    return content_signature(doc, root)


FRONT_ZH = ("title_zh", "abstract_label_zh", "abstract_zh", "keywords_zh")
FRONT_EN = ("title_en", "abstract_label_en", "abstract_en", "keywords_en")
BODY_START = ("preface", "h1", "h2", "h3", "h4")


def _role_of(doc, el):
    if el.tag == qn("w:tbl"):
        return "table"
    style = Paragraph(el, doc._body).style
    return next((k for k, v in ROLES.items() if style is not None and v == style.name), None)


def _section_break(doc, template_sectpr, number_format=None, start=None):
    """An empty 1 pt paragraph that ends a section (properties = the part before it)."""
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    p.append(ppr)
    sect = deepcopy(template_sectpr)
    for child in list(sect):
        if child.tag in (qn("w:headerReference"), qn("w:footerReference"), qn("w:pgNumType")):
            sect.remove(child)
    page = element(sect, "w:pgNumType", fmt=number_format or "decimal")
    if start is not None:
        page.set(qn("w:start"), str(start))
    ppr.append(sect)
    para = Paragraph(p, doc._body)
    para.style = ROLES["spacer"]
    para.paragraph_format.line_spacing = Pt(1)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    return p


def structure_paper(doc, has_cover, meta, additions):
    """Give an existing document the full template structure without touching its content.

    Original paragraphs/tables keep their order. Only these are inserted: section
    breaks between parts, a real table of contents, missing abstract-page titles or
    labels, and an English abstract the user asked for (additions).
    """
    body = doc.element.body
    final_sect = body.find(qn("w:sectPr"))
    items = [e for e in body if e.tag != qn("w:sectPr")]
    start = 0
    if has_cover:
        start = next(
            i + 1 for i, e in enumerate(items) if e.tag == qn("w:p") and e.xpath("./w:pPr/w:sectPr")
        )
    content = items[start:]
    for el in content:
        if el.tag == qn("w:p") and el.xpath("./w:pPr/w:sectPr"):
            fail("原稿正文里已有分节符，不能自动整理成完整论文结构；mapping.json 的 mode 改成 document。")
        if _role_of(doc, el) in ("toc_title", "toc1", "toc2", "toc3", "toc4"):
            fail("原稿已有目录：把旧目录各段的 role 改成 delete，脚本会生成新的目录。")
    parts = {"lead": [], "abstract_zh": [], "abstract_en": [], "body": []}
    order = ["lead", "abstract_zh", "abstract_en", "body"]
    current = "lead"
    for el in content:
        role = _role_of(doc, el)
        target = current
        if role in FRONT_ZH:
            target = "abstract_zh"
        elif role in FRONT_EN:
            target = "abstract_en"
        elif role in BODY_START or role == "reference_title":
            target = "body"
        elif current == "lead" and role not in ("spacer", "abstract_spacer"):
            fail("正文前有分不到摘要或正文里的段落：" + "".join(el.xpath(".//w:t/text()"))[:40])
        if order.index(target) < order.index(current):
            names = {"lead": "开头", "abstract_zh": "中文摘要", "abstract_en": "英文摘要", "body": "正文"}
            text = "".join(el.xpath(".//w:t/text()")).strip()[:20]
            fail(f"“{text}”（role={role}）排在{names[current]}后面，但它属于{names[target]}：摘要、英文摘要、正文的先后顺序和所选版式不同。"
                 "先看这一段的 role 是否标错；没标错就改用 mode=document，或先和用户确认。")
        current = target
        parts[current].append(el)
    if not parts["abstract_zh"]:
        fail("mode=paper 需要中文摘要；原稿没有摘要时把 mapping.json 的 mode 改成 document。")
    if not any(_role_of(doc, el) in BODY_START for el in parts["body"]):
        fail("mode=paper 需要正文标题（如“一、引言”）；没有时把 mode 改成 document。")
    first_body = next(el for el in parts["body"] if _role_of(doc, el) not in ("spacer",))
    if _role_of(doc, first_body) not in BODY_START:
        fail("正文第一段不是标题，目录无法定位正文起页；在正文开头加标题，或改用 mode=document。")
    for el in items[start:]:
        body.remove(el)

    def append(el):
        final_sect.addprevious(el)

    def ensure_front(elements, english):
        roles = [_role_of(doc, el) for el in elements]
        head = []
        title_role, label_role = ("title_en", "abstract_label_en") if english else ("title_zh", "abstract_label_zh")
        title = meta.get("title_en" if english else "title_zh", "")
        if title_role not in roles and title:
            head.append(Paragraph(OxmlElement("w:p"), doc._body))
            head[-1].add_run(title)
            head[-1].style = ROLES[title_role]
        label = []
        if label_role not in roles:
            label.append(Paragraph(OxmlElement("w:p"), doc._body))
            label[-1].add_run("Abstract" if english else "内 容 摘 要")
            label[-1].style = ROLES[label_role]
        # A generated label goes after the paper's own title, never above it.
        after = max((i + 1 for i, role in enumerate(roles) if role == title_role), default=0)
        sequence = [p._p for p in head] + elements[:after] + [p._p for p in label] + elements[after:]
        # Same blank lines as the template's abstract page: before the title and
        # between title / label / text / keywords, only where none exists yet.
        spaced, previous = [], None
        for el in sequence:
            role = _role_of(doc, el)
            boundary = role in (title_role, label_role, "keywords_zh", "keywords_en") or (
                role in ("abstract_zh", "abstract_en") and previous == label_role
            )
            if boundary and previous not in ("abstract_spacer", "spacer"):
                spacer = Paragraph(OxmlElement("w:p"), doc._body)
                spacer.style = ROLES["abstract_spacer"]
                spaced.append(spacer._p)
            spaced.append(el)
            previous = role
        return spaced

    if additions:
        if not SPEC["english_abstract"]:
            fail("mapping.json 的 add 只用于补英文摘要；本次版式不含英文摘要。")
        if parts["abstract_en"]:
            fail("原稿已经有英文摘要，不要再用 add 补。")
        missing = {"title_en", "abstract_en", "keywords_en"} - set(additions)
        if missing or set(additions) - {"title_en", "abstract_en", "keywords_en"}:
            fail("mapping.json 的 add 需要且只能有 title_en、abstract_en、keywords_en 三项。")
        mark = len([e for e in body if e.tag != qn("w:sectPr")])
        add_abstract(doc, additions, {"title_en": additions["title_en"]}, english=True)
        added = [e for e in body if e.tag != qn("w:sectPr")][mark:]
        for el in added:
            body.remove(el)
        parts["abstract_en"] = added
    if SPEC["english_abstract"] and not parts["abstract_en"]:
        fail("所选版式需要英文摘要，原稿没有。问用户：要你翻译补上（写进 mapping.json 的 add），还是不要英文摘要（job.json 的 format 写 \"english_abstract\": false）。")
    entries, bid = [], 900000
    for el in parts["body"]:
        role = _role_of(doc, el)
        if role in BODY_START and int(role[-1] if role != "preface" else 1) <= SPEC.get("heading_levels", 3):
            name = f"_xaufe_heading_{len(entries) + 1}"
            level = 1 if role == "preface" else int(role[-1])
        elif role == "reference_title":
            name, level = "_xaufe_references", 1
        else:
            continue
        p = Paragraph(el, doc._body)
        bookmark(p, name, bid)
        bid += 1
        entries.append((name, p.text, level))
    sequence = [x for x in SPEC["parts"] if x not in ("cover", "references")]
    abstract_started = False
    for part in sequence:
        if part == "toc":
            mark = len([e for e in body if e.tag != qn("w:sectPr")])
            add_toc(doc, entries)
            toc = [e for e in body if e.tag != qn("w:sectPr")][mark:]
            for el in toc:
                body.remove(el)
                append(el)
            append(_section_break(doc, final_sect, "decimal", 1))
        elif part.startswith("abstract_"):
            elements = parts[part]
            if part == "abstract_zh":
                elements = parts["lead"] + elements
            for el in ensure_front(elements, part == "abstract_en"):
                append(el)
            append(_section_break(doc, final_sect, "upperRoman", None if abstract_started else 1))
            abstract_started = True
        elif part == "body":
            for el in parts["body"]:
                append(el)
            element(final_sect, "w:pgNumType", fmt="decimal", start=1)
    kept = sum(1 for e in body if e in set(content))
    if kept != len(content):
        fail("整理论文结构时原稿内容数量不一致，已停止。")


def _delete_marked(doc, roles, mapping):
    """Remove paragraphs the user agreed to delete (role "delete") and the old TOC the
    template replaces (role "old_toc").  Returns the mapping id of each paragraph left (messages
    and roles keep using the ids the user saw in mapping.json) and the deleted texts."""
    old = list(doc.paragraphs)
    marked = [p for i, p in enumerate(old) if roles.get(f"p{i}") in ("delete", "old_toc")]
    if not marked:
        return [f"p{i}" for i in range(len(old))], []
    body = list(doc.element.body)
    end = mapping.get("replace_cover_before")
    shift = 0
    deleted = []
    for p in marked:
        if p._p.xpath("./w:pPr/w:sectPr"):
            fail("要删除的段落带有分节符，不能删除：" + p.text[:40])
        if isinstance(end, int) and not isinstance(end, bool) and body.index(p._p) < end:
            shift += 1
        deleted.append(p.text[:60])
    ids = [f"p{i}" for i, p in enumerate(old) if roles.get(f"p{i}") not in ("delete", "old_toc")]
    for p in marked:
        p._p.getparent().remove(p._p)
    if shift:
        mapping["replace_cover_before"] = end - shift
    return ids, deleted


def restyle(input_path, map_path, output, force=False, intake_path=None, cover_min_level=0):
    state = require_intake(intake_path)
    select_profile(state["template"], state.get("format_request"))
    SPEC["_cover_min_level"] = cover_min_level
    asset_check()
    m = load_json(map_path)
    m["metadata"] = merge_metadata(m.get("metadata", {}), state)
    out = safe_out(input_path, output, force)
    if not SPEC["english_abstract"] and any(
        x in ("title_en", "abstract_en", "abstract_label_en", "keywords_en")
        for x in m.get("roles", {}).values()
    ):
        fail(
            "原稿有英文摘要，但本次版式不含英文摘要。问用户：保留就在 job.json 的 format 写 \"english_abstract\": true；用户同意删除就把这些段落的 role 改成 delete。不要把它标成正文。"
        )
    if m.get("input_sha256") != hashlib.sha256(Path(input_path).read_bytes()).hexdigest():
        fail("映射文件与当前原稿哈希不一致。请重新 inspect，禁止使用旧段落编号。")
    doc = Document(input_path)
    # These require a dedicated OOXML preservation workflow; do not flatten them via python-docx.
    if any(x.tag not in (qn("w:p"), qn("w:tbl"), qn("w:sectPr")) for x in doc.element.body):
        fail("存在内容控件或其他复杂顶层对象。请采用保全OOXML的专门流程，不静默跳过。")
    tracked = doc.element.body.xpath(".//w:ins|.//w:del|.//w:altChunk")
    if tracked:
        from unsupported import where
        fail(f"原稿有修订或嵌入文档（{where(doc, tracked)}）：脚本不会替用户接受、拒绝或丢弃。"
             "请用户在 Word 的“审阅”里接受或拒绝这些修订后重新给原稿。")
    from unsupported import unsupported_reasons

    problems = unsupported_reasons(doc)
    if problems:
        fail("；".join(problems))
    roles = dict(m.get("roles", {}))
    default = m.get("default_role")
    if m.get("rebuild_structure"):
        # An earlier output made with the other template: drop its section breaks after the cover
        # so the parts can be laid out again in this template's order.
        start = m.get("replace_cover_before") or 0
        for el in list(doc.element.body)[start:]:
            for sect in el.xpath("./w:pPr/w:sectPr"):
                sect.getparent().remove(sect)
        # and the TOC bookmarks it put on the headings: the new TOC sets its own.
        for mark in doc.element.body.xpath('.//w:bookmarkStart[starts-with(@w:name, "_xaufe_")]'):
            ident = mark.get(qn("w:id"))
            for end in doc.element.body.xpath(f'.//w:bookmarkEnd[@w:id="{ident}"]'):
                end.getparent().remove(end)
            mark.getparent().remove(mark)
    ids, deleted = _delete_marked(doc, roles, m)
    spacing_changes = []
    # Only this explicit formatting whitelist is allowed to change visible ASCII spaces.
    for i, p in enumerate(doc.paragraphs):
        if roles.get(ids[i], default) in ("keywords_zh", "keywords_en"):
            try:
                change = normalize_keywords(p, m.get("keyword_items", {}).get(ids[i]))
            except ValueError as exc:
                fail(f"原稿 {ids[i]}“{p.text.strip()[:15]}”（role {roles.get(ids[i], default)}）：{exc}")
            if change:
                spacing_changes.append(dict(change, paragraph=ids[i]))
    # The template indents; a typed indent is layout, not text (the source comparison ignores it too).
    indent_changes = []
    for i, p in enumerate(doc.paragraphs):
        if roles.get(ids[i], default) in INDENTED_ROLES and strip_typed_indent(p):
            indent_changes.append(f"{ids[i]}“{p.text.strip()[:15]}”")
    frozen = freeze_list_labels(doc, roles, default, ids)
    for i, p in enumerate(doc.paragraphs):
        role = roles.get(ids[i], default)
        wanted = m["metadata"].get(role, "") if role in ("title_zh", "title_en") else None
        if wanted and re.sub(r"\s+", "", p.text) != re.sub(r"\s+", "", wanted):
            fail(f"原稿 {ids[i]} 的题目“{p.text.strip()[:30]}”和 job.json 的 info.{role}“{wanted[:30]}”不一致："
                 "封面和页眉用的是 job.json 里的题目，两处会对不上。问用户哪个是对的，改 job.json 的 info 或原稿后重新 run。")
    before = content_signature(doc)
    install_styles(doc)
    styled = None
    for i, p in enumerate(doc.paragraphs):
        role = roles.get(ids[i])
        if role is None:
            if not p.text and not p._p.xpath(".//w:drawing|.//m:oMath"):
                role = "spacer"
            elif default:
                role = default
            else:
                fail(f"mapping.json 里没有 {ids[i]}“{p.text.strip()[:30]}”：在 paragraphs 里给它一个 role（正文用 body）。")
        if role == "preserve":
            if styled is not None and p.text.strip() and not (p.style is not None and p.style.name.lower().startswith("toc ")):
                fail(f"原稿 {ids[i]}“{p.text.strip()[:15]}”的 role 是 preserve，但它在 {styled}（已经开始排版的内容）后面："
                     "preserve 只用于开头的旧封面，保留下来的段落用的是原稿样式，模板检查不认。"
                     "改成它实际的 role（正文 body、标题 h1/h2/h3 等）；旧目录用 old_toc。")
            continue
        if role != "spacer" and styled is None:
            styled = ids[i]
        try:
            style_paragraph(p, role, preserve_num=ids[i] not in frozen and keep_list_numbering(p))
        except ValueError as exc:
            fail(f"原稿 {ids[i]}“{p.text.strip()[:15]}”（role {role}）：{exc}")
        if role == "figure":
            resize_inline(
                p,
                WIDTH,
                SPEC["page_twips"]["height"]
                - SPEC["page_twips"]["top"]
                - SPEC["page_twips"]["bottom"]
                - 400,
            )
    for p in doc.paragraphs:
        nxt = p._p.getnext()
        if p.style.name == ROLES["figure"]:
            caption = (
                nxt is not None
                and nxt.tag == qn("w:p")
                and Paragraph(nxt, doc._body).style.name == ROLES["caption"]
            )
            p.paragraph_format.keep_with_next = caption
        elif p.style.name == ROLES["caption"]:
            # Restyling resets direct pPr. Rebind a table caption to its table;
            # a figure caption must not pull unrelated following prose along.
            p.paragraph_format.keep_with_next = nxt is not None and nxt.tag == qn("w:tbl")
    # A page break that only makes an empty page is layout, not content (the template's own
    # "start on a new page" headings often meet a typed page break).
    break_changes = [f"“{text.strip()[-15:]}”后面" if text.strip() else f"第 {i + 1} 段"
                     for i, text in drop_blank_page_breaks(doc)]
    if m.get("format_tables", False):
        for ti, t in enumerate(doc.tables):
            try:
                format_table(t)
            except ValueError as exc:
                first = " | ".join(c.text.strip() for c in t.rows[0].cells)[:20] if t.rows else ""
                fail(f"原稿第 {ti + 1} 个表格 t{ti}“{first}”：{exc}")
    if content_signature(doc) != before:
        fail("正文内容保真检查失败；未写入输出。")
    has_cover = m.get("has_cover", False)
    preserved_start = 0
    tail_before = None
    if m.get("replace_cover_before") is not None:
        end = m["replace_cover_before"]
        if (
            isinstance(end, bool)
            or not isinstance(end, int)
            or end < 0
            or end >= len(doc.element.body)
        ):
            fail("replace_cover_before 必须是有效的 body block_index。")
        tail_before = subtree_signature(doc, list(doc.element.body)[end:])
        for e in list(doc.element.body)[:end]:
            doc.element.body.remove(e)
        master = Document(ROOT / SPEC["master_file"])
        fill_cover(master, cover_metadata(m.get("metadata", {}), state))
        if "原模板封面基础" in doc.styles:
            doc.styles.element.remove(doc.styles["原模板封面基础"].element)
        doc.styles.element.append(deepcopy(master.styles["原模板封面基础"].element))
        rid, _ = doc.part.get_or_add_image(str(ROOT / SPEC["logo"]["file"]))
        cover = [deepcopy(p._p) for p in master.paragraphs]
        for e in cover:
            for blip in e.xpath(".//a:blip"):
                blip.set(qn("r:embed"), rid)
        ep = OxmlElement("w:p")
        pp = OxmlElement("w:pPr")
        ep.append(pp)
        sp = deepcopy(master.sections[0]._sectPr)
        pp.append(sp)
        for child in list(sp):
            if child.tag in (qn("w:headerReference"), qn("w:footerReference")):
                sp.remove(child)
        for e in reversed(cover + [ep]):
            doc.element.body.insert(0, e)
        # Same 1 pt section-break paragraph as a newly built cover.
        breaker = Paragraph(ep, doc._body)
        breaker.style = ROLES["spacer"]
        breaker.paragraph_format.line_spacing = Pt(1)
        breaker.paragraph_format.space_before = Pt(0)
        breaker.paragraph_format.space_after = Pt(0)
        has_cover = True
        preserved_start = len(cover) + 1
        if subtree_signature(doc, list(doc.element.body)[preserved_start:]) != tail_before:
            fail("封面替换后的正文保全检查失败。")
    # Existing cover content is untouched unless explicitly replaced. The date position
    # follows the current bottom-of-first-page rule even when preserving the cover.
    already_structured = any(
        p.style is not None and p.style.name == ROLES["toc_title"] for p in doc.paragraphs
    ) and len(doc.sections) > (2 if has_cover else 1)
    if m.get("mode") == "paper" and not already_structured:
        structure_paper(doc, has_cover, m.get("metadata", {}), m.get("add", {}))
    elif m.get("add"):
        fail("mapping.json 的 add 只能和 mode=paper 一起用。")
    finalize(doc, out, m.get("metadata", {}), has_cover, m.get("mode", "document"), state)
    return {
        "output": str(out),
        "deleted_paragraphs": deleted,
        "body_preservation": "passed; includes payload hashes and post-cover-replacement body check",
        "keyword_spacing_changes": spacing_changes,
        "typed_indent_removed": indent_changes,
        "blank_page_breaks_removed": break_changes,
        "list_labels_frozen": list(frozen.values()),
        "message": "重排初稿已生成；内容保全检查通过。",
    }
