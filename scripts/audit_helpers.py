"""Independent contract checks for paragraphs, runs, tables, and TOC relationships."""

from __future__ import annotations
import hashlib, json, re
from copy import deepcopy
from pathlib import Path
from docx.oxml.ns import qn
from style_model import paragraph_properties, run_properties, on
from keywords import LABEL, visible, spacing_issue
from table_layout import walk_tables, raw_rows, cell_margin

ROOT = Path(__file__).resolve().parents[1]


def page_contract(template, request=None):
    """Page box, margins and header/footer distances from the independent baseline plus the user's
    page_twips; a header or page numbers added to B without a distance get A's distances."""
    data = json.loads((ROOT / "assets/checks/audit-baseline.json").read_text(encoding="utf-8"))["pages"]
    request = request or {}
    page = dict(data[template])
    asked = request.get("page_twips") or {}
    page.update(asked)
    added = data["added_header_or_page_numbers_distance"]
    if (request.get("header") or {}).get("enabled") is True and "header" not in asked and page["header"] < 200:
        page["header"] = added["header"]
    if request.get("footer_page_numbers") and "footer" not in asked and page["footer"] < 200:
        page["footer"] = added["footer"]
    room = data["narrowed_margin"]
    for side, distance in (("top", "header"), ("bottom", "footer")):
        if side in asked and distance not in asked and page[distance] >= 200 \
                and page[distance] + room["line_room"] > page[side]:
            page[distance] = max(room["minimum_distance"], page[side] - room["line_room"])
    return page


def contract(template, request=None):
    data = json.loads((ROOT / "assets/checks/audit-baseline.json").read_text(encoding="utf-8"))
    result = deepcopy(data["roles"])
    for role, patch in data.get("overrides", {}).get(template, {}).items():
        for tag, attrs in patch.items():
            result[role]["pPr"][tag].update(attrs)
    # Apply explicit user deltas to an independent fixed baseline, not writer output.
    request = request or {}
    for role, patch in request.get("styles", {}).items():
        rule = result[role]
        pp = rule["pPr"]
        ft = rule["font"]
        if "cn" in patch:
            ft["eastAsia"] = patch["cn"]
        if "latin" in patch:
            ft["ascii"] = ft["hAnsi"] = patch["latin"]
        if "size" in patch:
            ft["size_half_points"] = int(patch["size"] * 2)
        if "bold" in patch:
            ft["bold_required"] = patch["bold"]
            ft["bold_exact"] = True
        if "align" in patch:
            pp["jc"]["val"] = patch["align"]
        keys = {
            "first": "firstLine",
            "chars": "firstLineChars",
            "left": "left",
            "left_chars": "leftChars",
            "hanging": "hanging",
            "right": "right",
        }
        if "first" in patch:
            pp["ind"].pop("firstLineChars", None)
            pp["ind"].pop("hanging", None)
        if "chars" in patch:
            pp["ind"].pop("hanging", None)
        if "hanging" in patch:
            pp["ind"].pop("firstLine", None)
            pp["ind"].pop("firstLineChars", None)
        if "left" in patch:
            pp["ind"].pop("leftChars", None)
        for key, target in keys.items():
            if key in patch:
                pp["ind"][target] = str(patch[key])
        keys = {
            "before": "before",
            "after": "after",
            "line": "line",
            "line_rule": "lineRule",
            "before_lines": "beforeLines",
            "after_lines": "afterLines",
        }
        for key in ("before", "after"):
            if key in patch and key + "_lines" not in patch:
                pp["spacing"].pop(key + "Lines", None)
        for key, target in keys.items():
            if key in patch:
                pp["spacing"][target] = str(patch[key])
        if "keep" in patch:
            pp["keepNext"]["val"] = pp["keepLines"]["val"] = str(int(patch["keep"]))
        if "break_before" in patch:
            pp["pageBreakBefore"]["val"] = str(int(patch["break_before"]))
        for key in ("italic", "color", "underline"):
            if key in patch:
                ft[key] = patch[key].upper() if key == "color" else patch[key]
    # Physical fallback must match the character-unit indent at the effective size.
    for role, rule in result.items():
        ind = rule["pPr"].get("ind", {})
        size = rule["font"]["size_half_points"] / 2
        for chars, absolute in (("firstLineChars", "firstLine"), ("leftChars", "left")):
            if chars in ind:
                ind[absolute] = str(round(size * 20 * int(ind[chars]) / 100))
        if role in ("keywords_zh", "keywords_en"):
            ft = deepcopy(rule["font"])
            ft["size_half_points"] = 28
            patch = request.get("keyword_labels", {}).get(role[-2:], {})
            if "size" in patch:
                ft["size_half_points"] = int(patch["size"] * 2)
            if "cn" in patch:
                ft["eastAsia"] = patch["cn"]
            if "latin" in patch:
                ft["ascii"] = ft["hAnsi"] = patch["latin"]
            if "bold" in patch:
                ft["bold_required"] = patch["bold"]
                ft["bold_exact"] = True
            for key in ("italic", "color", "underline"):
                if key in patch:
                    ft[key] = patch[key].upper() if key == "color" else patch[key]
            rule["label_font"] = ft
    return result


def paragraph_check(p, role, label, want, err, warn):
    for tag, expected in want["pPr"].items():
        actual = paragraph_properties(p, tag)
        for key, value in expected.items():
            got = actual.get(key)
            if tag in (
                "keepNext",
                "keepLines",
                "widowControl",
                "pageBreakBefore",
                "snapToGrid",
                "kinsoku",
                "overflowPunct",
            ):
                got = "1" if on(got) and got is not None else "0"
            # Empty section/page-break carriers are deliberately one-point tall.
            if (
                role == "spacer"
                and tag == "spacing"
                and key == "line"
                and not visible(p).strip()
                and got == "20"
            ):
                continue
            # A table caption must stay with its following table; a bare figure has no caption.
            if role == "caption" and tag == "keepNext":
                nxt = p._p.getnext()
                if nxt is not None and nxt.xpath('./w:pPr/w:pStyle[@w:val="XAUFETableUnit"]'):
                    nxt = nxt.getnext()  # caption, unit line, table: all three stay together
                value = "1" if nxt is not None and nxt.tag == qn("w:tbl") else "0"
            if role == "figure" and tag == "keepNext":
                nxt = p._p.getnext()
                if (
                    nxt is None
                    or nxt.tag != qn("w:p")
                    or not nxt.xpath('./w:pPr/w:pStyle[@w:val="XAUFECaption"]')
                ):
                    value = "0"
            if got != value:
                err("PARAGRAPH", f"{label}: {tag}/{key}={got!r}，应为{value!r}。")
        # Extra properties can override an otherwise correct indent or spacing.
        for key, value in actual.items():
            if key in expected:
                continue
            if tag in ("spacing", "ind"):
                harmless_zero = value in ("0", "false", "off")
                if key in ("beforeLines", "afterLines") and expected.get(key[:-5], "0") != "0":
                    harmless_zero = False
                if tag == "ind":
                    # Even zero-valued alternate-unit attributes can override nonzero indents.
                    paired = {
                        "leftChars": "left",
                        "startChars": "left",
                        "rightChars": "right",
                        "endChars": "right",
                        "firstLineChars": "firstLine",
                        "hangingChars": "hanging",
                        "start": "left",
                        "end": "right",
                    }
                    if key in paired and expected.get(paired[key], "0") != "0":
                        harmless_zero = False
                    if key in ("hanging", "hangingChars") and (
                        expected.get("firstLine", "0") != "0"
                        or expected.get("firstLineChars", "0") != "0"
                    ):
                        harmless_zero = False
                if not harmless_zero:
                    err("PARAGRAPH_EXTRA", f"{label}: 多余的{tag}/{key}={value}改变实际布局。")
    for tag in ("contextualSpacing", "mirrorIndents", "adjustRightInd"):
        actual = paragraph_properties(p, tag)
        if actual and on(actual.get("val")):
            err("PARAGRAPH_EXTRA", f"{label}: {tag}改变模板布局。")
    if role in ("keywords_zh", "keywords_en"):
        problem = spacing_issue(p)
        if problem:
            err("KEYWORD_SPACING", label + ": " + problem)
    if role in ("h1", "h2", "h3") and not p._p.xpath("./w:pPr/w:numPr"):
        patterns = {
            "h1": r"^\s*[一二三四五六七八九十百零〇]+、",
            "h2": r"^\s*[（(][一二三四五六七八九十百零〇]+[）)]",
            "h3": r"^\s*\d+[.．]",
        }
        if not re.match(patterns[role], visible(p)):
            warn(
                "HEADING_NUMBERING",
                f"标题“{visible(p)}”的编号和它的级别（{role}）对不上：一级用“一、”，二级用“（一）”，三级用“1.”。"
                "先看 mapping.json/content.json 里这一段的级别是否填错；原稿编号本来就是这样的，就在 review.json 的 warnings 里说明，不要改原文。",
            )


# Roles whose weight the template fixes: a bold run here is a format error, not emphasis.
FIXED_WEIGHT = {"h1", "h2", "h3", "h4", "preface", "reference_title", "caption", "title_zh", "title_en",
                "abstract_label_zh", "abstract_label_en", "toc_title", "toc1", "toc2", "toc3", "toc4", "header"}
# Running text may bold a few words, but a wholly bold paragraph is almost always a stray format.
RUNNING_TEXT = {"body", "reference", "abstract_zh", "abstract_en", "code"}


def run_check(p, role, label, want, err, warn=None):
    offset = 0
    bold_runs = plain_runs = 0
    match = LABEL.match(visible(p)) if role.startswith("keywords_") else None
    for ri, r in enumerate(p._p.xpath(".//w:r")):
        text = "".join(r.xpath("./w:t/text()"))
        if not text:
            continue
        try:
            props = run_properties(p, r)
        except ValueError as exc:
            err("STYLE_INHERITANCE", f"{label}: {exc}")
            continue
        ft = (
            want.get("label_font", dict(want["font"], size_half_points=28))
            if match and offset < match.end()
            else want["font"]
        )
        expected_size = ft["size_half_points"]
        if match and offset < match.end() < offset + len(text):
            err("KEYWORD_RUN", f"{label}: 标签和词项跨字号边界，需分离运行。")
        offset += len(text)
        if props.get("sz") != str(expected_size):
            err(
                "FONT_SIZE", f'{label}/r{ri}: 有效字号={props.get("sz")}半磅，应为{expected_size}。'
            )
        fonts = props["fonts"]
        for attr, code in [("eastAsia", "FONT_CN"), ("ascii", "FONT_EN"), ("hAnsi", "FONT_EN")]:
            if fonts.get(attr) != ft[attr] or fonts.get(attr + "Theme"):
                err(code, f"{label}/r{ri}: 有效{attr}字体={fonts.get(attr) or fonts.get(attr + 'Theme')}，应为{ft[attr]}。")
        if ft["bold_required"] and not props.get("b", False):
            err("FONT_BOLD", f"{label}/r{ri}: 加粗=否，应为是（要求加粗）。")
        if ft.get("bold_exact") and props.get("b", False) != ft["bold_required"]:
            err("FONT_BOLD", f"{label}/r{ri}: 加粗={'是' if props.get('b', False) else '否'}，应为{'是' if ft['bold_required'] else '否'}（用户明确的加粗开关）。")
        elif role in FIXED_WEIGHT and not ft["bold_required"] and props.get("b", False):
            err("FONT_BOLD", f"{label}/r{ri}: 加粗=是，应为否（模板这一级不加粗）。")
        if text.strip():
            bold_runs, plain_runs = bold_runs + bool(props.get("b", False)), plain_runs + (not props.get("b", False))
        if props.get("vanish") or props.get("webHidden") or props.get("specVanish"):
            err("HIDDEN_TEXT", f"{label}/r{ri}: 存在隐藏正文。")
        rpr = r.find(qn("w:rPr"))
        def shown(tag):
            node = rpr.find(qn("w:" + tag)) if rpr is not None else None
            if node is None:
                return False
            if tag == "shd":  # a clear pattern still paints its fill colour
                return node.get(qn("w:val")) not in ("clear", "nil") or (node.get(qn("w:fill")) or "auto").upper() not in ("AUTO", "FFFFFF")
            return node.get(qn("w:val")) not in ("none", "0", "false")

        looks = [name for tag, name in (("highlight", "高亮"), ("shd", "底纹"), ("emboss", "阳文"), ("imprint", "阴文"),
                                        ("outline", "空心"), ("shadow", "阴影"), ("effect", "动态效果")) if shown(tag)]
        if looks:
            err("FONT_DECORATION", f"{label}/r{ri}: 文字带{'、'.join(looks)}，模板里没有这种格式。")
        marks = [name for key, name in (("strike", "删除线"), ("dstrike", "双删除线"), ("caps", "全部大写"),
                                        ("smallCaps", "小型大写")) if props.get(key)]
        if rpr is not None and rpr.find(qn("w:em")) is not None and rpr.find(qn("w:em")).get(qn("w:val")) != "none":
            marks.append("着重号")
        if marks and warn:
            warn("FONT_MARKS", f"{label}/r{ri}: “{text[:20]}”带{'、'.join(marks)}，是原稿里的格式，成品保留了原样（模板没有这种格式）。")
        for attr, default in [("w", "100"), ("spacing", "0"), ("position", "0")]:
            if props.get(attr, default) != default:
                err("FONT_GEOMETRY", f"{label}/r{ri}: 字形宽度/间距/基线偏移{attr}未归一。")
        if props.get("color", "000000").upper() not in (
            {"000000", "AUTO"} if ft.get("color", "000000") == "000000" else {ft["color"]}
        ):
            err("FONT_COLOR", f"{label}/r{ri}: 文字颜色={props.get('color', '000000')}，应为{ft.get('color', '000000')}。")
        if "italic" in ft and (
            props.get("i", False) != ft["italic"] or props.get("iCs", False) != ft["italic"]
        ):
            err("FONT_ITALIC", f"{label}/r{ri}: 斜体不符合用户要求。")
        if "underline" in ft and props.get("u", "none") != ft["underline"]:
            err("FONT_UNDERLINE", f"{label}/r{ri}: 下划线不符合用户要求。")
        color = r.find("./" + qn("w:rPr") + "/" + qn("w:color"))
        if color is not None and any(
            k.split("}")[-1] in ("themeColor", "themeTint", "themeShade") for k in color.attrib
        ):
            err("FONT_COLOR", f"{label}/r{ri}: 主题色仍可能覆盖黑色。")
    if (warn and role in RUNNING_TEXT and bold_runs and not plain_runs and not want["font"].get("bold_exact")
            and len(visible(p).strip()) > 20):
        warn("FONT_BOLD_PARAGRAPH", f"{label}: 整段都是加粗，模板正文不加粗（只强调个别词可以）。")


# 三线表 (user rule 2026-10-26): top and bottom 1.5 pt (sz 12), under the header 0.25 pt (sz 2),
# no vertical rules and no other horizontal rules.  Written here from the rule, not from table_layout.
OUTER_RULE, HEADER_RULE = "12", "2"


def _visible(edge):
    return edge is not None and edge.get(qn("w:val")) not in (None, "nil", "none")


MIN_PADDING = 80  # twips; Word's usual cell margin is 108


def three_line_check(table, label, err):
    where = f"{label}（第 {int(label[1:]) + 1} 个表格）"
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    edges = {e: borders.find(qn("w:" + e)) if borders is not None else None
             for e in ("top", "bottom", "left", "right", "insideH", "insideV")}
    for e in ("top", "bottom"):
        if not _visible(edges[e]) or edges[e].get(qn("w:sz")) != OUTER_RULE:
            got = f"{int(edges[e].get(qn('w:sz'), '0')) / 8:g} 磅" if _visible(edges[e]) else "没有"
            err("TABLE_RULES", f"{where}: {'顶线' if e == 'top' else '底线'}应为 1.5 磅，实际{got}。")
    for e in ("left", "right", "insideH", "insideV"):
        if _visible(edges[e]):
            err("TABLE_RULES", f"{where}: 三线表不设竖线和其他横线，表格还有{ {'left': '左边框', 'right': '右边框', 'insideH': '内部横线', 'insideV': '内部竖线'}[e] }。")
    mar = table._tbl.tblPr.find(qn("w:tblCellMar"))
    for side in ("left", "right"):
        edge = mar.find(qn("w:" + side)) if mar is not None else None
        if edge is None or int(edge.get(qn("w:w"), "0")) < MIN_PADDING:
            err("TABLE_PADDING", f"{where}: 单元格{'左' if side == 'left' else '右'}边没有留白（应约 0.19 厘米），文字会贴着表格边缘。")
            break
    rows = table._tbl.tr_lst
    marked = [i for i, tr in enumerate(rows) if tr.xpath('./w:trPr/w:tblHeader[not(@w:val="0") and not(@w:val="false")]')]
    header = max(marked or [0])
    for ri, tr in enumerate(rows):
        for tc in tr.tc_lst:
            sides = {e.tag.split("}")[1]: e for e in tc.xpath("./w:tcPr/w:tcBorders/*")}
            for side, edge in sides.items():
                wanted = ri == header and side == "bottom"
                if _visible(edge) and not wanted:
                    err("TABLE_RULES", f"{where} 第 {ri + 1} 行: 单元格自带{side}边框，三线表只在表头下有一条线。")
                    return
            if ri == header and (not _visible(sides.get("bottom")) or sides["bottom"].get(qn("w:sz")) != HEADER_RULE):
                err("TABLE_RULES", f"{where}: 表头下面的分隔线应为 0.25 磅。")
                return


def tables_check(doc, width, err):
    count = 0
    for i, t in enumerate(doc.tables):
        try:
            for label, table, available in walk_tables(t, width, f"t{i}"):
                count += 1
                grid = [int(x.get(qn("w:w"), "0")) for x in table._tbl.tblGrid.gridCol_lst]
                total = sum(grid)
                if not grid or min(grid) <= 0:
                    err("TABLE_GRID", label + ": 列网格无效。")
                    continue
                if total > available:
                    err("TABLE_WIDTH", f"{label}: {total}twip超过可用{available}twip。")
                pp = table._tbl.tblPr
                if "/" not in label:
                    three_line_check(table, label, err)
                tw = pp.find(qn("w:tblW"))
                if tw is None or tw.get(qn("w:type")) != "dxa" or tw.get(qn("w:w")) != str(total):
                    err("TABLE_WIDTH", label + ": 表格首选宽度与网格宽度不一致。")
                alignment = pp.find(qn("w:jc"))
                if alignment is None or alignment.get(qn("w:val")) != "center":
                    err(
                        "TABLE_ALIGNMENT",
                        label + ": 表格须居中，避免左对齐兼容模式按单元格文字而非外边框定位。",
                    )
                indent = pp.find(qn("w:tblInd"))
                if indent is not None and int(indent.get(qn("w:w"), "0")) != 0:
                    err("TABLE_INDENT", label + ": 表格额外缩进未归一。")
                if pp.find(qn("w:tblpPr")) is not None:
                    err("TABLE_FLOAT", label + ": 浮动表格不在本检查器支持范围。")
                spacing = pp.find(qn("w:tblCellSpacing"))
                if spacing is not None and int(spacing.get(qn("w:w"), "0")):
                    err("TABLE_SPACING", label + ": 单元格外间距会额外扩宽表格。")
                for ri, tr, _, _, cells in raw_rows(table):
                    if tr.xpath('./w:trPr/w:trHeight[@w:hRule="exact"]'):
                        err(
                            "TABLE_FIXED_HEIGHT",
                            f"{label}/r{ri}: 固定行高可能裁字，应改自动或最小值。",
                        )
                    for col, span, cell in cells:
                        expected = sum(grid[col : col + span])
                        cw = cell._tc.tcPr.find(qn("w:tcW"))
                        if (
                            cw is None
                            or cw.get(qn("w:type")) != "dxa"
                            or cw.get(qn("w:w")) != str(expected)
                        ):
                            err("CELL_WIDTH", f"{label}/r{ri}/c{col}: 单元格宽度与合并网格不一致。")
                        for tag in ("noWrap", "tcFitText"):
                            node = cell._tc.tcPr.find(qn("w:" + tag))
                            if node is not None and on(node.get(qn("w:val"))):
                                err(
                                    "CELL_LAYOUT",
                                    f"{label}/r{ri}/c{col}: {tag}可能导致溢出或缩字。",
                                )
                        content = (
                            expected
                            - cell_margin(table, cell, "left")
                            - cell_margin(table, cell, "right")
                        )
                        if content < 210:
                            err("CELL_WIDTH", f"{label}/r{ri}/c{col}: 内边距后无法容纳五号文字。")
                        for paragraph in cell.paragraphs:
                            for inline in paragraph._p.xpath(".//wp:inline"):
                                if any(
                                    inline.get(key) != "0"
                                    for key in ("distT", "distB", "distL", "distR")
                                ):
                                    err(
                                        "CELL_IMAGE_SPACING",
                                        f"{label}/r{ri}/c{col}: 行内图片留白必须显式为0，避免兼容渲染默认间隔挤出边界。",
                                    )
                                ext = inline.find(qn("wp:extent"))
                                if ext is None:
                                    err("CELL_IMAGE_WIDTH", f"{label}/r{ri}/c{col}: 图片缺少尺寸。")
                                elif int(ext.get("cx", "0")) > content * 635:
                                    err(
                                        "CELL_IMAGE_WIDTH",
                                        f"{label}/r{ri}/c{col}: 图片超出单元格内容区。",
                                    )
        except (ValueError, TypeError, AttributeError) as exc:
            err("TABLE_STRUCTURE", f"t{i}: {exc}")
    return count


def toc_entries(doc):
    """Return actual anchor targets and fields; titles are roles, not substring guesses."""
    bookmarks = {}
    for p in doc.paragraphs:
        for el in p._p.xpath("./w:bookmarkStart"):
            name = el.get(qn("w:name"))
            if name in bookmarks:
                raise ValueError("书签名称重复：" + name)
            bookmarks[name] = p
    entries = []
    for p in doc.paragraphs:
        style_name = getattr(p.style, "name", None)
        if style_name not in ("TOC 1", "TOC 2", "TOC 3", "TOC 4"):
            continue
        fields = p._p.xpath("./w:fldSimple")
        pagefields = [
            f for f in fields if re.search(r"\bPAGEREF\b", f.get(qn("w:instr"), ""), re.I)
        ]
        if len(pagefields) != 1:
            raise ValueError("目录条目不是受支持的单个PAGEREF结构；不能伪造数字验收。")
        fld = pagefields[0]
        m = re.search(r'\bPAGEREF\s+(?:"([^"]+)"|([^\s\\]+))', fld.get(qn("w:instr"), ""), re.I)
        if not m:
            raise ValueError("目录PAGEREF缺少目标书签。")
        anchor = m.group(1) or m.group(2)
        target = bookmarks.get(anchor)
        if target is None:
            raise ValueError("目录目标书签不存在：" + anchor)
        expected_levels = {
            "Heading 1": 1,
            "Heading 2": 2,
            "Heading 3": 3,
            "Heading 4": 4,
            "XAUFE Preface": 1,
            "XAUFE References Title": 1,
        }
        target_style = getattr(target.style, "name", None)
        if target_style not in expected_levels:
            raise ValueError("目录不得指向摘要或非正文章节：" + anchor)
        level = int(style_name[-1])
        if level != expected_levels[target_style]:
            raise ValueError("目录层级与目标标题层级不一致：" + anchor)
        links = p._p.xpath("./w:hyperlink[@w:anchor]")
        if len(links) != 1 or links[0].get(qn("w:anchor")) != anchor:
            raise ValueError("目录链接与页引用不是同一目标。")
        text = "".join(links[0].xpath(".//w:t/text()"))
        if re.sub(r"\s+", "", text) != re.sub(r"\s+", "", visible(target)):
            raise ValueError("目录标题缓存与实际标题不一致：" + anchor)
        entries.append(
            {
                "anchor": anchor,
                "text": text,
                "level": level,
                "field": fld,
                "target": target,
                "paragraph": p,
                "cache": "".join(x.text or "" for x in fld.findall(".//" + qn("w:t"))),
            }
        )
    return entries


def toc_proof_check(path, entries, vv, final, err, warn, proof_path=None):
    if not entries:
        return
    report_path = Path(proof_path) if proof_path else Path(path).with_suffix(".toc-report.json")
    if not report_path.is_file():
        (err if final else warn)(
            "TOC_PROOF_REQUIRED", "缺少与当前DOCX绑定的目录渲染验证报告；刷新目录后再终审。"
        )
        return
    try:
        report = json.loads(report_path.read_text(encoding="utf-8"))
        if report.get("document_sha256") != hashlib.sha256(Path(path).read_bytes()).hexdigest():
            raise ValueError("DOCX已修改或报告不对应当前文件")
        if (
            not report.get("converged")
            or report.get("template") != vv.get("XAUFE_TEMPLATE")
            or report.get("job_id") != vv.get("XAUFE_INTAKE_ID")
        ):
            raise ValueError("报告模板、任务或收敛状态不符")
        expected = [
            {
                "anchor": e["anchor"],
                "text": re.sub(r"\s+", "", e["text"]),
                "logical_page": int(e["cache"]),
            }
            for e in entries
        ]
        if report.get("entries") != expected:
            raise ValueError("目录缓存与渲染记录不一致")
        if not report.get("rendered_pdf_sha256"):
            raise ValueError("没有实际渲染摘要")
        from render_preview import font_report

        if report.get("font_resolution") != font_report():
            (err if final else warn)(
                "TOC_RENDER_ENV_CHANGED",
                "当前字体解析与目录验证时不同；需在当前环境重新渲染验证页码。",
            )
    except (ValueError, KeyError, TypeError, OSError) as exc:
        (err if final else warn)("TOC_PROOF_STALE", "目录验证证据无效：" + str(exc))


CJK = re.compile(r"[㐀-䶿一-鿿豈-﫿　-〿＀-￯‘-‟…—·]")


def word_count(text):
    """Word's 字数统计: each CJK character or full-width mark counts 1, each other run without spaces counts 1."""
    text = text or ""
    return len(CJK.findall(text)) + len(CJK.sub(" ", text).split())


def lengths():
    return json.loads((ROOT / "assets/checks/audit-baseline.json").read_text(encoding="utf-8"))["lengths"]
