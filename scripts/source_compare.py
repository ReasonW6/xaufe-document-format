"""Compare the delivered Word with the user's original (format-only fidelity).

Text is compared paragraph by paragraph, in order.  Only layout whitespace is
ignored (see norm); full-width and half-width forms count as different, so a
changed digit, decimal point, sign, letter case or punctuation mark, and a dropped,
added, duplicated or moved paragraph are differences, reported with the exact
span.  Keyword separators may change (the template writes them its own way).

Objects are counted too: images, tables (with their row x column grid),
equations and link targets.

Parts the template generates (cover, TOC, labels such as 摘要/参考文献, the
reference numbers, the header) are not counted as added; the old cover and old
TOC the template replaces are not counted as missing; front-matter paragraphs
(title, abstract, keywords) may move to where the template puts them.
"""

from pathlib import Path
import difflib, re, unicodedata
from docx import Document
from docx.oxml.ns import qn
from textio import read_text_any

GENERATED_ROLES = {
    "toc_title", "toc1", "toc2", "toc3", "toc4", "abstract_label_zh", "abstract_label_en",
    "reference_title", "header", "spacer", "abstract_spacer",
}
MOVABLE_ROLES = {"title_zh", "title_en", "abstract_zh", "abstract_en", "keywords_zh", "keywords_en"}
NUMBER = re.compile(r"^\s*[\[［]\d+[\]］]\s*")
KEYWORDS = re.compile(r"^\s*(关键词|key\s*words?)\s*[:：]", re.I)
# Part labels the template prints itself; the original's own copy is replaced, not lost.
PART_LABELS = {"摘要", "内容摘要", "中文摘要", "abstract", "参考文献", "主要参考文献", "目录", "目錄"}
OBJECT_LABELS = (("images", "图片"), ("tables", "表格"), ("equations", "公式"), ("links", "链接"))


def norm(text, exact_spaces=False):
    """The text to keep, character for character (full-/half-width punctuation counts).

    Only layout noise is ignored: the reference number the template prints, line
    breaks and runs of spaces, and a space next to a Chinese character or a
    punctuation mark (hard-wrapped .txt lines).  A space between two Latin letters
    or digits is text ("Times New Roman").  Keyword lines compare their words only,
    because the template sets the separators.
    """
    text = NUMBER.sub("", unicodedata.normalize("NFC", str(text or "")))
    if KEYWORDS.match(text):
        return re.sub(r"[\s;,、；，]+", "", text.lower().replace("：", ":"))
    if exact_spaces:  # Word to Word: every space inside a paragraph is the user's text
        return text.strip(" \t\r\n\u3000")
    text = re.sub(r"\s+", " ", text).strip()
    return re.sub(r"(?<![A-Za-z0-9]) | (?![A-Za-z0-9])", "", text)


def _plain(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _text(element):
    return "".join(element.xpath(".//w:t/text()|.//m:t/text()"))


def _cells(table):
    return ["".join(_text(p) for p in cell.xpath("./w:p")) for cell in table.xpath(".//w:tc")]


def _has_cover(doc):
    return any(
        v.get(qn("w:name")) == "XAUFE_HAS_COVER" and v.get(qn("w:val")) == "1"
        for v in doc.settings.element.xpath("./w:docVars/w:docVar")
    )


def fragments(path):
    """[(location, text, block_index)] of the original in reading order; None for unreadable types."""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        doc = Document(path)
        # TOC entries are regenerated from the headings (their page numbers change), so they are not text to keep.
        toc = {st.style_id for st in doc.styles if st.name.lower().startswith("toc ")}
        result, pno, tno = [], 0, 0
        for index, el in enumerate(doc.element.body):
            if el.tag == qn("w:p"):
                if _text(el).strip() and not set(el.xpath("./w:pPr/w:pStyle/@w:val")) & toc:
                    result.append((f"p{pno}", _text(el), index))
                pno += 1
            elif el.tag == qn("w:tbl"):
                result += [(f"t{tno}", text, index) for text in _cells(el) if text.strip()]
                tno += 1
        return result
    if path.suffix.lower() not in (".md", ".txt"):
        return None
    text = read_text_any(path)
    result, lines, start, fenced = [], [], 1, False

    def flush():
        if lines:
            result.append((f"第 {start} 行", " ".join(lines), None))
        lines.clear()

    for index, line in enumerate(text.splitlines(), 1):
        stripped = line.strip()
        here = f"第 {index} 行"
        if stripped.startswith(("```", "~~~")):
            flush()
            fenced = not fenced
            continue
        if fenced:
            if stripped:
                result.append((here, stripped, None))
            continue
        if not stripped or stripped.startswith("![") or re.fullmatch(r"[-*_=\s]{3,}", stripped):
            flush()
            continue
        if stripped.startswith("|"):
            flush()
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells if c):
                result += [(here, c, None) for c in cells if c]
            continue
        if stripped.startswith("#"):
            flush()
            result.append((here, stripped.lstrip("#").strip(), None))
            continue
        title = re.match(r"^(?:English title|英文题目)\s*[:：]\s*(.+)$", stripped, re.I)
        if title or KEYWORDS.match(stripped):
            flush()
            result.append((here, title.group(1) if title else stripped, None))
            continue
        # Visible link labels and emphasised text, not Markdown syntax.
        stripped = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", stripped)
        stripped = re.sub(r"(?<!\\)(\*\*|__|`)", "", stripped)
        stripped = re.sub(r"^[-*+]\s+", "", stripped)
        if not lines:
            start = index
        lines.append(stripped)
    flush()
    return result


def objects(path, start=0):
    """Non-text content from body block ``start`` on: images, tables (grids), equations, link targets."""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        doc = Document(path)
        body = list(doc.element.body)[start:]
        links = sorted(
            doc.part.rels[rid].target_ref
            for el in body for rid in el.xpath(".//w:hyperlink/@r:id")
            if rid in doc.part.rels and doc.part.rels[rid].is_external
        )
        tables = [el for el in body if el.tag == qn("w:tbl")]
        return {
            "images": sum(len(el.xpath(".//a:blip|.//*[local-name()='imagedata']")) for el in body),
            "tables": [f"{len(t.xpath('./w:tr'))}x{max((len(r.xpath('./w:tc')) for r in t.xpath('./w:tr')), default=0)}"
                       for t in tables],
            "equations": sum(len(el.xpath(".//m:oMath")) for el in body),
            "links": links,
        }
    if path.suffix.lower() != ".md":
        return None
    text = re.sub(r"```.*?```|~~~.*?~~~", "", read_text_any(path), flags=re.S)
    tables, rows, cols = [], 0, 0
    for line in text.splitlines() + [""]:
        if line.strip().startswith("|"):
            cells = line.strip().strip("|").split("|")
            if not all(re.fullmatch(r"\s*:?-+:?\s*", c) for c in cells):
                rows, cols = rows + 1, max(cols, len(cells))
        elif rows:
            tables.append(f"{rows}x{cols}")
            rows = cols = 0
    return {
        "images": len(re.findall(r"!\[[^\]]*\]\([^)]*\)", text)),
        "tables": tables,
        "equations": None,  # Markdown formulas are text and are compared as text
        "links": sorted(re.findall(r"(?<!!)\[[^\]]+\]\((https?://[^)\s]+)\)", text)),
    }


def body_key(docx):
    """All visible text except the TOC (its page numbers may change), for "only formatting changed" checks."""
    doc = Document(docx)
    toc = {s.style_id for s in doc.styles if s.name.lower().startswith("toc ")}
    return norm("".join(
        _text(p) for p in doc.element.body.xpath(".//w:p")
        if not set(p.xpath("./w:pPr/w:pStyle/@w:val")) & toc
    ))


def first_text_difference(before, after):
    """(location in ``after``, was, now) of the first paragraph whose text differs; None when equal."""
    def paragraphs(docx):
        doc = Document(docx)
        toc = {s.style_id for s in doc.styles if s.name.lower().startswith("toc ")}
        return [(f"第 {i + 1} 段", _plain(_text(p))) for i, p in enumerate(doc.element.body.xpath(".//w:p"))
                if _text(p).strip() and not set(p.xpath("./w:pPr/w:pStyle/@w:val")) & toc]

    a, b = paragraphs(before), paragraphs(after)
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, [norm(t) for _, t in a], [norm(t) for _, t in b],
                                                      autojunk=False).get_opcodes():
        if op != "equal":
            was = a[i1][1] if i1 < i2 else ""
            now = b[j1][1] if j1 < j2 else ""
            where = b[j1][0] if j1 < j2 else (b[j1 - 1][0] + " 之后" if j1 else "开头")
            return (where,) + (_span(was, now) if was and now else (was, now))
    return None


def _after_cover(doc):
    """Index of the first body block after the template cover (0 when there is none)."""
    if not _has_cover(doc):
        return 0
    for i, el in enumerate(doc.element.body):
        if el.tag == qn("w:p") and el.xpath("./w:pPr/w:sectPr"):
            return i + 1
    return 0


def output_paragraphs(docx):
    """[(label, role, text)] of the output body in order, skipping the template's cover section."""
    from customization import ROLE_STYLES

    doc = Document(docx)
    roles = {name: role for role, name in ROLE_STYLES.items()}
    start = _after_cover(doc)
    result = []
    for number, el in enumerate(doc.element.body):
        if number < start:
            continue
        if el.tag == qn("w:p") and _text(el).strip():
            style = el.xpath("./w:pPr/w:pStyle/@w:val")
            name = doc.styles.get_by_id(style[0], 1).name if style else "Normal"
            result.append((f"成品第 {number + 1} 段", roles.get(name), _text(el)))
        elif el.tag == qn("w:tbl"):
            result += [(f"成品第 {number + 1} 块（表格）", "table_text", text) for text in _cells(el) if text.strip()]
    return result


def _covered(text, approved):
    """An approved change must contain the whole changed text, not just share a few characters."""
    t = norm(text)
    return not t or any(t in norm(a) for a in approved)


def _span(a, b):
    """The differing stretch of two strings, with a little context."""
    ops = [op for op in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes() if op[0] != "equal"]
    if not ops:
        return a, b
    return a[max(0, ops[0][1] - 8) : ops[-1][2] + 8], b[max(0, ops[0][3] - 8) : ops[-1][4] + 8]


def margin_texts(path):
    """Header and footer lines of a .docx; digits dropped so "第 3 页" and "第 1 页" are one line."""
    from docx import Document

    seen = []
    for section in Document(path).sections:
        for part in (section.header, section.first_page_header, section.even_page_header,
                     section.footer, section.first_page_footer, section.even_page_footer):
            if part.is_linked_to_previous:  # no own definition: nothing new to show (and reading would add one)
                continue
            for p in part.paragraphs:
                text = "".join(p._p.xpath(".//w:t/text()"))
                key = re.sub(r"[\s\d]+", "", text)
                if key and key not in [k for k, _ in seen]:
                    seen.append((key, text.strip()))
    return seen


def dropped_margins(sources, docx):
    """Header/footer text of a .docx original that the output's headers and footers no longer show."""
    kept = {k for k, _ in margin_texts(docx)}
    out = []
    for source in sources:
        path = Path(source["path"])
        if path.suffix.lower() == ".docx" and path.is_file():
            out += [text for key, text in margin_texts(path) if not any(key in k or k in key for k in kept)]
    return out


def has_page_numbers(path):
    """The .docx shows page numbers: a PAGE field in one of its headers or footers."""
    from zipfile import ZipFile

    with ZipFile(path) as z:
        return any(re.search(r"\bPAGE\b", z.read(n).decode("utf-8", "replace"))
                   for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml$", n))


def dropped_page_numbers(sources, docx):
    """A .docx original had page numbers and the output shows none."""
    return not has_page_numbers(docx) and any(
        Path(s["path"]).suffix.lower() == ".docx" and Path(s["path"]).is_file() and has_page_numbers(s["path"])
        for s in sources)


def _approved_count(approved, label):
    """How many objects of this kind the user agreed to add or drop, by the numbers next to the label."""
    count = 0
    for text in approved:
        for match in re.finditer(rf"(\d+)\s*[张个幅处条项份]?\s*{label}|{label}\s*(\d+)\s*[张个幅处条项份]", text):
            count += int(match.group(1) or match.group(2))
    return count


def compare(sources, docx, *, approved=(), generated=(), replaced=None, old_toc=()):
    """Differences between the readable originals and the output.

    ``approved``: texts the user agreed to add or drop (each must contain the whole change;
    for objects, a text naming 图片/表格/公式/链接); ``generated``: texts from job.json (cover
    values) or mapping.json "add" (agreed translations); ``replaced``/``old_toc``: the old
    cover block range and old TOC paragraph ids of a .docx original.
    """
    allowed = {norm(g) for g in generated if norm(g)}
    target = [(loc, role, text) for loc, role, text in output_paragraphs(docx)
              if role not in GENERATED_ROLES and norm(text) not in allowed]
    original, reports, wanted = [], [], []
    for source in sources:
        path = Path(source["path"])
        items = fragments(path)
        if items is None:
            reports.append({"source": str(path), "status": "manual_review", "reason": "此文件类型需按原件核对文字或对象。"})
            continue
        docx_source = path.suffix.lower() == ".docx"
        for loc, text, block in items:
            if docx_source and ((replaced is not None and block is not None and block < replaced) or loc in old_toc):
                continue
            if norm(text) and norm(text) not in allowed and norm(text).lower() not in PART_LABELS:
                original.append((f"{path.name} 的 {loc}", text))
        wanted.append(objects(path, (replaced or 0) if docx_source else 0))
        reports.append({"source": str(path), "status": "compared", "checked_fragments": len(items)})

    differences = []
    if not wanted:
        return {"sources": reports, "differences": differences}
    # A .txt/.md original may be hard-wrapped, so spaces beside Chinese are noise there; a Word original is exact.
    exact = all(Path(x["path"]).suffix.lower() == ".docx" for x in sources)
    a, b = [norm(t, exact) for _, t in original], [norm(t, exact) for _, _, t in target]
    removed, added = [], []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if op != "equal" and "".join(a[i1:i2]) != "".join(b[j1:j2]):  # equal joined = only re-paragraphed
            removed += range(i1, i2)
            added += range(j1, j2)
    # A paragraph that only moved: front matter may (the template orders it), body text may not.
    for i in list(removed):
        j = next((j for j in added if b[j] == a[i]), None)
        if j is not None:
            removed.remove(i)
            added.remove(j)
            if target[j][1] not in MOVABLE_ROLES and not _covered(target[j][2], approved):
                differences.append({"kind": "moved", "location": f"{original[i][0]} → {target[j][0]}", "original": original[i][1]})
    # What is left, in order: a changed paragraph when the two are similar, else dropped / added.
    while removed or added:
        i = removed.pop(0) if removed else None
        j = added.pop(0) if added else None
        if i is not None and j is not None and difflib.SequenceMatcher(None, a[i], b[j], autojunk=False).ratio() < 0.5:
            added.insert(0, j)  # unrelated: report the original paragraph as dropped, look at j again
            j = None
        before = original[i][1] if i is not None else ""
        after = target[j][2] if j is not None else ""
        if _covered(before, approved) and _covered(after, approved) and approved:
            continue
        if i is not None and j is not None:
            was, now = _span(_plain(before), _plain(after))
            differences.append({"kind": "changed", "location": f"{original[i][0]} → {target[j][0]}",
                                "original": was, "output": now})
        elif i is not None:
            differences.append({"kind": "missing", "location": original[i][0], "original": before})
        else:
            differences.append({"kind": "added", "location": target[j][0], "output": after})

    got = objects(docx, _after_cover(Document(docx)))
    for name, label in OBJECT_LABELS:
        values = [w[name] for w in wanted if w and w[name] is not None]
        if not values:
            continue
        if name in ("tables", "links"):
            want, have = sorted(x for v in values for x in v), sorted(got[name])
            same = want == have
            delta = len(set(want) ^ set(have)) or abs(len(want) - len(have))
            shown = (", ".join(want) or "无", ", ".join(have) or "无")
        else:
            want, have = sum(values), got[name]
            same, shown, delta = want == have, (str(want), str(have)), abs(want - have)
        # An approval covers objects only when it says how many ("删去图片 1 张"), so one approved
        # image does not hide every other image that went missing.
        if not same and delta != _approved_count(approved, label):
            differences.append({"kind": "objects", "location": label, "original": shown[0], "output": shown[1]})
    return {
        "sources": reports,
        "differences": differences,
        "notice": "逐段按顺序逐字比对文字（全角/半角标点也算；只忽略换行和中文旁边的空格），并核对图片、表格、公式、链接的数量。",
    }
