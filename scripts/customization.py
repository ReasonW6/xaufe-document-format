"""Task-local user instructions override defaults; never mutate bundled assets."""

from __future__ import annotations
from copy import deepcopy
from pathlib import Path
import hashlib, json, math
from profiles import load_profile, logo_info, ROOT

ROLE_STYLES = {
    "body": "XAUFE Body",
    "h1": "Heading 1",
    "h2": "Heading 2",
    "h3": "Heading 3",
    "h4": "Heading 4",
    "preface": "XAUFE Preface",
    "title_zh": "XAUFE Chinese Title",
    "title_en": "XAUFE English Title",
    "abstract_label_zh": "XAUFE Chinese Abstract Label",
    "abstract_label_en": "XAUFE English Abstract Label",
    "abstract_zh": "XAUFE Abstract Chinese",
    "abstract_en": "XAUFE Abstract English",
    "keywords_zh": "XAUFE Keywords Chinese",
    "keywords_en": "XAUFE Keywords English",
    "toc_title": "XAUFE TOC Title",
    "toc1": "TOC 1",
    "toc2": "TOC 2",
    "toc3": "TOC 3",
    "toc4": "TOC 4",
    "reference_title": "XAUFE References Title",
    "reference": "XAUFE Reference",
    "caption": "XAUFE Caption",
    "figure": "XAUFE Figure",
    "table_text": "XAUFE Table Text",
    "code": "XAUFE Code",
    "equation": "XAUFE Equation",
    "spacer": "XAUFE Spacer",
    "abstract_spacer": "XAUFE Abstract Spacer",
    "header": "XAUFE Header",
    "cover": "原模板封面基础",
}
STYLE_KEYS = {
    "cn",
    "latin",
    "size",
    "bold",
    "align",
    "first",
    "chars",
    "left",
    "before",
    "after",
    "line",
    "before_lines",
    "after_lines",
    "keep",
    "break_before",
    "line_rule",
    "hanging",
    "left_chars",
    "italic",
    "color",
    "underline",
    "right",
}
TOP_KEYS = {
    "logo",
    "logo_width_cm",
    "logo_alignment",
    "page_twips",
    "styles",
    "english_abstract",
    "parts",
    "header",
    "footer_page_numbers",
    "cover_label",
    "cover_title_label",
    "cover_supervisor",
    "cover_fonts",
    "keyword_labels",
    "heading_levels",
}


def fingerprint(value):
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def number(value, lo, hi, name, integer=False):
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
        or not lo <= value <= hi
        or (integer and int(value) != value)
    ):
        raise ValueError(f"{name}须是{lo}至{hi}之间的" + ("整数。" if integer else "数值。"))


def validate_style(patch, *, cover=False):
    if not isinstance(patch, dict) or set(patch) - (
        {"cn", "latin", "size", "bold", "italic", "color", "underline"} if cover else STYLE_KEYS
    ):
        raise ValueError(
            "这个格式要求脚本不支持；用户要求仍有效，按 references/special-cases.md 的“脚本做不到的要求”处理，不能退回默认格式。"
        )
    for key, val in patch.items():
        if key in ("cn", "latin"):
            if not isinstance(val, str) or not val.strip():
                raise ValueError("字体名称必须非空。")
        elif key in ("bold", "keep", "break_before", "italic"):
            if not isinstance(val, bool):
                raise ValueError(key + "须为布尔值。")
        elif key == "color":
            if (
                not isinstance(val, str)
                or __import__("re").fullmatch(r"[0-9A-Fa-f]{6}", val) is None
            ):
                raise ValueError("颜色须为六位十六进制，如0055AA。")
        elif key == "underline":
            if val not in ("none", "single", "double", "dotted", "dash", "wave"):
                raise ValueError("下划线类型无效。")
        elif key == "align":
            if val not in ("left", "center", "right", "both"):
                raise ValueError("对齐方式须为left/center/right/both。")
        elif key == "line_rule":
            if val not in ("auto", "exact", "atLeast"):
                raise ValueError("行距类型无效。")
        elif key == "size":
            number(val, 1, 96, key)
        else:
            number(val, 0, 50000, key, True)
    if "first" in patch and "chars" in patch:
        raise ValueError("首行缩进不能同时指定绝对值和字符值。")
    if "left" in patch and "left_chars" in patch:
        raise ValueError("左缩进不能同时指定绝对值和字符值。")
    if patch.get("size", 1) * 2 % 1:
        raise ValueError("Word字号精度为半磅，不能静默四舍五入。")
    if "hanging" in patch and ("first" in patch or "chars" in patch):
        raise ValueError("首行缩进与悬挂缩进不能同时设置。")


def effective_profile(template, request=None):
    spec = load_profile(template)
    request = deepcopy({} if request is None else request)
    if not isinstance(request, dict) or set(request) - TOP_KEYS:
        raise ValueError(
            "format 里有不认识的字段；只能用 references/job.md 列出的字段。脚本不支持的要求见 references/special-cases.md，不能忽略用户要求。"
        )
    for key in ("english_abstract", "footer_page_numbers", "cover_supervisor"):
        if key in request and not isinstance(request[key], bool):
            raise ValueError(key + "须为布尔值。")
    for key in ("cover_label", "cover_title_label"):
        if key in request and (not isinstance(request[key], str) or not request[key].strip()):
            raise ValueError(key + "须为非空文本。")
    if "logo" in request:
        # Only the image changes; the chosen layout keeps its own logo width.
        base = spec["logo"]
        spec["logo"] = dict(
            logo_info(request["logo"]),
            width_emu=base["width_emu"],
            alignment=base["alignment"],
        )
        if request["logo"] == base["color"]:
            spec["logo"]["height_emu"] = base["height_emu"]
        else:
            source = spec["logo"]
            spec["logo"]["height_emu"] = round(
                base["width_emu"] * source["pixel_height"] / source["pixel_width"]
            )
    if "logo_width_cm" in request:
        number(request["logo_width_cm"], 0.5, 25, "logo_width_cm")
        spec["logo"]["width_emu"] = round(request["logo_width_cm"] * 360000)
        spec["logo"]["height_emu"] = round(
            spec["logo"]["width_emu"] * spec["logo"]["pixel_height"] / spec["logo"]["pixel_width"]
        )
    if "logo_alignment" in request:
        if request["logo_alignment"] not in ("left", "center", "right"):
            raise ValueError("Logo对齐无效。")
        spec["logo"]["alignment"] = request["logo_alignment"]
    pages = request.get("page_twips", {})
    if not isinstance(pages, dict) or set(pages) - set(spec["page_twips"]):
        raise ValueError("未知页面参数。")
    for key, val in pages.items():
        number(val, 0, 50000, key, True)
    spec["page_twips"].update(pages)
    m = spec["page_twips"]
    width = m["width"] - m["left"] - m["right"]
    if width < 1200 or m["height"] - m["top"] - m["bottom"] < 2000:
        raise ValueError("页面剩余空间不足，需重新安排布局，不能静默缩字。")
    if spec["logo"]["width_emu"] > width * 635:
        raise ValueError("Logo超过当前版心，请按用户要求确定等比尺寸。")
    styles = request.get("styles", {})
    if not isinstance(styles, dict) or set(styles) - (set(ROLE_STYLES) - {"cover"}):
        raise ValueError("未知段落角色。")
    for role, patch in styles.items():
        validate_style(patch)
        spec.setdefault("style_overrides", {}).setdefault(ROLE_STYLES[role], {}).update(patch)
    if "english_abstract" in request:
        enabled = request["english_abstract"]
        spec["english_abstract"] = enabled
        spec["parts"] = [p for p in spec["parts"] if p != "abstract_en"]
        if enabled:
            spec["parts"].insert(spec["parts"].index("abstract_zh") + 1, "abstract_en")
    if "parts" in request:
        parts = request["parts"]
        allowed = {"cover", "abstract_zh", "abstract_en", "toc", "body", "references"}
        if (
            not isinstance(parts, list)
            or any(not isinstance(x, str) or x not in allowed for x in parts)
            or len(parts) != len(set(parts))
        ):
            raise ValueError("parts须为不重复的已知部分名称数组。")
        if (
            "body" not in parts
            or "abstract_zh" not in parts
            or "toc" not in parts
            or not parts
            or parts[-1] != "references"
        ):
            raise ValueError(
                "parts 必须包含正文、中文摘要、目录，并以参考文献结尾；其他结构请在 content.json 用 mode=document。"
            )
        if "cover" in parts and parts[0] != "cover":
            raise ValueError("封面只可为第一部分。")
        if ("abstract_en" in parts) != spec["english_abstract"]:
            raise ValueError("parts与english_abstract不一致。")
        if parts.index("body") < max(parts.index(x) for x in ("abstract_zh", "toc")):
            raise ValueError("摘要和目录要排在正文之前；其他顺序请在 content.json 用 mode=document。")
        spec["parts"] = parts
    header = request.get("header", {})
    if not isinstance(header, dict) or set(header) - {"enabled", "left", "right", "border"}:
        raise ValueError("未知页眉参数。")
    for key, val in header.items():
        if key in ("enabled", "border") and not isinstance(val, bool):
            raise ValueError("页眉开关须为布尔值。")
        if key in ("left", "right") and not isinstance(val, str):
            raise ValueError("页眉文字须为字符串。")
    spec["header"].update(header)
    # A header-free template has no usable default header distance. Adding one
    # should not inherit a zero/near-zero distance that places text at the paper edge.
    # An explicit user distance always wins; no default B layout is changed.
    if header.get("enabled") is True and "header" not in pages and m["header"] < 200:
        m["header"] = load_profile("A")["page_twips"]["header"]
    spec["footer"]["page_numbers"] = request.get("footer_page_numbers", False)
    if "cover_label" in request:
        spec["cover"]["document_label"] = request["cover_label"]
    if "cover_title_label" in request:
        spec["cover"]["title_label"] = request["cover_title_label"]
    if "cover_supervisor" in request:
        fields = spec["cover"]["fields"]
        required = spec["cover"]["required_fields"]
        if request["cover_supervisor"]:
            fields["supervisor"] = 13
            if "supervisor" not in required:
                required.insert(required.index("date"), "supervisor")
        else:
            fields.pop("supervisor", None)
            if "supervisor" in required:
                required.remove("supervisor")
    cf = request.get("cover_fonts", {})
    if not isinstance(cf, dict) or set(cf) - {"label", "title", "fields", "date", "course"}:
        raise ValueError("未知封面字体区域。")
    for patch in cf.values():
        validate_style(patch, cover=True)
    levels = request.get("heading_levels", 3)
    if isinstance(levels, bool) or levels not in (3, 4):
        raise ValueError("heading_levels只支持3或4；第四级需用户明确启用。")
    if any(role in styles for role in ("h4", "toc4")) and levels != 4:
        raise ValueError("四级样式需要先明确heading_levels=4。")
    spec["heading_levels"] = levels
    spec["toc"]["indents_twips"] = list(spec["toc"]["indents_twips"])
    if len(spec["toc"]["indents_twips"]) == 3:
        vals = spec["toc"]["indents_twips"]
        vals.append(vals[-1] + (vals[-1] - vals[-2]))
    labels = request.get("keyword_labels", {})
    if not isinstance(labels, dict) or set(labels) - {"zh", "en"}:
        raise ValueError("keyword_labels仅支持zh/en标签字体参数。")
    for patch in labels.values():
        validate_style(patch, cover=True)
    spec["_format_request"] = request
    return spec


def profile_for(state):
    return effective_profile(state["template"], state.get("format_request", {}))


def normalize_style_options(options, patch):
    """An explicit alternate-unit indent replaces, not coexists with, the default."""
    result = dict(options)
    if "first" in patch:
        result.pop("chars", None)
        result.pop("hanging", None)
    if "chars" in patch:
        result.pop("first", None)
        result.pop("hanging", None)
    if "left_chars" in patch:
        result.pop("left", None)
    if "hanging" in patch:
        result.pop("chars", None)
        result.pop("first", None)
    if "left" in patch:
        result.pop("left_chars", None)
    for key in ("before", "after"):
        if key in patch and key + "_lines" not in patch:
            result.pop(key + "_lines", None)
    result.update(patch)
    return result


def apply_cover(doc, spec):
    """Apply only authorized differences; master bytes and default behavior stay intact."""
    request = spec.get("_format_request", {})
    if not request:
        return
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from copy import deepcopy
    from docx import Document
    from engine import font

    ps = doc.paragraphs
    c = spec["cover"]
    if any(k == "logo" or k.startswith("logo_") for k in request):
        p = ps[c["logo_paragraph"]]
        logo = spec["logo"]
        rid, _ = doc.part.get_or_add_image(str(ROOT / logo["file"]))
        old_ids = {b.get(qn("r:embed")) for b in p._p.xpath(".//a:blip")}
        for b in p._p.xpath(".//a:blip"):
            b.set(qn("r:embed"), rid)
        from engine import drop_unused_images

        drop_unused_images(doc, old_ids - {rid})
        for ex in p._p.xpath(".//wp:inline/wp:extent|.//a:xfrm/a:ext"):
            ex.set("cx", str(logo["width_emu"]))
            ex.set("cy", str(logo["height_emu"]))
        pp = p._p.get_or_add_pPr()
        jc = pp.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            pp.append(jc)
        jc.set(qn("w:val"), logo["alignment"])
    if "cover_label" in request:
        ps[c["label_paragraph"]].runs[0].text = c["document_label"]
    if "cover_title_label" in request:
        ps[c["title_paragraph"]].runs[0].text = c["title_label"]
    if "cover_supervisor" in request:
        # Slot 13 is an existing optional line; no indexing changes to date/cover.
        source = "A" if request["cover_supervisor"] else "B"
        source_p = Document(ROOT / load_profile(source)["master_file"]).paragraphs[13]
        old = ps[13]._p
        replacement = deepcopy(source_p._p)
        # Reapply only if the slot has not already been initialized and filled.
        replaced = (
            request["cover_supervisor"]
            and not ps[13].text.startswith("指导教师")
            or not request["cover_supervisor"]
            and ps[13].text.strip()
        )
        if replaced:
            old.getparent().replace(old, replacement)
        ps = doc.paragraphs
        if request["cover_supervisor"] and replaced:
            # The label comes from A, but its columns belong to the chosen base
            # cover. B uses different left indents and tab stops. Copied once,
            # before the cover text is fitted, so later fitting is kept.
            reference = ps[c["fields"]["student_name"]]._p.get_or_add_pPr()
            target = ps[13]._p.get_or_add_pPr()
            for tag in ("ind", "tabs", "jc", "spacing"):
                for node in list(target.findall(qn("w:" + tag))):
                    target.remove(node)
                node = reference.find(qn("w:" + tag))
                if node is not None:
                    target.append(deepcopy(node))
    for area, patch in request.get("cover_fonts", {}).items():
        indices = {
            "label": [c["label_paragraph"]],
            "title": [c["title_paragraph"]],
            "fields": list(c["fields"].values()),
            "date": [c["date_paragraph"]],
            "course": [] if c.get("course_paragraph") is None else [c["course_paragraph"]],
        }[area]
        for pi in indices:
            runs = ps[pi].runs[1:2] if area == "title" else ps[pi].runs
            for run in runs:
                from style_model import run_properties

                old = run_properties(ps[pi], run._r)
                font(
                    run._r.get_or_add_rPr(),
                    patch.get("cn", old["fonts"].get("eastAsia", "宋体")),
                    patch.get("size", int(old.get("sz", "32")) / 2),
                    patch.get("bold", old.get("b", False)),
                    latin=patch.get("latin", old["fonts"].get("ascii", "Times New Roman")),
                )
                from engine import apply_text_options

                apply_text_options(run._r.get_or_add_rPr(), patch)


def merge_request(current, patch):
    """Deep, non-mutating merge. null restores a key; arrays replace; {} is a no-op.

    Clearing sibling indent units is deliberate: a new absolute indent supersedes
    a prior character indent and vice versa. The final effective profile is validated.
    """
    if not isinstance(current, dict) or not isinstance(patch, dict):
        raise ValueError("格式差异必须为JSON对象。")
    if set(patch) - TOP_KEYS:
        raise ValueError("未知格式差异字段。")
    dictionaries = {
        "styles": set(ROLE_STYLES) - {"cover"},
        "header": {"enabled", "left", "right", "border"},
        "page_twips": set(load_profile("A")["page_twips"]),
        "cover_fonts": {"label", "title", "fields", "date", "course"},
        "keyword_labels": {"zh", "en"},
    }
    font_keys = {"cn", "latin", "size", "bold", "italic", "color", "underline"}
    for key, allowed in dictionaries.items():
        section = patch.get(key)
        if section is None:
            continue
        if not isinstance(section, dict) or set(section) - allowed:
            raise ValueError("未知或无效的" + key + "差异。")
        if key in ("styles", "cover_fonts", "keyword_labels"):
            for value in section.values():
                if value is not None and (
                    not isinstance(value, dict)
                    or set(value) - (STYLE_KEYS if key == "styles" else font_keys)
                ):
                    raise ValueError("未知样式参数。")

    def merge(old, delta):
        result = deepcopy(old)
        for key, value in delta.items():
            if value is None:
                result.pop(key, None)
            elif isinstance(value, dict):
                result[key] = merge(
                    result.get(key, {}) if isinstance(result.get(key, {}), dict) else {}, value
                )
                if not result[key]:
                    result.pop(key, None)
            else:
                result[key] = deepcopy(value)
        return result

    value = merge(current, patch)
    for role, delta in (patch.get("styles") or {}).items():
        if not isinstance(delta, dict):
            continue
        target = value.get("styles", {}).get(role, {})
        conflicts = {
            "first": ("chars", "hanging"),
            "chars": ("first", "hanging"),
            "hanging": ("first", "chars"),
            "left": ("left_chars",),
            "left_chars": ("left",),
            "before": ("before_lines",),
            "before_lines": ("before",),
            "after": ("after_lines",),
            "after_lines": ("after",),
        }
        for key, others in conflicts.items():
            if key in delta and delta[key] is not None:
                for other in others:
                    if other not in delta:
                        target.pop(other, None)
    return value


def keyword_label_options(spec, role):
    from engine import role_format, role_latin

    cn, size, bold = role_format(role)
    default = {"cn": cn, "latin": role_latin(role), "size": 14, "bold": bold}
    base = spec.get("style_overrides", {}).get(ROLE_STYLES[role], {})
    default.update({k: base[k] for k in ("italic", "color", "underline") if k in base})
    default.update(spec.get("_format_request", {}).get("keyword_labels", {}).get(role[-2:], {}))
    return default
