"""Turn acceptance failures into issues a model can act on.

Each issue says where (page / paragraph / text / which input to edit), what was
expected and what was found, the likely cause, and a fix that edits the inputs
(job.json, content.json, mapping.json), never the generated Word.
"""

from __future__ import annotations
import json, re
from collections import Counter
from pathlib import Path
from textio import read_json

MAX_STALLED = 3  # failed runs in a row that fixed nothing: stop and report
MAX_ROUNDS = 8  # failed runs in a row in total, even while some issues still get fixed


class AcceptanceFailed(ValueError):
    """The document did not pass an acceptance check; ``issues`` explain each failure."""

    def __init__(self, stage, issues):
        self.stage = stage
        self.issues = issues
        super().__init__(f"{stage}未通过（{len(issues)} 个问题）：\n" + render_text(issues))


_STALE = ("任务记录和成品对不上，通常是改了 job.json 后没有重新 run。", "重新运行 run；不要手改 WORK 里的文件。")
_DIRECT = (
    "这段在原稿里带着直接格式（或者角色选错了），排版后仍和模板不一致。",
    "先看 mapping.json 里这一段的 role 是否正确（标题、正文、摘要、参考文献）；角色没错仍失败，说明是脚本不支持的原稿格式，按 references/special-cases.md 的“脚本做不到的要求”处理并告诉用户。",
)
_USER_STYLE = (
    "实际格式和“模板默认值 + job.json 里的特别要求”不一致。",
    "核对 job.json 的 format.styles 写法（references/job.md）；用户没要求改的项不要写。",
)
HINTS = {
    # static audit
    "PARAGRAPH": _DIRECT,
    "PARAGRAPH_EXTRA": _DIRECT,
    "STYLE_INHERITANCE": _DIRECT,
    "FONT_SIZE": _USER_STYLE,
    "FONT_CN": _USER_STYLE,
    "FONT_EN": _USER_STYLE,
    "FONT_BOLD": _USER_STYLE,
    "FONT_COLOR": _USER_STYLE,
    "FONT_ITALIC": _USER_STYLE,
    "FONT_UNDERLINE": _USER_STYLE,
    "FONT_GEOMETRY": _DIRECT,
    "HIDDEN_TEXT": ("原稿里有隐藏文字。", "问用户这段隐藏文字要不要；要就在 Word 里取消隐藏后重新给，不要就经用户同意把那段 role 设为 delete。"),
    "ROLE_MISMATCH": ("这一段的角色标错了，格式检查会按错的角色放行，排出来却不对。", "改 mapping.json 里这一段的 role（正文 body、标题 h1/h2/h3、图表标题 caption、图片 figure、空行 spacer）；content.json 时把这一块的 type/level 改对。"),
    "UNKNOWN_STYLE": ("有段落没有分到任何角色。", "在 mapping.json 里给这一段一个 role（正文用 body）。"),
    "MISPLACED_COVER_STYLE": ("封面以外的段落用了封面样式。", "把 mapping.json 里这一段改成正确的 role。"),
    "HEADING_DEPTH": ("出现了第四级标题，但没有开启四级。", "用户要四级标题就在 job.json 写 \"heading_levels\": 4；否则把这一段改成三级以内。"),
    "KEYWORD_SPACING": ("关键词之间的分隔不对。", "content.json 用 keywords_zh 列表写关键词；已有 Word 时核对这一段的 role 是 keywords_zh。"),
    "KEYWORD_RUN": ("“关键词：”标签和词项混在一个文字段里。", "重新 run；仍失败就把原稿这一段改成“关键词：词1 词2”的形式。"),
    "PAGE": ("页面尺寸或页边距和本次有效设置不一致。", "核对 job.json 的 page_twips（1 厘米 = 567）；用户没要求就删掉这一项。"),
    "HEADER": ("页眉缺失或内容不对。", "用户自定义了页眉就核对 job.json 的 format.header（enabled、left、right 不能是空文字）；没有自定义时，A 版页眉用专业和题目，确认 info 里有 major、title_zh；B 版默认没有页眉。"),
    "HEADER_TEXT": ("页眉文字和封面信息或 header 设置不一致。", "核对 job.json 的 major、title_zh 和 format.header.left/right。"),
    "HEADER_BORDER": ("页眉横线和设置不一致。", "用户不要横线写 \"header\": {\"border\": false}；否则删掉这项。"),
    "RED_HEADER": ("不该有页眉的地方出现了页眉。", "B 版默认无页眉；用户要页眉才写 format.header.enabled=true。"),
    "COVER_HEADER": ("封面出现了页眉。", "重新 run；仍有就检查原稿是不是自带页眉的旧封面（mapping 的 replace_cover_before）。"),
    "FOOTER": ("页脚或页码和设置不一致。", "要页码写 \"footer_page_numbers\": true，不要就删掉这一项。"),
    "FOOTER_FIELD": ("页脚里残留了页码域。", "重新 run；原稿自带页脚时检查是否带进来了。"),
    "COVER": ("封面结构不完整。", "重新 run；不要手改封面。"),
    "COVER_VALUE": ("封面上的信息和 job.json 不一致。", "核对 job.json 的 info 和 blank；改完重新 run。"),
    "COVER_DATE": ("封面日期和要求不一致。", "日期写成 \"date\": \"2026-10-07\"；要留空就把 date 放进 blank。"),
    "COVER_LAYOUT": ("封面版式和母版不一致。", "重新 run；仍失败说明封面适配有问题，请把题目或课程名缩短后再试，或问用户能否缩小封面字号（format.cover_fonts）。"),
    "COVER_RUNS": ("封面出现了母版以外的文字段。", "重新 run；不要手改封面。"),
    "COVER_FONT": ("封面字体和母版不一致。", "用户要改封面字体就写 format.cover_fonts；否则删掉这一项后重新 run。"),
    "COVER_EFFECTIVE_FONT": ("封面字体被其他样式覆盖。", "同上：核对 format.cover_fonts。"),
    "COVER_TITLE": ("封面题目标签不对。", "核对 format.cover_title_label。"),
    "COVER_LABEL": ("封面文种（如“期末大作业”）不对。", "核对 format.cover_label。"),
    "COVER_BREAK": ("封面没有单独成节。", "重新 run；原稿的旧封面要在 mapping 里用 replace_cover_before 替换。"),
    "COVER_DATE_POSITION": ("封面日期段落缺失。", "重新 run。"),
    "COVER_DATE_LAYOUT": ("封面日期的位置属性不对。", "重新 run；不要手改封面。"),
    "COVER_DATE_BREAK": ("日期段落带了分页。", "重新 run。"),
    "COVER_FRAME": ("封面以外出现了定位框（浮动段落）。", "原稿有浮动段落：请用户改成普通段落，或经同意删掉。"),
    "LOGO": ("校徽图片不是选定的原图。", "job.json 的 logo 只能是 red 或 green；不要替换图片文件。"),
    "LOGO_SIZE": ("校徽尺寸不对。", "用户要改大小就写 format.logo_width_cm；否则删掉这一项。"),
    "LOGO_FLOAT": ("校徽变成了浮动图片。", "重新 run；不要手改封面。"),
    "LOGO_CROP": ("校徽被裁切。", "重新 run；不要手改封面。"),
    "LOGO_ALIGN": ("校徽对齐方式不对。", "用户要居中写 format.logo_alignment；否则删掉这一项。"),
    "TABLE_WIDTH": ("表格比版心宽。", "已有 Word 时确认 mapping 的 format_tables 是 true；仍超宽就请用户减少列数或拆表。"),
    "TABLE_GRID": ("表格列宽信息无效。", "原稿表格结构异常：请用户在 Word 里重新插入这个表格。"),
    "TABLE_FLOAT": ("原稿有浮动表格。", "请用户把表格改成“嵌入文本行中”（取消文字环绕）后再给。"),
    "TABLE_ALIGNMENT": ("表格没有居中。", "确认 mapping 的 format_tables 是 true 后重新 run。"),
    "TABLE_INDENT": ("表格有额外缩进。", "确认 format_tables 是 true 后重新 run。"),
    "TABLE_SPACING": ("表格单元格外间距会扩宽表格。", "确认 format_tables 是 true；仍失败请用户去掉单元格间距。"),
    "TABLE_FIXED_HEIGHT": ("表格行是固定行高，文字可能被裁掉。", "确认 format_tables 是 true；仍失败请用户把行高改成“最小值”。"),
    "TABLE_STRUCTURE": ("表格结构脚本读不懂。", "请用户简化这个表格（拆掉嵌套或复杂合并）。"),
    "CELL_WIDTH": ("单元格宽度和表格网格不一致或太窄。", "确认 format_tables 是 true；列太多就请用户拆表。"),
    "CELL_LAYOUT": ("单元格设置了不换行或压缩文字。", "确认 format_tables 是 true 后重新 run。"),
    "CELL_IMAGE_WIDTH": ("表格里的图片比单元格宽。", "请用户把这张图改小，或从表格里移出来。"),
    "CELL_IMAGE_SPACING": ("表格里的图片留白会撑出边界。", "确认 format_tables 是 true 后重新 run。"),
    "IMAGE_WIDTH": ("图片比版心宽。", "content.json 的图片会自动缩放；已有 Word 时确认这一段 role 是 figure。"),
    "LINE_TOO_TIGHT": ("固定行距比字号还小，打印出来上下两行会叠在一起。", "改 job.json 里这一角色的 styles：把 line 调到不小于字号（固定行距 line = 磅数×20，例如五号 10.5 磅至少 210），或改用 \"line_rule\": \"auto\"。这是会让文字重叠的缺陷，不能交付：用户坚持要这么挤，就告诉用户会重叠、请用户换一种要求（比如缩小字号）。"),
    "PAGE_NUMBERING": ("某一节的页码格式或起始页和规范不一致（完整论文：摘要、目录用罗马数字从 I 开始，正文用阿拉伯数字从 1 开始、一直接到最后；普通文档：封面后从 1 开始、一直接到最后）。", "重新 run，脚本会按规范重设各节页码；仍失败说明成品被改动过或 mapping 的 mode/role 不对：完整论文核对摘要、目录、正文的 role，普通文档不用改 mode。"),
    "OBJECT_HEIGHT": ("大图片放在固定行距的段落里，会被裁掉。", "把这一段的 role 改成 figure。"),
    "PART_MISSING": ("完整论文结构缺了一部分（摘要、目录、正文或参考文献）。", "先看 mapping.json 是不是把那部分标成了别的 role；原稿真的没有就补全（只排版也要补）：摘要、关键词根据正文起草，参考文献只补 refs 核实通过的真实文献，写进 job.json 的 corrections；正文缺了没法补，问用户。"),
    "EMPTY_CONTENT": ("摘要、关键词或参考文献在成品里是空的。", "先排除 role 标错；原稿真的没有就补全：摘要、关键词根据正文起草，参考文献只补 refs 核实通过的真实文献，写进 job.json 的 corrections，交付时告诉用户。用户明确说不要补，就把原话写进 job.json 的 keep_as_is。"),
    "TITLE_TOO_LONG": ("论文题目超过 24 字。", "保留原意精简题目：job.json 的 info.title_zh 和原稿题目段（corrections 里 part 为 title）一起改，交付时告诉用户原题和新题；用户明确要保留原题，就把原话写进 keep_as_is。"),
    "ABSTRACT_LENGTH": ("中文摘要不在 150～300 字之间。", "少了按正文已有内容补写，多了精简（保留目的、方法、结论），写进 job.json 的 corrections（part 为 abstract），交付时告诉用户；用户明确要保留原文，就把原话写进 keep_as_is。"),
    "PART_ORDER": ("各部分顺序和版式不一致：多半是某一段的角色标错了（摘要或目录里的段落标成了正文/标题，或反过来），也可能是 format.parts 写错。", "看“位置”里那一段：已有 Word 就改 mapping.json 里这一段的 role；content.json 就检查 blocks 的顺序和 type。角色都对时，删掉 format.parts 用默认顺序，或按用户要求核对它的写法。"),
    "PART_BREAKS": ("封面、摘要、目录、正文没有各自起页。", "重新 run；mapping 的 mode 应为 paper。"),
    "PLACEHOLDER": ("模板里的说明或示例文字留在了成品里。", "从 content.json / 原稿里删掉这段占位文字（只排版时先问用户）。"),
    "TOC_SCOPE": ("目录条目和标题对不上。", "检查标题的 role/level 是否正确后重新 run。"),
    "TOC_UNRESOLVED": ("目录页码没算出来（没有实际分页）。", "有办公软件就重新 run；没有就是静态交付，如实告诉用户。"),
    "TOC_PROOF_STALE": ("目录页码和实际分页不一致。", "重新 run。"),
    "TOC_FONT": _USER_STYLE,
    "TOC_FONT_SIZE": _USER_STYLE,
    "TOC_EMPHASIS": _USER_STYLE,
    "TOC_COLOR": _USER_STYLE,
    "TOC_GEOMETRY": _DIRECT,
    "TOC_TABS": ("目录点引导线或右侧页码位置不对。", "重新 run；改过 page_twips 时核对写法。"),
    "TOC_STYLE_MISSING": ("目录样式缺失。", "重新 run。"),
    "UNSUPPORTED_STRUCTURE": ("原稿里有脚本不能安全排版的对象（文本框、修订、批注、内容控件等）。", "按 references/special-cases.md 的“输入文件”一节处理：请用户先接受修订、把文本框改成普通文字，再重新给。"),
    "EMBEDDED_FONT": ("原稿嵌入了字体文件。", "请用户在 Word“选项 → 保存”里取消“嵌入字体”后重新给。"),
    "AUTO_FIELD_UPDATE": _STALE,
    "CUSTOMIZATION_HASH": _STALE,
    "CUSTOMIZATION_STALE": _STALE,
    "METADATA_STALE": _STALE,
    "TEMPLATE_MISMATCH": _STALE,
    "JOB_MISMATCH": _STALE,
    "WORKFLOW": _STALE,
    "INTAKE": _STALE,
    "INTAKE_REQUIRED": _STALE,
    "DOCUMENT_SCOPE": ("有没有封面和用户要求的不一致。", "核对 job.json 的 cover。"),
    "RED_ENGLISH_ABSTRACT": ("出现了英文摘要，但本次设置不要英文摘要。", "用户要英文摘要就写 \"english_abstract\": true；不要就去掉英文摘要（已有 Word 时问用户后把这些段落设为 delete）。"),
    "OPEN": ("成品打不开或不是本工具生成的。", "重新 run。"),
    "AUDIT_ABORTED": ("检查器读不懂文档结构。", "重新 run；仍失败就把错误原文告诉用户，不要绕过检查交付。"),
    # rendered-page audit
    "PAGE_SIZE": ("办公软件按别的纸张大小导出了。", "核对 job.json 的 page_twips；仍不对就换一个办公软件（--office word/wps/libreoffice）。"),
    "OFF_PAGE": ("内容超出纸张，打印会被裁掉：通常是超宽表格、长网址或图片；也可能是 job.json 的 page_twips 把纸张设得太小或页边距设成了负数/零。", "先核对 job.json 的 page_twips（单位 twip，1 厘米 = 567，A4 宽 11906、高 16838）；没问题就拆表、给网址加空格断行，或缩小图片，改 content.json / 原稿后重新 run。"),
    "CAPTION_POSITION": ("表题不在表格上方，或图题不在图片下方。", "已有 Word 的图表题脚本会自动挪；仍出现说明图表题和图表之间隔着别的段落，或 caption 标在了别处：看 mapping.json 里这一段和相邻段落的 role。"),
    "TABLE_NOTE_LABEL": ("表注以“注”或“附注”开头，学校不允许。", "已有 Word 的表注（role table_note）脚本会自动去掉；content.json 的表格 note 里也会去掉。仍出现说明这段没标成 table_note：改 mapping.json 的 role。"),
    "TABLE_UNIT_WORD": ("表格单位里写了“单位”二字。", "表格单位段（role table_unit，或 content.json 表格的 unit）脚本会自动去掉“单位：”；仍出现说明这段没标成 table_unit。"),
    "EQUATION_NUMBER": ("公式行末的编号缺了或不按顺序。", "公式段（role equation）脚本会自动按全文顺序编号；仍出现说明这一段没标成 equation，或公式写在了正文段里：改 mapping.json 的 role 或 content.json 的 type。"),
    "EQUATION_LAYOUT": ("公式行的排法不对（制表位、虚线、加粗或字号）。", "脚本会给公式段设好居中和右对齐制表位；加粗或改字号是原稿里公式自带的，告诉用户在原稿里去掉，或在 Word 公式编辑器里改。"),
    "OBJECT_SPACING": ("图或表和上下文之间没有空一行。", "脚本会在图、表组（含图题、表题、单位、表注）前后各加一个空行；仍出现说明这一份是手工处理过的，或图表题的 role 标错了：看 mapping.json。"),
    "TABLE_PADDING": ("表格单元格左右没有留白。", "已有 Word 和 content.json 的表格脚本都会设好留白，仍出现说明这一份是手工处理过的，重新 run。"),
    "TABLE_RULES": ("表格不是三线表（顶线、底线 1.5 磅，表头下 0.25 磅，没有其他线）。", "已有 Word 的表格脚本会自动改成三线表，仍出现说明这一份是手工处理过的，重新 run；content.json 的表格不用写边框。"),
    "TEXT_OUTSIDE_MARGINS": ("文字伸进了页边距：常见于不能断行的长英文、网址、公式或代码，或 page_twips 的页边距设得太小。", "在 content.json 里给长串加断行（代码块用 code 类型），或请用户同意缩短；封面就缩短这一项或问用户能否缩小封面字号；改过 page_twips 就核对它（1 厘米 = 567）。"),
    "IMAGE_OFF_PAGE": ("图片超出纸张。", "把图片改小，或确认这一段 role 是 figure。"),
    "IMAGE_OUTSIDE_MARGINS": ("图片比版心宽，或比一页的版心还高、伸进了上下页边距。", "把图片改小（太高的图请用户裁开或缩小），或确认这一段 role 是 figure（脚本会按版心缩放宽度）。"),
    "TEXT_OVER_IMAGE": ("文字压在图片上：常见于原稿的浮动图片、环绕方式，或图片所在段落是固定行距。", "已有 Word 时把这张图所在段落的 role 设为 figure；仍压字就请用户把图片改成“嵌入型”后再给。"),
    "HEADING_AT_PAGE_BOTTOM": ("标题和它后面的正文被分页拆开了。模板给标题设了“与下段同页”，通常是标题前后有手动加的空行或分页符，或原稿的段落设置盖过了模板。", "去掉这个标题前后多余的空行或分页符（content.json 里的 page_break 块，或 mapping.json 里设成 delete 的空段）后重新 run；仍然这样就告诉用户这一页的情况。"),
    "CAPTION_SEPARATED": ("图题或表题和它的图、表被分页拆开了。通常是图表太大放不进剩下的空间，或题注前后有手动空行。", "去掉题注前后的空行；表格很长时可以拆成两张表（续表），图片太高就请用户同意缩小（content.json 里 image 块写 \"width_cm\": 厘米数）；改完重新 run。"),
    "COVER_TITLE_MISALIGNED": ("封面题目折成几行后对不齐：通常是渲染时没有模板字体（宋体/黑体），替代字体的字宽不同，或题目里有不能断开的长英文。", "先看有没有 FONT_SUBSTITUTED 提示：有就按缺字体处理（换装有字体的 Word/WPS 渲染，或问用户）；题目里有长英文串就请用户同意在合适的位置断开或缩短。"),
    "BODY_IN_MARGIN": ("正文文字落到了上边距或下边距里（页眉页脚区），打印时会和页眉页码挤在一起：常见于原稿的浮动对象或文本框、段落的负间距，或 job.json 改小了 page_twips 的上下边距、改大了行距。", "已有 Word 时核对这一段在 mapping.json 里的 role，有浮动对象就请用户改成嵌入型；改过 page_twips 或 styles 的行距/段距就核对它们，然后重新 run。"),
    "INDENT_RENDERED": ("渲染出来的首行缩进或左缩进和模板要求不一样：通常是原稿这一段带着自己的缩进设置（或用空格、Tab 顶出来的缩进），或这一段的 role 标错了。", "已有 Word 时核对这一段在 mapping.json 里的 role；段首用空格、全角空格或 Tab 打出来的缩进脚本会自动去掉，仍偏就是原稿段落自带缩进或 role 不对；用户要别的缩进就写 job.json 的 format.styles.<role>.chars。"),
    "NOT_CENTERED": ("应该居中的标题或图表标题显示出来没有居中：通常是原稿这一段带着自己的缩进或对齐，或 role 标错了。", "核对这一段在 mapping.json 里的 role；改过 format.styles 的 align/left/chars 就核对它们，然后重新 run。"),
    "COVER_OUTSIDE_MARGINS": ("封面信息栏的下划线伸出了版心：通常是 job.json 的 page_twips 把左右边距改大了，或封面信息太长。", "核对 job.json 的 format.page_twips 左右边距（1 厘米=567）；封面信息太长就问用户能否缩短，然后重新 run。"),
    "SIZE_RENDERED": ("标题、图表题等显示出来的字号和模板（或 job.json 的要求）不同：通常是原稿这一段带着自己的字号，或 role 标错了。", "核对这一段在 mapping.json / content.json 里的级别和 role；用户要别的字号就写 job.json 的 format.styles.<role>.size，然后重新 run。"),
    "WEIGHT_RENDERED": ("标题、英文题目等显示出来的加粗和模板要求不同（只看得出英文字母的粗细）。", "核对这一段的 role；用户要加粗或不加粗就写 job.json 的 format.styles.<role>.bold，然后重新 run。"),
    "LINE_PITCH_RENDERED": ("正文显示出来的行距和要求的固定行距不同：通常是原稿段落自带行距设置盖过了模板，或 job.json 的 line/line_rule 写法不对。", "核对 job.json 的 format.styles.body.line 和 line_rule（固定 20 磅是 line 400 + exact）；已有 Word 时核对 mapping.json 里这些段的 role，然后重新 run。"),
    "TEXT_AT_PAPER_EDGE": ("页眉或页脚距离（page_twips.header / footer）几乎是 0，文字贴着纸边。B 版模板本来没有页眉页脚，加页码或页眉后最容易这样。", "在 job.json 的 format 里设 \"page_twips\": {\"footer\": 992}（页眉太靠上就设 header，单位 twip，992 约 1.75 厘米），然后重新 run。"),
    "CAPTION_DETACHED": ("表题（在表格上方）或图题（在图片下方）和它的表格/图片之间夹着空段落，常见于原稿里多敲的回车或分页符。", "把夹在中间的空段落删掉：mapping.json 里把它的 role 设为 \"delete\"（content.json 里就删掉那个空块），然后重新 run。"),
    "FONT_BOLD_PARAGRAPH": ("整段正文、摘要或参考文献都是加粗：通常是原稿整段手动加粗，或 content.json 给整段写了 bold。", "只排版：问用户这段加粗是不是有意的，不是就在 review.json 说明并请用户确认；写作：去掉整段的 bold，只给需要强调的几个词加粗。"),
    "HEADING_LEVEL": ("标题的级别和编号对不上（用的是模板另一级的编号，或同样的编号在别处是另一级，或标题被当成了正文），通常是 mapping.json 或 content.json 里这一段的级别填错了。", "按编号改这段的级别：mapping.json 的 role（h1/h2/h3）或 content.json 的 level/type，然后重新 run；不要改标题文字。用户确认原稿就按自己的方式编号时，在 job.json 的 format 写 \"heading_numbering\": \"own\" 再 run。"),
    "FONT_DECORATION": ("文字带高亮、底纹或阴影一类的装饰：手工处理的副本或原稿里带来的，脚本排版时会去掉。", "用 run 重新生成就会去掉；这是 adopt 的手工副本时，在 Word 里清除这些格式后重新 adopt。"),
    "FONT_MARKS": ("原稿里有删除线、全部大写、着重号这类可能有含义的标记（全部大写会让显示的字母和原文的大小写不同）。", "不要自己去掉：在 review.json 的 warnings 里写明，并在回复里告诉用户成品保留了这些标记；用户要去掉就请他在原稿里清除后重新给。"),
    "BODY_SIZE_RENDERED": ("Word 文件里写的正文字号没有按预期显示：通常是段落里的直接格式或字符样式盖过了正文样式，或渲染软件对样式继承的理解不同。", "看页图确认正文字号；已有 Word 的原稿里有整段的直接字号就重新 inspect 后 run；仍不对就如实告诉用户这台机器的渲染结果，建议用 Word 打开核对。"),
    "HEADER_WRAPPED": ("页眉文字太长，一行放不下折成了两行。A 版页眉右侧默认放论文题目，长题目最容易这样。", "在 job.json 的 format 里给页眉写一个短的题目：\"header\": {\"right\": \"简短题目\"}（先问用户用哪个简称；左侧太长就改 left）；或经用户同意缩小页眉字号（styles.header.size）。"),
    "HEADER_OVERFLOW": ("页眉文字太长，折成多行后伸出了上边距，把正文往下挤。A 版页眉右侧默认放论文题目，长题目最容易这样。", "在 job.json 的 format 里给页眉写一个短的题目：\"header\": {\"right\": \"简短题目\"}（先问用户用哪个简称；左侧太长就改 left）；或经用户同意缩小页眉字号（styles.header.size）。"),
    "TABLE_OUTSIDE_MARGINS": ("表格或图形伸出了版心。", "已有 Word 时确认 format_tables 是 true；列太多就请用户拆表；浮动图形请用户改成嵌入式。"),
    "BLANK_PAGE": ("出现了多余的空白页：常见于连续分页符、段落自带的“段前分页”，或分节后又手动分页。", "content.json：删掉多余的 page_break 块。已有 Word：在 mapping.json 里把位置里那些只有分页符的空段落 role 设为 delete（空段落没有文字，删掉不算改原文；设为 spacer 不会去掉分页）。"),
    "HEADER_MISSING": ("应有页眉的页面没有页眉。", "核对 job.json 的 format.header 和封面信息（A 版页眉要 major、title_zh）。"),
    "HEADER_UNEXPECTED": ("不该有页眉的页面出现了页眉文字。", "B 版默认无页眉；检查原稿是否带页眉，以及 format.header 写法。"),
    "COVER_HEADER_VISIBLE": ("封面上出现了页眉。", "重新 run；原稿有旧封面时用 mapping 的 replace_cover_before 替换。"),
    "COVER_FOOTER_VISIBLE": ("封面上出现了页脚或页码。", "重新 run；仍出现就是原稿封面自带页脚，请用户删掉后再给。"),
    "PAGE_NUMBER_MISSING": ("用户要页码，但渲染结果里页脚没有页码。", "确认 job.json 写了 \"footer_page_numbers\": true；换办公软件再 run 一次排除渲染差异。"),
    "FOOTER_UNEXPECTED": ("页脚应该空白，渲染结果里有文字。", "用户没要页码时删掉 footer_page_numbers；原稿自带页脚文字时请用户删掉。"),
    "TEXT_OVERLAP": ("两行文字在页面上叠在一起（上下两行的基线距离不到字号的 0.9 倍）：常见于过小的固定行距、文字位置被提升或降低、负的段前段后，或浮动对象。", "打开这一页的页图看是哪两行；检查 job.json 里这类段落的 line/line_rule（固定行距磅数×20 不能小于字号×20）和 size；原稿里文字设了“位置提升/降低”或有浮动对象，就把那一段 role 重新标对后再 run，仍重叠请用户在 Word 里取消这些设置。"),
    "PAGE_NUMBER_SEQUENCE": ("页码跳号、重复，或在不该重新开始的地方重新从 1 开始。脚本会把各节页码设成连续的，所以多半是页脚里另有手写的页码文字，或某一节的页脚不是页码域。", "看这一页的页图和 Word 页脚：原稿页脚里有手打的数字就请用户删掉（只排版时先问）；再重新 run。仍不对就把页码问题和页图如实告诉用户。"),
    "WORD_BLANK_PAGE": ("两处“从新页开始”之间什么都没有：连着两个分页符、分页符后面紧跟分节符或设了“段前分页”的段落"
                        "（模板的一级标题等本来就从新页开始）。Word 会多出一页空白，LibreOffice 有时会把它们合并、页图上看不出来。",
                        "content.json：删掉多余的 page_break 块（标题本来就从新页开始时，它前面不需要 page_break）。"
                        "已有 Word：脚本排版时会自动去掉这种分页符；仍出现时把位置里只有分页符的空段落在 mapping.json 里设为 delete。"),
    "RENDERER_FALLBACK": ("run --office 指定的软件没装、不能自动调用或导出失败，脚本换了另一个能用的软件。", "在 review.json 的 warnings 里写明实际用的是哪个软件，交付时告诉用户；用户坚持要那个软件，就先让它能用再重新 run。"),
    "FONT_SUBSTITUTED": ("渲染机器上没有文档要求的字体。", "不用改文档。在 review.json 的 warnings 里写明页图是用替代字体看的，交付时告诉用户分页以装有这些字体的 Word 为准。"),
}

_LABEL = re.compile(r"^(?P<label>(?:p\d+|t\d+(?:/[rcpt]\d+)*|header\d+/\w+/p\d+)(?:/r\d+)?)\s*[:：]\s*(?P<rest>.*)$")
_SHOWN = re.compile(r"^(?P<label>成品[^:：，。]*?)\s*[:：]\s*(?P<rest>.*)$")


def output_place(label):
    """An output paragraph label ("p39", "t0/r1/c2/p0", "header1/default/p0") as people and agents read it:
    "成品第40段".  Never "p39", which reads like a mapping.json id (those number the original)."""
    label = re.sub(r"/r\d+$", "", label)
    m = re.fullmatch(r"p(\d+)", label)
    if m:
        return f"成品第{int(m.group(1)) + 1}段"
    m = re.fullmatch(r"t(\d+)/r(\d+)/c(\d+)(?:/p\d+)?", label)
    if m:
        t, r, c = (int(x) + 1 for x in m.groups())
        return f"成品第{t}个表格第{r}行第{c}格"
    m = re.fullmatch(r"header(\d+)/\w+/p\d+", label)
    if m:
        return f"成品第{int(m.group(1)) + 1}节的页眉"
    return f"成品 {label}"


def _shown_message(label, message):
    """The message with output paragraph labels in words (the leading label and any "p12" inside)."""
    rest = _LABEL.match(message).group("rest")
    rest = re.sub(r"(?<![\w/])p(\d+)(?![\w/])", lambda m: output_place(m.group(0)), rest)
    return f"{output_place(label)}：{rest}"
_EXPECT = re.compile(r"(?P<what>[\w/]+)=(?P<actual>[^，,]+?)(?:半磅)?，应为(?P<expected>[^。]+)。?")


def _paragraph(doc, label):
    from table_layout import walk_paragraphs

    base = re.sub(r"/r\d+$", "", label)
    if re.fullmatch(r"p\d+", base):
        index = int(base[1:])
        return doc.paragraphs[index] if index < len(doc.paragraphs) else None
    if base.startswith("t"):
        top = int(re.match(r"t(\d+)", base).group(1))
        if top < len(doc.tables):
            for found, p in walk_paragraphs(doc.tables[top], f"t{top}"):
                if found == base:
                    return p
    m = re.fullmatch(r"header(\d+)/(\w+)/p(\d+)", base)
    if m and int(m.group(1)) < len(doc.sections):
        section = doc.sections[int(m.group(1))]
        part = {"default": section.header, "first": section.first_page_header, "even": section.even_page_header}.get(m.group(2))
        if part is not None and int(m.group(3)) < len(part.paragraphs):
            return part.paragraphs[int(m.group(3))]
    return None


def _flat(text):
    return re.sub(r"\s+", "", str(text or ""))


def _sources(root):
    """(pointer, text) for every piece of input the user's text came from.  Text the template
    generates (cover, TOC, header, labels) is not here, so it is never pinned on an input."""
    content, mapping = Path(root) / "content.json", Path(root) / "mapping.json"
    found = []
    if content.is_file():
        data = read_json(content)
        for i, block in enumerate(data.get("blocks", [])):
            kind = block.get("type")
            body = block.get("text") or "".join(r.get("text", "") for r in block.get("runs") or [])
            found.append((f"content.json 的 blocks[{i}]（{kind}）", body))
            if block.get("caption"):
                found.append((f"content.json 的 blocks[{i}]（{kind} 的 caption）", block["caption"]))
            for row in [block.get("header") or []] + (block.get("rows") or []):
                found += [(f"content.json 的 blocks[{i}]（表格单元格）", str(c)) for c in row]
        for field in ("abstract_zh", "abstract_en"):
            if data.get(field):
                found.append((f"content.json 的 {field}", data[field]))
        for field in ("keywords_zh", "keywords_en"):
            if data.get(field):
                found.append((f"content.json 的 {field}", "".join(map(str, data[field]))))
        found += [(f"content.json 的 references[{i}]", str(r)) for i, r in enumerate(data.get("references") or [])]
    elif mapping.is_file():
        rows = read_json(mapping).get("paragraphs", [])
        full = Path(root) / "paragraph-text.json"
        texts = read_json(full) if full.is_file() else {}
        found += [(f"mapping.json 里 id 为 {r.get('id')} 的段落（当前 role：{r.get('role')}）", texts.get(r.get("id"), r.get("text", "")))
                  for r in rows]
    return [(where, _flat(text)) for where, text in found if _flat(text)]


def _input_hint(root, text):
    """Which input holds this text: the one content.json block / mapping.json paragraph that contains it.
    Too short or found in several places gives no pointer rather than a wrong one."""
    key = _flat(text)
    if len(key) < 4 or root is None:
        return None
    try:
        hits = list(dict.fromkeys(where for where, body in _sources(root) if key in body))
    except (OSError, ValueError):
        return None
    if len(hits) == 1:
        return hits[0]
    if 1 < len(hits) <= 3:
        return "以下之一（这段文字出现了不止一次）：" + "；".join(hits)
    return None


def explain(errors, docx=None, root=None, stage="格式检查"):
    """Static-audit errors -> actionable issues."""
    doc = None
    if docx:
        try:
            from docx import Document

            doc = Document(docx)
        except Exception:
            doc = None
    issues = []
    for e in errors:
        code, message = e.get("code", "?"), e.get("message", "")
        cause, fix = HINTS.get(code, ("", "看 message 的说明修改 job.json / content.json / mapping.json 后重新 run。"))
        item = {"code": code, "stage": stage, "message": message, "cause": cause, "fix": fix}
        m = _LABEL.match(message)
        if m:
            item["message"] = _shown_message(m.group("label"), message)
            item["where"] = {"location": output_place(m.group("label"))}
            p = _paragraph(doc, m.group("label")) if doc is not None else None
            if p is not None:
                text = "".join(p._p.xpath(".//w:t/text()")).strip()
                item["where"]["text"] = text[:40] or "（空段落）"
                item["where"]["style"] = p.style.name if p.style is not None else ""
                source = _input_hint(root, text)
                styled = _job_style(root, p.style.name if p.style is not None else "") if code in STYLE_CODES else None
                if styled:  # the paragraph's look comes from job.json, not from its words
                    item["where"]["edit"] = styled + (f"（问题出在这一段：{source}）" if source else "")
                elif source:
                    item["where"]["edit"] = source
                gap = _mapping_gap(root, text) if code == "CAPTION_DETACHED" and root else []
                if gap:
                    item["where"]["edit"] = f"mapping.json 里 id 为 {'、'.join(gap)} 的空段落（role 改为 \"delete\"）"
        if code == "WORD_BLANK_PAGE" and root:
            previous = re.search(r"前面最后的文字是“(.+?)”", message)
            blocks = _content_breaks(root, previous.group(1)) if previous else []
            gap = _mapping_gap(root, previous.group(1)) if previous else []
            if blocks:
                item.setdefault("where", {})["edit"] = "content.json 的 " + "、".join(blocks) + "（page_break）"
            elif gap:
                item.setdefault("where", {})["edit"] = f"mapping.json 里 id 为 {'、'.join(gap)} 的空段落（role 改为 \"delete\"）"
        x = _EXPECT.search(message)
        if x:
            item["expected"] = x.group("expected").strip()
            item["actual"] = x.group("actual").strip()
        issues.append(item)
    return with_job_suspects(root, issues)


# Static codes about how a paragraph looks: when job.json restyles its role, that is the input to edit.
STYLE_CODES = {"LINE_TOO_TIGHT", "FONT_SIZE", "FONT_CN", "FONT_EN", "FONT_BOLD", "FONT_COLOR", "FONT_ITALIC",
               "FONT_UNDERLINE", "PARAGRAPH", "PARAGRAPH_EXTRA", "OBJECT_HEIGHT"}


def _job_style(root, style_name):
    """'job.json 的 format.styles.body（line=200、size=16）' when job.json restyles this paragraph's role."""
    from customization import ROLE_STYLES

    role = next((r for r, name in ROLE_STYLES.items() if name == style_name), None)
    try:
        patch = (read_json(Path(root) / "job.json").get("format") or {}).get("styles", {}).get(role) if root and role else None
    except (OSError, ValueError, AttributeError):
        return None
    if not isinstance(patch, dict) or not patch:
        return None
    return f"job.json 的 format.styles.{role}（" + "、".join(f"{k}={v}" for k, v in patch.items()) + "）"


def _mapping_gap(root, previous):
    """mapping.json ids of the empty paragraphs right after the text ``previous`` (existing Word)."""
    path = Path(root) / "mapping.json" if root else None
    if not path or not path.is_file():
        return []
    try:
        rows = read_json(path).get("paragraphs", [])
        full = path.parent / "paragraph-text.json"
        texts = read_json(full) if full.is_file() else {}
    except (OSError, ValueError):
        return []
    key = re.sub(r"\s+", "", previous)[-12:]
    start = next((i for i, r in enumerate(rows)
                  if key and key in re.sub(r"\s+", "", texts.get(r.get("id"), r.get("text", "")))), None)
    found = []
    for r in rows[start + 1 : start + 8] if start is not None else []:
        if r.get("text", "").strip():
            break
        found.append(str(r.get("id")))
    return found


def _content_breaks(root, previous):
    """content.json page_break blocks right after the block holding the text ``previous`` (new document)."""
    path = Path(root) / "content.json" if root else None
    if not path or not path.is_file():
        return []
    try:
        blocks = read_json(path).get("blocks", [])
    except (OSError, ValueError):
        return []
    key = re.sub(r"\s+", "", previous)[-12:]
    text = lambda b: re.sub(r"\s+", "", b.get("text") or "".join(r.get("text", "") for r in b.get("runs") or []))
    start = next((i for i, b in enumerate(blocks) if key and key in text(b)), None)
    found = []
    for i in range(start + 1, min(start + 8, len(blocks))) if start is not None else []:
        if blocks[i].get("type") != "page_break":
            break
        found.append(f"blocks[{i}]")
    return found


def _page_breaks(docx, previous):
    """Output paragraphs that only carry a page break (or start a page) right after the text ``previous``."""
    from docx import Document

    try:
        paragraphs = Document(docx).paragraphs
    except Exception:
        return []
    key = re.sub(r"\s+", "", previous)[-12:]
    start = next((i for i, p in enumerate(paragraphs) if key and key in re.sub(r"\s+", "", p.text)), None)
    if start is None:
        return []
    found = []
    for i in range(start + 1, min(start + 8, len(paragraphs))):
        p = paragraphs[i]
        if p.text.strip():
            break
        if p._p.xpath('.//w:br[@w:type="page"]|./w:pPr/w:pageBreakBefore|./w:pPr/w:sectPr'):
            found.append(f"p{i}")
    return found


# Failures that a changed page or text geometry in job.json can cause.
GEOMETRY = {"OFF_PAGE", "TEXT_OUTSIDE_MARGINS", "TEXT_AT_PAPER_EDGE", "IMAGE_OFF_PAGE", "IMAGE_OUTSIDE_MARGINS",
            "TABLE_OUTSIDE_MARGINS", "TEXT_OVERLAP", "TEXT_OVER_IMAGE", "LINE_TOO_TIGHT", "HEADER_OVERFLOW",
            "HEADER_WRAPPED", "BLANK_PAGE", "HEADING_AT_PAGE_BOTTOM", "CAPTION_SEPARATED", "PAGE_NUMBER_MISSING",
            "TABLE_WIDTH", "IMAGE_WIDTH", "CELL_WIDTH", "OBJECT_HEIGHT", "BODY_IN_MARGIN", "INDENT_RENDERED",
            "NOT_CENTERED", "LINE_PITCH_RENDERED", "SIZE_RENDERED", "WEIGHT_RENDERED",
            "COVER_OUTSIDE_MARGINS", "COVER_FONTS_MISSING", "COVER_TOO_LONG"}


def _job_changes(root):
    """What job.json changed from the template layout: page_twips and style values, as "key=value"."""
    try:
        fmt = read_json(Path(root) / "job.json").get("format") or {}
    except (OSError, ValueError, AttributeError):
        return []
    found = []

    def walk(prefix, value):
        if isinstance(value, dict):
            for key, inner in value.items():
                walk(f"{prefix}.{key}", inner)
        elif not isinstance(value, (list, str)) or prefix.endswith(("line_rule", "align")):
            found.append(f"{prefix}={value}")

    for key in ("page_twips", "styles", "cover_fonts"):
        walk(f"format.{key}", fmt.get(key) or {})
    if "logo_width_cm" in fmt:
        found.append(f"format.logo_width_cm={fmt['logo_width_cm']}")
    return found


def with_job_suspects(root, issues):
    """Name the job.json changes first when a layout failure could come from them: the template's
    own values never cause these, so a changed margin, size or spacing is the likeliest cause."""
    changes = _job_changes(root) if root else []
    if not changes:
        return issues
    listed = "、".join(changes[:8]) + ("等" if len(changes) > 8 else "")
    for item in issues:
        if item.get("code") in GEOMETRY or (item.get("code") == "INPUT_REJECTED"
                                           and re.search(r"封面文字太长|放不下|校徽|Logo|版心", item.get("message", ""))):
            item["cause"] = (f"job.json 改过模板的版面设置：{listed}。模板默认值不会出这个问题，先查这些改动"
                             "（数值是否离谱、单位是否写错：page_twips 单位是 twip，1 厘米=567）。其次才是：" + (item.get("cause") or ""))
            item["fix"] = ("先把不合理的值改回合理范围，或删掉这一项恢复模板默认（用户明确要求的改动要先问用户），再重新 run；"
                           "改动都合理时再按下面处理：" + (item.get("fix") or ""))
    return issues


def _indent_settings(root):
    """"body（left、first）" for each role whose indent the job changed, joined; "" when none."""
    try:
        styles = read_json(Path(root) / "job.json").get("format", {}).get("styles", {})
    except (OSError, ValueError, AttributeError):
        return ""
    keys = ("left", "left_chars", "right", "first", "chars", "hanging")
    found = [f"{role}（{'、'.join(k for k in v if k in keys)}）"
             for role, v in styles.items() if isinstance(v, dict) and set(v) & set(keys)]
    return "、".join(found)


def from_render(audit, root=None, docx=None):
    """Rendered-page audit errors -> actionable issues (they already carry page/expected/actual)."""
    issues = []
    for e in audit.get("errors", []):
        cause, fix = HINTS.get(e["code"], ("", ""))
        item = {"code": e["code"], "stage": "渲染后版面检查", "message": e["message"], "cause": cause, "fix": fix,
                "where": {"page": e.get("page")}}
        quoted = re.search(r"“(.+)”", e["where"]) if e.get("where") else None
        if e.get("where"):
            item["where"]["text"] = e["where"]
            source = _input_hint(root, quoted.group(1) if quoted else e["where"])
            if source:
                item["where"]["edit"] = source
        if e.get("bbox_pt"):
            item["where"]["bbox_pt"] = e["bbox_pt"]
        if e["code"] == "TEXT_OUTSIDE_MARGINS" and root is not None:
            indents = _indent_settings(root)
            if indents:
                item["cause"] += f" 这次 job.json 的 format.styles 改过缩进：{indents}；缩进加上文字超出了版心，先核对这几项。"
        if e["code"] == "BLANK_PAGE" and docx and quoted:
            breaks = _page_breaks(docx, quoted.group(1))
            rows = _mapping_gap(root, quoted.group(1))
            blocks = _content_breaks(root, quoted.group(1))
            if blocks:
                item["where"]["edit"] = "content.json 的 " + "、".join(blocks) + "（page_break）"
            if breaks:
                shown = "、".join(output_place(b) for b in breaks)
                item["where"]["location"] = shown
                item["cause"] = f"{shown}是只有分页符（或分节符）的空段落，连在一起就多出一页空白。"
            if rows:
                item["where"]["edit"] = f"mapping.json 里 id 为 {'、'.join(rows)} 的空段落"
        for key in ("expected", "actual"):
            if key in e:
                item[key] = e[key]
        issues.append(item)
    return with_job_suspects(root, issues)


# Input errors raised while laying out (ValueError text) -> cause and fix.
_MESSAGE_HINTS = (
    (r"模板这一级不加粗", "标题、图表标题、目录或页眉被加粗了：原稿里手动加粗，或 content.json 的 runs 里写了 bold。",
     "已有 Word：这是格式问题，脚本按模板统一后会去掉，若仍出现说明是手工处理的副本，在 Word 里取消加粗后 adopt；"
     "content.json：去掉这段 runs 里的 bold。用户明确要加粗就在 job.json 的 format.styles 里写 bold: true。"),
    (r"的题目“.*”和 job\.json 的 info\.title", "原稿里的题目和 job.json 封面信息里的题目不是同一个。",
     "问用户哪个对：job.json 写错了就改 info.title_zh / title_en；原稿要改就请用户确认后改原稿。不要自己挑一个。"),
    (r"放不下一个五号汉字|列太多|太窄", "表格列太多或某几列太窄，按版心缩放后一列连一个汉字都放不下。",
     "content.json 里给这个表格写 widths（各列宽度比例，长文字列给大、编号列给小），或减少列数、缩短表头、把表拆成两个；"
     "已有 Word 里的表格：请用户减少列数或拆表。用户要求横向页面时脚本做不到，按 references/special-cases.md 的“脚本做不到的要求”处理。"),
    (r"封面文字太长", "封面上的题目、课程名或某一项太长，自动换行、收紧空白和行距后第一页仍放不下。",
     "问用户：缩短题目、课程名或过长的封面项（改 job.json 的 info），或同意缩小封面字号（format.cover_fonts）。"),
    (r"校徽", "校徽比版心宽。", "按报错里的数值把 job.json 的 format.logo_width_cm 改小，或按用户要求调整 page_twips 的左右边距。"),
    (r"隐藏文字", "原稿里有隐藏文字，脚本不会擅自显示或删除。", "问用户这段隐藏文字要不要：要就在 Word 里取消隐藏后重新给，不要就经用户同意把那段 role 设为 delete。"),
    (r"关键词", "关键词段的写法脚本读不懂，或这一段的 role 标成了关键词但其实不是。",
     "核对 mapping.json 里这一段的 role（只有以“关键词：”/“Key words:”开头的段落才是 keywords_zh/keywords_en）；"
     "content.json 用 keywords_zh/keywords_en 列表写关键词，每项一个词。"),
    (r"摘要和目录要排在正文之前|parts", "job.json 的 format.parts 顺序或内容不合规。",
     "删掉 format.parts 用默认顺序；用户要的顺序脚本不支持时，在 content.json 用 mode=document 自己安排并告诉用户。"),
    (r"先后顺序和所选版式不同", "这一段的 role 可能标错了（比如正文段标成了摘要或关键词），或者原稿本来就是这个顺序。",
     "role 标错就改 mapping.json 里这一段的 role；原稿顺序本来如此，就把 mapping.json 的 mode 改成 document，或先问用户。"),
    (r"英文摘要", "原稿有没有英文摘要和本次版式设置不一致。", "版式要英文摘要、原稿没有：把中文题目、摘要、关键词译成英文补上（已有 Word 写进 mapping.json 的 add，文本写进 content.json 的 abstract_en/keywords_en/metadata.title_en），交付时告诉用户是补的；用户明确不要，才在 job.json 的 format 写 \"english_abstract\": false。原稿有、版式不要：保留就写 \"english_abstract\": true，用户同意删除才把那些段落设为 delete。"),
    (r"role 是 preserve", "preserve 只用于开头的旧封面，正文里的段落用了它就保持原稿样式，模板检查不认。",
     "把这一段的 role 改成它实际的角色（正文 body、标题 h1/h2/h3 等），旧目录用 old_toc，然后重新 run。"),
    (r"自动编号.*认不出", "这一段的自动编号用了脚本不认识的编号格式，没法转成同样的文字。",
     "请用户在 Word 里把这几段的自动编号转成普通文字后重新给原稿，或经用户同意把这一段的 role 改成 body（编号照原样保留）。"),
    (r"未标记角色|未映射|mapping\.json 里没有 p\d+", "mapping.json 里有段落没有可用的 role。", "按 references/content.md 的“已有 Word”一节给这一段写 role。"),
    (r"内容控件|文本框|修订|批注|嵌入", "原稿里有脚本不能安全排版的对象。", "按 references/special-cases.md 的“输入文件”一节处理，请用户先改成普通文字/嵌入式图片或接受修订。"),
)


def from_message(message, root=None, stage="排版"):
    """A layout ValueError -> located issues (one per line), counted and fixed like any other failure."""
    if "\n" in message.strip():
        return [item for line in message.strip().splitlines() if line.strip()
                for item in from_message(line.strip(), root, stage)]
    cause, fix = "排版脚本拒绝了这份输入。", "按上面的说明改 job.json / content.json / mapping.json 后重新 run。"
    for pattern, c, f in _MESSAGE_HINTS:
        if re.search(pattern, message):
            cause, fix = c, f
            break
    item = {"code": "INPUT_REJECTED", "stage": stage, "message": message, "cause": cause, "fix": fix}
    where = {}
    block = re.search(r"blocks\[(\d+)\]", message)
    label = re.search(r"\b(p\d+|t\d+)\b", message)
    quoted = re.search(r"[“《](.{2,20}?)[”》]", message)
    if block:
        where["edit"] = f"content.json 的 blocks[{block.group(1)}]"
    elif label:
        where["location"] = label.group(1)
        hint = _input_hint(root, quoted.group(1)) if quoted else None
        where["edit"] = hint or f"mapping.json 里 id 为 {label.group(1)} 的段落"
    elif quoted and root is not None:
        hint = _input_hint(root, quoted.group(1))
        if hint:
            where["edit"] = hint
    if quoted:
        where["text"] = quoted.group(1)
    if where:
        item["where"] = where
    return with_job_suspects(root, [item])


def warning_key(w):
    """Key for review.json "warnings": the code, plus the page when the warning has one."""
    return f"{w['code']}（第{w['page']}页）" if w.get("page") else w["code"]


def warning_text(w):
    """One line for the agent: code, page/where, message."""
    place = "，".join(x for x in (f"第{w['page']}页" if w.get("page") else "", str(w.get("where") or "")) if x)
    return w["code"] + (f"（{place}）" if place else "") + f"：{w['message']}"


def _same_kind(item):
    """Issues that differ only in page or paragraph get one entry: same code, message, expected, actual."""
    message = item.get("message", "")
    m = _LABEL.match(message) or _SHOWN.match(message)
    where = item.get("where") or {}
    return (item["code"], m.group("rest") if m else message, item.get("expected"), item.get("actual"),
            where.get("text") if where.get("page") else None)


def group(issues):
    """Merge repeats so a 30-page header problem reads as one problem on 30 pages."""
    merged, order = {}, []
    for item in issues:
        key = _same_kind(item)
        if key not in merged:
            merged[key] = dict(item, where=dict(item.get("where") or {}), pages=[], locations=[])
            order.append(key)
        where = item.get("where") or {}
        if where.get("page") and where["page"] not in merged[key]["pages"]:
            merged[key]["pages"].append(where["page"])
        if where.get("location"):
            merged[key]["locations"].append(where["location"])
    return [merged[k] for k in order]


def _span(values, label, unit=""):
    if len(values) <= 1:
        return ""
    shown = "、".join(f"{label}{v}{unit}" for v in values[:6])
    return shown + (f" 等共 {len(values)} 处" if len(values) > 6 else f"（共 {len(values)} 处）")


def render_text(issues, limit=12):
    """Readable summary for the agent; the full list is in WORK/issues.json."""
    lines, fixes = [], {}
    issues = group(issues)
    for n, item in enumerate(issues[:limit], 1):
        where = item.get("where") or {}
        place = "，".join(
            x for x in (
                _span(item.get("pages", []), "第", "页") or (f"第{where['page']}页" if where.get("page") else ""),
                _span(item.get("locations", []), "") or where.get("location", ""),
                (where["text"] if "“" in where["text"] else f"“{where['text']}”") if where.get("text") else "",
            ) if x
        )
        lines.append(f"{n}. [{item['code']}] {item['message']}")
        if place:
            lines.append(f"   位置：{place}")
        if where.get("edit"):
            lines.append(f"   要改的输入：{where['edit']}")
        shown = [f"{name}：{item[key]}" for key, name in (("actual", "实际"), ("expected", "应为")) if key in item]
        if shown:
            lines.append("   " + "；".join(shown))
        same = fixes.get((item.get("cause"), item.get("fix")))
        if same:
            lines.append(f"   可能原因和怎么改：同第 {same} 条。")
            continue
        fixes[(item.get("cause"), item.get("fix"))] = n
        if item.get("cause"):
            lines.append(f"   可能原因：{item['cause']}")
        if item.get("fix"):
            lines.append(f"   怎么改：{item['fix']}")
    if len(issues) > limit:
        lines.append(f"……另有 {len(issues) - limit} 个问题，见 issues.json。")
    return "\n".join(lines)


_PLACE = re.compile(r"成品第\d+(?:段|个表格第\d+行第\d+格|节的页眉)|第\s*\d+\s*页|(?<![\w/])[pt]\d+(?:/[rcpt]\d+)*")


def _fingerprint(item):
    """One issue as it stands: the code, the input to edit, the text and the measured value.  Where
    it sits (page, output paragraph number) is left out: an unrelated edit that only moves the problem
    to another page or paragraph number is not progress."""
    where = item.get("where") or {}
    return "|".join((item["code"], _PLACE.sub("#", str(item.get("message"))), _PLACE.sub("#", str(where.get("text"))),
                     str(where.get("edit")), str(item.get("actual"))))


def _attempts(root):
    """(counters, session state).  The counters are kept in attempts.json and in the session
    marker, so deleting attempts.json does not reset the stop rule."""
    try:
        data = json.loads((Path(root) / "attempts.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        data = {}
    try:
        state = json.loads((Path(root) / ".xaufe-session.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        state = None
    kept = (state or {}).get("attempts") or {}
    if int(kept.get("failed_rounds", 0)) > int(data.get("failed_rounds", 0)) \
            or int(kept.get("stalled", 0)) > int(data.get("stalled", 0)):
        data = dict(kept)
    return data, state


def _save_attempts(root, data, state):
    (Path(root) / "attempts.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    if state is not None:
        state["attempts"] = data
        (Path(root) / ".xaufe-session.json").write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def stopped(root):
    """Why automatic retries must stop now (None while they may go on)."""
    data, _ = _attempts(root)
    if int(data.get("stalled", 0)) >= MAX_STALLED:
        return f"连续 {data['stalled']} 次修改都没有解决任何一个问题"
    if int(data.get("failed_rounds", 0)) >= MAX_ROUNDS:
        return f"已经连续 {data['failed_rounds']} 次没通过"
    return None


def reset_rounds(root, user_said):
    """The user asked for another try: count from zero, keep their words with the attempts."""
    data, state = _attempts(root)
    data.setdefault("user_retries", []).append(str(user_said).strip())
    data.update(failed_rounds=0, stalled=0)
    _save_attempts(root, data, state)


def record_round(root, issues):
    """Count failed runs in a row and how many of them fixed nothing.

    Returns (rounds, stalled, repeated).  A run makes progress when at least one
    issue of the previous run is gone; ``stalled`` counts runs in a row without
    that, ``repeated`` means exactly the same issues as last time.
    """
    data, state = _attempts(root)
    now = sorted(_fingerprint(i) for i in issues)  # a list: two blank pages, one fixed, is progress
    before = data.get("last_issues")
    rounds = int(data.get("failed_rounds", 0)) + 1 if issues else 0
    progress = not before or bool(Counter(before) - Counter(now))  # after a pass, a new failure is not a stall
    stalled = 0 if not issues or progress else int(data.get("stalled", 0)) + 1
    repeated = bool(issues) and now == before
    _save_attempts(root, {"failed_rounds": rounds, "stalled": stalled, "last_issues": now,
                          "user_retries": data.get("user_retries", [])}, state)
    return rounds, stalled, repeated
