#!/usr/bin/env python3
"""Profile-aware static audit. No implicit template; no visual-pass claims."""

from __future__ import annotations
import argparse, hashlib, json, re, sys
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
import engine as E
from intake import require as require_intake, digest
from audit_helpers import (
    contract,
    page_contract,
    paragraph_check,
    run_check,
    tables_check,
    toc_entries,
    toc_proof_check,
)
from table_layout import walk_paragraphs
from style_model import paragraph_properties, run_properties
from cover_date_audit import static as check_cover_date
from customization import apply_cover, fingerprint


def attrs(e):
    return {} if e is None else dict(e.attrib)


def shape(e):
    if e is None:
        return None
    return (e.tag, tuple(sorted(e.attrib.items())), tuple(shape(x) for x in e))


def inherited(element, style, tag):
    pp = style.element.find(qn("w:pPr"))
    result = attrs(pp.find(qn("w:" + tag))) if pp is not None else {}
    pp = element.find(qn("w:pPr"))
    if pp is not None:
        result.update(attrs(pp.find(qn("w:" + tag))))
    return result


def _cover_values(doc, spec):
    """Read back the cover strings, ignoring wrapping breaks, tabs and placeholders."""

    def clean(text):
        text = text.replace("\n", "").replace("\t", "").strip()
        return "" if not text.strip("　") else text

    cover = spec["cover"]
    ps = doc.paragraphs
    values = {key: clean(ps[i].runs[1].text) for key, i in cover["fields"].items()}
    values["title_zh"] = clean(ps[cover["title_paragraph"]].runs[1].text)
    if cover.get("course_paragraph") is not None:
        values["course_name"] = clean(ps[cover["course_paragraph"]].text).strip("《》").strip()
        values["course_name"] = clean(values["course_name"])
    return values


def text_of(p):
    return "".join(p._p.xpath(".//w:t/text()"))


def visible_part(part):
    return bool(
        part._element.xpath(".//w:t|.//w:instrText|.//w:fldSimple|.//w:drawing|.//w:pict|.//w:pBdr")
    )


def line_check(p, label, err):
    """A fixed line height below the font size prints lines on top of each other."""
    try:
        spacing = paragraph_properties(p, "spacing")
        if spacing.get("lineRule") != "exact" or not str(spacing.get("line", "")).isdigit():
            return
        sizes = [
            int(run_properties(p, r).get("sz") or 21)
            for r in p._p.xpath(".//w:r")
            if "".join(r.xpath("./w:t/text()")).strip()
        ]
    except ValueError:  # broken style chains are reported by paragraph_check
        return
    line = int(spacing["line"]) / 20
    if sizes and line < max(sizes) / 2:
        err(
            "LINE_TOO_TIGHT",
            f"{label}: 固定行距={line:g}磅，应为不小于字号{max(sizes) / 2:g}磅。",
        )


CAPTION = re.compile(r"^\s*(?:图|表|续表|Figure|Fig\.?|Table)\s*[\dA-Za-z一二三四五六七八九十]", re.I)
ROLE_NAMES = {"spacer": "空行", "abstract_spacer": "空行", "figure": "图片段", "caption": "图表标题",
              "h1": "一级标题", "h2": "二级标题", "h3": "三级标题", "h4": "四级标题", "preface": "序言标题", "body": "正文"}


HEADING_LIKE = re.compile(r"^(?:[一二三四五六七八九十]+、|[（(][一二三四五六七八九十]+[）)]|第[一二三四五六七八九十\d]+[章节])\s*\S")


def _blank(el):
    """An empty paragraph (at most a page or line break): it carries nothing to keep a caption with."""
    return el.tag == qn("w:p") and not el.xpath(".//w:t[normalize-space()]|.//w:drawing|.//w:pict|.//m:oMath")


def caption_check(p, label, err):
    """A caption cut off from its table (below it) or figure (above it) by empty paragraphs:
    keep-with-next then ties it to the blank, and the page can break right between them."""
    for step, wanted, name in ((lambda e: e.getnext(), lambda e: e.tag == qn("w:tbl"), "表格"),
                               (lambda e: e.getprevious(), lambda e: bool(e.xpath(".//w:drawing|.//w:pict")), "图片")):
        el, skipped = step(p._p), 0
        while el is not None and _blank(el):
            el, skipped = step(el), skipped + 1
        if skipped and el is not None and wanted(el):
            err("CAPTION_DETACHED", f"{label}: 图表标题和它的{name}之间隔着 {skipped} 个空段落，分页时会被拆开。")
            return


HEADING_STYLE_NUMBER = {
    "h1": re.compile(r"^[一二三四五六七八九十百零〇]+、"),
    "h2": re.compile(r"^[（(][一二三四五六七八九十百零〇]+[）)]"),
    "h3": re.compile(r"^\d+[.．](?!\d)"),
}


def heading_levels_check(paragraphs, err, own=False):
    """Heading levels that contradict the numbering.  A heading numbered the template's way for
    another level ("（一）" set as h1) is blocked, unless the user approved numbering the document
    its own way (format "heading_numbering": "own").  Even then it is blocked when the document
    itself shows the intended level: the same numbering sits at another level elsewhere, the level
    is already used by a different numbering, or a short unpunctuated body line carries a numbering
    that the document uses for headings."""
    name = {"h1": "一级", "h2": "二级", "h3": "三级"}
    style_of = lambda text: next((r for r, rx in HEADING_STYLE_NUMBER.items() if rx.match(text)), None)
    headings = [(label, role, text, style_of(text)) for label, role, text in paragraphs if role != "body"]
    home = {style for _, role, _, style in headings if style == role}
    for label, role, text, style in headings:
        if not style or style == role:
            continue
        if style in home:
            err("HEADING_LEVEL", f"{label}: 标题“{text[:20]}”按编号是{name[style]}标题（用同样编号的其他标题是 {style}），这里却设成了 {role}。")
        elif role in home:
            err("HEADING_LEVEL", f"{label}: 标题“{text[:20]}”的编号是{name[style]}的写法，但设成了 {role}，"
                                 f"而文档里 {role} 用的是另一种编号。")
        elif not own:
            err("HEADING_LEVEL", f"{label}: 标题“{text[:20]}”用的是模板{name[style]}标题的编号，却设成了 {role}"
                                 f"（模板：一级“一、”，二级“（一）”，三级“1.”）。")
    # A list ("1.准确性" "2.完整性" …) runs on numbered lines; a lost heading is followed by plain text.
    for before, (label, role, text), after in zip([(None, None, "")] + paragraphs, paragraphs, paragraphs[1:] + [(None, None, "")]):
        style = style_of(text) if role == "body" else None
        alone = after[1] == "body" and not style_of(after[2]) and not (before[1] == "body" and style_of(before[2]))
        if style in home and alone and len(text) <= 30 and not re.search(r"[。；;：:，,！!？?]$", text):
            err("HEADING_LEVEL", f"{label}: “{text[:20]}”像{name[style]}标题（编号和文档里的 {style} 标题一样，短且没有句末标点），却设成了正文。")


def role_check(p, role, label, visible, err):
    """Catch a paragraph whose role (from mapping.json or content.json) does not fit its content:
    the formatting would then be "correct" for the wrong role and pass unnoticed."""
    text = visible.strip()
    problem = None
    if role in ("spacer", "abstract_spacer") and text:
        problem = "空行里有文字"
    elif role == "figure" and text and not p._p.xpath(".//w:drawing|.//w:pict"):
        problem = "图片段里没有图片，只有文字"
    elif role == "caption" and text and not CAPTION.match(text):
        problem = "图表标题应以“图1”“表2”“Figure 3”等开头"
    elif role in ("h1", "h2", "h3", "h4", "preface") and (len(text) > 80 or re.search(r"[。；;]$", text)):
        problem = "标题太长或以句号/分号结尾，像是一段正文"
    elif role == "body" and len(text) <= 30 and HEADING_LIKE.match(text) and not re.search(r"[。；;，,：:]$", text):
        problem = "它是“一、”“（一）”“第一章”这样开头的短行，像是标题（标成正文会从目录里漏掉）"
    if problem:
        err("ROLE_MISMATCH", f"{label}: 这一段标成了{ROLE_NAMES[role]}（{role}），但{problem}：“{text[:20]}”。")


FRONT = {"title_zh", "title_en", "abstract_label_zh", "abstract_label_en", "abstract_zh", "abstract_en",
         "keywords_zh", "keywords_en", "toc_title", "toc1", "toc2", "toc3", "toc4"}
MAIN = {"preface", "h1", "h2", "h3", "h4", "body", "reference_title", "reference"}


def numbering_check(doc, reverse, err):
    """Front matter I, II, ... from I; body 1, 2, ... from 1 to the end (format-spec.md)."""
    from docx.text.paragraph import Paragraph

    roles, index, front, main = set(), 0, False, False
    for el in doc.element.body:
        if el.tag == qn("w:p"):
            roles.add(reverse.get(Paragraph(el, doc._body).style.name))
            sect = el.find(qn("w:pPr") + "/" + qn("w:sectPr"))
        else:
            sect = el if el.tag == qn("w:sectPr") else None
        if sect is None:
            continue
        index += 1
        page = sect.find(qn("w:pgNumType"))
        fmt = (page.get(qn("w:fmt")) if page is not None else None) or "decimal"
        start = page.get(qn("w:start")) if page is not None else None
        if main or roles & MAIN:
            want = ("decimal", None if main else "1")
            what = "正文" if not main else "正文之后"
            main = True
        elif roles & FRONT:
            want = ("upperRoman", None if front else "1")
            what = "摘要/目录"
            front = True
        else:
            roles = set()
            continue
        if (fmt, start) != want:
            show = lambda f, s: ("罗马数字" if f == "upperRoman" else "阿拉伯数字" if f == "decimal" else f) + (
                f"从{s}开始" if s else "接着上一节"
            )
            err("PAGE_NUMBERING", f"第{index}节（{what}）页码格式={show(fmt, start)}，应为{show(*want)}。")
        roles = set()


def plain_numbering_check(doc, has_cover, err):
    """Without paper structure: arabic from 1 after the cover, then continuing (no restarts)."""
    for i, section in enumerate(doc.sections):
        if has_cover and i == 0:
            continue
        page = section._sectPr.find(qn("w:pgNumType"))
        fmt = (page.get(qn("w:fmt")) if page is not None else None) or "decimal"
        start = page.get(qn("w:start")) if page is not None else None
        want = ("decimal", "1" if i == (1 if has_cover else 0) else None)
        if (fmt, start) != want:
            err("PAGE_NUMBERING", f"第{i + 1}节页码格式={fmt}{'从' + start + '开始' if start else '接着上一节'}，"
                f"应为阿拉伯数字{'从1开始' if want[1] else '接着上一节'}（{'封面后的第一节' if want[1] else '正文中途不重新编号'}）。")


def _check(path, final=False, intake_path=None, toc_report=None, validation_scope="full"):
    if validation_scope not in ("full", "static-only"):
        raise ValueError("未知验收范围。")
    errors = []
    warnings = []

    def err(code, msg):
        errors.append({"code": code, "message": msg})

    def warn(code, msg):
        warnings.append({"code": code, "message": msg})

    try:
        doc = Document(path)
        vv = E.get_vars(doc)
        request = json.loads(vv.get("XAUFE_FORMAT_REQUEST", "{}"))
        E.select_profile(vv.get("XAUFE_TEMPLATE"), request)
        E.asset_check()
    except Exception as exc:
        return {
            "passed": False,
            "errors": [{"code": "OPEN", "message": str(exc)}],
            "warnings": [],
            "stats": {},
        }
    spec = E.SPEC
    expected = contract(spec["id"], request)
    if doc.settings.element.xpath('./w:updateFields[@w:val="true" or @w:val="1"]'):
        err(
            "AUTO_FIELD_UPDATE",
            "交付稿不得请求打开时自动更新域；关闭文档级更新请求，保留手动更新能力，不通过修改 Word 全局提示设置规避。",
        )
    if "XAUFE_FORMAT_REQUEST" in vv and vv.get("XAUFE_FORMAT_REQUEST_HASH") != fingerprint(request):
        err("CUSTOMIZATION_HASH", "任务格式差异摘要无效；请重新排版，不手改成品声明。")
    from unsupported import unsupported_reasons

    for reason in unsupported_reasons(doc):
        err("UNSUPPORTED_STRUCTURE", reason)
    has_cover = vv.get("XAUFE_HAS_COVER") == "1"
    mode = vv.get("XAUFE_MODE", "unknown")
    count = spec["cover"]["paragraph_count"]
    state = None
    if mode != "template":
        if vv.get("XAUFE_WORKFLOW") != "ready" or not vv.get("XAUFE_INTAKE_ID"):
            err("WORKFLOW", "缺少本次模板选择及封面信息已解决的执行记录。")
        if intake_path:
            try:
                state = require_intake(intake_path)
                if state.get("format_request", {}) != request:
                    err("CUSTOMIZATION_STALE", "成品没有采用本次用户的特别要求，或任务差异已改变。")
                if state["template"] != spec["id"]:
                    err("TEMPLATE_MISMATCH", "输出模板与用户选择不同。")
                if "include_cover" in state and state["include_cover"] != has_cover:
                    err("DOCUMENT_SCOPE", "有无封面与用户明确的文档范围不一致。")
                if state["job_id"] != vv.get("XAUFE_INTAKE_ID"):
                    err("JOB_MISMATCH", "输出关联了其他任务的交互记录。")
                if digest(state["metadata"]) != vv.get("XAUFE_COVER_METADATA_HASH") or sorted(
                    json.loads(vv.get("XAUFE_COVER_BLANK_FIELDS", "[]"))
                ) != sorted(state.get("blank_fields", [])):
                    err("METADATA_STALE", "封面信息、留空设置或默认日期已变化，请重新生成和检查。")
            except Exception as exc:
                err("INTAKE", str(exc))
        elif final:
            err("INTAKE_REQUIRED", "最终检查需 --intake 以核对真实选择、封面信息和留空设置。")
        else:
            warn("INTAKE_NOT_RECHECKED", "本次静态检查未提供交互记录；最终检查必须提供。")
    if has_cover and mode != "template" and len(doc.sections) < 2:
        err("COVER_BREAK", "封面必须独立分节，与后续部分分开。")
    for role, patch in (request.get("styles") or {}).items():
        # Allowed but self-defeating: the user may not have meant it, so they are told before delivery.
        color = str(patch.get("color") or "")
        if len(color) == 6:
            r, g, b = (int(color[i:i + 2], 16) for i in (0, 2, 4))
            if 0.2126 * r + 0.7152 * g + 0.0722 * b > 230:
                warn("STYLE_HARD_TO_READ", f"{ROLE_NAMES.get(role, role)}（{role}）的文字颜色是 {color}，在白纸上几乎看不见。"
                     "确认是用户要的就在回复里提醒用户；不是就改 job.json 的 format.styles。")
        if isinstance(patch.get("size"), (int, float)) and patch["size"] < 5:
            warn("STYLE_HARD_TO_READ", f"{ROLE_NAMES.get(role, role)}（{role}）的字号只有 {patch['size']:g} 磅，打印出来几乎看不清。"
                 "确认是用户要的就在回复里提醒用户；不是就改 job.json 的 format.styles。")
    page_rule = page_contract(spec["id"], request)  # independent of assets/*/layout.json
    for i, s in enumerate(doc.sections):
        for attr, key in [
            ("page_width", "width"),
            ("page_height", "height"),
            ("top_margin", "top"),
            ("bottom_margin", "bottom"),
            ("left_margin", "left"),
            ("right_margin", "right"),
            ("header_distance", "header"),
            ("footer_distance", "footer"),
            ("gutter", "gutter"),
        ]:
            actual = getattr(s, attr)
            if actual is None or actual.twips != page_rule[key]:
                err("PAGE", f"第{i+1}节 {attr} 与本次有效格式（默认值加用户覆盖）不一致。")
            elif i == 0 and key in ("top", "bottom", "left", "right") and actual.twips < 284:
                warn("MARGIN_TOO_SMALL", f"页边距 {key} 只有 {actual.twips / 567:.2f} 厘米：多数打印机印不到离纸边 0.5 厘米以内的地方，"
                     "边上的文字可能被裁掉。是用户要求的就在回复里提醒用户；不是就删掉 job.json 的 page_twips 这一项。")
        for variant, f in [
            ("default", s.footer),
            ("first", s.first_page_footer),
            ("even", s.even_page_footer),
        ]:
            if spec["footer"].get("page_numbers") and not (i == 0 and has_cover):
                fields = f._element.xpath(".//w:fldSimple/@w:instr")
                if len(fields) != 1 or fields[0].strip() != "PAGE":
                    err("FOOTER", "用户要求的页码域缺失或混入其他字段。")
            elif visible_part(f):
                err("FOOTER", f"第{i+1}节 {variant} 页脚不是空白。")
        for variant, h in [
            ("default", s.header),
            ("first", s.first_page_header),
            ("even", s.even_page_header),
        ]:
            if not spec["header"]["enabled"] or (i == 0 and has_cover):
                if visible_part(h):
                    err(
                        "RED_HEADER" if not spec["header"]["enabled"] else "COVER_HEADER",
                        f"第{i+1}节 {variant} 不应有页眉、横线或字段。",
                    )
            elif not "".join(h._element.xpath(".//w:t/text()")).strip():
                err("HEADER", f"第{i+1}节 {variant} 页眉缺失。")
            elif state:
                want = (
                    spec["header"].get(
                        "left", (state["metadata"].get("major") or "＿＿") + "专业本科学年论文"
                    )
                    + "\t"
                    + spec["header"].get(
                        "right",
                        state["metadata"].get("title_zh") or doc.core_properties.title or "",
                    )
                )
                actual = [p.text.replace("\n", "") for p in h.paragraphs]
                if actual != [want] and actual != want.split("\t", 1):
                    err(
                        "HEADER_TEXT",
                        "页眉文字与本次有效设置或已核实的专业／题目不一致；以header覆盖后的左右文字为准，不按校徽颜色判断。",
                    )
    if spec["header"]["enabled"]:
        for si, section in enumerate(doc.sections):
            if si == 0 and has_cover:
                continue
            for variant, h in [
                ("default", section.header),
                ("first", section.first_page_header),
                ("even", section.even_page_header),
            ]:
                bottoms = h._element.xpath(".//w:pBdr/w:bottom")
                if not spec["header"].get("border", True):
                    if bottoms:
                        err("HEADER_BORDER", "用户要求无页眉横线。")
                    continue
                if len(bottoms) != 1 or any(
                    bottoms[0].get(qn("w:" + key)) != value
                    for key, value in {
                        "val": "single",
                        "sz": "6",
                        "space": "1",
                        "color": "000000",
                    }.items()
                ):
                    err("HEADER_BORDER", f"第{si+1}节 {variant} 页眉横线缺失或改变。")
    if has_cover:
        check_cover_date(doc, spec, err)
        if len(doc.paragraphs) < count:
            err("COVER", "封面字段段落不完整。")
        else:
            p = doc.paragraphs[0]
            blips = p._p.xpath(".//a:blip")
            if len(blips) != 1:
                err("LOGO", "封面首段须包含本次选定的原校徽；校徽来源与基础版式分别核对。")
            else:
                rid = blips[0].get(qn("r:embed"))
                if not rid or rid not in doc.part.related_parts:
                    err("LOGO", "校徽不是有效嵌入式图片。")
                elif (
                    hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest()
                    != spec["logo"]["sha256"]
                ):
                    err(
                        "LOGO",
                        "校徽图像与本次选定来源不一致，或原图字节已改变。A/B均可配红／绿校徽；按logo核对，不按基础版式推断颜色。",
                    )
            ext = p._p.xpath(".//wp:inline/wp:extent")
            if (
                len(ext) != 1
                or ext[0].get("cx") != str(spec["logo"]["width_emu"])
                or ext[0].get("cy") != str(spec["logo"]["height_emu"])
            ):
                err("LOGO_SIZE", "校徽显示尺寸或嵌入方式改变。")
            if p._p.xpath(".//wp:anchor"):
                err("LOGO_FLOAT", "校徽不得使用浮动锚点。")
            if any(
                any(v not in ("0", "") for v in x.attrib.values())
                for x in p._p.xpath(".//a:srcRect")
            ):
                err("LOGO_CROP", "校徽不得裁切。")
            if p._p.xpath("./w:pPr/w:jc/@w:val") != [spec["logo"]["alignment"]]:
                names = {"left": "左对齐", "start": "左对齐", "center": "居中", "right": "右对齐", "end": "右对齐"}
                want = spec["logo"]["alignment"]
                err("LOGO_ALIGN", f"校徽应为{names.get(want, want)}（模板默认或用户要求），"
                    f"现在是{names.get((p._p.xpath('./w:pPr/w:jc/@w:val') or ['left'])[0], '其他')}。")
            if not doc.paragraphs[spec["cover"]["title_paragraph"]].text.startswith(
                spec["cover"]["title_label"]
            ):
                err("COVER_TITLE", "封面题目标签与所选模板不符。")
            if (
                doc.paragraphs[spec["cover"]["label_paragraph"]].text
                != spec["cover"]["document_label"]
            ):
                err("COVER_LABEL", "封面文种与所选模板不符。")
            master = Document(E.ROOT / spec["master_file"])
            apply_cover(master, spec)
            import cover_fit

            cover_values = _cover_values(doc, spec)
            try:
                fit_level = int(vv.get("XAUFE_COVER_FIT", "0"))
            except ValueError:
                fit_level = -1
            if not 0 <= fit_level <= cover_fit.MAX_LEVEL:
                err("COVER_LAYOUT", "封面适配级别记录无效；请重新生成文档。")
                fit_level = 0
            # Expected cover = master + user overrides + the same long-text fitting.
            cover_fit.apply(master, spec, cover_fit.plan(master, spec, cover_values, level=fit_level))
            for pi in range(count):
                actual = doc.paragraphs[pi]
                baseline = master.paragraphs[pi]
                for tag in ["jc", "ind", "tabs", "spacing"]:
                    a = actual._p.pPr.find(qn("w:" + tag)) if actual._p.pPr is not None else None
                    b = (
                        baseline._p.pPr.find(qn("w:" + tag))
                        if baseline._p.pPr is not None
                        else None
                    )
                    if shape(a) != shape(b):
                        err("COVER_LAYOUT", f"封面p{pi}的{tag}与母版不符。")
                for ri, r in enumerate(actual.runs):
                    if not r.text.strip():
                        continue
                    if ri >= len(baseline.runs):
                        err("COVER_RUNS", f"封面p{pi}出现母版外运行；须确认布局。")
                        continue
                    rr = baseline.runs[ri]
                    effective = run_properties(actual, r._r)
                    original = run_properties(baseline, rr._r)
                    for key in (
                        "fonts",
                        "sz",
                        "b",
                        "i",
                        "iCs",
                        "u",
                        "color",
                        "vanish",
                        "webHidden",
                        "specVanish",
                        "spacing",
                        "w",
                        "position",
                    ):
                        default = (
                            False
                            if key in ("b", "i", "iCs", "vanish", "webHidden", "specVanish")
                            else None
                        )
                        if effective.get(key, default) != original.get(key, default):
                            err("COVER_EFFECTIVE_FONT", f"封面p{pi}/r{ri}的有效{key}被样式覆盖。")
                    for tag in ["rFonts", "sz", "b"]:
                        want = (
                            attrs(rr._r.rPr.find(qn("w:" + tag))) if rr._r.rPr is not None else {}
                        )
                        got = attrs(r._r.rPr.find(qn("w:" + tag))) if r._r.rPr is not None else {}
                        if got != want:
                            err("COVER_FONT", f"封面p{pi}/r{ri}的{tag}与母版不符。")
            if state:
                cm = E.cover_metadata(state["metadata"], state)
                for key, pi in spec["cover"]["fields"].items():
                    got = cover_values.get(key, "")
                    if got != cm.get(key, ""):
                        err("COVER_VALUE", f"{key} 未按已核实信息或明确留空设置填写。")
                title = (
                    doc.paragraphs[spec["cover"]["title_paragraph"]]
                    .runs[1]
                    .text.strip()
                    .replace("\n", "")
                )
                if title != cm.get("title_zh", ""):
                    err("COVER_VALUE", "封面题目未按已核实信息／留空设置处理。")
                cp = spec["cover"].get("course_paragraph")
                if cp is not None:
                    val = cover_values.get("course_name", "")
                    if val != cm.get("course_name", "").strip().strip("《》"):
                        err("COVER_VALUE", "课程名称与本次信息不符。")
                date = cm.get("date", "")
                dp = doc.paragraphs[spec["cover"]["date_paragraph"]].text.replace("完成日期：", "")
                if date:
                    y, m, d = map(int, date.split("-"))
                    want = f"{y}年{m}月{d}日"
                    if re.sub(r"\s+", "", dp) != want:
                        err("COVER_DATE", "日期不是指定日期或本次运行日。")
                elif re.search(r"\d", dp):
                    err("COVER_DATE", "用户要求留空，却保留了数字日期。")
    reverse = {v: k for k, v in E.ROLES.items()}
    allp = [(f"p{i}", p) for i, p in enumerate(doc.paragraphs)]
    for ti, t in enumerate(doc.tables):
        allp.extend(walk_paragraphs(t, f"t{ti}"))
    if spec["header"]["enabled"]:
        for si, sec in enumerate(doc.sections):
            if si == 0 and has_cover:
                continue
            for variant, h in [
                ("default", sec.header),
                ("first", sec.first_page_header),
                ("even", sec.even_page_header),
            ]:
                allp.extend((f"header{si}/{variant}/p{pi}", p) for pi, p in enumerate(h.paragraphs))
    checked, headings = 0, []
    for label, p in allp:
        if has_cover and label.startswith("p") and int(label[1:]) < count:
            continue
        role = reverse.get(p.style.name)
        visible = text_of(p)
        if not role:
            if visible.strip():
                err("UNKNOWN_STYLE", f"{label}: 未映射的样式 {p.style.name}。")
            continue
        if role == "cover":
            err("MISPLACED_COVER_STYLE", f"{label} 在封面范围外使用封面样式。")
            continue
        if not spec["english_abstract"] and role in (
            "title_en",
            "abstract_label_en",
            "abstract_en",
            "keywords_en",
        ):
            err(
                "RED_ENGLISH_ABSTRACT",
                "本次有效设置english_abstract=false，但文档含英文摘要角色；检查用户要求是否已记录，不按A/B或校徽颜色禁用。",
            )
        if role in ("h4", "toc4") and spec.get("heading_levels", 3) < 4:
            err("HEADING_DEPTH", "第四级结构未被本次用户要求启用。")
        checked += 1
        rule = expected[role]
        if role == "header" and label.startswith("header") and len(p._parent.paragraphs) == 2:
            from copy import deepcopy

            rule = deepcopy(rule)
            rule["pPr"]["jc"]["val"] = "right" if label.endswith("/p1") else "left"
        paragraph_check(p, role, label, rule, err, warn)
        run_check(p, role, label, expected[role], err, warn)
        line_check(p, label, err)
        role_check(p, role, label, visible, err)
        if role == "caption" and label.startswith("p"):
            caption_check(p, label, err)
        if (role in HEADING_STYLE_NUMBER or role == "body") and label.startswith("p"):
            headings.append((label, role, visible.strip()))
        if p._p.xpath(".//wp:inline") and role != "figure":
            maxh = max((int(x.get("cy", "0")) for x in p._p.xpath(".//wp:extent")), default=0)
            if maxh > 254000 and paragraph_properties(p, "spacing").get("lineRule") == "exact":
                err("OBJECT_HEIGHT", "大图片不能放在固定20磅的普通段落里。")
    heading_levels_check(headings, err, own=request.get("heading_numbering") == "own")
    with ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith("word/fonts/"):
                err("EMBEDDED_FONT", "不得嵌入并分发字体文件。")
            if (
                not spec["footer"].get("page_numbers")
                and name.startswith("word/footer")
                and name.endswith(".xml")
                and re.search(r"\b(?:PAGE|NUMPAGES|SECTIONPAGES)\b", z.read(name).decode("utf-8"))
            ):
                err("FOOTER_FIELD", f"{name} 残留页码域。")
    for p in doc.paragraphs:
        if re.fullmatch(r"\s*(?:空一行|本科毕业论文（设计）模板)\s*", p.text) or "※※※" in p.text:
            err("PLACEHOLDER", "模板说明或示例占位泄漏到成品。")
    unresolved = []
    for fld in doc.element.body.xpath(".//w:fldSimple"):
        if "PAGEREF" in fld.get(qn("w:instr"), ""):
            text = "".join(x.text or "" for x in fld.findall(".//" + qn("w:t")))
            if not text.isdigit():
                unresolved.append(text)
    if unresolved:
        (err if final and validation_scope == "full" else warn)(
            "TOC_UNRESOLVED", "目录页码还没有计算（需要实际分页）。"
        )
    entries = []
    try:
        entries = toc_entries(doc)
    except ValueError as exc:
        err("TOC_SCOPE", str(exc))
    if mode == "paper" and doc.core_properties.title:
        title_key = re.sub(r"\s+", "", doc.core_properties.title)
        if any(re.sub(r"\s+", "", entry["text"]) == title_key for entry in entries):
            warn(
                "DOCUMENT_TITLE_IN_TOC",
                "全篇题目“"
                + doc.core_properties.title
                + "”同时被列为正文章节。默认全篇题目放在元数据供封面/摘要使用，不重复进入正文和目录；请检查原始材料的文档标题是否误映射进blocks。仅在用户明确要求时保留该特殊范围。",
            )
    if validation_scope == "static-only":
        from render_availability import pending_check

        pending_check(doc, err)
        warn(
            "RENDER_NOT_VERIFIED",
            "仅做静态格式检查；未完成实际分页、目录页位与逐页视觉检查，必须向用户告知。",
        )
    else:
        if final and vv.get("XAUFE_VALIDATION_SCOPE") == "static-only":
            err("RENDER_REQUIRED", "此稿尚未实际分页，不能当作完整检查；请重新运行 run。")
        if entries:
            toc_proof_check(path, entries, vv, final, err, warn, toc_report)
    from toc_format import audit as audit_toc

    audit_toc(doc, spec, err)
    table_count = tables_check(doc, E.WIDTH, err)
    if mode == "document":
        if any(reverse.get(p.style.name) in ("toc_title", "abstract_label_zh", "abstract_label_en") for p in doc.paragraphs):
            numbering_check(doc, reverse, err)  # an earlier output of this skill, revised
        else:
            plain_numbering_check(doc, has_cover, err)
    if mode == "paper":
        front = []
        for part in spec["parts"]:
            front += {
                "toc": ["toc_title"],
                "abstract_zh": ["title_zh", "abstract_label_zh"],
                "abstract_en": ["title_en", "abstract_label_en"],
            }.get(part, [])
        roles = [reverse.get(p.style.name) for p in doc.paragraphs]
        names = {"toc_title": "“目录”标题", "title_zh": "中文摘要页的题目", "abstract_label_zh": "“摘要”标题",
                 "title_en": "英文摘要页的题目", "abstract_label_en": "“Abstract”标题"}
        quote = lambda i: "“" + doc.paragraphs[i].text.strip()[:20] + "”"
        pos, found = [], []
        for role in front:
            if role not in roles:
                err("PART_MISSING", f"缺少{names.get(role, role)}（{role}）。")
            else:
                pos.append(roles.index(role))
                found.append(role)
        for i in range(1, len(pos)):
            if pos[i] < pos[i - 1]:
                err("PART_ORDER", f"p{pos[i]}: {names[found[i]]}{quote(pos[i])}排在了{names[found[i - 1]]}"
                    f"{quote(pos[i - 1])}（p{pos[i - 1]}）前面，和所选版式的前置部分顺序不一致。")
                break
        # An abstract tagged as body sits between its label and the keywords: name it, so the
        # misleading "body before the front matter" / "keywords inside the body" errors stay away.
        misfiled = set()
        for label_role, text_role, keywords_role in (("abstract_label_zh", "abstract_zh", "keywords_zh"),
                                                     ("abstract_label_en", "abstract_en", "keywords_en")):
            start = roles.index(label_role) if label_role in roles else None
            end = next((i for i in range(start + 1, len(roles)) if roles[i] == keywords_role), None) if start is not None else None
            for i in range(start + 1, end) if end is not None else ():
                if roles[i] in ("preface", "h1", "h2", "h3", "h4", "body") and doc.paragraphs[i].text.strip():
                    misfiled.add(i)
                    err("ROLE_MISMATCH", f"p{i}: 这一段{quote(i)}在{names[label_role]}和关键词之间，应该是摘要正文"
                        f"（{text_role}），现在标成了 {roles[i]}。")
        body = next(
            (i for i, x in enumerate(roles) if x in ("preface", "h1", "h2", "h3", "h4", "body") and i not in misfiled),
            None,
        )
        if body is None:
            err("PART_MISSING", "缺少正文。")
        elif pos and body <= max(pos):
            last = max(pos)
            err("PART_ORDER", f"p{body}: 正文段落{quote(body)}（{roles[body]}）排在了{names[roles[last]]}"
                f"{quote(last)}（p{last}）前面：前置部分里有段落被标成了正文或标题。")
        if "reference_title" not in roles:
            err("PART_MISSING", "缺少参考文献标题。")
        elif body is not None and roles.index("reference_title") <= body:
            at = roles.index("reference_title")
            err("PART_ORDER", f"p{at}: 参考文献标题{quote(at)}排在了第一段正文{quote(body)}（p{body}）前面。")
        ref_at = roles.index("reference_title") if "reference_title" in roles else len(roles)
        for i, x in enumerate(roles):
            if x == "reference" and i < ref_at and doc.paragraphs[i].text.strip():
                err("ROLE_MISMATCH", f"p{i}: 这一段{quote(i)}标成了参考文献条目（reference），但它在参考文献标题前面：正文段落应为 body。")
                break
        late = {"abstract_zh", "abstract_en", "keywords_zh", "keywords_en", "abstract_label_zh", "abstract_label_en"}
        for i in range(body + 1 if body is not None else len(roles), len(roles)):
            if roles[i] in late and doc.paragraphs[i].text.strip():
                err("ROLE_MISMATCH", f"p{i}: 正文开始以后的这一段{quote(i)}标成了{roles[i]}（摘要/关键词）：正文里的段落应为 body 或标题。")
                break
        if "reference_title" in roles:
            for i in range(roles.index("reference_title") + 1, len(roles)):
                if roles[i] in ("h1", "h2", "h3", "h4", "preface"):
                    break
                if roles[i] == "body" and doc.paragraphs[i].text.strip():
                    err("ROLE_MISMATCH", f"p{i}: 参考文献标题后面的这一段标成了正文（body）{quote(i)}：参考文献条目应为 reference。")
                    break
        numbering_check(doc, reverse, err)
        minimum = (5 if spec["english_abstract"] else 4) - (0 if has_cover else 1)
        if len(doc.sections) < minimum:
            err("PART_BREAKS", "封面／摘要／目录／正文没有独立分节起页。")
    if mode == "paper":
        missing_content = []
        for role in ("abstract_zh", "keywords_zh", "reference") + (
            ("abstract_en", "keywords_en") if spec["english_abstract"] else ()
        ):
            paragraphs = [p for p in doc.paragraphs if reverse.get(p.style.name) == role]
            nonempty = any(
                re.sub(r"^(关键词\s*[:：]|Key\s*words\s*[:：])", "", text_of(p), flags=re.I).strip()
                for p in paragraphs
            )
            if not nonempty:
                missing_content.append(role)
        if missing_content:
            after = {"abstract_zh": "abstract_label_zh", "abstract_en": "abstract_label_en", "reference": "reference_title"}
            prefix = {"keywords_zh": r"^\s*关键词", "keywords_en": r"^\s*key\s*words?"}
            paragraphs = list(enumerate(doc.paragraphs))
            candidates = []
            for role in missing_content:
                if role in after:
                    starts = [i for i, q in paragraphs if reverse.get(q.style.name) == after[role]]
                    found = [(i, q) for i, q in paragraphs[starts[0] + 1 : starts[0] + 4] if q.text.strip()] if starts else []
                else:
                    found = [(i, q) for i, q in paragraphs if re.match(prefix[role], q.text, re.I)]
                candidates += [f"p{i}“{q.text.strip()[:15]}”（现在是 {reverse.get(q.style.name) or q.style.name}）" for i, q in found[:2]]
            labels = {"abstract_zh": "中文摘要", "keywords_zh": "中文关键词", "reference": "参考文献条目",
                      "abstract_en": "英文摘要", "keywords_en": "英文关键词"}
            warn(
                "EMPTY_CONTENT",
                "这些部分在成品里是空的："
                + "、".join(labels[r] for r in missing_content)
                + "。先排除角色标错：已有 Word 时打开 mapping.json，看原稿里对应的文字是不是被标成了 body 或别的角色"
                "（例如摘要正文应为 abstract_zh，参考文献条目应为 reference），改对后重新 run。"
                + ("可能是这几段：" + "；".join(candidates) + "。" if candidates else "")
                + "原稿本来就没有这部分时：写作模式补写；只排版模式不补，在 review.json 说明原稿就是空的。",
            )
    for i, img in enumerate(doc.inline_shapes):
        if i == 0 and has_cover:
            continue
        if img.width > E.Twips(E.WIDTH):
            err("IMAGE_WIDTH", "图片宽度超过所选模板版心。")
    warn(
        "VISUAL_REQUIRED",
        "静态检查不等于看图检查；有看图能力就逐页打开页图。",
    )
    stats = {
        "template": spec["id"],
        "mode": mode,
        "cover": has_cover,
        "sections": len(doc.sections),
        "checked_paragraphs": checked,
        "tables": len(doc.tables),
        "tables_including_nested": table_count,
        "inline_images": len(doc.inline_shapes),
        "toc_unresolved": len(unresolved),
    }
    return {
        "passed": not errors,
        "document_sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(),
        "errors": errors,
        "warnings": warnings,
        "stats": stats,
        "validation_scope": validation_scope,
        "fully_validated": False,
    }


def check(path, final=False, intake_path=None, toc_report=None, validation_scope="full"):
    try:
        return _check(path, final, intake_path, toc_report, validation_scope)
    except Exception as exc:
        return {
            "passed": False,
            "errors": [{"code": "AUDIT_ABORTED", "message": "文档结构无法安全检查：" + str(exc)}],
            "warnings": [],
            "stats": {},
        }


def main():
    p = argparse.ArgumentParser()
    p.add_argument("docx")
    p.add_argument("--report", required=True)
    p.add_argument("--intake")
    p.add_argument("--final", action="store_true")
    p.add_argument("--toc-report")
    p.add_argument("--validation-scope", choices=["full", "static-only"], default="full")
    a = p.parse_args()
    result = check(a.docx, a.final, a.intake, a.toc_report, a.validation_scope)
    Path(a.report).parent.mkdir(parents=True, exist_ok=True)
    Path(a.report).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "passed": result["passed"],
                "errors": len(result["errors"]),
                "warnings": len(result["warnings"]),
                "report": a.report,
            },
            ensure_ascii=False,
        )
    )
    return 0 if result["passed"] else 2


if __name__ == "__main__":
    sys.exit(main())
