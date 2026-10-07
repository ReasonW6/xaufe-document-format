"""Conservative fixed-grid table layout with nested-cell bounds and no text loss."""

from __future__ import annotations
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


def setel(parent, tag, **attrs):
    el = parent.find(qn("w:" + tag))
    if el is None:
        el = OxmlElement("w:" + tag)
        parent.append(el)
    for key, val in attrs.items():
        el.set(qn("w:" + key), str(val))
    return el


def raw_rows(table):
    """Yield real XML cells once, including vMerge continuations (row.cells aliases)."""
    grid = table._tbl.tblGrid
    if grid is None:
        raise ValueError("表格缺少网格定义；需专门保全处理。")
    n = len(grid.gridCol_lst)
    for ri, tr in enumerate(table._tbl.tr_lst):
        before = tr.xpath("./w:trPr/w:gridBefore/@w:val")
        after = tr.xpath("./w:trPr/w:gridAfter/@w:val")
        col = int(before[0]) if before else 0
        start = col
        cells = []
        for tc in tr.tc_lst:
            if tc.xpath("./w:tcPr/w:hMerge"):
                raise ValueError("旧式 hMerge 水平合并需先无损转换为 gridSpan；未写入输出。")
            spans = tc.xpath("./w:tcPr/w:gridSpan/@w:val")
            span = int(spans[0]) if spans else 1
            if span < 1 or col + span > n:
                raise ValueError("表格合并跨度超出网格；未写入输出。")
            cells.append((col, span, _Cell(tc, table)))
            col += span
        end = int(after[0]) if after else 0
        if col + end != n:
            raise ValueError("表格行与列网格不一致；需专门保全处理。")
        yield ri, tr, start, end, cells


def cell_margin(table, cell, side):
    """Use explicit cell > table margins > Word default 108twip side padding."""
    alias = {"left": "start", "right": "end"}[side]
    for parent in (cell._tc.tcPr, table._tbl.tblPr):
        if parent is None:
            continue
        name = "tcMar" if parent is cell._tc.tcPr else "tblCellMar"
        mar = parent.find(qn("w:" + name))
        if mar is None:
            continue
        el = mar.find(qn("w:" + alias))
        if el is None:
            el = mar.find(qn("w:" + side))
        if el is not None:
            if el.get(qn("w:type"), "dxa") not in ("dxa", "nil"):
                raise ValueError("不支持百分比单元格内边距。")
            return 0 if el.get(qn("w:type")) == "nil" else max(0, int(el.get(qn("w:w"), "0")))
    return 108


def resize_inline(p, max_width_twips, max_height_twips):
    for inline in p._p.xpath(".//wp:inline"):
        # Missing distances acquire renderer-dependent defaults (about 9pt
        # horizontally in the tested LibreOffice). They can clip a correctly
        # sized image inside a cell. Body images are inline: use explicit zero.
        for key in ("distT", "distB", "distL", "distR"):
            inline.set(key, "0")
        ext = inline.find(qn("wp:extent"))
        if ext is None:
            raise ValueError("图片缺少尺寸。")
        cx, cy = int(ext.get("cx", "0")), int(ext.get("cy", "0"))
        if cx <= 0 or cy <= 0:
            raise ValueError("图片尺寸无效。")
        factor = min(1, max_width_twips * 635 / cx, max_height_twips * 635 / cy)
        if factor < 1:
            x, y = max(1, int(cx * factor)), max(1, int(cy * factor))
            ext.set("cx", str(x))
            ext.set("cy", str(y))
            for transform in inline.xpath(".//a:xfrm/a:ext"):
                transform.set("cx", str(x))
                transform.set("cy", str(y))


def fit_table(table, max_width, style_paragraph, *, repeat_header=False, max_height=13958):
    if max_width < 426:
        raise ValueError("表格可用宽度不足以容纳五号文字及边距；不能擅自缩小字号。")
    if table._tbl.xpath("./w:tblPr/w:tblpPr"):
        raise ValueError("浮动表格需专门核对锚点；本路径只处理行内表格。")
    rows = list(raw_rows(table))
    widths = [int(c.get(qn("w:w"), "0")) for c in table._tbl.tblGrid.gridCol_lst]
    if not widths or min(widths) <= 0:
        raise ValueError("表格列宽缺失或非正数；需先核对原网格。")
    old_total = sum(widths)
    target = min(old_total, int(max_width))
    # Cumulative rounding keeps the exact total and original column proportions.
    cumulative = [0]
    for width in widths:
        cumulative.append(cumulative[-1] + width)
    endpoints = [round(x * target / old_total) for x in cumulative]
    fitted = [endpoints[i + 1] - endpoints[i] for i in range(len(widths))]
    for _, _, _, _, cells in rows:
        for col, span, cell in cells:
            content = (
                sum(fitted[col : col + span])
                - cell_margin(table, cell, "left")
                - cell_margin(table, cell, "right")
            )
            if content < 210:
                raise ValueError(
                    "缩放后的单元格放不下一个五号汉字；需改横向页面或人工调整列，不缩字号。"
                )
    table.autofit = False
    pp = table._tbl.tblPr
    setel(pp, "tblW", w=target, type="dxa")
    setel(pp, "jc", val="center")
    setel(pp, "tblInd", w=0, type="dxa")
    setel(pp, "tblCellSpacing", w=0, type="dxa")
    for col, width in zip(table._tbl.tblGrid.gridCol_lst, fitted):
        col.set(qn("w:w"), str(width))
    for ri, tr, start, end, cells in rows:
        trpr = tr.get_or_add_trPr()
        for h in trpr.findall(qn("w:trHeight")):
            # Keep a requested minimum but never impose a clipping ceiling.
            if int(h.get(qn("w:val"), "0")) > max_height:
                raise ValueError("表格最小行高超过整页版心；需要明确调整。")
            if h.get(qn("w:hRule")) == "exact":
                h.set(qn("w:hRule"), "atLeast")
        header = (
            repeat_header
            and ri == 0
            or bool(tr.xpath('./w:trPr/w:tblHeader[not(@w:val="0") and not(@w:val="false")]'))
        )
        setel(trpr, "cantSplit", val=int(header))
        if repeat_header and ri == 0:
            setel(trpr, "tblHeader", val=1)
        if start:
            setel(trpr, "wBefore", w=sum(fitted[:start]), type="dxa")
        if end:
            setel(trpr, "wAfter", w=sum(fitted[-end:]), type="dxa")
        for col, span, cell in cells:
            width = sum(fitted[col : col + span])
            tcpr = cell._tc.get_or_add_tcPr()
            setel(tcpr, "tcW", w=width, type="dxa")
            setel(tcpr, "noWrap", val=0)
            setel(tcpr, "tcFitText", val=0)
            directions = cell._tc.xpath("./w:tcPr/w:textDirection/@w:val")
            if directions and directions[0] not in ("lrTb", "lr"):
                raise ValueError("竖排单元格需单独核对；不自动改写方向。")
            available = width - cell_margin(table, cell, "left") - cell_margin(table, cell, "right")
            for p in cell.paragraphs:
                if p._p.xpath(".//wp:anchor|.//w:pict|.//w:txbxContent"):
                    raise ValueError("单元格含浮动图片或文本框；需专门保全处理。")
                role = (
                    "equation"
                    if p._p.xpath(".//m:oMath")
                    else "figure" if p._p.xpath(".//wp:inline") else "table_text"
                )
                style_paragraph(p, role, preserve_num=bool(p._p.xpath("./w:pPr/w:numPr")))
                if role == "figure":
                    resize_inline(p, available, max_height - 400)
                    p.paragraph_format.keep_with_next = False
                if (
                    not p._p.xpath(".//w:t|.//w:drawing|.//m:oMath|.//w:fldSimple|.//w:instrText")
                    and cell.tables
                ):
                    style_paragraph(p, "spacer")
                    setel(
                        p._p.get_or_add_pPr(),
                        "spacing",
                        line=20,
                        lineRule="exact",
                        before=0,
                        after=0,
                    )
            for nested in cell.tables:
                fit_table(nested, available, style_paragraph, max_height=max_height)
    return {"before_twips": old_total, "after_twips": target, "columns": len(widths)}


def walk_tables(table, max_width, label="t0"):
    """Read-only walk used for reporting, with each nested table's actual available width."""
    yield label, table, max_width
    for ri, _, _, _, cells in raw_rows(table):
        for col, span, cell in cells:
            widths = [int(c.get(qn("w:w"), "0")) for c in table._tbl.tblGrid.gridCol_lst]
            available = (
                sum(widths[col : col + span])
                - cell_margin(table, cell, "left")
                - cell_margin(table, cell, "right")
            )
            for ni, nested in enumerate(cell.tables):
                yield from walk_tables(nested, available, f"{label}/r{ri}/c{col}/t{ni}")


def walk_paragraphs(table, label="t0"):
    for ri, _, _, _, cells in raw_rows(table):
        for col, _, cell in cells:
            for pi, p in enumerate(cell.paragraphs):
                yield f"{label}/r{ri}/c{col}/p{pi}", p
            for ti, nested in enumerate(cell.tables):
                yield from walk_paragraphs(nested, f"{label}/r{ri}/c{col}/t{ti}")
