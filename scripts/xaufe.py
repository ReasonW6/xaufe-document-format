#!/usr/bin/env python3
"""The only command an agent needs. Run from the skill folder:

    python -X utf8 scripts/xaufe.py compare
    python -X utf8 scripts/xaufe.py read --docx 用户的文件.docx
    python -X utf8 scripts/xaufe.py new --out "成品.docx" [--source 原稿]
    python -X utf8 scripts/xaufe.py check --work WORK
    python -X utf8 scripts/xaufe.py inspect --work WORK --docx 原稿.docx
    python -X utf8 scripts/xaufe.py inspect --work WORK --text 原稿.txt   # .txt/.md 原稿逐字导入 content.json
    python -X utf8 scripts/xaufe.py refs --work WORK [--min 篇数]
    python -X utf8 scripts/xaufe.py run --work WORK [--office none] [--office-note "原因"]
    python -X utf8 scripts/xaufe.py finish --work WORK
    python -X utf8 scripts/xaufe.py preview --work WORK --docx 文件.docx
    python -X utf8 scripts/xaufe.py adopt --work WORK --docx 手工改好的.docx --note "改了什么"
    python -X utf8 scripts/xaufe.py cancel --work WORK

Output is plain Chinese text meant to be read and followed step by step.
"""

from __future__ import annotations
import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
ROOT = HERE.parent
import feedback as F  # noqa: E402  (needs the scripts folder on sys.path)

EXIT_ASK_USER = 3
EXIT_NOT_ACCEPTED = 5


def _read(path, what):
    path = Path(path)
    if not path.is_file():
        raise ValueError(f"找不到 {what}：{path}")
    from textio import read_text_any

    text = read_text_any(path)
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        line = text.splitlines()[exc.lineno - 1] if 0 < exc.lineno <= len(text.splitlines()) else ""
        near = line[max(0, exc.colno - 21):exc.colno + 20].strip()
        raise ValueError(
            f"{what} 不是合法的 JSON（第 {exc.lineno} 行第 {exc.colno} 列）：{exc.msg}。"
            + (f"那里是：{near}。" if near else "")
            + "最常见的原因是文字里有英文双引号 \"：写成 \\\"（前面加反斜杠）；"
            "不要换成中文引号，那会改掉用户的原文。其次是漏了逗号或多了结尾逗号。"
        ) from exc


def _write(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def _work(value):
    import delivery as D

    root, state = D.load(value)
    return root, state


# ---------------------------------------------------------------- compare
def cmd_compare(_args):
    print((ROOT / "references/comparison.md").read_text(encoding="utf-8").strip())


# ---------------------------------------------------------------- new
def _previous_job(source):
    """job.json of an earlier output of this skill, read back from the file itself: template, logo,
    format request, cover on/off, cover values and blanks, date.  None for any other file."""
    import docx
    import engine as E
    from check_format import _cover_values
    from customization import effective_profile

    if Path(source).suffix.lower() != ".docx":
        return None
    doc = docx.Document(source)
    vv = E.get_vars(doc)
    template = vv.get("XAUFE_TEMPLATE")
    if template not in ("A", "B"):
        return None
    recorded = json.loads(vv.get("XAUFE_FORMAT_REQUEST", "{}"))
    request = dict(recorded)
    job = {"task": "format", "template": template}
    if "logo" in request:
        job["logo"] = request.pop("logo")
    if vv.get("XAUFE_HAS_COVER") == "1":
        values = _cover_values(doc, effective_profile(template, recorded))
        job["info"] = {k: v for k, v in values.items() if v}
        if vv.get("XAUFE_EFFECTIVE_DATE"):
            job["info"]["date"] = vv["XAUFE_EFFECTIVE_DATE"]
        blank = [k for k, v in values.items() if not v]
        if blank:
            job["blank"] = blank
    else:
        job["cover"] = False
    if request:
        job["format"] = request
    if vv.get("XAUFE_KEEP_AS_IS"):
        job["keep_as_is"] = vv["XAUFE_KEEP_AS_IS"]
    return job


def cmd_new(args):
    import delivery as D

    job = D.create(args.out, sources=args.source)
    work = job["workspace"]
    print(f"WORK = {work}")
    print(f"最终 Word 将保存到：{job['destination']}")
    for source in args.source or []:
        try:
            previous = _previous_job(source)
        except Exception:  # an unreadable file is the inspect step's business, not new's
            previous = None
        if previous:
            _write(Path(work) / "job.json", previous)
            print(f"原稿是本技能以前生成的成品：已按它记录的设置恢复 {Path(work) / 'job.json'}（版式、校徽、封面信息、全部格式要求）。"
                  "这是修改旧成品：只改用户这次说的那一项，其余保持不动；task 默认为 format，要改内容就改成 write。")
            print(json.dumps(previous, ensure_ascii=False, indent=1))
            manual = D.previous_manual(source)
            if manual:
                print(f"必须处理：上一版交付前有一项是手工处理的（“{manual.get('note', '')}”），它不在上面的设置里，"
                      "这次重新排版不会自动带上：写得进 job.json 的（字号、对齐、缩进、行距等）就写进 format.styles，"
                      "写不进的在新成品上重新手工处理后 adopt。run 会把它列进 review.json 的 warnings。")
            break
    print("下一步：")
    print(f"  1. 写 {Path(work) / 'job.json'}（写法见 references/job.md）")
    print(f"  2. 运行 check --work \"{work}\" 看还缺什么封面信息")
    texts = [s for s in args.source or [] if Path(s).suffix.lower() in (".txt", ".md")]
    if texts:
        print(f"  3. 原稿是文本文件：运行 inspect --work \"{work}\" --text \"{texts[0]}\"（逐字导入成 content.json 草稿）；"
              f"新写文档就直接写 {Path(work) / 'content.json'}")
    else:
        print(
            f"  3. 新写文档：写 {Path(work) / 'content.json'}；已有 Word：运行 inspect --work \"{work}\" --docx 原稿.docx"
        )


# ---------------------------------------------------------------- check
def _intake(root):
    import intake as I

    job = _read(root / "job.json", "job.json")
    if isinstance(job, dict) and job.get("task") not in ("write", "format"):
        raise ValueError(
            'job.json 缺 task：这次要写内容就写 "task": "write"，只排版就写 "task": "format"（判断方法见 SKILL.md 的“需要用户决定的事”）。'
        )
    return I.from_job(job, root / "intake.json")


def _task(root):
    return _read(root / "job.json", "job.json").get("task")


# ---------------------------------------------------------------- writing checks
NUMBER = re.compile(r"\d[\d,.]*\s*(?:%|％|‰|元|万|亿|人|名|倍|个百分点|户|家)|百分之"
                    # Chinese numerals with a magnitude ("三千万人", "十几万") or a multiple ("增长一倍")
                    r"|[一二两三四五六七八九十][一二两三四五六七八九十几余多]*[百千万亿]+[余多]?(?:人|名|元|户|家|吨)?"
                    r"|[一二两三四五六七八九十]+倍")
CITE = re.compile(r"[\[［]\d+(?:\s*[,，\-–]\s*\d+)*[\]］]")
CN_DIGITS = "一二三四五六七八九"
HEADING_NUMBER = {
    1: re.compile(r"^([一二三四五六七八九十]+)\s*、"),
    2: re.compile(r"^[（(]([一二三四五六七八九十]+)[)）]"),
    3: re.compile(r"^(\d+)\s*[.．]"),
}


def _cn_number(text):
    if text.isdigit():
        return int(text)
    tens, ten, ones = text.partition("十")
    if not ten:
        return CN_DIGITS.index(text) + 1 if len(text) == 1 else -1
    return (CN_DIGITS.index(tens) + 1 if tens else 1) * 10 + (CN_DIGITS.index(ones) + 1 if ones else 0)


def _ref_key(text):
    return re.sub(r"\s+", "", re.sub(r"^\s*[\[［]\d+[\]］]", "", str(text)))


def _sentences(text):
    """Sentences of a paragraph; a citation right after the full stop belongs to the sentence before it."""
    parts = re.findall(r"[^。！？!?；;]+(?:[。！？!?；;]\s*(?:\[[\d,，\-–—\s]+\])?|$)", str(text))
    return [p for p in parts if p.strip()]


def _cited_numbers(text):
    """[1] / [1,3] / [2-4] -> [1] / [1, 3] / [2, 3, 4], in reading order."""
    numbers = []
    for match in CITE.finditer(text):
        for part in re.split(r"\s*[,，]\s*", match.group(0)[1:-1]):
            bounds = [int(x) for x in re.split(r"\s*[-–]\s*", part) if x.strip().isdigit()]
            if len(bounds) == 2 and bounds[0] <= bounds[1] <= bounds[0] + 50:
                numbers.extend(range(bounds[0], bounds[1] + 1))
            elif bounds:
                numbers.append(bounds[0])
    return numbers


def _citation_problems(content):
    """Every [n] points to a reference, every reference is cited, numbered by first citation."""
    references = content.get("references") or []
    cited = []
    for block in content.get("blocks", []):
        runs = block.get("runs") or []
        parts = [block.get("text") or "", "".join(r.get("text", "") for r in runs), str(block.get("caption") or "")]
        if block.get("type") == "table":
            parts += [str(c) for row in [block.get("header") or []] + (block.get("rows") or []) for c in row]
        cited += _cited_numbers(" ".join(parts))
    problems = []
    if not references:
        if cited:
            problems.append(f"正文引用了 [{cited[0]}]，但 references 是空的：要么补上核对过的文献，要么删掉引用标记。")
        return problems
    beyond = sorted({n for n in cited if n < 1 or n > len(references)})
    if beyond:
        problems.append(f"正文引用了 {', '.join(f'[{n}]' for n in beyond)}，但参考文献只有 {len(references)} 条。")
    uncited = [i for i in range(1, len(references) + 1) if i not in cited]
    if uncited:
        problems.append(
            "参考文献第 " + "、".join(map(str, uncited)) + " 条在正文里没有被引用：在支撑的句子后标上序号，或删掉这条。"
        )
    first = list(dict.fromkeys(n for n in cited if 1 <= n <= len(references)))
    if first != sorted(first):
        problems.append(
            "参考文献要按正文第一次引用的顺序编号：现在第一次出现的顺序是 "
            + "、".join(f"[{n}]" for n in first[:10])
            + "。调整 references 的顺序并同步改正文里的序号。"
        )
    return problems


class Problem(str):
    """A writing-check message; ``edit`` says which input to change."""

    edit = None


def _problem(message, edit=None):
    item = Problem(message)
    item.edit = edit
    return item


def _material_text(root):
    """Text of the user's registered material (new --source / inspect), for figures the user gave."""
    import delivery as D
    from source_compare import fragments

    try:
        sources = D.load(root)[1].get("source_documents", [])
    except (OSError, ValueError, KeyError):
        return ""
    texts = []
    for source in sources:
        if source.get("role") == "original":
            items = fragments(Path(source["path"])) or []
            texts += [text for _, text, _ in items]
    return re.sub(r"\s+", "", "".join(texts)).replace("％", "%")


def _from_material(sentence, material):
    """Every figure in the sentence also appears in the user's material."""
    if not material:
        return False
    for match in NUMBER.finditer(sentence):
        token = re.sub(r"\s+", "", sentence[match.start(): match.end() + (4 if match.group(0) == "百分之" else 0)])
        if token.replace("％", "%") not in material:
            return False
    return True


def _writing_problems(root, content):
    """What a written document must satisfy before it is laid out (task = write)."""
    import refcheck as R

    problems = []
    references = content.get("references") or []
    result_path = root / "refs-result.json"
    if references and not result_path.is_file():
        problems.append(_problem("参考文献还没有核对：按 references/citations.md 写 refs.json 并运行 refs。", "refs.json"))
    elif references and _refs_result_forged(root):
        problems.append(_problem("refs-result.json 不是 refs 命令写出的那一份（被手写或改过）：删掉它，重新运行 refs。"
                                 "核对结果只能来自 refs，不能自己填。", "refs-result.json"))
    elif references:
        saved = _read(result_path, "refs-result.json")
        refs_json = root / "refs.json"
        current = hashlib.sha256(refs_json.read_bytes()).hexdigest() if refs_json.is_file() else None
        if saved.get("refs_sha256") != current:
            problems.append(_problem("refs.json 改过以后还没有重新运行 refs。", "refs.json"))
        results = {_ref_key(r["text"]): r for r in saved.get("results", [])}
        accepted = _read(root / "job.json", "job.json").get("accept_unverified_refs") is True
        for index, text in enumerate(references, 1):
            edit = (f"mapping.json 里 id 为 {content['_ref_ids'][index - 1]} 的段落" if "_ref_ids" in content
                    else f"content.json 的 references[{index - 1}]")
            r = results.get(_ref_key(text))
            if r is None or r["status"] not in (R.PASS, R.PENDING, R.USER):
                problems.append(_problem(f"参考文献第 {index} 条不在 refs 核对可用的条目里（要用 refs.json 里的 text 原样写）：{str(text)[:40]}", edit))
            elif r["status"] == R.USER and not r.get("provenance"):
                problems.append(_problem(
                    f"参考文献第 {index} 条标成用户提供（from_user），但原稿里没有这一条，refs.json 也没写 user_said（用户给它时的原话）："
                    "确实是用户给的就补上 user_said 后重新 refs；不是就去掉 from_user 正常核对。" + str(text)[:30], "refs.json"))
            elif r.get("unreachable") and not accepted:
                problems.append(_problem(
                    f"参考文献第 {index} 条因为连不上核对服务根本没核对：联网后重新运行 refs；确实没法联网，"
                    "先问用户是否接受未核实的文献，用户同意后在 job.json 写 \"accept_unverified_refs\": true：" + str(text)[:30], "job.json"))
            elif r["status"] == R.PENDING and not accepted:
                problems.append(_problem(
                    f"参考文献第 {index} 条无法验证（{R.plain_reason(r)[:40]}）：换一条能核实的文献；"
                    "确实要用，先问用户是否接受这条未核实的文献，用户同意后在 job.json 写 \"accept_unverified_refs\": true：" + str(text)[:30], edit))
            elif r["status"] != R.USER and r.get("support") == R.NO_QUOTE:
                problems.append(_problem(
                    f"参考文献第 {index} 条没有写 quote（支撑正文那句话的原文原句）：自己找的文献要从摘要或正文照抄一句写进 refs.json 的 quote，"
                    "找不到能支撑的原句就换文献或删掉那句话：" + str(text)[:30], "refs.json"))
    material = _material_text(root)
    for i, block in enumerate(content.get("blocks", [])):
        edit = f"mapping.json 里 id 为 {block['_id']} 的段落" if "_id" in block else f"content.json 的 blocks[{i}]"
        if block.get("type") == "paragraph":
            runs = block.get("runs") or []
            # Each sentence with a figure needs its own citation; one [n] elsewhere in the paragraph is not enough.
            # Figures the user supplied (they appear in the registered material) need no reference.
            for sentence in _sentences(block.get("text") or "".join(r.get("text", "") for r in runs)):
                if NUMBER.search(sentence) and not CITE.search(sentence) and not _from_material(sentence, material):
                    problems.append(_problem("这句话有具体数字但没有标出处（删掉数字，或在这句末尾标上核对过的文献序号；"
                                             f"数字是用户给的，就把用户的资料用 new --source 登记）：{sentence.strip()[:40]}", edit))
        elif block.get("type") == "table":
            cells = [str(c) for row in [block.get("header") or []] + (block.get("rows") or []) for c in row]
            text = " ".join(cells)
            if NUMBER.search(text) and not CITE.search(str(block.get("caption") or "")) and not _from_material(text, material):
                problems.append(_problem(f"表格里有具体数字但表题没有标出处（在表题末尾标上核对过的文献序号）：{text[:40]}", edit))
    problems += _citation_problems(content)
    if references and result_path.is_file():
        problems += _unbacked_figures(content, {_ref_key(r["text"]): r for r in _read(result_path, "refs-result.json").get("results", [])})
    # Heading numbers are put right by the layout step itself (heading_numbers); what writing must
    # get right is the prose around figures and tables: cited before, used after, same chapter.
    blocks = content.get("blocks", [])
    chapter, chapters = 0, []
    for block in blocks:
        chapter += block.get("type") == "heading" and block.get("level") == 1
        chapters.append(chapter)
    prose = [(i, block.get("text") or "".join(r.get("text", "") for r in block.get("runs") or []))
             for i, block in enumerate(blocks) if block.get("type") == "paragraph"]
    for i, block in enumerate(blocks):
        m = re.match(r"^\s*(图|表)\s*(\d+(?:[-－.．]\d+)?)", str(block.get("caption") or "")) \
            if block.get("type") in ("table", "image") else None
        if not m:
            continue
        kind, number = m.group(1), m.group(2)
        mention = re.compile(rf"{kind}\s*{re.escape(number)}(?![\d\-－.．]\d)")
        edit = f"mapping.json 里 id 为 {block['_id']} 的段落" if "_id" in block else f"content.json 的 blocks[{i}]"
        if not any(j < i and chapters[j] == chapters[i] and mention.search(t) for j, t in prose):
            problems.append(_problem(f"“{kind}{number}”前面同一章的正文没有引用它：先文后{kind}，在它前面写上“见{kind}{number}”之类的话。", edit))
        if not any(j > i and chapters[j] == chapters[i] and mention.search(t) for j, t in prose):
            problems.append(_problem(f"“{kind}{number}”后面同一章的正文没有交待它：写一句“由{kind}{number}可见……”说明它说明了什么。", edit))
    return problems


def _citing_sentences(content):
    """{n: every sentence of the text that cites [n], in reading order}."""
    found = {}
    for block in content.get("blocks", []):
        runs = block.get("runs") or []
        for sentence in _sentences(block.get("text") or "".join(r.get("text", "") for r in runs)):
            for n in _cited_numbers(sentence):
                found.setdefault(n, []).append(sentence.strip())
    return found


def _figures(text):
    """The numbers a sentence states as facts ("85%", "1.2万人" -> "85", "1.2")."""
    return {m.group(1).replace(",", "") for m in re.finditer(r"(\d[\d,]*(?:\.\d+)?)\s*(?:%|％|‰|元|万|亿|人|名|倍|个百分点|户|家)", text)}


def _unbacked_figures(content, results):
    """Sentences whose figures appear in none of the excerpts of the references they cite."""
    import refcheck as R

    problems = []
    references = content.get("references") or []
    for i, block in enumerate(content.get("blocks", [])):
        runs = block.get("runs") or []
        for sentence in _sentences(block.get("text") or "".join(r.get("text", "") for r in runs)):
            figures = _figures(CITE.sub("", sentence))
            cited = [n for n in _cited_numbers(sentence) if 1 <= n <= len(references)]
            if not figures or not cited:
                continue
            entries = [results.get(_ref_key(references[n - 1]), {}) for n in cited]
            if any(e.get("status") == R.USER and not e.get("quote") for e in entries):
                continue  # the user's own source: they vouch for its numbers
            quoted = set().union(*(set(re.findall(r"\d+(?:\.\d+)?", str(e.get("quote") or "").replace(",", "")))
                                   for e in entries))
            missing = sorted(f for f in figures if f not in quoted)
            if missing:
                edit = f"mapping.json 里 id 为 {block['_id']} 的段落" if "_id" in block else f"content.json 的 blocks[{i}]"
                problems.append(_problem(
                    f"这句话的数字 {'、'.join(missing)} 在它引用的文献 {'、'.join(f'[{n}]' for n in cited)} 的 quote（原文摘录）里找不到："
                    "数字要能在原文里找到——把原文里写着这个数字的句子作为 quote 写进 refs.json 后重新 refs，"
                    f"原文里没有就删掉或改正这个数字：{sentence.strip()[:40]}", edit))
    return problems


COMMON = set("研究 本文 分析 方法 系统 我们 进行 通过 基于 问题 可以 以及 发展 影响 提出 结果 表明 认为 使用 应用 方面 主要 相关 "
             "the and for with that this from are was were study paper using based".split())


def _terms(text):
    """Content words of a text: CJK character pairs and Latin words, minus the commonest ones."""
    text = str(text).lower()
    pairs = {a + b for chunk in re.findall(r"[\u4e00-\u9fff]+", text) for a, b in zip(chunk, chunk[1:])}
    return (pairs | set(re.findall(r"[a-z][a-z0-9-]{2,}", text))) - COMMON


def _off_topic_lines(root, content):
    """[n] placed on sentences that share no content word with reference n's claim, excerpt or entry:
    the number most likely points at the wrong reference.  Every citing sentence is checked."""
    if not (root / "refs-result.json").is_file():
        return []
    results = {_ref_key(r["text"]): r for r in _read(root / "refs-result.json", "refs-result.json").get("results", [])}
    citing = {}
    for block in content.get("blocks", []):
        runs = block.get("runs") or []
        for sentence in _sentences(block.get("text") or "".join(r.get("text", "") for r in runs)):
            for n in _cited_numbers(sentence):
                citing.setdefault(n, []).append(CITE.sub("", sentence).strip())
    lines = []
    for index, text in enumerate(content.get("references") or [], 1):
        r = results.get(_ref_key(text), {})
        about = _terms(" ".join([str(text), str(r.get("supports") or ""), str(r.get("quote") or "")]))
        off = [s for s in citing.get(index, []) if not _terms(s) & about]
        if off:
            lines.append((index, f"[{index}] 引用句“{off[0][:30]}”和这条文献的题目、supports、quote 没有一个共同的词"))
    return lines


def _off_topic_numbers(root, content):
    return {index for index, _ in _off_topic_lines(root, content)}


def _off_topic(root, content):
    """REF_CITATION_MISMATCH warning for the run output and review.json."""
    lines = [line for _, line in _off_topic_lines(root, content)]
    if not lines:
        return None
    return {"code": "REF_CITATION_MISMATCH",
            "message": "这些引用号标的句子和对应文献看起来说的不是一回事，可能标错了序号：" + "；".join(lines)
                       + "。逐条核对：序号标错就改正文里的 [n]；确实是这篇文献支撑的，在 review.json 里说明它支撑了哪一点。"}


def _weak_support(root, content):
    """References whose quoted excerpt was not found in what could be read: the claim rests on the agent's word."""
    import refcheck as R

    if not (root / "refs-result.json").is_file():
        return None
    results = {_ref_key(r["text"]): r for r in _read(root / "refs-result.json", "refs-result.json").get("results", [])}
    citing = _citing_sentences(content)
    lines = []
    for index, text in enumerate(content.get("references") or [], 1):
        r = results.get(_ref_key(text), {})
        if r.get("status") != R.USER and r.get("support") in (R.NOT_FOUND, R.UNCHECKED):
            lines.append(f"[{index}]（文献核对：{r.get('status')}）引用句“{(citing.get(index) or [''])[0][:30]}”"
                         f"／摘录“{str(r.get('quote') or '')[:30]}”（{r['support']}）")
    if not lines:
        return None
    return {"code": "REF_SUPPORT_UNCONFIRMED",
            "message": "这些引用的原文摘录没能在读得到的原文里对上，支撑关系只有你的判断：" + "；".join(lines)
                       + "。逐条核对引用句和摘录说的是不是同一件事；对不上就换文献或改写那句话。清单会把这几条标成未确认。"}


def _readable(text):
    return re.sub(r"^\s*[\[［]\d+[\]］]\s*", "", str(text)).strip()


def _written(root):
    """The written text in content.json shape, also when it comes from a Word draft (mapping.json):
    the writing checks then apply the same way, and each block remembers its mapping id."""
    if (root / "content.json").is_file():
        return _read(root / "content.json", "content.json")
    if not (root / "mapping.json").is_file():
        return None
    rows = _read(root / "mapping.json", "mapping.json").get("paragraphs", [])
    blocks, references, ref_ids = [], [], []
    for r in rows:
        role, text = r.get("role"), r.get("text", "")
        if role == "reference" and text.strip():
            references.append(text)
            ref_ids.append(r.get("id"))
        elif role in ("h1", "h2", "h3"):
            blocks.append({"type": "heading", "level": int(role[1]), "text": text, "_id": r.get("id")})
        elif role == "body":
            blocks.append({"type": "paragraph", "text": text, "_id": r.get("id")})
    return {"blocks": blocks, "references": references, "_ref_ids": ref_ids}


def _document_references(root):
    """Reference lines of the document being laid out: content.json references or mapping rows,
    plus the references job.json "corrections" adds to an existing Word."""
    if (root / "content.json").is_file():
        return [str(r) for r in _read(root / "content.json", "content.json").get("references") or []]
    if (root / "mapping.json").is_file():
        rows = _read(root / "mapping.json", "mapping.json").get("paragraphs", [])
        return [r.get("text", "") for r in rows if r.get("role") == "reference" and r.get("text", "").strip()] \
            + _added_references(root)
    return []


def _added_references(root):
    """References the agent adds or rewrites through job.json "corrections"."""
    return [str(c["after"]) for c in _corrections(root) if c.get("part") == "references" and str(c.get("after") or "").strip()]


def _added_reference_problems(root):
    """A reference the agent adds to the user's paper must be found and confirmed real first (PASS);
    one that cannot be confirmed is left out and the user is told, never added on trust."""
    import refcheck as R

    added = _added_references(root)
    if not added:
        return []
    result_path = root / "refs-result.json"
    if not result_path.is_file():
        return [_problem("job.json 的 corrections 补了参考文献，但还没有核对：运行 refs --from-document（会带上补的条目），"
                         "全部“通过”才能补；核实不了的那条去掉，交付时告诉用户。", "refs.json")]
    if _refs_result_forged(root):
        return [_problem("refs-result.json 不是 refs 命令写出的那一份（被手写或改过）：删掉它，重新运行 refs。", "refs-result.json")]
    saved = _read(result_path, "refs-result.json")
    refs_json = root / "refs.json"
    if saved.get("refs_sha256") != (hashlib.sha256(refs_json.read_bytes()).hexdigest() if refs_json.is_file() else None):
        return [_problem("refs.json 改过以后还没有重新运行 refs。", "refs.json")]
    results = {_ref_key(r["text"]): r for r in saved.get("results", [])}
    problems = []
    for text in added:
        r = results.get(_ref_key(text))
        if r is None:
            problems.append(_problem("corrections 补的这条参考文献不在 refs.json 里：删掉 refs.json 后运行 refs --from-document，"
                                     "或把它原样写进 refs.json 再运行 refs：" + text[:40], "refs.json"))
        elif r["status"] != R.PASS:
            reason = "连不上核对服务，没能核对" if r.get("unreachable") else R.plain_reason(r)[:40]
            problems.append(_problem(f"corrections 补的这条参考文献没有核实通过（{reason}）：按用户的要求，核实不了的不补——"
                                     "从 job.json 的 corrections 里去掉它，换一条能核实的真实文献，或交付时告诉用户少补了一条：" + text[:40],
                                     "job.json"))
    return problems


def _save_refs_result(root, data):
    """refs-result.json plus its fingerprint in the workspace record: a result file the refs command
    did not write (or one edited afterwards) is then caught at run and finish."""
    import delivery as D

    _write(root / "refs-result.json", data)
    root, state = D.load(root)
    state["refs_result_sha256"] = hashlib.sha256((root / "refs-result.json").read_bytes()).hexdigest()
    D.write_json(root / D.MARKER, state)


def _refs_result_forged(root):
    """True when refs-result.json is not the file the refs command last wrote."""
    import delivery as D

    path = root / "refs-result.json"
    if not path.is_file():
        return False
    recorded = D.load(root)[1].get("refs_result_sha256")
    return recorded != hashlib.sha256(path.read_bytes()).hexdigest()


def _auto_original_refs(root):
    """Format only: check the original's references once (report only, the text is never changed),
    so finish can tell the user what was found.  Skipped when the agent already ran refs."""
    import os
    import refcheck as R

    if os.environ.get("XAUFE_SKIP_REF_CHECK") or (root / "refs.json").exists() or (root / "refs-result.json").exists() \
            or (root / "refs-offline.json").exists():
        return None
    lines = _document_references(root)
    if not lines:
        return None
    if not R._online():
        _write(root / "refs-offline.json", {"references": len(lines)})
        return "原文参考文献：这台机器连不上文献核对服务，没有核对（finish 会如实告诉用户；联网后可以运行 refs --from-document 再核对）。"
    entries = [R.parse_line(line) for line in lines]
    _write(root / "refs.json", entries)
    results = R.check_all(entries)
    _save_refs_result(root, {"refs_sha256": hashlib.sha256((root / "refs.json").read_bytes()).hexdigest(),
                             "min": 0, "results": results})
    return ("原文参考文献已自动核对（只报告，不改原文；finish 会把结果给用户）：\n"
            + R.report(results, entries, 0, report_only=True))


def _original_reference_report(root):
    """Format only: what the refs check found about the user's own references (nothing was changed)."""
    import refcheck as R

    if not (root / "refs-result.json").is_file():
        try:
            count = len(_document_references(root))
        except (OSError, ValueError):
            count = 0
        why = "这次连不上文献核对服务，" if (root / "refs-offline.json").is_file() else ""
        return (f"这是只排版：文中的 {count} 条参考文献按原文保留，只统一了格式；{why}没有核对它们是否真实存在、"
                "信息是否准确。需要的话可以再帮你逐条核对。") if count else None
    saved = _read(root / "refs-result.json", "refs-result.json")
    added = {_ref_key(t) for t in _added_references(root)}
    lines = ["| 序号 | 文献 | 核对结果 | 说明 |", "|---|---|---|---|"]
    for index, r in enumerate(saved.get("results", []), 1):
        reason = "" if r["status"] == R.PASS else R.plain_reason(r)[:80]
        mark = "（补充）" if _ref_key(r["text"]) in added else ""
        lines.append(f"| {index} | {_readable(r['text'])[:60]}{mark} | {r['status']} | {reason} |")
    lines.append(("标“补充”的是原稿缺少、按要求补上并逐条核实过的文献；其余是原文的，没有改动。" if added else
                  "这是只排版：原文的参考文献没有改动。") + "“信息不一致”“排除”的条目建议你自己核对后再决定要不要改。")
    return "\n".join(lines)


def _original_text(root):
    """All text of the registered originals, normalised for "is this line in there" checks."""
    import delivery as D
    from source_compare import fragments, norm

    texts = []
    for item in D.load(root)[1].get("source_documents", []):
        found = fragments(item["path"]) if Path(item["path"]).is_file() else None
        texts += [t for _, t, _ in found or []]
    return norm(" ".join(texts))


def _user_provenance(entry, originals):
    """Where a from_user reference came from: the user's own words, or a registered original."""
    from source_compare import norm

    if isinstance(entry.get("user_said"), str) and entry["user_said"].strip():
        return "用户原话"
    line = norm(_readable(entry.get("text", "")))
    return "原稿" if len(line) >= 8 and line in originals else None


def _reference_checklist(root):
    """Ready-made table for the reply to the user (task = write)."""
    import refcheck as R

    written = _written(root)
    if written is None:
        return None
    references = written.get("references") or []
    saved = _read(root / "refs-result.json", "refs-result.json") if (root / "refs-result.json").is_file() else {}
    minimum = saved.get("min", R.DEFAULT_MIN)
    by_key = {_ref_key(r["text"]): r for r in saved.get("results", [])}
    lines = []
    if references:
        citing = _citing_sentences(written)
        doubtful = _off_topic_numbers(root, written)
        lines += ["| 序号 | 文献 | 真实性 | 支撑的论点 | 正文引用句 | 观点支撑 |", "|---|---|---|---|---|---|"]
        for index, text in enumerate(references, 1):
            r = by_key.get(_ref_key(text), {})
            status = {R.PASS: "已验证"}.get(r.get("status"))
            if r.get("status") == R.USER:
                status = f"你提供的（自动核对：{r.get('checked', '未核对')}）"
            if status is None:
                status = "无法验证：" + R.plain_reason(r)[:40]
            supports = r.get("supports") or ""
            if supports and r.get("relevance") == "中":
                supports += "（弱支撑）"
            backing = r.get("support") or R.NO_QUOTE
            if index in doubtful:
                backing += "；引用号可能标错"
            said = " ／ ".join(x[:40] for x in citing.get(index, []))[:160]
            lines.append(f"| [{index}] | {_readable(text)[:60]} | {status} | {supports} | {said} | {backing} |")
        lines.append("“真实性”只说明文献存在、信息对得上；“观点支撑”才说明你引用的那句话在原文里找到了依据。")
    else:
        lines.append("这次没有找到能核实的参考文献，请提供你读过的文献。")
    if 0 < len(references) < minimum:
        lines.append(f"参考文献现在 {len(references)} 篇，比要求的 {minimum} 篇少 {minimum - len(references)} 篇，请补充你读过的文献。")
    stats = _body_length(root)
    if stats and stats[0] < 5000:
        lines.append(f"正文现在约 {stats[0]} 字，少于默认的 5000 字：用户另说过字数就以用户的为准，没说过就如实告诉用户。")
    lines.append("请在提交前再核对一遍全部参考文献。")
    return "\n".join(lines)


def cmd_check(args):
    import intake as I

    root, _ = _work(args.work)
    state = _intake(root)
    print(I.summary(state))
    if state["stage"] != "ready":
        print("\n→ 先问用户上面缺的项（可以让用户选“留空自己填”），把回答写进 job.json 后再运行 check。")
        return EXIT_ASK_USER
    print("\n→ 封面信息齐了，可以准备内容并运行 run。")
    return 0


# ---------------------------------------------------------------- inspect
H1 = re.compile(r"^(第[一二三四五六七八九十百\d]+章|[一二三四五六七八九十]+\s*[、.．])")
H2 = re.compile(r"^[（(][一二三四五六七八九十]+[)）]")
H3 = re.compile(r"^\d+\s*[.．、]\s*\S")
NUM2 = re.compile(r"^\d+\.\d+(?!\.\d)\s*\S")
NUM3 = re.compile(r"^\d+\.\d+\.\d+\s*\S")
COVER_WORDS = ("学号", "学生姓名", "指导教师", "班级", "完成日期", "学  号", "学　　号")
FIGURE_LISTS = ("图目录", "表目录", "插图目录", "插图清单", "附图目录", "附表目录", "图表目录", "图表清单", "附表清单")
COVER_FIELDS = ("学号", "姓名", "指导教师", "班级", "专业", "学院", "完成日期", "提交日期", "课程名称", "任课教师")


def _plain(text):
    return re.sub(r"[\s　]+", "", text)


def _heading_level(block):
    style = (block.get("style") or "").lower()
    match = re.search(r"(?:heading|标题)\s*(\d)", style)
    if match and int(match.group(1)) <= 3:
        return int(match.group(1))
    text = block["text"].strip()
    if len(text) > 40 or text.endswith(("。", "；", "，")):
        return None
    if H1.match(text):
        return 1
    if H2.match(text) or NUM2.match(text):
        return 2
    if NUM3.match(text) or (H3.match(text) and len(text) <= 30):
        return 3
    return None


def guess_roles(blocks):
    """Best-effort role for every top-level paragraph; the agent must review it.

    Paragraphs that already use this skill's styles (an earlier output) are
    mapped exactly; other documents are guessed from text patterns.
    """
    from customization import ROLE_STYLES
    from engine import TABLE_NOTE, TABLE_UNIT

    known = {name: role for role, name in ROLE_STYLES.items() if role != "cover"}
    paragraphs = [b for b in blocks if b["kind"] == "paragraph"]
    cover_end = None
    for b in blocks[:60]:
        if b["kind"] == "paragraph" and b["section_break"]:
            before = "".join(x["text"] for x in blocks if x["block_index"] <= b["block_index"])
            if any(word in before for word in COVER_WORDS):
                cover_end = b["block_index"] + 1
            break
        if b["kind"] == "paragraph" and b.get("page_break"):
            # A cover ended by a page break: at least two cover fields ("学号" "指导教师" …) on that first page.
            end = b["block_index"] + (b["page_break"] == "after")
            before = _plain("".join(x["text"] for x in blocks if x["block_index"] < end))
            if sum(word in before for word in COVER_FIELDS) >= 2:
                cover_end = end
            break
    roles, area, seen_heading = {}, None, False
    for b in paragraphs:
        text = b["text"].strip()
        plain = _plain(text)
        style = (b.get("style") or "").lower()
        if cover_end is not None and b["block_index"] < cover_end:
            role = "preserve"
        elif b.get("style") in known:
            role = known[b["style"]]
            seen_heading = seen_heading or role in ("h1", "h2", "h3", "h4", "preface")
        elif not text and b["images"]:
            role = "figure"
        elif not text and b["equations"]:
            role = "equation"
        elif not text:
            role = "spacer"
        elif plain in ("目录", "目錄"):
            role, area = "toc_title", "toc"
        elif plain in FIGURE_LISTS:  # the school's TOC has no list of figures or tables: it is dropped
            role, area = "old_toc", "figure_list"
        elif area == "figure_list" and re.search(r"\d+\s*$", text) and re.match(r"^\s*(图|表|Figure|Table)", text):
            role = "old_toc"
        elif style.startswith("table of figures"):
            role = "old_toc"
        elif style.startswith("toc") or (area == "toc" and re.search(r"\d+\s*$", text) and not seen_heading):
            role = "preserve"
        elif plain in ("摘要", "内容摘要", "中文摘要"):
            role, area = "abstract_label_zh", "abstract_zh"
        elif not seen_heading and INLINE_ABSTRACT_ZH.match(plain):  # "摘要：本文……" in one paragraph
            role, area = "abstract_zh", "abstract_zh"
        elif not seen_heading and INLINE_ABSTRACT_EN.match(text):
            role, area = "abstract_en", "abstract_en"
        elif re.match(r"^关键词\s*[:：]", plain):
            role, area = "keywords_zh", None
        elif plain.lower() == "abstract":
            role, area = "abstract_label_en", "abstract_en"
        elif re.match(r"^key\s*words?\s*[:：]", text, re.I):
            role, area = "keywords_en", None
        elif plain in ("参考文献", "主要参考文献"):
            role, area = "reference_title", "references"
        elif area == "references":
            role = "reference"
        elif area in ("abstract_zh", "abstract_en"):
            role = area
        elif plain in ("序言", "引言", "前言"):
            role, area, seen_heading = "preface", None, True
        elif TABLE_NOTE.match(text) and (_neighbour(blocks, b, -1).get("kind") == "table"
                                         or roles.get(_neighbour(blocks, b, -1).get("id")) == "table_note"):
            role = "table_note"
        elif TABLE_UNIT.match(text) and _neighbour(blocks, b, 1).get("kind") == "table":
            role = "table_unit"
        elif re.match(r"^(图|表|Figure|Table)\s*\d", text) and len(text) <= 60:
            role = "caption"
        elif b.get("numbered"):  # a Word list item (its number is not a typed heading number)
            role = "body"
        elif _heading_level(b):
            role, area, seen_heading = "h" + str(_heading_level(b)), None, True
        elif (
            not seen_heading
            and area is None
            and len(text) <= 200
            and re.search(r"[A-Za-z]", text)
            and not re.search(r"[一-鿿]", text)
            and not text.endswith(".")
            and not roles_has(roles, "title_en")
        ):
            role = "title_en"
        elif not seen_heading and area is None and len(text) <= 40 and not text.endswith(("。", "：", ":")) \
                and not roles_has(roles, "title_zh"):
            role = "title_zh"
        else:
            role, area = "body", (None if area == "toc" else area)
        roles[b["id"]] = role
    return roles, cover_end


def _neighbour(blocks, b, step):
    """The nearest non-blank block before (step -1) or after (step 1) block b."""
    at = next(i for i, x in enumerate(blocks) if x is b) + step
    while 0 <= at < len(blocks):
        x = blocks[at]
        if x["kind"] == "table" or x["text"].strip() or x.get("images") or x.get("equations"):
            return x
        at += step
    return {}


INLINE_ABSTRACT_ZH = re.compile(r"^(?:内容|中文)?摘\s*要\s*[:：]\s*\S")
INLINE_ABSTRACT_EN = re.compile(r"^abstract\s*[:：]\s*\S", re.I)


def roles_has(roles, role):
    return role in roles.values()


def _inspect_text(root, s, source):
    """A .txt/.md original -> content.json with its words verbatim and a guessed structure."""
    import delivery as D
    from text_import import to_content

    if source.suffix.lower() not in (".txt", ".md") or not source.is_file():
        raise ValueError("--text 要指向原稿的 .txt 或 .md 文件（Word 原稿用 --docx）。")
    for name in ("content.json", "mapping.json"):
        if (root / name).exists():
            raise ValueError(f"WORK 里已经有 {name}：要从原稿重新导入就先删掉它。")
    content, notes, changes = to_content(source, root)
    s["import_corrections"] = changes
    _write(root / "content.json", content)
    if not any(x["path"] == str(source) for x in s["source_documents"]):
        s["source_documents"].append({"path": str(source), "sha256": D.sha(source), "role": "original"})
    D.write_json(root / D.MARKER, s)
    kinds = {}
    for b in content["blocks"]:
        kinds[b["type"]] = kinds.get(b["type"], 0) + 1
    print(f"已把原稿逐字导入 {root / 'content.json'}（mode={content['mode']}）："
          + "，".join(f"{k} {v} 块" for k, v in sorted(kinds.items()))
          + (f"，参考文献 {len(content['references'])} 条" if content.get("references") else "") + "。")
    for note in notes:
        print("  注意：" + note)
    print("→ 打开 content.json 核对结构（题目、标题级别、摘要、关键词、参考文献、表题/图题），只改结构不改字；然后运行 run。")
    return 0


def cmd_inspect(args):
    import delivery as D
    import engine as E
    import intake as I

    root, s = _work(args.work)
    if args.text:
        return _inspect_text(root, s, Path(args.text).expanduser().resolve())
    source = Path(args.docx).expanduser().resolve()
    if source.suffix.lower() != ".docx":
        raise ValueError("只能读取 .docx。旧版 .doc 请先用 Word/WPS“另存为 .docx”。")
    state = _intake(root)
    inv = E.inventory(source)
    roles, cover_end = guess_roles(inv["blocks"])
    wants_cover = state.get("include_cover", True)
    own_output = any((b.get("style") or "").startswith("XAUFE ") for b in inv["blocks"])
    values = set(roles.values())
    # A full paper (abstract + headings) gets the template's structure: its own TOC,
    # abstract pages and page numbering. An earlier output of this skill already has it.
    # An earlier output of this skill made with the other template has that template's part order
    # (A: TOC before the abstract, B: after): take its structure apart and build the new one.
    previous = E.get_vars(__import__("docx").Document(source)).get("XAUFE_TEMPLATE") if own_output else None
    rebuild = bool(previous) and previous != state.get("template") and "abstract_zh" in values
    if rebuild:
        for b in inv["blocks"]:
            if b["kind"] == "paragraph" and roles.get(b["id"]) in ("toc_title", "toc1", "toc2", "toc3", "toc4"):
                roles[b["id"]] = "old_toc"
    # The school's paper structure is the default: parts the original lacks are completed, not skipped.
    paper = rebuild or (
        not own_output
        and bool(values & {"h1", "h2", "h3", "preface"})
        and not any(b.get("section_break") for b in inv["blocks"][(cover_end or 0) :] if b["kind"] == "paragraph")
    )
    if paper:
        for b in inv["blocks"]:
            if b["kind"] == "paragraph" and (roles.get(b["id"]) == "toc_title" or (
                roles.get(b["id"]) == "preserve" and (cover_end is None or b["block_index"] >= cover_end)
            )):
                roles[b["id"]] = "old_toc"
    mapping = {
        "source": str(source),
        "input_sha256": inv["input_sha256"],
        "mode": "paper" if paper else "document",
        **({"rebuild_structure": True} if rebuild else {}),
        "format_tables": True,
        "replace_cover_before": (cover_end if cover_end is not None else 0) if wants_cover else None,
        "paragraphs": [
            {"id": b["id"], "role": roles[b["id"]], "text": (b.get("list_label") or "") + b["text"][:60]}
            for b in inv["blocks"]
            if b["kind"] == "paragraph"
        ],
    }
    old_cover = [b["text"].strip() for b in inv["blocks"] if b["kind"] == "paragraph" and cover_end is not None
                 and b["block_index"] < cover_end and b["text"].strip()]
    if old_cover and wants_cover:
        mapping["old_cover"] = old_cover
    _write(root / "mapping.json", mapping)
    # mapping.json shows the first 60 characters; failures are located against the whole paragraph.
    _write(root / "paragraph-text.json", {b["id"]: (b.get("list_label") or "") + b["text"] for b in inv["blocks"] if b["kind"] == "paragraph"})
    if not any(x["path"] == str(source) for x in s["source_documents"]):
        s["source_documents"].append(
            {"path": str(source), "sha256": D.sha(source), "role": "original"}
        )
        D.write_json(root / D.MARKER, s)
    counts = {}
    for role in roles.values():
        counts[role] = counts.get(role, 0) + 1
    print(f"已生成 {root / 'mapping.json'}，共 {len(roles)} 段。自动判断结果：")
    print("  " + "，".join(f"{k} {v} 段" for k, v in sorted(counts.items())))
    if cover_end is not None and wants_cover:
        print(f"  检测到原稿开头有旧封面（到第 {cover_end} 个块为止），会整页换成所选模板的封面：旧封面上的文字、图片、"
              "声明之类模板封面没有的内容都会去掉（交付时告诉用户）。")
        found = [line for line in old_cover if re.search(r"[:：]\s*\S", line)]
        if found:
            print("  旧封面上写着：" + "；".join(found[:8]) + "。job.json 的 info 里还没有的项，可以从这里取"
                  "（和用户确认过再写；看不清的不要猜）。")
    elif cover_end is not None:
        print("  注意：原稿开头有封面，但用户不要封面。先和用户确认；脚本不会自动删除封面。")
    elif wants_cover:
        print("  原稿没有封面，会在最前面加上所选模板的封面。")
    if rebuild:
        print(f"  原稿是这个技能之前用 {previous} 版排的，这次换成 {state.get('template')} 版："
              "旧目录和分节会拆掉，按新版式重新排摘要、目录的顺序和页码。")
    if paper:
        print("  原稿有摘要和标题：按完整论文结构排（mode=paper），会自动生成目录、摘要单独成页。")
        if "old_toc" in roles.values():
            print("  原稿里的旧目录已标为 old_toc，会换成自动生成的目录。")
        texts = {b["id"]: b["text"].strip() for b in inv["blocks"] if b["kind"] == "paragraph"}
        for pid, role in roles.items():
            if role in ("abstract_zh", "abstract_en") and (INLINE_ABSTRACT_ZH.match(_plain(texts.get(pid, "")))
                                                          or INLINE_ABSTRACT_EN.match(texts.get(pid, ""))):
                print(f"  注意：{pid} 以“{texts[pid][:3]}”开头。模板会在它上面单独放一行摘要标题，这段原样保留，"
                      "读起来会有两个“摘要”。脚本不会改这段文字：在给用户的回复里说明，"
                      "用户想去掉就请用户在原稿里删掉段首这几个字，再从 new 重新排。")
        from customization import profile_for

        missing = [name for role, name in (("abstract_zh", "中文摘要"), ("keywords_zh", "关键词"), ("reference", "参考文献"))
                   if role not in values]
        if missing:
            print(f"  原稿缺{'、'.join(missing)}：只排版也要补全。摘要（150～300 字）和关键词根据正文起草；参考文献只补"
                  "联网查到并用 refs 核实通过的真实文献，核实不了的不补。都写进 job.json 的 corrections（只写 after 和 why，"
                  "见 references/job.md），交付时会逐条告诉用户是补的。")
        if profile_for(state)["english_abstract"] and "abstract_en" not in values:
            print(
                "  原稿没有英文摘要，所选版式需要：把中文题目、摘要、关键词译成英文写进 mapping.json 的 add"
                "（见 references/content.md），交付时会告诉用户是补的；用户明确不要英文摘要，才在 job.json 的 format 写 "
                "\"english_abstract\": false。"
            )
    elif not own_output:
        print("  原稿没有章节标题，按普通文档排（mode=document），不加目录和摘要页。论文缺章节标题没法自动补："
              "在回复里告诉用户，用户给出章节划分后再按论文结构排。")
    print("→ 打开 mapping.json，逐段核对 role（尤其是标题级别和正文），改错的项；然后运行 run。")
    print("  可用 role 见 references/content.md 的“已有 Word”一节。")
    return 0


# ---------------------------------------------------------------- run
def _corrections(root):
    """job.json "corrections", checked for shape: the fixes the agent made to the user's text."""
    import engine as E

    if not (root / "job.json").is_file():
        return []
    items = _read(root / "job.json", "job.json").get("corrections") or []
    if not isinstance(items, list):
        raise ValueError("job.json 的 corrections 要写成列表：[{\"part\": \"abstract\", \"before\": \"原文整段\", "
                         "\"after\": \"改后整段\", \"why\": \"为什么改\"}]。")
    for n, c in enumerate(items):
        where = f"job.json 的 corrections[{n}]"
        if not isinstance(c, dict):
            raise ValueError(f"{where} 要写成 {{\"part\": …, \"before\": …, \"after\": …, \"why\": …}}。")
        if c.get("part") not in E.CORRECTION_PARTS:
            raise ValueError(f"{where} 的 part 只能是：" + "、".join(f"{k}（{v}）" for k, v in E.CORRECTION_PARTS.items()) + "。")
        if not str(c.get("why") or "").strip():
            raise ValueError(f"{where} 缺 why：写清为什么改（交付时原样告诉用户），例如“摘要只有 120 字，按规定补到 150 字以上”。")
        if not str(c.get("before") or "").strip() and not str(c.get("after") or "").strip():
            raise ValueError(f"{where} 的 before 和 after 都是空的：改一段就两样都写，补一段只写 after，删一段只写 before。")
        if not str(c.get("before") or "").strip() and c.get("part") == "cover":
            raise ValueError(f"{where}：封面由模板生成，不能补写；封面信息写在 job.json 的 info 里。")
        extra = set(c) - {"part", "before", "after", "why", "insert_after"}
        if extra:
            raise ValueError(f"{where} 有不认识的项“{sorted(extra)[0]}”；只能有 part、before、after、why、insert_after。")
    return items


def _resolve_mapping(root):
    mapping = _read(root / "mapping.json", "mapping.json")
    rows = mapping.get("paragraphs")
    if not isinstance(rows, list):
        raise ValueError("mapping.json 缺少 paragraphs 列表；请重新运行 inspect。")
    roles = {}
    for row in rows:
        if not isinstance(row, dict) or not row.get("id") or not row.get("role"):
            raise ValueError("mapping.json 的每一段都要有 id 和 role。")
        roles[row["id"]] = row["role"]
    resolved = {
        "input_sha256": mapping.get("input_sha256"),
        "mode": mapping.get("mode", "document"),
        "format_tables": bool(mapping.get("format_tables", True)),
        "has_cover": False,
        "roles": roles,
    }
    if mapping.get("replace_cover_before") is not None:
        resolved["replace_cover_before"] = mapping["replace_cover_before"]
    if mapping.get("add"):
        resolved["add"] = mapping["add"]
    if mapping.get("rebuild_structure"):
        resolved["rebuild_structure"] = True
    corrections = _corrections(root)
    if corrections:
        resolved["corrections"] = corrections
    _write(root / "mapping.resolved.json", resolved)
    return Path(mapping["source"]), root / "mapping.resolved.json"


def _build(root, level):
    import delivery as D
    import engine as E

    content = root / "content.json"
    mapping = root / "mapping.json"
    draft = root / "draft.docx"
    if content.is_file() and mapping.is_file():
        raise ValueError("WORK 里同时有 content.json 和 mapping.json，只能保留一个：新写文档用 content.json，已有 Word 用 mapping.json。")
    if content.is_file():
        result = E.build(content, draft, force=True, intake_path=root / "intake.json", cover_min_level=level) or {}
        applied = _declared_corrections(root, content)
    elif mapping.is_file():
        source, resolved = _resolve_mapping(root)
        result = E.restyle(
            source, resolved, draft, force=True, intake_path=root / "intake.json", cover_min_level=level
        )
        if result.get("deleted_paragraphs") and level == 0:
            print(f"按 mapping 删除了 {len(result['deleted_paragraphs'])} 段（role 为 delete 或 old_toc）：")
            for text in result["deleted_paragraphs"]:
                print("  - " + (text or "（空段）"))
        if result.get("typed_indent_removed") and level == 0:
            shown = result["typed_indent_removed"]
            print(f"去掉了 {len(shown)} 段段首用空格/全角空格/Tab 打出来的缩进（模板自带首行缩进，文字没动）："
                  + "、".join(shown[:5]) + ("等" if len(shown) > 5 else ""))
        if result.get("list_labels_frozen") and level == 0:
            shown = result["list_labels_frozen"]
            print(f"{len(shown)} 段标题/文献的自动编号转成了同样的文字（看到的编号不变，目录和编号检查要读文字）："
                  + "、".join(shown[:5]) + ("等" if len(shown) > 5 else ""))
        if result.get("blank_page_breaks_removed") and level == 0:
            shown = result["blank_page_breaks_removed"]
            print(f"去掉了 {len(shown)} 个只会在 Word 里多出空白页的分页符（文字没动）：" + "、".join(shown[:5])
                  + ("等" if len(shown) > 5 else ""))
        applied = [dict(zip(("part", "before", "after", "why"), c)) for c in result.get("corrections_applied", [])]
        for c in result.get("keyword_spacing_changes", []):
            if not c.get("punctuation"):
                continue
            own = next((a for a in applied if a.get("after") == c["before"]), None)
            if own:  # a keyword line the agent added: report it as it now reads
                own["after"] = c["after"]
            else:
                applied.append({"part": "keywords", "before": c["before"], "after": c["after"],
                                "why": "关键词之间按模板用空格分开，去掉了原来的标点"})
        rows = _read(mapping, "mapping.json")
        if rows.get("old_cover") and rows.get("replace_cover_before"):
            applied.append({"part": "cover", "before": "；".join(rows["old_cover"]), "after": "",
                            "why": "旧封面整页换成了模板封面，上面模板没有的文字、图片和声明都去掉了"})
        lists = [r["text"] for r in rows.get("paragraphs", []) if r.get("role") == "old_toc"
                 and re.sub(r"\s+", "", r.get("text", "")) in FIGURE_LISTS]
        if lists:
            applied.append({"part": "text", "before": "、".join(lists), "after": "",
                            "why": "学校要求目录只用标题目录，原稿的图表目录连同条目删去了"})
        added = rows.get("add") or {}
        for key in ("title_en", "abstract_en", "keywords_en"):
            if added.get(key):
                text = "; ".join(map(str, added[key])) if isinstance(added[key], list) else str(added[key])
                applied.append({"part": key, "before": "", "after": text, "why": "原稿没有英文摘要，按版式要求译自中文补上"})
    else:
        raise ValueError("还没有内容：新写文档请写 content.json；已有 Word 请先运行 inspect。")
    applied += [dict(zip(("part", "before", "after", "why"), c)) for c in result.get("auto_corrections", [])]
    r, st = D.load(root)
    if content.is_file():  # what reading the .txt/.md original already put the school's way
        applied = st.get("import_corrections", []) + applied
    st["corrections"] = applied
    D.write_json(r / D.MARKER, st)
    if applied and level == 0:
        print(f"按版式要求改了 {len(applied)} 处文字（交付时会逐条告诉用户）：")
        for c in applied:
            print("  - " + D.correction_line(c))
    return draft


def _declared_corrections(root, content):
    """content.json is written by the agent, so its corrections are already in it: check each one is."""
    items = _corrections(root)
    if not items:
        return []

    def flat(value):
        if isinstance(value, dict):
            return "".join(flat(v) for v in value.values())
        if isinstance(value, list):
            return "".join(flat(v) for v in value)
        return value if isinstance(value, str) else ""

    squash = lambda t: re.sub(r"\s+", "", t or "")
    text = squash(flat(_read(content, "content.json")))
    for n, c in enumerate(items):
        after = squash(c.get("after"))
        if after and after not in text:
            raise ValueError(f"job.json 的 corrections[{n}] 写了改后的文字，但 content.json 里没有它：content.json 是你写的，"
                             f"改动要直接写进 content.json，corrections 只是登记（交付时告诉用户）。没找到：“{c['after'][:30]}”。")
    return [{k: c.get(k, "") for k in ("part", "before", "after", "why")} for c in items]


def _body_length(root):
    """Characters of body paragraphs in content.json (CJK characters plus Latin words, no punctuation)."""
    content = root / "content.json"
    if not content.is_file():
        return None
    data = _read(content, "content.json")
    total = 0
    for block in data.get("blocks", []):
        if block.get("type") != "paragraph":
            continue
        text = block.get("text") or "".join(r.get("text", "") for r in block.get("runs", []))
        total += len(re.findall(r"[\u4e00-\u9fff]", text)) + len(re.findall(r"[A-Za-z0-9]+", text))
    return total, len(data.get("references") or [])


def _review_template(result):
    warnings = {F.warning_key(w): "请填写：这条提示怎么处理的——" + F.warning_text(w)[:100] for w in result["warnings"]}
    static = result["validation_scope"] == "static-only"
    return {
        "vision": not static,
        "pages": {
            str(p["page"]): f"请填写：第{p['page']}页实际看到的情况（有无缺字、重叠、裁切、空白页，日期/目录/页眉是否正常）"
            for p in result["pages"]
        },
        "content_check": "请填写：需求与内容保全——对照用户要求和原文，正文、图表、公式、链接是否齐全，封面信息和留空是否正确",
        "layout_check": "请填写：格式与版面——字体字号、标题层级、目录与页码、分页、表格图片是否正常",
        "delivery_check": "请填写：最终交付——发现的问题是否都已改好，交付的是不是用户要的这一个 Word",
        "warnings": warnings,
    }


def cmd_run(args):
    import docx
    import delivery as D
    import engine as E
    import intake as I
    from cover_date_audit import CoverOverflow
    from cover_fit import MAX_LEVEL

    root, _ = _work(args.work)
    stopped = F.stopped(root)
    if stopped and not str(getattr(args, "user_says_retry", None) or "").strip():
        print(f"{stopped}，自动重试已经停下。先把还没解决的问题（{root / 'issues.json'}）如实告诉用户、问用户怎么处理；"
              "用户要求再试时，用 run --user-says-retry \"用户原话\" 重新开始计数。")
        return EXIT_ASK_USER
    if stopped:
        F.reset_rounds(root, args.user_says_retry)
    state = _intake(root)
    if state["stage"] != "ready":
        print(I.summary(state))
        print("\n→ 封面信息还不全。先问用户（可选“留空自己填”），写进 job.json 后重新运行 run。")
        return EXIT_ASK_USER
    writing = _task(root) == "write" and _written(root) is not None
    if writing or _added_references(root):
        problems = (_writing_problems(root, _written(root)) if writing else []) + _added_reference_problems(root)
        if problems:
            def fix(text):
                where = (getattr(text, "edit", None) or "content.json / refs.json").split(" ")[0]
                return f"按上面的说明改 {where} 后重新 run；不要为了通过检查编造数字或出处。"

            cause = "写作模式（task 为 write）的硬性要求没满足。" if writing else "补进论文的参考文献必须先核实是真实存在的。"
            issues = [
                dict({"code": "WRITING_CHECK", "stage": "写作检查", "message": str(text),
                      "cause": cause, "fix": fix(text)},
                     **({"where": {"edit": text.edit}} if getattr(text, "edit", None) else {}))
                for text in problems
            ]
            return _not_accepted(root, F.AcceptanceFailed("写作检查", issues))
    if args.office in ("word", "libreoffice", "wps", "wps-manual"):
        engine = "wps" if args.office.startswith("wps") else args.office
        D.configure_office(root, engine, "manual" if args.office == "wps-manual" else "auto")
    if (root / "content.json").is_file() == (root / "mapping.json").is_file():
        raise ValueError(
            "WORK 里同时有 content.json 和 mapping.json，只能保留一个：新写文档用 content.json，已有 Word 用 mapping.json。"
            if (root / "content.json").is_file()
            else "还没有内容：新写文档请写 content.json；已有 Word 请先运行 inspect。"
        )
    level = 0
    while True:
        try:
            draft = _build(root, level)
        except ValueError as exc:
            return _not_accepted(root, F.AcceptanceFailed("排版", F.from_message(str(exc), root)))
        try:
            result = D.prepare(
                root, draft, root / "intake.json", args.office_note, "none" if args.office == "none" else "auto"
            )
            break
        except F.AcceptanceFailed as exc:
            return _not_accepted(root, exc)
        except CoverOverflow as exc:
            used = int(E.get_vars(docx.Document(draft)).get("XAUFE_COVER_FIT", "0"))
            if exc.too_high:
                pages = ((_read(root / "job.json", "job.json").get("format") or {}).get("page_twips") or {})
                changed = "、".join(f"{k}={v}" for k, v in pages.items())
                return _not_accepted(root, F.AcceptanceFailed("封面", [{
                    "code": "COVER_DATE_TOO_HIGH", "stage": "封面",
                    "message": "封面的完成日期在第一页上，但比底部保留区高：" + str(exc),
                    "where": dict({"location": "第 1 页（封面）"},
                                  **({"edit": f"job.json 的 format.page_twips（{changed}）"} if changed else {})),
                    "expected": "日期在第一页版心最底部", "actual": "日期浮在保留区上面",
                    "cause": ("不是封面太长：页面下方可用的高度变小了。" + (
                        f"job.json 改了页边距（{changed}），页脚放不进这么窄的下边距，正文区的底边被顶上去了。"
                        if changed else "页脚比下边距高，或渲染软件的页面设置和文件不同。")),
                    "fix": ("先核对 page_twips：页脚距离 footer 要比下边距 bottom 至少小 300（一行页脚的高度），"
                            "不是用户要的值就删掉这项恢复模板默认；用户要的就是这么窄的边距，就问用户是把 footer 改小"
                            "（不小于 280，再小打印机印不到）还是把 bottom 放宽。不要缩短封面内容，也不用换字体。"),
                }]))
            if used >= MAX_LEVEL and exc.missing_fonts:
                return _not_accepted(root, F.AcceptanceFailed("封面", F.with_job_suspects(root, [{
                    "code": "COVER_FONTS_MISSING", "stage": "封面",
                    "message": "封面在这台机器上放不下：" + str(exc),
                    "where": {"location": "第 1 页（封面）"},
                    "expected": "封面信息和完成日期都在第一页",
                    "actual": "渲染时没有" + "、".join(exc.missing_fonts) + "，用的是替代字体（"
                              + "、".join(exc.rendered_fonts)[:120] + "）",
                    "cause": "替代字体的字宽和行高和要求的字体不同，封面被撑出第一页，很可能不是内容太长。Word 文件本身的字体设置没有变。",
                    "fix": "不要自己装软件或字体，也不要先让用户缩短封面。问用户选一个：① 电脑上有 Word 或 WPS 就用它渲染（run --office word 或 run --office wps）；"
                           "② 用户自己装上" + "、".join(exc.missing_fonts) + "后再 run；"
                           "③ 用户同意这次不做实际分页检查，按静态交付（run --office none，会如实告诉用户哪些检查没做）。"
                           "装好字体或换了软件后仍放不下，才问用户缩短封面内容。",
                }])))
            if used >= MAX_LEVEL:
                values = {I.LABELS.get(k, k): str(v) for k, v in state.get("metadata", {}).items() if k != "date" and v}
                longest = sorted(values.items(), key=lambda kv: -len(kv[1]))[:2]
                return _not_accepted(root, F.AcceptanceFailed("封面", F.with_job_suspects(root, [{
                    "code": "COVER_TOO_LONG", "stage": "封面",
                    "message": "封面已经收紧到最紧，第一页仍放不下：" + str(exc),
                    "where": {"location": "第 1 页（封面）", "edit": "job.json 的 info"},
                    "expected": "封面信息和完成日期都在第一页",
                    "actual": "最长的是" + "、".join(f"{k}（{len(v)} 字：“{v[:20]}”）" for k, v in longest),
                    "cause": "封面字段太长，空白和行距已经收到最紧。",
                    "fix": "问用户：缩短这几项，或同意缩小封面字号（job.json 的 format.cover_fonts）。不要自己删改封面内容。",
                }])))
            level = used + 1
            print(f"封面太长，日期被挤出第一页；自动收紧封面空白（第 {level} 级）后重试……")
    F.record_round(root, [])
    if _task(root) == "write" and _written(root) is not None:
        content = _written(root)
        notes = [n for n in (_weak_support(root, content), _off_topic(root, content)) if n]
        if notes:
            result["warnings"].extend(notes)
            D.add_review_warnings(root, notes)
    _write(root / "review.json", _review_template(result))
    print(I.summary(state).splitlines()[0])
    stats = _body_length(root)
    if stats:
        print(
            f"正文约 {stats[0]} 字（不含标点、标题、摘要、参考文献），参考文献 {stats[1]} 条。"
            "写作时用户没提要求的，默认要正文不少于 5000 字、参考文献至少 6 篇；只排版时不用管。"
        )
    if result["validation_scope"] == "static-only":
        print("\n排版完成，但没有做实际分页和页面预览（原因见下）。")
        print("需要在最终回复里告诉用户：" + result["user_notice"])
    else:
        used = D._renderer_name(root)
        shown = dict(D.RENDERER_NAMES, word="Word").get(used.lower(), used) or "未知软件"
        print(f"\n排版完成，共 {len(result['pages'])} 页（分页和页图用 {shown} 生成）。请逐页打开下面的图片查看：")
        for page in result["pages"]:
            print(f"  第{page['page']}页：{page['path']}")
    if result["warnings"]:
        print("\n需要处理的提示：")
        for w in result["warnings"]:
            print("  - " + F.warning_text(w))
    if _task(root) == "format":
        checked = _auto_original_refs(root)
        if checked:
            print("\n" + checked)
    review = root / "review.json"
    if result["validation_scope"] == "static-only":
        print(f"\n→ 下一步：这次没有页图。核对内容和格式后，把 {review} 里所有“请填写”换成你的实际检查结果，然后运行 finish。")
    else:
        print(f"\n→ 下一步：看完每一页后，把 {review} 里所有“请填写”换成你的实际检查结果，然后运行 finish。")
    print("  发现问题就改 job.json / content.json / mapping.json 后重新 run，不要直接改 WORK 里的 Word。")
    return 0


def _not_accepted(root, exc):
    """Print located, actionable issues; count rounds so the agent knows when to stop."""
    _write(root / "issues.json", {"stage": exc.stage, "issues": exc.issues})
    rounds, stalled, repeated = F.record_round(root, exc.issues)
    kinds = len(F.group(exc.issues))
    count = f"{len(exc.issues)} 处，归成 {kinds} 类" if kinds != len(exc.issues) else f"{kinds} 个问题"
    print(f"{exc.stage}没通过，共 {count}（完整列表：{root / 'issues.json'}）：\n")
    print(F.render_text(exc.issues))
    print("\n→ 改 job.json / content.json / mapping.json 后重新 run；不要直接改 WORK 里的 Word，也不要为了通过检查改动用户原文。")
    if stalled >= F.MAX_STALLED or rounds >= F.MAX_ROUNDS:
        why = f"连续 {stalled} 次修改都没有解决任何一个问题" if stalled >= F.MAX_STALLED else f"已经连续 {rounds} 次没通过"
        print(
            f"\n{why}。停止自动重试：把上面还没解决的问题如实告诉用户"
            "（哪一页、哪一段、差在哪里、你试过什么），问用户怎么处理；不要交付没通过检查的文件。"
        )
    elif repeated:
        print(f"\n注意：问题和上一次完全一样（这次修改没有效果）。换一种改法；连续 {F.MAX_STALLED} 次没有进展就停下来报告用户。")
    return EXIT_NOT_ACCEPTED


# ---------------------------------------------------------------- finish / cancel
def cmd_finish(args):
    import delivery as D

    root, _ = _work(args.work)
    D.attest(root, _read(root / "review.json", "review.json"))
    task = _task(root)
    checklist = _reference_checklist(root) if task == "write" else _original_reference_report(root)
    result = D.publish(root)
    print("已交付：" + result["output"])
    print("工作区已清理。")
    if result["user_notice"]:
        print("\n必须在给用户的回复里说明：" + result["user_notice"])
    else:
        print("\n已完成排版、分页和逐页检查。")
    if checklist and task == "write":
        print("\n必须在给用户的回复里附上这张参考文献核实清单（原样照抄）：\n" + checklist)
    elif checklist and checklist.startswith("这是只排版"):
        print("\n必须在给用户的回复里说明：" + checklist)
    elif checklist:
        print("\n必须在给用户的回复里附上原文参考文献的核对结果（原样照抄）：\n" + checklist)
    return 0


def cmd_refs(args):
    import refcheck as R

    root, _ = _work(args.work)
    if args.from_document:
        if (root / "refs.json").exists():
            raise ValueError("WORK 里已经有 refs.json：要从文档重新生成就先删掉它，否则去掉 --from-document 直接核对。")
        lines = _document_references(root)
        if not lines:
            raise ValueError("文档里没有参考文献条目：content.json 的 references 为空，或 mapping.json 里没有 role 为 reference 的段落。")
        _write(root / "refs.json", [R.parse_line(line) for line in lines])
        print(f"已按文档里的 {len(lines)} 条参考文献生成 {root / 'refs.json'}（字段是从原文猜的，text 是原文原样）。\n")
    entries = _read(root / "refs.json", "refs.json")
    if not isinstance(entries, list):
        raise ValueError("refs.json 应该是一个列表：[{...}, {...}]，写法见 references/citations.md")
    results = R.check_all(entries)
    _keep_earlier_verdicts(root, results, entries)
    originals = None
    for result, entry in zip(results, entries):
        if isinstance(entry, dict):
            result["supports"] = entry.get("supports", "")
            result["relevance"] = entry.get("relevance", "")
            if entry.get("from_user"):
                originals = _original_text(root) if originals is None else originals
                result["provenance"] = _user_provenance(entry, originals)
                if not result["provenance"]:
                    result["notes"].append("标了 from_user，但登记的原稿里找不到这一条，也没写 user_said（用户给这条文献时的原话）")
    _save_refs_result(
        root,
        {"refs_sha256": hashlib.sha256((root / "refs.json").read_bytes()).hexdigest(), "min": args.min, "results": results},
    )
    only_report = (root / "job.json").is_file() and _task(root) == "format"
    print(R.report(results, entries, args.min, report_only=only_report))
    return 0


def _entry_sha(entry):
    return hashlib.sha256(json.dumps(entry, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _keep_earlier_verdicts(root, results, entries):
    """A refs run that could not reach the checking services says nothing new: an unchanged entry
    keeps the verdict an earlier refs run reached (a 排除 must not turn back into 连不上)."""
    earlier = {}
    if (root / "refs-result.json").is_file() and not _refs_result_forged(root):
        earlier = {r.get("entry_sha256"): r for r in _read(root / "refs-result.json", "refs-result.json").get("results", [])
                   if r.get("entry_sha256") and not r.get("unreachable")}
    for i, (result, entry) in enumerate(zip(results, entries)):
        sha = _entry_sha(entry)
        if result.get("unreachable") and sha in earlier:
            results[i] = {**earlier[sha], "notes": list(earlier[sha].get("notes", []))
                          + ["这次连不上核对服务，沿用上次核对的结论"]}
        results[i]["entry_sha256"] = sha


def cmd_read(args):
    """Print the text of a .docx so the agent can read the user's file."""
    import docx
    from docx.oxml.ns import qn

    source = Path(args.docx).expanduser().resolve()
    if source.suffix.lower() != ".docx":
        raise ValueError("只能读取 .docx。旧版 .doc 请先用 Word/WPS“另存为 .docx”。")
    doc = docx.Document(source)
    body = doc.element.body
    print(f"《{source.name}》的文字内容（按顺序）：")
    for element in body.iterchildren():
        if element.tag == qn("w:p"):
            text = "".join(element.xpath(".//w:t/text()")).strip()
            if text:
                print(text)
        elif element.tag == qn("w:tbl"):
            for row in element.xpath("./w:tr"):
                cells = ["".join(c.xpath(".//w:t/text()")).strip() for c in row.xpath("./w:tc")]
                print("[表格] " + " | ".join(cells))
    images = len(body.xpath(".//w:drawing"))
    if images:
        print(f"[另有 {images} 张图片]")
    return 0


def cmd_preview(args):
    """Page images of any .docx (e.g. one edited by hand), kept inside WORK."""
    import shutil
    import pymupdf as fitz
    from office_backend import export_pdf, OfficeUnavailable

    root, _ = _work(args.work)
    source = Path(args.docx).expanduser().resolve()
    out = root / "extra-preview" / source.stem
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    pdf = out / "preview.pdf"
    try:
        export_pdf(source, pdf)
    except OfficeUnavailable as exc:
        raise ValueError("没有可用的办公软件，生不成页图：" + str(exc)) from exc
    with fitz.open(pdf) as rendered:
        print(f"共 {len(rendered)} 页，请逐页打开查看：")
        for index, page in enumerate(rendered, 1):
            image = out / f"page-{index:03d}.png"
            page.get_pixmap(dpi=110).save(image)
            print(f"  第{index}页：{image}")
    try:
        from cover_date_audit import verify_pdf

        verify_pdf(source, pdf)
    except Exception as exc:
        print("注意（封面检查）：" + str(exc))
    from render_audit import audit

    rendered = audit(source, pdf)
    _write(out / "render-audit.json", rendered)
    if rendered["errors"]:
        print("\n渲染后版面检查发现问题：\n" + F.render_text(F.from_render(rendered)))
    for w in rendered["warnings"]:
        print("注意：" + F.warning_text(w))
    return 0


def cmd_adopt(args):
    """A requirement the scripts cannot meet, handled by hand on a copy of the accepted output."""
    import delivery as D

    root, _ = _work(args.work)
    try:
        result = D.adopt(root, args.docx, args.note)
    except F.AcceptanceFailed as exc:
        return _not_accepted(root, exc)
    _write(root / "review.json", _review_template(result))
    print(f"手工处理过的 Word 已通过渲染后检查，共 {len(result['pages'])} 页。请逐页打开查看：")
    for page in result["pages"]:
        print(f"  第{page['page']}页：{page['path']}")
    print("\n需要处理的提示（每条都要在 review.json 的 warnings 里写怎么处理的）：")
    for w in result["warnings"]:
        print("  - " + F.warning_text(w))
    print(f"\n→ 下一步：看完每一页后填写 {root / 'review.json'}，然后运行 finish。交付时会把手工处理的内容告诉用户。")
    return 0


def cmd_cancel(args):
    import delivery as D

    D.clean(args.work)
    print("已删除本次工作区，没有交付文件。")
    return 0


def main(argv=None):
    parser = argparse.ArgumentParser(description="西安财经大学文档格式：唯一入口")
    sub = parser.add_subparsers(dest="cmd", required=True)
    sub.add_parser("compare", help="打印 A/B 对比表（用户还没选版式时发给用户）")
    a = sub.add_parser("new", help="开始一个任务，创建工作区 WORK")
    a.add_argument("--out", required=True, help="最终 Word 的保存路径（.docx，不能已存在）")
    a.add_argument("--source", action="append", default=[], help="用户提供的原稿文件，可写多次")
    for name, text in (
        ("check", "检查 job.json，告诉你还缺什么"),
        ("run", "排版、分页、生成页图"),
        ("finish", "读取 review.json，交付 Word 并清理"),
        ("cancel", "放弃任务并清理"),
    ):
        a = sub.add_parser(name, help=text)
        a.add_argument("--work", required=True)
        if name == "run":
            a.add_argument(
                "--office",
                default="auto",
                choices=["auto", "word", "libreoffice", "wps", "wps-manual", "none"],
                help="none = 用户要求不启动办公软件",
            )
            a.add_argument("--office-note", help="办公软件找到了却无法调用时，写实际原因")
            a.add_argument("--user-says-retry", metavar="用户原话",
                           help="自动重试已经停下后，用户要求再试时写用户的原话，计数从头开始")
    a = sub.add_parser("read", help="打印用户 Word 文件里的文字（看原稿内容用）")
    a.add_argument("--docx", required=True)
    a = sub.add_parser("preview", help="给任意 .docx 生成页图（手工处理过的文档用）")
    a.add_argument("--work", required=True)
    a.add_argument("--docx", required=True)
    a = sub.add_parser("adopt", help="脚本做不到的要求：验收在成品副本上手工处理过的 Word")
    a.add_argument("--work", required=True)
    a.add_argument("--docx", required=True)
    a.add_argument("--note", required=True, help="手工改了什么，会告诉用户")
    a = sub.add_parser("inspect", help="读取已有 Word，生成可核对的 mapping.json；或把 .txt/.md 原稿逐字导入 content.json")
    a.add_argument("--work", required=True)
    given = a.add_mutually_exclusive_group(required=True)
    given.add_argument("--docx")
    given.add_argument("--text", help="只排版、原稿是 .txt/.md：逐字导入 content.json")
    a = sub.add_parser("refs", help="核对 WORK/refs.json 里的参考文献（需要联网）")
    a.add_argument("--work", required=True)
    a.add_argument("--min", type=int, default=6, help="至少几篇；用户另有要求时填用户的数")
    a.add_argument("--from-document", action="store_true",
                   help="按文档里已有的参考文献生成 refs.json 再核对（只排版时用：只报告，不改原文）")
    args = parser.parse_args(argv)
    handler = {
        "adopt": cmd_adopt,
        "compare": cmd_compare,
        "new": cmd_new,
        "check": cmd_check,
        "inspect": cmd_inspect,
        "preview": cmd_preview,
        "read": cmd_read,
        "refs": cmd_refs,
        "run": cmd_run,
        "finish": cmd_finish,
        "cancel": cmd_cancel,
    }[args.cmd]
    try:
        return handler(args) or 0
    except Exception as exc:  # every failure is reported as one readable line
        from wps_exchange import ManualExportRequired

        if isinstance(exc, ManualExportRequired):
            print("需要在 WPS 里手动导出 PDF，请求文件：" + str(exc.request), file=sys.stderr)
            print("做法见 references/special-cases.md 的 WPS 一节。", file=sys.stderr)
            return 4
        if args.cmd == "run" and isinstance(exc, ValueError):
            # An input the scripts refuse is a failed round too: located if possible, and counted.
            try:
                root = _work(args.work)[0]
            except Exception:
                root = None
            if root is not None:
                return _not_accepted(root, F.AcceptanceFailed("输入检查", F.from_message(str(exc), root, "输入检查")))
        print("出错了：" + str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    # Without -X utf8 a GBK/ASCII console must not crash on a character it cannot show.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")
    sys.exit(main())
